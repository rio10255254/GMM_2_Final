"""Verify that each fitted disk uses the P-pillar SEM cell for its image ID."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
PROJECT_ROOT = Path(__file__).resolve().parents[3]


PROJECT = Path(str(PROJECT_ROOT))
ROOT = Path(__file__).resolve().parent
IDS = [[36, 35, 34], [30, 29, 28], [24, 23, 22]]


def main():
    source_workbook = PROJECT / "Thesis_MOSR_Evidence" / "SEM_data" / "via_pilar.xlsx"
    workbook = load_workbook(source_workbook, read_only=True, data_only=True)
    sheet = workbook["SEM data"]
    source_mapping = PROJECT / "Thesis_MOSR_Evidence" / "SEM_data" / "ordering_result_20260729" / "SEM_center36_file_mapping.csv"
    with source_mapping.open(newline="", encoding="utf-8-sig") as f:
        full_map = {int(r["global_pillar_id"]): r for r in csv.DictReader(f)}
    with (ROOT / "sem_9_mapping.csv").open(newline="", encoding="utf-8") as f:
        fitted_map = list(csv.DictReader(f))
    assert len(fitted_map) == 9

    matlab_source = PROJECT / "PSF_Fitting" / "center_9pillar_sem_radius_nm.m"
    source_text = matlab_source.read_text(encoding="utf-8")
    block = re.search(r"radius_nm\s*=\s*\[(.*?)\];", source_text, re.S)
    assert block is not None
    hardcoded = np.array([float(x) for x in re.findall(r"\d+\.\d+", block.group(1))]).reshape(3, 3)
    physical = np.array([[float(sheet.cell(row, col).value) for col in range(10, 13)] for row in range(8, 11)])
    assert np.allclose(physical, hardcoded, atol=1e-9, rtol=0)

    checked = []
    for index, row in enumerate(fitted_map):
        image_row, image_col = divmod(index, 3)
        expected_id = IDS[image_row][image_col]
        cell = sheet.cell(10 - image_row, 10 + image_col)
        historical = full_map[expected_id]
        assert int(row["image_row"]) == image_row + 1
        assert int(row["image_col"]) == image_col + 1
        assert int(row["sem_id_working"]) == expected_id
        assert row["workbook_cell"] == f"SEM data!{cell.coordinate}"
        assert abs(float(row["sem_radius_nm"]) - float(cell.value)) < 1e-9
        assert abs(float(row["sem_diameter_nm"]) - 2 * float(cell.value)) < 1e-9
        assert int(historical["grid_row_1_to_6"]) == cell.row - 4
        assert int(historical["grid_col_1_to_6"]) == cell.column - 9
        assert row["pillar_sem_file"] == historical["pillar_sem_file"]
        assert row["pillar_sem_file"].startswith("P-")
        assert row["sem_local_corner"] == historical["local_corner"]
        checked.append({"image_row": image_row + 1, "image_col": image_col + 1,
                        "sem_id": expected_id, "cell": cell.coordinate,
                        "P_file": row["pillar_sem_file"],
                        "diameter_nm": float(row["sem_diameter_nm"])})

    current_config_path = PROJECT / "Thesis_MOSR_Evidence" / "Final_Data_Figures" / "01_CURRENT_THESIS_RESULTS_V42_CENTER9" / "config" / "sequence_center9_50x20x.json"
    current = json.loads(current_config_path.read_text(encoding="utf-8"))["sequences"]["20x_1um"]
    assert [int(c[0]) for c in current["centers"]] == [item for row in IDS for item in row]
    transfer = json.loads((ROOT / "current_20x_identity_transfer.json").read_text(encoding="utf-8"))
    assert transfer["max_center_residual_px"] < 1.5
    archive = Path(r"H:\From_2TB_D_disk\13_thesis\SEM_data")
    archive_names = ["via_pilar.xlsx", "P-01.tiff", "P-02.tiff", "P-05.tiff", "P-06.tiff"]
    def digest(path):
        h = hashlib.sha256()
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()
    for name in archive_names:
        assert digest(PROJECT / "Thesis_MOSR_Evidence" / "SEM_data" / name) == digest(archive / name)
    result = {
        "status": "PASS",
        "P_pillar_data_count": 9,
        "SEM_workbook": str(source_workbook),
        "SEM_file_mapping": str(source_mapping),
        "current_20x_ID_source": str(current_config_path),
        "current_to_new_max_center_residual_px": transfer["max_center_residual_px"],
        "sem_ids_image_rows": IDS,
        "archive_files_identical_by_sha256": archive_names,
        "checked": checked,
    }
    (ROOT / "sem_mapping_audit.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "checked"}, indent=2))


if __name__ == "__main__":
    main()
