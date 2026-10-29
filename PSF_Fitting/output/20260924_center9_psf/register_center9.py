"""Register the 25 September frames to the existing July center-nine calibration.

This only locates the physical pillars. The PSF fit itself uses the existing
MATLAB center-nine pair builder and BADS fitter.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import cv2
import numpy as np
from openpyxl import load_workbook
PROJECT_ROOT = Path(__file__).resolve().parents[3]


ROOT = Path(str(PROJECT_ROOT))
CAPTURE = ROOT / "20260924"
OUT = Path(__file__).resolve().parent
REFERENCE_JSON = (
    ROOT / "PSF_Fitting" / "output" / "pillar_via_junction_psf_20260713"
    / "20x" / "selected_frame.json"
)
PILLAR_IDS = np.array([[36, 35, 34], [30, 29, 28], [24, 23, 22]])
CURRENT_20X_CONFIG = (
    ROOT / "Thesis_MOSR_Evidence" / "Final_Data_Figures"
    / "01_CURRENT_THESIS_RESULTS_V42_CENTER9" / "config"
    / "sequence_center9_50x20x.json"
)


def sift_features(paths: list[Path]):
    sift = cv2.SIFT_create(nfeatures=8500, contrastThreshold=0.015)
    result = []
    for path in paths:
        img = cv2.imread(str(path), cv2.IMREAD_GRAYSCALE)
        if img is None:
            raise RuntimeError(f"Cannot read {path}")
        k, d = sift.detectAndCompute(img, None)
        result.append((img, k, d))
    return result


def match_affine(a, b):
    _, ka, da = a
    _, kb, db = b
    matcher = cv2.FlannBasedMatcher(dict(algorithm=1, trees=5), dict(checks=80))
    matches = matcher.knnMatch(da, db, k=2)
    good = [m for m, n in matches if m.distance < 0.7 * n.distance]
    src = np.float32([ka[m.queryIdx].pt for m in good])
    dst = np.float32([kb[m.trainIdx].pt for m in good])
    affine, inliers = cv2.estimateAffinePartial2D(
        src, dst, method=cv2.RANSAC, ransacReprojThreshold=2,
        maxIters=3000, confidence=0.999,
    )
    n = int(inliers.sum()) if inliers is not None else 0
    if affine is None or n < 70:
        raise RuntimeError(f"Insufficient feature match: {len(good)} good, {n} inliers")
    return affine, len(good), n


def centroid_refine(img: np.ndarray, expected: np.ndarray):
    center = expected.copy()
    background = 0.0
    contrast = 0.0
    for _ in range(3):
        x0, y0 = np.floor(center).astype(int)
        if x0 < 15 or y0 < 15 or x0 + 15 >= img.shape[1] or y0 + 15 >= img.shape[0]:
            raise RuntimeError(f"Pillar too close to image boundary: {center}")
        patch = img[y0 - 14:y0 + 15, x0 - 14:x0 + 15].astype(np.float32)
        yy, xx = np.mgrid[-14:15, -14:15]
        rr = np.hypot(xx, yy)
        background = float(np.median(patch[(rr >= 11) & (rr <= 14)]))
        weights = np.maximum(background - patch, 0)
        weights[rr > 10] = 0
        if weights.sum() < 1:
            raise RuntimeError(f"No dark pillar near {expected}")
        delta = np.array([(weights * xx).sum(), (weights * yy).sum()]) / weights.sum()
        center = np.array([x0, y0], dtype=float) + delta
        contrast = background - float(np.mean(patch[rr <= 5]))
    return center, contrast, background


def main():
    cv2.setRNGSeed(0)
    OUT.mkdir(parents=True, exist_ok=True)
    files = sorted(CAPTURE.glob("*.bmp"))
    if len(files) != 25:
        raise RuntimeError(f"Expected 25 BMP frames, found {len(files)}")
    ref = json.loads(REFERENCE_JSON.read_text(encoding="utf-8"))
    old_path = Path(ref["selected_image"])
    old_features = sift_features([old_path])[0]
    data = sift_features(files)

    # Adjacent overlapping views form a connected 5 x 5 registration graph.
    edges = []
    for i in range(25):
        for j in (i + 1, i + 5):
            if j >= 25 or (j == i + 1 and i % 5 == 4):
                continue
            affine, good, n = match_affine(data[i], data[j])
            scale = float(np.hypot(affine[0, 0], affine[0, 1]))
            if abs(scale - 1) > 0.015:
                raise RuntimeError(f"Unexpected scale between {i + 1} and {j + 1}: {scale}")
            midpoint = np.array([1032.0, 772.0, 1.0])
            shift = affine @ midpoint - midpoint[:2]
            edges.append((i, j, shift, n, good, scale))

    # Weighted least squares uses all loops, rather than one fragile path.
    design = np.zeros((len(edges) + 1, 25))
    target = np.zeros((len(edges) + 1, 2))
    weight = np.ones(len(edges) + 1)
    for row, (i, j, shift, n, _, _) in enumerate(edges):
        design[row, i] = -1
        design[row, j] = 1
        target[row] = shift
        weight[row] = np.sqrt(n)
    design[-1, 12] = 1
    weight[-1] = 1000
    transforms, *_ = np.linalg.lstsq(design * weight[:, None], target * weight[:, None], rcond=None)
    residual = (design[:-1] @ transforms) - target[:-1]
    residual_norm = np.linalg.norm(residual, axis=1)
    if residual_norm.max() > 3:
        raise RuntimeError(f"Mosaic loop disagreement: {residual_norm.max():.3f} px")

    ref_affine, ref_good, ref_inliers = match_affine(old_features, data[12])
    ref_scale = float(np.hypot(ref_affine[0, 0], ref_affine[0, 1]))
    if abs(ref_scale - 1) > 0.025:
        raise RuntimeError(f"July to September scale mismatch: {ref_scale}")
    old_centers = np.asarray(ref["pillar_centers_xy"], dtype=float) - 1
    center_frame_centers = (ref_affine[:, :2] @ old_centers.T).T + ref_affine[:, 2]

    records = []
    previews = []
    for i, (path, (img, _, _)) in enumerate(zip(files, data)):
        predicted = center_frame_centers + transforms[i]
        corrected = []
        for p_idx, expected in enumerate(predicted):
            actual, contrast, background = centroid_refine(img, expected)
            correction = float(np.linalg.norm(actual - expected))
            if correction > 5.0 or contrast < 12:
                raise RuntimeError(
                    f"Bad center match: {path.name} pillar {p_idx + 1}, "
                    f"correction={correction:.2f}, contrast={contrast:.1f}"
                )
            corrected.append(actual)
            records.append({
                "image_index": i + 1,
                "image": path.name,
                "pillar_row": p_idx // 3 + 1,
                "pillar_col": p_idx % 3 + 1,
                "sem_id_working": int(PILLAR_IDS.flat[p_idx]),
                # OpenCV uses zero-based pixel coordinates; MATLAB interp2
                # and imfindcircles use one-based coordinates.
                "x": f"{actual[0] + 1:.6f}",
                "y": f"{actual[1] + 1:.6f}",
                "registration_x": f"{expected[0] + 1:.6f}",
                "registration_y": f"{expected[1] + 1:.6f}",
                "center_correction_px": f"{correction:.6f}",
                "dark_contrast_gray": f"{contrast:.3f}",
                "background_gray": f"{background:.3f}",
            })
        corrected = np.asarray(corrected)
        xmin = max(0, int(np.floor(corrected[:, 0].min() - 16)))
        ymin = max(0, int(np.floor(corrected[:, 1].min() - 16)))
        xmax = min(img.shape[1], int(np.ceil(corrected[:, 0].max() + 16)))
        ymax = min(img.shape[0], int(np.ceil(corrected[:, 1].max() + 16)))
        crop = cv2.cvtColor(img[ymin:ymax, xmin:xmax], cv2.COLOR_GRAY2BGR)
        for p_idx, xy in enumerate(corrected):
            pos = np.round(xy - np.array([xmin, ymin])).astype(int)
            cv2.circle(crop, tuple(pos), 10, (0, 0, 255), 1)
            cv2.putText(crop, str(PILLAR_IDS.flat[p_idx]), tuple(pos + [7, -7]),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.27, (0, 0, 255), 1)
        crop = cv2.resize(crop, (180, 180))
        cv2.putText(crop, f"{i+1:02d} {path.stem}", (4, 16),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.42, (0, 0, 255), 1)
        previews.append(crop)

    with (OUT / "center9_locations.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    montage = np.vstack([np.hstack(previews[i:i + 5]) for i in range(0, 25, 5)])
    cv2.imwrite(str(OUT / "center9_25_frames.png"), montage)
    current_config = json.loads(CURRENT_20X_CONFIG.read_text(encoding="utf-8"))["sequences"]["20x_1um"]
    current_ids = [int(row[0]) for row in current_config["centers"]]
    if current_ids != PILLAR_IDS.ravel().tolist():
        raise RuntimeError("Current 20x SEM ordering differs from expected image rows")
    current_raw = Path(ROOT / current_config["raw_dir"]) / current_config["first_frame"]
    current_affine, current_good, current_inliers = match_affine(sift_features([current_raw])[0], data[12])
    current_centers_zero_based = np.asarray([row[1:3] for row in current_config["centers"]], dtype=float) - 1
    transferred = (current_affine[:, :2] @ current_centers_zero_based.T).T + current_affine[:, 2]
    actual_center_frame = np.asarray([[float(r["x"]), float(r["y"])] for r in records[12 * 9:13 * 9]]) - 1
    transfer_errors = np.linalg.norm(transferred - actual_center_frame, axis=1)
    if transfer_errors.max() > 0.2 * 5 / (3.45 / 20):
        raise RuntimeError(f"Current 20x identity transfer differs by {transfer_errors.max():.2f} pixels")
    identity_evidence = {
        "current_20x_config": str(CURRENT_20X_CONFIG),
        "current_20x_raw_image": str(current_raw),
        "new_image": str(files[12]),
        "inlier_matches": current_inliers,
        "good_matches": current_good,
        "affine_current_to_new": current_affine.tolist(),
        "sem_ids_image_rows": PILLAR_IDS.tolist(),
        "transferred_centers_zero_based_xy": transferred.tolist(),
        "located_centers_zero_based_xy": actual_center_frame.tolist(),
        "center_residual_px": transfer_errors.tolist(),
        "max_center_residual_px": float(transfer_errors.max()),
    }
    (OUT / "current_20x_identity_transfer.json").write_text(
        json.dumps(identity_evidence, indent=2) + "\n", encoding="utf-8"
    )
    manifest = {
        "capture_count": len(files),
        "locations_count": len(records),
        "reference_image": str(old_path),
        "reference_json": str(REFERENCE_JSON),
        "july_to_september_center_frame_good_matches": ref_good,
        "july_to_september_center_frame_inliers": ref_inliers,
        "july_to_september_affine": ref_affine.tolist(),
        "adjacent_edges": len(edges),
        "min_adjacent_inliers": min(e[3] for e in edges),
        "max_loop_error_px": float(residual_norm.max()),
        "mean_loop_error_px": float(residual_norm.mean()),
        "max_center_correction_px": max(float(r["center_correction_px"]) for r in records),
        "min_dark_contrast_gray": min(float(r["dark_contrast_gray"]) for r in records),
        "sem_id_display_rows_working": PILLAR_IDS.tolist(),
        "sem_id_note": "Directly transferred current locked 20x row-major IDs from sequence_center9_50x20x.json; physical SEM cells J8:L10 are flipped vertically into this image orientation.",
        "current_20x_identity_transfer": str(OUT / "current_20x_identity_transfer.json"),
        "center_coordinates": "MATLAB one-based image coordinates",
    }
    (OUT / "registration_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    workbook = load_workbook(
        ROOT / "Thesis_MOSR_Evidence" / "SEM_data" / "via_pilar.xlsx",
        data_only=True, read_only=True,
    )
    sheet = workbook["SEM data"]
    mapping_file = (ROOT / "Thesis_MOSR_Evidence" / "SEM_data" / "ordering_result_20260729"
                    / "SEM_center36_file_mapping.csv")
    with mapping_file.open(newline="", encoding="utf-8-sig") as f:
        by_id = {int(row["global_pillar_id"]): row for row in csv.DictReader(f)}
    sem_rows = []
    # Current locked 20x sequence flips the physical J8:L10 block vertically.
    for image_row in range(3):
        for image_col in range(3):
            sem_id = int(PILLAR_IDS[image_row, image_col])
            worksheet_row = 10 - image_row
            worksheet_col = 10 + image_col
            cell = sheet.cell(worksheet_row, worksheet_col)
            radius = float(cell.value)
            mapping = by_id[sem_id]
            if not mapping["pillar_sem_file"].startswith("P-"):
                raise RuntimeError(f"SEM ID {sem_id} is not linked to a P image")
            sem_rows.append({
                "image_row": image_row + 1,
                "image_col": image_col + 1,
                "sem_id_working": sem_id,
                "workbook_cell": f"SEM data!{cell.coordinate}",
                "sem_radius_nm": f"{radius:.12f}",
                "sem_diameter_nm": f"{2 * radius:.12f}",
                "pillar_sem_file": mapping["pillar_sem_file"],
                "sem_local_corner": mapping["local_corner"],
            })
    with (OUT / "sem_9_mapping.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(sem_rows[0]))
        writer.writeheader()
        writer.writerows(sem_rows)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
