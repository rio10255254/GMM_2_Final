from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


INPUT_DOCX = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節"
    r"\第六章 實驗結果與分析_6.1.2方法與銳利度驗證完整版.docx"
)
OUTPUT_DOCX = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節"
    r"\第六章 實驗結果與分析_6.1.2門檻與指標定義完整版.docx"
)


def find_paragraph(doc: Document, text: str, *, startswith: bool = False) -> Paragraph:
    matches = []
    for paragraph in doc.paragraphs:
        candidate = paragraph.text.strip()
        if (startswith and candidate.startswith(text)) or (not startswith and candidate == text):
            matches.append(paragraph)
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph for {text!r}; found {len(matches)}")
    return matches[0]


def copy_paragraph_properties(source: Paragraph, target_element) -> None:
    if source._p.pPr is not None:
        target_element.insert(0, deepcopy(source._p.pPr))


def copy_run_properties(source: Paragraph, target_run) -> None:
    if source.runs and source.runs[0]._r.rPr is not None:
        target_run._r.insert(0, deepcopy(source.runs[0]._r.rPr))


def replace_text(paragraph: Paragraph, text: str) -> None:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def insert_paragraph_after(
    reference: Paragraph,
    text: str,
    prototype: Paragraph,
) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    copy_paragraph_properties(prototype, element)
    run = paragraph.add_run(text)
    copy_run_properties(prototype, run)
    return paragraph


def insert_definition_after(
    reference: Paragraph,
    label: str,
    definition: str,
    prototype: Paragraph,
) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    copy_paragraph_properties(prototype, element)
    paragraph.paragraph_format.keep_together = True

    label_run = paragraph.add_run(label)
    copy_run_properties(prototype, label_run)
    label_run.bold = True

    body_run = paragraph.add_run(definition)
    copy_run_properties(prototype, body_run)
    body_run.bold = False
    return paragraph


def main() -> None:
    if not INPUT_DOCX.exists():
        raise FileNotFoundError(INPUT_DOCX)

    doc = Document(INPUT_DOCX)
    body_prototype = find_paragraph(doc, "每顆 pillar 的量測以九宮格定位所得的中心為起點", startswith=True)
    heading_prototype = find_paragraph(doc, "1.1 CD 與邊緣寬度的量測方式")
    method_caption = find_paragraph(doc, "CD 與 10-90% 邊緣過渡寬度之實際量測流程")

    intro = find_paragraph(doc, "為評估擬合所得之點擴散函數是否具有實際的影像修復能力", startswith=True)
    replace_text(
        intro,
        "為評估擬合所得之點擴散函數是否具有實際的影像修復能力，本研究分別以 20 倍與 50 倍影像所擬合得到的 PSF 作為 Wiener 反卷積核，處理各自的 pillar 九宮格影像。Wiener 反卷積利用 PSF 反推原本被模糊的影像，同時以雜訊參數限制高頻雜訊被過度放大。後續先定義 CD 邊界、銳利度與誤差指標，再比較反卷積前後的結果，最後檢查影像銳化是否伴隨環狀假影。",
    )

    nsr_paragraph = find_paragraph(doc, "反卷積前先估計影像中的緩慢背景變化", startswith=True)
    replace_text(
        nsr_paragraph,
        "反卷積前先估計影像中的緩慢背景變化，再將暗色 pillar 轉換為正向結構後進行處理，最後還原為原本的灰階方向。Wiener 反卷積中的 NSR 為 noise-to-signal power ratio，表示雜訊功率與訊號功率之比；數值越小，反卷積越積極，但也越容易放大雜訊與振鈴，數值越大則會加強抑制而保留較多模糊。本研究的 20 倍與 50 倍影像共同使用 NSR = 0.0565，以避免逐倍率或逐顆 pillar 分別調整造成比較規則不一致。",
    )

    anchor = method_caption
    anchor = insert_paragraph_after(anchor, "1.2 對比門檻的意義", heading_prototype)
    anchor.paragraph_format.keep_with_next = True
    anchor = insert_paragraph_after(
        anchor,
        "本節所稱的 40%、50% 與 60% 門檻，是徑向灰階輪廓由 pillar 中心過渡至周圍背景時的相對對比位置，不是顯示影像時調整的亮度門檻、二值化機率，也不是 Wiener 反卷積的強度參數。設中心灰階為 Icore、背景灰階為 Ibg，比例 f 的灰階門檻定義為 T_f = Icore + f(Ibg - Icore)，其中 f 分別取 0.4、0.5 與 0.6。",
        body_prototype,
    )
    anchor = insert_paragraph_after(
        anchor,
        "由於本研究的 pillar 中心較暗、背景較亮，正規化對比 C(r) 會由中心的 0 隨半徑向外增加至背景的 1。因此，40% 交點較靠近柱中心，通常得到較小的半徑與 CD；60% 交點較靠近外側背景，通常得到較大的半徑與 CD；50% 則位於中心與背景灰階的中間位置，作為本節的主要半對比邊界。",
        body_prototype,
    )
    anchor = insert_paragraph_after(
        anchor,
        "三種門檻使用完全相同的原始影像或反卷積影像、pillar 中心、背景估計與尺寸換算，只改變 f 的數值。每條徑向輪廓均由柱中心向外尋找第一個由低於門檻轉為高於門檻的交點，再以相鄰樣點內插位置。因此，40%、50% 與 60% 的比較是在檢查 CD 對邊界定義的敏感度，而不是重新調整或重跑三種不同強度的反卷積。",
        body_prototype,
    )

    anchor = insert_paragraph_after(anchor, "1.3 評估指標與判讀方式", heading_prototype)
    anchor.paragraph_format.keep_with_next = True
    anchor = insert_definition_after(
        anchor,
        "CD 與平均偏差：",
        "單顆 pillar 的 CD 為 180 組對向 50% 邊界直徑的中位數。對第 i 顆 pillar，帶正負號的偏差定義為 b_i = CD_opt,i - CD_SEM,i；b_i < 0 表示光學量測低估，b_i > 0 表示高估。平均偏差為九顆 b_i 的平均，可判斷整體偏向低估或高估，但正負誤差可能互相抵消，因此不能單獨代表準確度。",
        body_prototype,
    )
    anchor = insert_definition_after(
        anchor,
        "單顆絕對誤差、MAE 與最大絕對誤差：",
        "單顆絕對誤差為 e_i = |CD_opt,i - CD_SEM,i|。平均絕對誤差（mean absolute error, MAE）為九顆 e_i 的平均，用來描述整體尺寸誤差；最大絕對誤差為九顆 e_i 中的最大值，用來檢查最差位置。兩者越小越好，且不會因高估與低估相互抵消。",
        body_prototype,
    )
    anchor = insert_definition_after(
        anchor,
        "改善比例與改善數量：",
        "某項誤差指標 E 的改善比例定義為 (E_before - E_after) / E_before × 100%。正值表示誤差下降，負值表示誤差增加。逐顆改善數量則統計反卷積後 e_i 小於反卷積前 e_i 的 pillar 數量；例如 7/9 表示九顆中有七顆的 CD 絕對誤差下降。",
        body_prototype,
    )
    anchor = insert_definition_after(
        anchor,
        "10-90% 邊緣過渡寬度與邊緣集中度：",
        "每條有效徑向輪廓的邊緣寬度為 w_θ = r_90(θ) - r_10(θ)，單顆結果取 360 個方向的中位數，倍率摘要再取九顆的平均。w 越小表示灰階變化集中於較窄範圍。輔助指標 S = 1/w 稱為邊緣集中度，S 越大表示邊緣越集中；它只能量化銳利度變化，不能取代 CD 準確度或光學解析度。",
        body_prototype,
    )
    anchor = insert_definition_after(
        anchor,
        "振鈴比例與最大振鈴比例：",
        "理想的暗柱徑向輪廓應由中心向背景大致單調上升。若反卷積後出現先升後降的波動，便視為振鈴。振鈴比例 R 定義為輪廓向外過程中所有反向下降量的總和，除以所有正向上升量的總和；R = 0 代表沒有可量測的反向波動，數值越大表示環狀假影越明顯。本文的最大振鈴比例是九顆 pillar 中最大的 R，用來呈現最嚴重的局部假影，而不是九顆的平均。",
        body_prototype,
    )
    anchor = insert_definition_after(
        anchor,
        "中心亮暗反轉：",
        "原始結構應滿足 Icore < Ibg，也就是 pillar 中心比背景暗。若處理後出現 Icore >= Ibg，便記為中心亮暗反轉，代表主要明暗關係已失真。未發生反轉只能證明中心結構仍保持正確方向，不能排除邊界外圍的振鈴。",
        body_prototype,
    )
    insert_definition_after(
        anchor,
        "逐顆實體 pillar 留出驗證：",
        "每次將同一顆實體 pillar 在 20 倍與 50 倍下的兩筆結果同時排除，只使用其餘八顆選擇共同 NSR，再回到被排除的兩筆資料計算誤差。此結果用來檢查改善是否完全依賴該顆 pillar 參與選參；它仍屬於相同視野與相同樣品內的驗證，不等同於跨視野或跨樣品測試。",
        body_prototype,
    )

    threshold_intro = find_paragraph(doc, "pillar 在光學影像中的邊界是一段連續的灰階過渡區域", startswith=True)
    replace_text(
        threshold_intro,
        "如前述定義，pillar 在光學影像中的邊界是一段連續的灰階過渡區域，因此 40%、50% 與 60% 相對對比門檻會得到不同的邊界半徑與 CD。若改善只在單一門檻下出現，代表結果可能高度依賴邊界定義。",
    )

    threshold_method = find_paragraph(doc, "本研究比較 40%、50% 與 60% 對比門檻", startswith=True)
    replace_text(
        threshold_method,
        "本研究在同一組原始與反卷積影像上分別套用 f = 0.4、0.5 與 0.6，其他中心、背景、方向與尺寸換算均不改變。此分析不是重新選擇反卷積參數，而是用來判斷主要結論能否在合理的邊界位置變動下維持。",
    )

    threshold_interpretation = find_paragraph(doc, "因此，反卷積的 CD 改善主要出現在", startswith=True)
    replace_text(
        threshold_interpretation,
        "因此，反卷積的 CD 改善主要出現在較靠近柱中心的 40% 邊界與中間位置的 50% 邊界，不能宣稱在所有門檻下都穩定。60% 門檻位於較靠近外側背景的位置，較容易受到反卷積後的過衝與環狀振鈴影響；本次 60% 結果明顯惡化，表示外側邊界的位置對灰階波動特別敏感。",
    )

    threshold_choice = find_paragraph(doc, "本節以 50% 半對比邊界作為主要結果", startswith=True)
    replace_text(
        threshold_choice,
        "本節以 50% 半對比邊界作為主要結果，是因為 f = 0.5 位於中心與背景灰階的中點，具有固定且對稱的定義，而不是因為它產生最大的改善。40% 與 60% 結果則作為敏感度檢查，用來揭示 CD 結論隨邊界定義改變的程度。",
    )

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)

    reopened = Document(OUTPUT_DOCX)
    text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    required = (
        "1.2 對比門檻的意義",
        "不是顯示影像時調整的亮度門檻",
        "40% 交點較靠近柱中心",
        "1.3 評估指標與判讀方式",
        "平均絕對誤差（mean absolute error, MAE）",
        "本文的最大振鈴比例是九顆 pillar 中最大的 R",
        "60% 門檻位於較靠近外側背景的位置",
    )
    for phrase in required:
        if phrase not in text:
            raise AssertionError(f"Missing required phrase: {phrase}")

    forbidden = (
        "較外側的 40% 邊界",
        "60% 門檻位於較接近 pillar 內部的位置",
    )
    for phrase in forbidden:
        if phrase in text:
            raise AssertionError(f"Obsolete threshold direction remains: {phrase}")

    print(OUTPUT_DOCX)
    print(f"Paragraphs: {len(reopened.paragraphs)}; tables: {len(reopened.tables)}")


if __name__ == "__main__":
    main()
