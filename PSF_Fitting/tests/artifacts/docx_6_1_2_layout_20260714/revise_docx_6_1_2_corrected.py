from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image
PROJECT_ROOT = Path(__file__).resolve().parents[4]


INPUT_DOCX = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節\第六章 實驗結果與分析_6.1.2圖表排版版.docx"
)
OUTPUT_DOCX = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節\第六章 實驗結果與分析_6.1.2圖表與數據校正版.docx"
)
FIGURE_ROOT = Path(
    str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/Final_Data_Figures/03_PSF_EXPERIMENT_RESULTS/01_global_bads_fit/figures/center9_pillar_fit_20x_50x')
)


PARAGRAPH_REPLACEMENTS = {
    3: "為比較不同物鏡倍率下 pillar 結構的成像特性，本研究分別使用 50 倍與 20 倍物鏡進行影像拍攝，並計算各影像的二維頻譜。兩種倍率的頻譜均採用相同的絕對尺度呈現，未進行逐張峰值正規化或額外強化，因此圖中的強度差異可直接反映各影像實際的空間頻率分布。",
    18: "本研究分別使用 Olympus MPLN-50X 與 Olympus MPLN-20X 物鏡所拍攝的 pillar 九宮格影像進行點擴散函數擬合。兩組影像皆選取 pillar 與 via 交界附近的九顆 pillar 作為擬合目標。SEM 提供的九個數值為半徑，範圍為 1299.376-1338.398 nm；換算為 CD 直徑後，範圍為 2598.752-2676.797 nm，平均為 2629.737 nm。Ground Truth 依各 pillar 的實際半徑與物體空間取樣尺寸繪製，並根據 50 倍與 20 倍影像的拍攝方向調整排列，使影像位置與同一顆實體 pillar 正確對應。",
    20: "圖 6-X 呈現 Olympus MPLN-50X 與 Olympus MPLN-20X 物鏡的點擴散函數擬合結果。兩組 PSF 的能量皆主要集中於中心，且水平與垂直方向的分布十分接近；以二階矩估計的長短軸比，50 倍為 1.001，20 倍為 1.012，顯示本次選定區域的模糊近似圓形對稱。由於兩種倍率的物體空間取樣間距不同，圖中的像素寬度不能直接代表實際光學擴散範圍，仍須換算至物體空間比較。",
    21: "Olympus MPLN-50X 物鏡的物體空間取樣間距為 0.069 µm/pixel，PSF 水平與垂直半高全寬分別為 1.298 µm 與 1.298 µm；Olympus MPLN-20X 物鏡的取樣間距為 0.1725 µm/pixel，水平與垂直半高全寬分別為 1.309 µm 與 1.311 µm。兩種倍率的物體空間半高全寬非常接近，50 倍結果僅略小約 0.012 µm，因此較適合解釋為兩組擬合 PSF 具有相近的中心擴散尺度，而非存在明顯的解析能力差距。",
    29: "為評估擬合所得之點擴散函數是否具有實際的影像修復能力，本研究分別以 20 倍與 50 倍影像所擬合得到的 PSF 作為 Wiener 反卷積核，處理各自的 pillar 九宮格影像，並比較反卷積前後的 10-90% 邊緣過渡寬度與臨界尺寸（critical dimension, CD）。本節的判讀重點為：邊緣是否變得較集中、量測 CD 是否整體朝 SEM 參考尺寸靠近，以及影像銳化是否同時引入過度的環狀假影。",
    30: "各倍率所使用的 PSF 皆由同一個 pillar 九宮格區域擬合而得，之後再回用於該區域。因此，本節屬於擬合區域內的一致性與可行性驗證，而不是對全新視野或不同樣品的完全獨立驗證。此設計可用來確認擬合 PSF 是否具備可量化的修復潛力，但不能直接推論其在其他視野位置也會得到相同效果。",
    31: "反卷積前先估計影像中的緩慢背景變化，再將暗色 pillar 轉換為正向結構後進行處理，最後還原為原本的灰階方向。20 倍與 50 倍影像共同使用雜訊對訊號比參數 0.0565，以避免逐倍率或逐顆 pillar 分別調整造成比較規則不一致。",
    32: "參數 0.0565 是以兩種倍率合併後的 CD 平均絕對誤差、最差情況與邊緣寬度改善共同選定，因此 SEM 尺寸確實參與了整體反卷積強度的校準；但未針對個別 pillar 或單一倍率另外微調。為降低直接配合 SEM 數值所造成的偏差，本研究另進行逐顆實體 pillar 留出驗證：每次排除同一顆 pillar 在兩種倍率下的量測，再利用其餘資料選擇參數，最後檢查被排除資料的誤差。",
    33: "反卷積前後均使用相同的半對比邊界判定規則：先估計 pillar 中心與周圍背景的亮度，再沿 360 個方向尋找 50% 對比位置，並以相對方向的邊界距離得到 CD。SEM 提供的九個數值為半徑，建立 Ground Truth 與比較 CD 時均先乘以 2 轉換為直徑；20 倍與 50 倍資料也依各自實際拍攝方位重新排列，使同一影像位置對應到相同的實體 pillar。",
    34: "由於影像看起來較銳利並不等於真實邊界已被完整恢復，本研究除主要的 50% 對比門檻外，也比較 40% 與 60% 門檻，並同時檢查邊緣過渡寬度、CD 平均絕對誤差、最大絕對誤差、個別 pillar 的改善數量及環狀假影。",
    37: "邊緣過渡寬度是指影像亮度由 pillar 內部轉換至周圍背景時，從 10% 變化至 90% 變化所跨越的距離。受到繞射、離焦與其他成像因素影響，實際邊界通常不是理想的階梯，而是一段連續的灰階漸變。",
    38: "在相同量測規則下，過渡寬度越小，代表亮暗變化集中於較窄的範圍，可視為邊緣較集中的證據。由於兩種倍率的物體空間取樣分別為 0.1725 µm/pixel 與 0.069 µm/pixel，跨倍率比較時必須先換算為實際尺寸。",
    39: "20 倍影像的平均邊緣過渡寬度由 4.185 pixels 降低至 3.060 pixels，縮小 26.88%。換算至物體空間後，約由 0.722 µm 降低至 0.528 µm。",
    40: "50 倍影像的平均邊緣過渡寬度由 8.998 pixels 降低至 7.242 pixels，縮小 19.51%；換算至物體空間後，約由 0.621 µm 降低至 0.500 µm。",
    41: "50 倍影像以像素表示的過渡寬度大於 20 倍影像，主要是因為每個像素所代表的實際尺寸較小。換算後可看出，50 倍影像在反卷積前後的實際過渡距離均略小於 20 倍影像，因此不能直接用像素數判斷兩種倍率的邊界品質。",
    42: "兩種倍率在反卷積後皆呈現明顯的邊緣寬度下降，說明擬合 PSF 能使 pillar 與背景之間的亮暗轉換更加集中。然而，此結果也可能同時受到對比重新分配與環狀假影影響，因此仍須配合 CD 誤差及振鈴分析共同判讀。",
    47: "SEM 擬合所得的九個 pillar 數值為半徑，乘以 2 後得到的直徑範圍為 2598.752-2676.797 nm，平均 CD 為 2629.737 nm。以此作為光學影像的比較基準時，20 倍與 50 倍原始影像的平均 CD 分別為 2588.895 nm 與 2596.666 nm，均略低於 SEM 平均直徑。",
    48: "因此，本次修正後的主要現象不是原始光學 CD 過度放大，而是半對比邊界落在較內側，使平均 CD 略為低估。反卷積的合理結果應是讓量測邊界整體向外移動並接近 SEM，而不是單純追求 CD 變大或變小。所有原始與反卷積結果皆使用相同的背景估計、中心亮度、360 方向邊界與像素尺寸換算方式。",
    50: "在 50% 對比門檻下，20 倍原始影像的平均 CD 為 2588.895 nm，反卷積後增加至 2625.655 nm，平均增加 36.761 nm。相對於 SEM 平均 CD 2629.737 nm，其平均偏差由 -40.843 nm 縮小至 -4.082 nm，顯示整體平均值明顯朝 SEM 靠近。",
    51: "20 倍影像的平均絕對誤差由 40.843 nm 降低至 22.928 nm，改善 43.86%；九顆 pillar 中有七顆的絕對誤差下降，另外兩顆則小幅增加。這表示平均誤差的改善具有多數位置支持，但不能描述為九顆全部改善。",
    52: "20 倍影像的最大絕對誤差由 87.468 nm 降低至 83.099 nm，改善 4.99%。最大誤差的下降幅度小於平均誤差，顯示整體量測雖更接近 SEM，但仍有少數位置保留較大的局部偏差。",
    53: "綜合平均值、平均絕對誤差與個別 pillar 結果，20 倍影像可支持反卷積具有整體 CD 修正能力；不過，其證據應表述為多數 pillar 改善與整體誤差下降，而不是所有位置都被完整還原。",
    57: "20 倍影像反卷積前後之實際半對比 CD 量測邊界",
    59: "20 倍影像反卷積前後之平均絕對 CD 誤差",
    62: "50 倍原始影像的平均 CD 為 2596.666 nm，反卷積後增加至 2635.838 nm，平均增加 39.172 nm。相對於 SEM 平均 CD 2629.737 nm，其平均偏差由 -33.071 nm 變為 +6.101 nm，表示平均值在接近 SEM 後出現輕微高估。",
    63: "50 倍影像的平均絕對誤差由 33.071 nm 降低至 23.966 nm，改善 27.53%；九顆 pillar 中有六顆的絕對誤差下降，三顆增加。因此，50 倍結果同樣呈現整體改善，但個別位置的一致性低於 20 倍影像。",
    64: "50 倍影像的最大絕對誤差由 80.884 nm 降低至 52.219 nm，改善 35.44%。雖然平均 CD 略微超過 SEM 平均值，但最大誤差仍明顯下降，顯示不能只以平均偏差判斷修復品質，仍需保留逐顆 pillar 的誤差比較。",
    65: "整體而言，50 倍影像的平均絕對誤差與最大絕對誤差均下降，但九顆中仍有三顆未改善。較適當的結論是反卷積降低了整體 CD 偏差與最差位置的誤差，而非所有位置皆得到相同程度的修正。",
    70: "50 倍影像反卷積前後之實際半對比 CD 量測邊界",
    72: "50 倍影像反卷積前後之平均絕對 CD 誤差",
    75: "兩種倍率共得到 18 筆量測，代表同一組九顆實體 pillar 分別在 20 倍與 50 倍影像中的結果，而不是 18 顆不同的 pillar。此配對關係也用於確認兩種倍率的 SEM 尺寸排列與影像方向一致。",
    76: "在使用全部資料校準共同參數 0.0565 的主要結果中，合併平均絕對誤差由 36.957 nm 降低至 23.447 nm，改善 36.56%；18 筆量測中有 13 筆改善。這表示整體誤差下降並非只由單一倍率主導，但仍存在五筆個別結果未改善。",
    77: "逐顆實體 pillar 留出驗證得到較保守的結果：被排除資料的平均絕對誤差由 36.957 nm 降低至 28.164 nm，改善 23.79%，18 筆中有 11 筆改善。此結果顯示共同參數的修正趨勢在未參與當次選參的 pillar 上仍可部分保留，但效果弱於使用全部資料時的結果。",
    78: "因此，本研究可主張擬合 PSF 在目前九宮格區域內具有整體邊緣修復與 CD 誤差降低的潛力；但由於結果並非每一顆 pillar 都改善，且留出驗證仍使用相同視野與同一組實體 pillar，本節不將其延伸解釋為跨視野或跨樣品的普遍校正能力。",
    80: "pillar 在光學影像中的邊界是一段連續的灰階過渡區域，因此不同對比門檻會改變邊界位置與 CD。若改善只在單一門檻下出現，代表結果可能高度依賴邊界定義。",
    81: "本研究比較 40%、50% 與 60% 對比門檻，並在每個門檻下使用完全相同的反卷積結果與 CD 計算流程。此分析用來判斷主要結論能否在合理的邊界位置變動下維持。",
    82: "20 倍影像在 40% 門檻下，平均絕對誤差改善 29.08%，九顆均改善；在主要的 50% 門檻下改善 43.86%，七顆改善；提高至 60% 時，平均絕對誤差反而增加 29.34%，僅一顆改善。",
    83: "50 倍影像在 40% 門檻下，平均絕對誤差改善 16.70%，九顆均改善；在 50% 門檻下改善 27.53%，六顆改善；在 60% 門檻下，平均絕對誤差增加 96.51%，且九顆皆未改善。",
    84: "合併兩種倍率後，40% 門檻共有 18/18 筆改善，50% 門檻共有 13/18 筆改善；60% 門檻則只有 1/18 筆改善，且兩種倍率的平均誤差均增加。",
    85: "因此，反卷積的 CD 改善主要出現在較外側的 40% 邊界與標準半對比的 50% 邊界，不能宣稱在所有門檻下都穩定。60% 門檻位於較接近 pillar 內部的位置，對反卷積後的局部灰階重新分配與環狀結構特別敏感。",
    86: "本節以 50% 半對比邊界作為主要結果，是因為它在中心與背景亮度之間具有明確且對稱的定義，而不是因為它產生最大的改善。敏感度分析則顯示 CD 修正仍會受到邊界定義影響，因此絕對尺寸結論必須保守解讀。",
    87: "不同對比門檻下之 CD 平均絕對誤差變化與改善數量",
    91: "反卷積可使邊緣變得集中，但處理強度提高時，也可能在高對比邊界周圍形成不屬於原始結構的亮環或暗環。這類環狀假影可能讓輪廓看起來更清楚，卻不一定代表真實結構被正確恢復。",
    92: "因此，邊緣寬度下降與 CD 誤差改善必須和最大振鈴比例一併判讀。本研究同時檢查 pillar 中心是否發生亮暗反轉，以排除明顯的結構失真。",
    93: "20 倍影像的最大振鈴比例由 0.0316 增加至 0.1947。雖然平均邊緣寬度縮小 26.88%、CD 平均絕對誤差降低 43.86%，但環狀假影也明顯增加，因此不能把全部銳化效果都視為真實邊界恢復。",
    94: "50 倍影像的最大振鈴比例由 0.0155 增加至 0.2241，增幅同樣明顯。其邊緣寬度與 CD 誤差雖有改善，但反卷積後的外圍光暈比原始影像更突出，顯示目前設定偏向較積極的銳化。",
    95: "兩種倍率在反卷積後均未出現 pillar 中心亮暗反轉，表示九宮格的主要明暗結構仍被保留；然而，中心未反轉不代表外圍假影可以忽略。尤其在 60% 對比門檻下，CD 誤差明顯惡化，與振鈴增加的結果相互呼應。",
    96: "綜合而言，目前設定適合用來證明擬合 PSF 具有使邊緣變窄並降低主要 CD 誤差的能力，但不適合宣稱已無失真地恢復完整影像。較保守的判讀應同時呈現銳化收益與環狀假影代價。",
    102: "綜合主要的 50% 半對比結果，20 倍與 50 倍影像的平均邊緣過渡寬度分別縮小 26.88% 與 19.51%，CD 平均絕對誤差分別降低 43.86% 與 27.53%。20 倍有七顆、50 倍有六顆 pillar 的 CD 誤差改善；合併全部資料後平均絕對誤差降低 36.56%，逐顆實體 pillar 留出驗證則降低 23.79%。",
    106: "本節仍有三項主要限制。第一，反卷積影像與 PSF 擬合影像來自同一九宮格區域，因此只能支持擬合區域內的可行性。第二，共同參數 0.0565 是利用 SEM 誤差與邊緣條件校準所得，雖未針對個別 pillar 微調，仍不屬於完全未知答案的盲測。第三，兩種倍率在反卷積後的最大振鈴比例均明顯增加，且 60% 對比門檻下的 CD 結果惡化，顯示較積極的銳化伴隨邊界假影與門檻敏感性。",
    107: "反卷積後的平均 CD 已接近 SEM 平均直徑：20 倍平均偏差為 -4.082 nm，50 倍為 +6.101 nm。然而，平均值接近不代表每顆 pillar 都已準確還原；20 倍與 50 倍仍分別存在 83.099 nm 與 52.219 nm 的最大絕對誤差，且部分 pillar 的誤差反而增加。因此，本結果不能解釋為完成絕對 CD 校準。",
    108: "整體而言，擬合 PSF 在目前九宮格區域內能使兩種倍率的 pillar 邊緣更加集中，並在 40% 與 50% 對比邊界下降低整體 CD 誤差。此結果可支持 PSF 具有實際影像修復潛力，但其效果仍受個別位置、邊界門檻與環狀假影影響；後續仍需以未參與擬合的新視野或不同樣品進行外部驗證。",
}


TABLE_VALUES = {
    55: [
        ["誤差指標", "原始影像量測結果", "反卷積影像量測結果"],
        ["平均絕對誤差（nm）", "40.843", "22.928"],
        ["最大絕對誤差（nm）", "87.468", "83.099"],
    ],
    67: [
        ["誤差指標", "原始影像量測結果", "反卷積影像量測結果"],
        ["平均絕對誤差（nm）", "33.071", "23.966"],
        ["最大絕對誤差（nm）", "80.884", "52.219"],
    ],
    88: [
        ["物鏡倍率", "對比門檻", "平均絕對誤差變化", "改善數量"],
        ["20 倍", "40%", "下降 29.08%", "9 / 9"],
        ["20 倍", "50%", "下降 43.86%", "7 / 9"],
        ["20 倍", "60%", "增加 29.34%", "1 / 9"],
        ["50 倍", "40%", "下降 16.70%", "9 / 9"],
        ["50 倍", "50%", "下降 27.53%", "6 / 9"],
        ["50 倍", "60%", "增加 96.51%", "0 / 9"],
    ],
    104: [
        ["物鏡倍率", "邊緣寬度（µm）", "CD 平均絕對誤差（nm）", "最大振鈴比例"],
        ["20 倍", "0.722 → 0.528", "40.843 → 22.928", "0.0316 → 0.1947"],
        ["50 倍", "0.621 → 0.500", "33.071 → 23.966", "0.0155 → 0.2241"],
    ],
}


IMAGE_REPLACEMENTS = {
    "rId6": FIGURE_ROOT / "50x" / "05_original_pillar_3x3.png",
    "rId7": FIGURE_ROOT / "50x" / "07_original_pillar_3x3_spectrum.png",
    "rId8": FIGURE_ROOT / "20x" / "05_original_pillar_3x3.png",
    "rId9": FIGURE_ROOT / "20x" / "07_original_pillar_3x3_spectrum.png",
    "rId11": FIGURE_ROOT / "50x" / "03_fitted_psf.png",
    "rId12": FIGURE_ROOT / "20x" / "03_fitted_psf.png",
    "rId15": FIGURE_ROOT / "20x" / "14_edge_width_before_after.png",
    "rId16": FIGURE_ROOT / "50x" / "14_edge_width_before_after.png",
    "rId17": FIGURE_ROOT / "20x" / "10_cd_measurement_overlay_captured.png",
    "rId18": FIGURE_ROOT / "20x" / "11_cd_measurement_overlay_deconvolved.png",
    "rId19": FIGURE_ROOT / "20x" / "13_cd_error_before_after.png",
    "rId20": FIGURE_ROOT / "50x" / "10_cd_measurement_overlay_captured.png",
    "rId21": FIGURE_ROOT / "50x" / "11_cd_measurement_overlay_deconvolved.png",
    "rId22": FIGURE_ROOT / "50x" / "13_cd_error_before_after.png",
    "rId23": FIGURE_ROOT / "20x" / "05_original_pillar_3x3.png",
    "rId24": FIGURE_ROOT / "20x" / "09_fitted_psf_wiener_deconvolution_pillar_3x3.png",
    "rId25": FIGURE_ROOT / "50x" / "05_original_pillar_3x3.png",
    "rId26": FIGURE_ROOT / "50x" / "09_fitted_psf_wiener_deconvolution_pillar_3x3.png",
}


def body_elements(doc: Document):
    return list(doc.element.body.iterchildren())


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def replace_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    replace_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def update_table(table: Table, values: list[list[str]]) -> None:
    if len(table.rows) != len(values):
        raise ValueError(f"Unexpected table row count: {len(table.rows)} != {len(values)}")
    for row, row_values in zip(table.rows, values):
        if len(row.cells) != len(row_values):
            raise ValueError(f"Unexpected table column count: {len(row.cells)} != {len(row_values)}")
        for cell, value in zip(row.cells, row_values):
            replace_cell_text(cell, value)


def set_image_aspect_ratio(doc: Document, relationship_id: str, image_path: Path) -> None:
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    for blip in doc.element.body.iter(qn("a:blip")):
        if blip.get(qn("r:embed")) != relationship_id:
            continue
        ancestor = blip
        while ancestor is not None and ancestor.tag != qn("wp:inline"):
            ancestor = ancestor.getparent()
        if ancestor is None:
            continue
        extent = ancestor.find(qn("wp:extent"))
        if extent is None:
            continue
        width_emu = int(extent.get("cx"))
        height_emu = round(width_emu * height_px / width_px)
        extent.set("cy", str(height_emu))
        for xfrm_extent in ancestor.iter(qn("a:ext")):
            xfrm_extent.set("cx", str(width_emu))
            xfrm_extent.set("cy", str(height_emu))


def main() -> None:
    if not INPUT_DOCX.exists():
        raise FileNotFoundError(INPUT_DOCX)
    for image_path in IMAGE_REPLACEMENTS.values():
        if not image_path.exists():
            raise FileNotFoundError(image_path)

    doc = Document(INPUT_DOCX)
    elements = body_elements(doc)

    for body_index, new_text in PARAGRAPH_REPLACEMENTS.items():
        element = elements[body_index]
        if element.tag != qn("w:p"):
            raise TypeError(f"Body element {body_index} is not a paragraph")
        replace_paragraph_text(Paragraph(element, doc._body), new_text)

    for body_index, values in TABLE_VALUES.items():
        element = elements[body_index]
        if element.tag != qn("w:tbl"):
            raise TypeError(f"Body element {body_index} is not a table")
        update_table(Table(element, doc._body), values)

    for relationship_id, image_path in IMAGE_REPLACEMENTS.items():
        image_part = doc.part.related_parts[relationship_id]
        image_part._blob = image_path.read_bytes()
        set_image_aspect_ratio(doc, relationship_id, image_path)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
