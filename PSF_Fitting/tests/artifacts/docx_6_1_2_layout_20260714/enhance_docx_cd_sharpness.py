from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
PROJECT_ROOT = Path(__file__).resolve().parents[4]


INPUT_DOCX = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節"
    r"\第六章 實驗結果與分析_6.1.2圖表與數據校正版.docx"
)
OUTPUT_DOCX = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節"
    r"\第六章 實驗結果與分析_6.1.2方法與銳利度驗證完整版.docx"
)
FIGURE_ROOT = Path(
    str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/Final_Data_Figures/03_PSF_EXPERIMENT_RESULTS/01_global_bads_fit/figures/center9_pillar_fit_20x_50x/comparison')
)
METHOD_FIGURE = FIGURE_ROOT / "04_cd_measurement_method.png"
SHARPNESS_FIGURE = FIGURE_ROOT / "05_sharpness_validation.png"
CD_FIGURE = FIGURE_ROOT / "06_cd_pointwise_validation.png"


def find_paragraph(doc: Document, text: str, *, startswith: bool = False) -> Paragraph:
    matches = []
    for paragraph in doc.paragraphs:
        candidate = paragraph.text.strip()
        if (startswith and candidate.startswith(text)) or (not startswith and candidate == text):
            matches.append(paragraph)
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph for {text!r}; found {len(matches)}")
    return matches[0]


def copy_run_format(source: Paragraph, target_run) -> None:
    if source.runs and source.runs[0]._r.rPr is not None:
        target_run._r.insert(0, deepcopy(source.runs[0]._r.rPr))


def replace_text(paragraph: Paragraph, text: str) -> None:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def insert_paragraph_after(
    reference: Paragraph,
    text: str,
    prototype: Paragraph,
) -> Paragraph:
    new_element = OxmlElement("w:p")
    reference._p.addnext(new_element)
    paragraph = Paragraph(new_element, reference._parent)
    if prototype._p.pPr is not None:
        new_element.insert(0, deepcopy(prototype._p.pPr))
    run = paragraph.add_run(text)
    copy_run_format(prototype, run)
    return paragraph


def insert_picture_after(reference: Paragraph, image_path: Path, width_cm: float) -> Paragraph:
    new_element = OxmlElement("w:p")
    reference._p.addnext(new_element)
    paragraph = Paragraph(new_element, reference._parent)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    return paragraph


def insert_picture_before(reference: Paragraph, image_path: Path, width_cm: float) -> Paragraph:
    new_element = OxmlElement("w:p")
    reference._p.addprevious(new_element)
    paragraph = Paragraph(new_element, reference._parent)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    return paragraph


def remove_table_immediately_before(paragraph: Paragraph) -> None:
    previous = paragraph._p.getprevious()
    if previous is None or previous.tag != qn("w:tbl"):
        raise ValueError("Expected the old two-panel image table immediately before the caption")
    previous.getparent().remove(previous)


def image_hashes_in_docx(path: Path) -> set[str]:
    with ZipFile(path) as archive:
        return {
            sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/")
        }


def main() -> None:
    for path in (INPUT_DOCX, METHOD_FIGURE, SHARPNESS_FIGURE, CD_FIGURE):
        if not path.exists():
            raise FileNotFoundError(path)

    doc = Document(INPUT_DOCX)
    body_prototype = find_paragraph(
        doc, "反卷積前後均使用相同的半對比邊界判定規則", startswith=True
    )
    subheading_prototype = find_paragraph(doc, "3.1 20 倍影像")
    caption_prototype = find_paragraph(
        doc, "20 倍與 50 倍影像反卷積前後之邊緣過渡寬度比較"
    )

    replace_text(
        body_prototype,
        "反卷積前後均採用相同的半對比邊界規則，且只更換輸入影像；pillar 中心位置、徑向取樣方向、尺寸換算與 SEM 對應順序均保持固定。SEM 提供的九個數值為半徑，建立 Ground Truth 與比較 CD 時均先乘以 2 轉換為直徑；20 倍與 50 倍資料也依各自實際拍攝方位重新排列，使同一影像位置對應到相同的實體 pillar。",
    )

    anchor = body_prototype
    anchor = insert_paragraph_after(
        anchor, "1.1 CD 與邊緣寬度的量測方式", subheading_prototype
    )
    anchor = insert_paragraph_after(
        anchor,
        "每顆 pillar 的量測以九宮格定位所得的中心為起點。設相鄰 pillar 的平均間距為 p，20 倍與 50 倍影像中的 p 分別為 28.986 pixels 與 72.464 pixels。由中心向外建立 360 條等角度放射線，徑向取樣間隔為 0.05 pixel，最遠取至 0.38p，以涵蓋柱邊界與背景，同時避免進入鄰近 pillar。",
        body_prototype,
    )
    anchor = insert_paragraph_after(
        anchor,
        "中心灰階 Icore 取自 r ≤ 0.07p 範圍內所有樣點的中位數，背景灰階 Ibg 則取自 0.30p ≤ r ≤ 0.36p 環帶內所有樣點的中位數。每條徑向輪廓依 C(r) = [I(r) - Icore] / [Ibg - Icore] 轉換為相對對比，並以標準差 6 個徑向樣點的平滑處理降低單一像素雜訊；由於取樣間隔為 0.05 pixel，此平滑尺度相當於 0.30 pixel。接著由柱中心向外尋找第一個 10%、50% 與 90% 對比交點，交點位置以相鄰兩個樣點線性內插。",
        body_prototype,
    )
    anchor = insert_paragraph_after(
        anchor,
        "對第 θ 個方向取得的 50% 邊界半徑記為 r50(θ)，再與相反方向的 r50(θ + 180°) 相加，形成 180 組對向直徑。單顆 pillar 的 CD 取這 180 組直徑的中位數，可降低中心定位略有偏移所造成的單邊誤差。10%、50% 與 90% 交點皆要求至少 324 條方向有效，亦即有效比例不得低於 90%。",
        body_prototype,
    )
    anchor = insert_paragraph_after(
        anchor,
        "像素直徑最後依物體空間取樣換算為實際尺寸：20 倍為 172.5 nm/pixel，50 倍為 69.0 nm/pixel。邊緣過渡寬度則取所有有效方向中 r90 - r10 的中位數。反卷積前後完全沿用此流程，因此後續疊圖中的黃色線是 360 個實際半對比交點連成的輪廓，不是以單一平均半徑畫出的理想圓。",
        body_prototype,
    )
    picture = insert_picture_after(anchor, METHOD_FIGURE, 15.7)
    insert_paragraph_after(
        picture,
        "CD 與 10-90% 邊緣過渡寬度之實際量測流程",
        caption_prototype,
    )

    sharpness_heading = find_paragraph(doc, "2. 反卷積前後的邊緣過渡寬度")
    replace_text(sharpness_heading, "2. 銳利度驗證：反卷積前後的邊緣過渡寬度")

    replacements = {
        "邊緣過渡寬度是指影像亮度由 pillar 內部轉換至周圍背景時":
            "本研究不以影像顯示對比或主觀觀察作為銳利度判準，而是使用同一組 360 方向輪廓的 10-90% 邊緣過渡寬度。此寬度表示影像亮度由 pillar 內部轉換至周圍背景時，從 10% 變化至 90% 變化所跨越的距離；數值越小，代表亮暗轉換集中於較窄範圍。",
        "在相同量測規則下，過渡寬度越小":
            "為同時檢查平均輪廓與個別 pillar，本研究先將每條徑向輪廓的 50% 交點對齊，再對九顆 pillar 的有效輪廓取中位數；另保留每顆 pillar 的 r90 - r10 中位寬度。圖中上排用來比較反卷積前後的平均邊緣斜率，下排則顯示九顆 pillar 是否一致變窄。跨倍率比較時，所有寬度均依 20 倍 0.1725 µm/pixel 與 50 倍 0.069 µm/pixel 換算至物體空間。",
        "20 倍影像的平均邊緣過渡寬度由 4.185 pixels":
            "20 倍影像的平均邊緣過渡寬度由 4.185 pixels 降低至 3.060 pixels，縮小 26.88%；換算至物體空間後，由 0.722 µm 降低至 0.528 µm。",
        "50 倍影像的平均邊緣過渡寬度由 8.998 pixels":
            "50 倍影像的平均邊緣過渡寬度由 8.998 pixels 降低至 7.242 pixels，縮小 19.51%；換算至物體空間後，由 0.621 µm 降低至 0.500 µm。",
        "50 倍影像以像素表示的過渡寬度大於 20 倍影像":
            "逐顆結果顯示，20 倍與 50 倍各九顆 pillar 的邊緣寬度在反卷積後均下降，合計 18/18 筆皆變窄。因此，平均寬度的下降並不是由少數柱子主導，而是在目前九宮格內具有一致方向。50 倍以像素表示的寬度雖較大，僅是因為每個像素代表的實際尺寸較小，不能直接以像素數比較倍率。",
        "兩種倍率在反卷積後皆呈現明顯的邊緣寬度下降":
            "若以 S = 1/w 作為輔助的邊緣集中度指標，其中 w 為物體空間的平均 10-90% 邊緣寬度，20 倍由 1.385 µm⁻¹ 增加至 1.895 µm⁻¹，提升 36.76%；50 倍由 1.611 µm⁻¹ 增加至 2.001 µm⁻¹，提升 24.24%。此指標只表示邊緣變化更集中，不能取代光學解析度或真實 CD 的判定，因此後續仍需與 SEM 誤差及振鈴現象共同解讀。",
    }
    for prefix, text in replacements.items():
        replace_text(find_paragraph(doc, prefix, startswith=True), text)

    sharpness_caption = find_paragraph(
        doc, "20 倍與 50 倍影像反卷積前後之邊緣過渡寬度比較"
    )
    remove_table_immediately_before(sharpness_caption)
    insert_picture_before(sharpness_caption, SHARPNESS_FIGURE, 15.7)
    replace_text(
        sharpness_caption,
        "20 倍與 50 倍反卷積前後之平均邊緣輪廓與逐顆邊緣寬度",
    )

    combined_conclusion = find_paragraph(
        doc, "因此，本研究可主張擬合 PSF 在目前九宮格區域內", startswith=True
    )
    detail = insert_paragraph_after(
        combined_conclusion,
        "為避免平均值掩蓋局部差異，逐顆 CD 比較圖使用相同的縱軸範圍，同時列出 SEM、原始影像與反卷積結果。20 倍有 7/9 顆的絕對誤差下降，50 倍有 6/9 顆下降；圖中也可直接看出未改善位置仍被保留，未以平均值或顯示尺度加以隱藏。",
        body_prototype,
    )
    cd_picture = insert_picture_after(detail, CD_FIGURE, 15.7)
    insert_paragraph_after(
        cd_picture,
        "20 倍與 50 倍逐顆 pillar CD 與 SEM 參考值比較",
        caption_prototype,
    )

    summary = find_paragraph(doc, "綜合主要的 50% 半對比結果", startswith=True)
    replace_text(
        summary,
        "綜合主要的 50% 半對比結果，20 倍與 50 倍影像的平均邊緣過渡寬度分別縮小 26.88% 與 19.51%，且兩種倍率的 18 筆逐顆邊緣寬度皆下降；輔助的邊緣集中度指標分別提升 36.76% 與 24.24%。CD 平均絕對誤差則分別降低 43.86% 與 27.53%，其中 20 倍有七顆、50 倍有六顆 pillar 的 CD 誤差改善；合併全部資料後平均絕對誤差降低 36.56%，逐顆實體 pillar 留出驗證則降低 23.79%。",
    )

    final_conclusion = find_paragraph(
        doc, "整體而言，擬合 PSF 在目前九宮格區域內", startswith=True
    )
    replace_text(
        final_conclusion,
        "整體而言，平均輪廓與 18/18 筆逐顆邊緣寬度皆支持擬合 PSF 能使目前九宮格區域的 pillar 邊緣更加集中；在 40% 與 50% 對比邊界下，整體 CD 誤差也隨之下降。這些結果可支持 PSF 具有實際影像修復潛力，但個別 CD 並非全部改善，且效果仍受邊界門檻與環狀假影影響；後續仍需以未參與擬合的新視野或不同樣品進行外部驗證。",
    )

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)

    reopened = Document(OUTPUT_DOCX)
    full_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    required_phrases = (
        "1.1 CD 與邊緣寬度的量測方式",
        "360 條等角度放射線",
        "至少 324 條方向有效",
        "合計 18/18 筆皆變窄",
        "20 倍與 50 倍逐顆 pillar CD 與 SEM 參考值比較",
    )
    for phrase in required_phrases:
        if phrase not in full_text:
            raise AssertionError(f"Missing required text: {phrase}")

    embedded_hashes = image_hashes_in_docx(OUTPUT_DOCX)
    for figure in (METHOD_FIGURE, SHARPNESS_FIGURE, CD_FIGURE):
        if sha256(figure.read_bytes()).hexdigest() not in embedded_hashes:
            raise AssertionError(f"Figure was not embedded exactly: {figure}")

    print(OUTPUT_DOCX)
    print(f"Paragraphs: {len(reopened.paragraphs)}; inline shapes: {len(reopened.inline_shapes)}")


if __name__ == "__main__":
    main()
