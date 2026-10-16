from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v20_量測位置補圖版.docx")
OUTPUT = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v21_位移標準差表格版.docx")


CAPTION = (
    "表 6-11 四組位移序列之逐次位移量標準差。單位為nm；"
    "改善率為MOSR相對Raw或Wiener之降低比例。"
)

HEADERS = [
    "Condition",
    "Raw SD (nm)",
    "Wiener SD (nm)",
    "MOSR SD (nm)",
    "MOSR reduction vs Raw",
    "MOSR reduction vs Wiener",
]

ROWS = [
    ["50×, 0.5 µm", "7.04", "4.94", "3.61", "48.72%", "26.82%"],
    ["50×, 1.0 µm", "7.69", "4.36", "3.72", "51.60%", "14.67%"],
    ["20×, 0.5 µm", "16.25", "15.19", "14.73", "9.35%", "3.02%"],
    ["20×, 1.0 µm", "13.52", "7.46", "6.94", "48.69%", "6.95%"],
    ["Mean across four conditions", "11.13", "7.99", "7.25", "34.83%", "9.20%"],
]


def find_paragraph(doc: Document, startswith: str, style: str | None = None) -> Paragraph:
    matches = [
        p
        for p in doc.paragraphs
        if p.text.strip().startswith(startswith)
        and (style is None or p.style.name == style)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph starting with {startswith!r}, found {len(matches)}"
        )
    return matches[0]


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    paragraph.text = text


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_geometry(table: Table, widths: list[int]) -> None:
    if len(table.columns) != len(widths):
        raise RuntimeError("Column-width count mismatch")

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = sum(widths)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "0")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid_cols = table._tbl.tblGrid.gridCol_lst
    for grid_col, width in zip(grid_cols, widths):
        grid_col.set(qn("w:w"), str(width))

    for row in table.rows:
        row.height = None
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def build_table_from_existing(doc: Document) -> Table:
    # Reuse the thesis's established six-column table treatment.
    template = doc.tables[99]
    tbl_el = deepcopy(template._tbl)
    while len(tbl_el.tr_lst) > 6:
        tbl_el.remove(tbl_el.tr_lst[-1])
    table = Table(tbl_el, template._parent)

    for c, text in zip(table.rows[0].cells, HEADERS):
        set_cell_text(c, text, bold=True, size=9.0)
    for row, values in zip(table.rows[1:], ROWS):
        for c, text in zip(row.cells, values):
            set_cell_text(c, text, bold=False, size=9.5)

    # Wider comparison columns keep the English headers readable.
    set_table_geometry(table, [1500, 1150, 1200, 1200, 1625, 1625])
    return table


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if OUTPUT.exists():
        OUTPUT.unlink()

    doc = Document(SOURCE)

    # 1) Insert the new list-of-tables entry and shift later table numbers.
    old_list_611 = find_paragraph(doc, "表 6-11", style="toc 1")
    new_list_el = deepcopy(old_list_611._p)

    # Renumber from the end so each source number is still unique.
    for old_no, new_no in [(13, 14), (12, 13), (11, 12)]:
        p = find_paragraph(doc, f"表 6-{old_no}", style="toc 1")
        set_paragraph_text(p, p.text.replace(f"表 6-{old_no}", f"表 6-{new_no}", 1))
    old_list_611._p.addprevious(new_list_el)
    new_list = Paragraph(new_list_el, old_list_611._parent)
    set_paragraph_text(new_list, CAPTION + "\t200")

    # 2) Shift the later body-table numbers before inserting the new 6-11.
    for old_no, new_no in [(13, 14), (12, 13), (11, 12)]:
        p = find_paragraph(doc, f"表 6-{old_no}", style="Table Caption")
        set_paragraph_text(p, p.text.replace(f"表 6-{old_no}", f"表 6-{new_no}", 1))

    # 3) Insert the caption and English table immediately after Figure 6-48.
    fig_caption = find_paragraph(doc, "圖 6-48", style="Figure Caption")
    caption_template = find_paragraph(doc, "表 6-10", style="Table Caption")
    new_caption_el = deepcopy(caption_template._p)
    fig_caption._p.addnext(new_caption_el)
    new_caption = Paragraph(new_caption_el, fig_caption._parent)
    set_paragraph_text(new_caption, CAPTION)
    # The source caption begins a fresh page in its original location; this
    # inserted caption should instead use the open space below Figure 6-48.
    new_caption.paragraph_format.page_break_before = False
    new_caption.paragraph_format.keep_with_next = True

    new_table = build_table_from_existing(doc)
    new_caption._p.addnext(new_table._tbl)

    # 4) Add the exact standard-deviation comparison to the surrounding analysis.
    analysis = find_paragraph(doc, "四組序列合併後，Raw、Wiener與MOSR")
    set_paragraph_text(
        analysis,
        "逐次位移量標準差的逐條件數值列於表6-11。MOSR在四組條件皆為最低；"
        "相較Raw的降幅依序為48.72%、51.60%、9.35%與48.69%，相較Wiener則為"
        "26.82%、14.67%、3.02%與6.95%。四條件平均標準差由Raw的11.13 nm與"
        "Wiener的7.99 nm降至MOSR的7.25 nm。四組序列合併後，Raw、Wiener與MOSR"
        "相對PI平台實際回授位移的平均絕對偏差為9.54、7.46與6.74 nm；九顆pillar"
        "位移離散度則為17.63、7.92與6.47 nm。MOSR相較Raw分別降低29.36%與"
        "63.29%，且改善不是由改變整段位移比例所造成。",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
