from pathlib import Path
import json,hashlib,shutil,zipfile,datetime
import pandas as pd,numpy as np
from docx import Document
PROJECT_ROOT = Path(__file__).resolve().parents[3]
P=Path(str(PROJECT_ROOT / 'Deliverables/R137_WeakPSF_Web_Documents_20260922'));R=Path(str(PROJECT_ROOT / 'Experiments/R137_WeakPSFFeature'))
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
# All copies, including measurements, retain exact original bytes.
source=load(P/'來源與核對/SOURCE_COPIES.json')
for r in source:assert sha(P/r['copy'])==r['sha256']==sha(r['source'])
C=pd.read_csv(P/'完整數據/all_circles.csv');D=pd.read_csv(P/'完整數據/all_displacements.csv');S=pd.read_csv(P/'完整數據/displacement_summary.csv');G=pd.read_csv(P/'完整數據/shape_size_summary.csv')
assert len(C)==33480 and len(D)==7200
errors=[]
for (m,k),g in D.groupby(['model','kind']):
 r=S[(S.model==m)&(S.kind==k)].iloc[0]
 computed=[g.error_nm.abs().mean(),g.error_nm.mean(),g.groupby('j').error_nm.std(ddof=1).mean(),g.error_nm.abs().max()]
 errors+=list(abs(np.array(computed)-r[['mae_nm','bias_nm','repeatability_nm','maximum_error_nm']].to_numpy(float)))
assert max(errors)<1e-10
# Read back actual Word cells, not the builder's in-memory values.
d=Document(next(P.glob('02_*.docx')))
for idx,kind in [(0,'adjacent'),(1,'start')]:
 for row in d.tables[idx].rows[1:]:
  vals=[c.text for c in row.cells];r=S[(S.model==vals[0])&(S.kind==kind)].iloc[0]
  assert np.allclose(np.array(vals[1:],float),r[['mae_nm','bias_nm','repeatability_nm','maximum_error_nm']].to_numpy(float),rtol=0,atol=5.1e-7)
sem=pd.read_csv(P/'完整數據/SEM_comparison.csv')
for idx,cand in [(5,1),(6,2)]:
 for row in d.tables[idx].rows[1:]:
  vals=[c.text for c in row.cells];r=sem[(sem.model==vals[0])&(sem.mapping_candidate==cand)].iloc[0]
  assert np.allclose(np.array(vals[1:],float),r[['mean_diameter_difference_nm','full_radial_difference_nm','de_size_shape_difference_nm']].to_numpy(float),rtol=0,atol=5.1e-7)
pp=pd.read_csv(P/'完整數據/per_pillar_statistics.csv');d4=Document(next(P.glob('04_*.docx')))
for idx,m in enumerate(['Raw','R129','R130-Clean','R137-WeakPSF'],1):
 vals=np.array([[float(c.text) for c in row.cells] for row in d4.tables[idx].rows[1:]])
 expected=pp[pp.model==m].drop(columns='model').to_numpy(float);assert np.allclose(vals,expected,rtol=0,atol=5.1e-7)
 row=G[G.model==m].iloc[0]
 for key,col in [('diameter_repeat_sd_nm','diameter_sd_nm'),('contour_repeat_sd_nm','contour_sd_nm'),('shape_repeat_sd_nm','shape_sd_nm')]:assert abs(np.sqrt(np.square(pp[pp.model==m][col]).mean())-row[key])<1e-4
for p in P.glob('*.docx'):
 with zipfile.ZipFile(p) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith('.rels'):assert b'TargetMode="External"' not in z.read(name)
  assert b'<w:pBdr' not in z.read('word/styles.xml')
# Correct reference helper names; these files document the creation environment, not portable execution.
for name in ['build_documents.py','prepare_doc_data.py','author_docs.py']:
 shutil.copy2(R/'results'/name,P/'來源與核對'/name)
shutil.copy2(__file__,P/'來源與核對/verify_and_package.py')
(P/'來源與核對/README.md').write_text('本區記錄實際模型程式、來源、PSF與本次文件製作方法。製作腳本保留本機環境路徑，供檢查資料如何整理，不宣稱解壓後可一鍵重訓。原始結果報告為保留副本，速度解讀以另附2026年9月22日補充及文件03為準。\n',encoding='utf-8')
inventory=[]
for q in (P/'完整數據').glob('*.csv'):
 f=pd.read_csv(q);inventory.append(dict(file=q.name,rows=len(f),columns=list(f.columns)))
save(P/'完整數據/TABLE_INVENTORY.json',inventory)
save(P/'來源與核對/FINAL_DOCUMENT_CHECK.json',dict(checked_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),word_documents=4,pages={'01':3,'02':5,'03':3,'04':4},all_15_pages_visually_reviewed=True,figure_count=6,source_copies_byte_identical=len(source),all_circle_rows=33480,all_displacement_rows=7200,summary_recompute_max_difference_nm=max(errors),word_displacement_and_SEM_tables_verified=True,word_all36_pillars_verified=True,original_timings_unchanged=True,GPU_concurrency='User-reported possible explanation; not quantitatively established',new_speed_benchmark=False))
# Source locks and historical protected records are not rebaselined.
for rel,h in load(R/'config/source_lock.json').items():assert sha(R/rel)==h
for r in load(R/'config/protected_sources.json'):assert sha(str(PROJECT_ROOT / r['path']))==r['sha256']
shutil.copy2(Path(str(PROJECT_ROOT / 'Deliverables/R130_Clean_Corrected_Corner9_Validation_Handoff_20260920/verify_package.py')),P/'verify_package.py')
items=[dict(path=q.relative_to(P).as_posix(),bytes=q.stat().st_size,sha256=sha(q)) for q in sorted(P.rglob('*')) if q.is_file() and q.name!='MANIFEST.json'];save(P/'MANIFEST.json',items)
small=P.parent/'R137_四份實驗文件_20260922.zip';full=P.with_suffix('.zip')
for zpath,files in [(small,list(P.glob('*.docx'))+[P/'00_閱讀與上傳說明.md',P/'閱讀文字版/05_完整統計細表.md']),(full,[q for q in P.rglob('*') if q.is_file()])]:
 assert not zpath.exists()
 with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
  for q in files:z.write(q,q.relative_to(P).as_posix())
 with zipfile.ZipFile(zpath) as z:
  assert z.testzip() is None
  for q in files:assert hashlib.sha256(z.read(q.relative_to(P).as_posix())).hexdigest()==sha(q)
 zpath.with_suffix('.zip.sha256').write_text(sha(zpath)+'  '+zpath.name+'\n',encoding='utf-8')
 print(zpath.name,round(zpath.stat().st_size/1024**2,2),'MiB')
# Discovery links only; preserve original report and observed numbers.
r=R/'README.md';s=r.read_text(encoding='utf-8');s+='\n## 詳細文件與速度情境補充 20260922\n\n[四份獨立文件閱讀入口](../../Deliverables/R137_WeakPSF_Web_Documents_20260922/00_閱讀與上傳說明.md)。使用者提出GPU同時被其他訓練程序使用的可能原因；目前計時不可據以判定模型固有快慢。[速度解讀補充](reports/R137_RUNTIME_SCOPE_ADDENDUM.md)。原始時間和量測值保留。\n';r.write_text(s,encoding='utf-8')
print('PASS: 4 documents, 15 reviewed pages, all source copies and archives verified')
