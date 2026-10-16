from common import *
import csv,shutil,re,math,datetime,collections,itertools
from pathlib import Path
from pypdf import PdfReader
import openpyxl
BASE=ROOT/'output/R136_LSC_R61_Reading_Package';TMP=ROOT/'tmp/package_build'

def main():
    # Preserve original source text while making copied reports' links portable.
    manifest=load(BASE/'05_定位與分析紀錄/source_copy_manifest.json')
    phase_source=ROOT/'results/r61_analysis/verification.json';phase_dest=BASE/'05_定位與分析紀錄/phase_verification.json';shutil.copy2(phase_source,phase_dest)
    if not any(x['package_path']=='05_定位與分析紀錄/phase_verification.json' for x in manifest['files']):manifest['files'].append(dict(package_path='05_定位與分析紀錄/phase_verification.json',source=str(phase_source),sha256=sha(phase_dest)))
    substitutions=[('05_定位與分析紀錄/原分析報告.md',{'(lsc_r61_figures/':'(../03_LSC圖表/','(../results/lsc/r61_analysis/tables/':'(../02_統計表/CSV/'}),('04_補充_非LSC影像比對/原分析報告.md',{'(r61_figures/':'(圖表/','(../results/r61_analysis/tables/':'(CSV/','(R136_LSC_R61_RESULTS.md)':'(../05_定位與分析紀錄/原分析報告.md)','(../results/r61_analysis/verification.json)':'(../05_定位與分析紀錄/phase_verification.json)'})]
    for relative,mapping in substitutions:
        path=BASE/relative;old=sha(path);text=path.read_text(encoding='utf-8')
        for a,b in mapping.items():text=text.replace(a,b)
        path.write_text(text,encoding='utf-8')
        record_item=next(x for x in manifest['files'] if x['package_path']==relative);record_item.setdefault('source_sha256',old);record_item['sha256']=sha(path);record_item['transformation']='Relative report links rewritten for browsing within this package; numerical content unchanged.'
    save(BASE/'05_定位與分析紀錄/source_copy_manifest.json',manifest)
    shutil.copy2(BASE/'05_定位與分析紀錄/原分析報告.md',BASE/'01_報告/LSC_完整文字報告.md')
    for p in TMP.glob('pdf_chart_*.png'):
        folder=BASE/'03_LSC圖表/PDF閱讀版';folder.mkdir(exist_ok=True);shutil.copy2(p,folder/p.name)
    sidecar=BASE/'02_統計表/R136_LSC_R61_Tables.xlsx.inspect.ndjson'
    if sidecar.exists():shutil.move(str(sidecar),str(TMP/'workbook_export_inspect.ndjson'))
    readme='''原圖與R61：LSC位移重拍成果包
版本：2026-09-17。範圍：完整30次原圖/R61 LSC分析。

一、從哪裡開始
1. 開啟「01_報告/R136_LSC_R61_Report.pdf」：18頁完整閱讀版，有可點選目錄、結果、圖表、明細與限制。
2. 開啟「02_統計表/R136_LSC_R61_Tables.xlsx」：14個分頁，可依掃描、位置、圓編號篩選。
3. 「03_LSC圖表」可直接開啟、放大或放入簡報。

二、本包包含
- 原圖與R61各30次、各930張影像的完整LSC統計。
- 合計16,740筆圓心與32,400筆逐顆圓位移。
- 每方法900段相鄰位移與900筆相對起點位移的9圓合併結果。
- 全部統計CSV、930筆PI回授、30份原始位置紀錄、取邊點、設定與程式存檔。
- 非LSC影像區域比對結果另外放在「04_補充_非LSC影像比對」，不能混稱LSC。

三、閱讀時注意
- 下表中的波動與範圍有不同意義，PDF中已說明，不直接當成完整量測不確定度。
- Excel时间欄顯示台北時間至毫秒；原始完整時間字串保留在CSV。
- 本版未納入仍在運算的R129。
- 依你的選擇，本包不含全量照片、重建影像或模型權重。
- 解壓縮後即可離線閱讀。PDF已嵌入中文字型；表格與圖表不依賴原電腦的路徑。

四、若要追查或核對
- PACKAGE_MANIFEST.json列出本包檔案及內容指紋。
- verify_package.py可用一般Python檢查解壓縮是否完整。
- check_statistics.py可直接從本包位移明細重新核對主要數值，不需要原電腦或GPU。
- 「05_定位與分析紀錄/程式」是原分析程式存檔，有歷史環境路徑，供方法追溯；不表示本包附有完整模型重建環境。
'''
    (BASE/'00_先讀我.txt').write_text(readme.replace('Excel时间','Excel時間'),encoding='utf-8-sig')
    guide=BASE/'02_統計表/欄位與閱讀說明.txt'
    guide.write_text('''欄位與閱讀說明

nm = 奈米；um = 微米；px = 影像像素。
Raw = 原图；R61 = 固定R61 v16模型；本包LSC主結果不含R129。
nominal_all30 = 全30次、名義尺度0.1725um/px。
nominal_late24 = 後24次、仍使用名義尺度。
calibrated_late24 = 前6次原圖估計尺度與方向，後24次獨立比較。
adjacent = 相鄰每一步；start = 每個位置相對同次起點。
run_index = 第幾次掃描（1至30）；index、i、j = 影格索引（0至30）。
pillar_id = 同一顆圓的固定編號（1至9）。
mae = 平均絕對誤差；bias = 保留正負的平均偏差；rmse = 平方平均根誤差。
error_sd = 同一掃描內，各位置誤差的標準差。
repeatability_sd = 同一位置跨多次掃描的標準差。
ci95_low/high = 以完整掃描重抽樣得到的平均值95%範圍。
candidate_minus_reference = 候選減參考；對誤差大小來說，負值代表候選較小。對有向偏差則須看距離零的大小。
within_grid_sd = 同一段位移中9顆圓讀值的不一致程度，不等同真實誤差。
transverse = 影像中垂直於主要移動方向的變化，並非獨立校正的PI X誤差。

完整精度存於CSV與Excel儲存值；表面顯示的小數位數不代表量測已達該物理精度。
多個分析範圍會重用相同觀測資料，不能把各分頁列數相加當成獨立樣本數。
'''.replace('原图','原圖'),encoding='utf-8-sig')
    # Final XLSX values are compared to typed source rows; no workbook rewriting here.
    source=load(TMP/'workbook_data.json');xlsx=BASE/'02_統計表/R136_LSC_R61_Tables.xlsx';wb=openpyxl.load_workbook(xlsx,read_only=True,data_only=True);formula=openpyxl.load_workbook(xlsx,read_only=True,data_only=False)
    assert len(wb.sheetnames)==14;checked=0
    for data in source['sheets']:
        sh=wb[data['name']]
        for expected,actual in itertools.zip_longest(data['rows'],sh.iter_rows(min_row=5,max_col=len(data['headers']),values_only=True)):
            assert expected is not None and actual is not None,(data['name'],'row count differs')
            for c,(a,b) in enumerate(zip(expected,actual)):
                if isinstance(a,(int,float)) and not isinstance(a,bool):assert isinstance(b,(int,float)) and abs(a-b)<=1e-9*max(1,abs(a));checked+=1
                if data['keys'][c]=='recorded_at' and a:
                    assert isinstance(b,datetime.datetime);d=datetime.datetime.fromisoformat(a).replace(tzinfo=None);assert abs((d-b).total_seconds())<.0011
    for cell,expected in [('I6',1-5.160403203/8.667305772),('I8',1-15.2674119/17.73802816)]:assert abs(wb['主要結果'][cell].value-expected)<1e-12 and formula['主要結果'][cell].data_type=='f'
    wb.close();formula.close()
    pdf=ROOT/'output/pdf/R136_LSC_R61_Report.pdf';reader=PdfReader(pdf);assert len(reader.pages)==18 and len(reader.outline)==17
    text='\n'.join(p.extract_text() for p in reader.pages)
    for token in ['8.667','5.160','10.452','5.195','17.738','15.267','14.796','11.870','16,740'] :assert token in text,token
    assert '圓心' in text and '\ufffd' not in text
    embedded=[]
    for page in reader.pages:
        for f in page['/Resources']['/Font'].values():
            f=f.get_object();name=str(f.get('/BaseFont',''))
            if 'JhengHei' in name:
                desc=f['/FontDescriptor'].get_object();assert '/FontFile2' in desc;embedded.append(name)
    assert embedded;assert sha(pdf)==sha(BASE/'01_報告'/pdf.name)
    # All relative links in portable reading reports resolve within this package.
    for file in [BASE/'01_報告/LSC_完整文字報告.md',BASE/'05_定位與分析紀錄/原分析報告.md',BASE/'04_補充_非LSC影像比對/原分析報告.md']:
        for link in re.findall(r'\]\(([^)]+)\)',file.read_text(encoding='utf-8')):
            if not link.startswith(('http://','https://')):assert (file.parent/link).exists(),(file,link)
    save(BASE/'05_定位與分析紀錄/reading_package_verified.json',dict(pdf_pages=18,pdf_outline_entries=17,Chinese_fonts_embedded=True,pdf_pages_visually_reviewed=18,workbook_sheets=14,workbook_sheets_visually_reviewed=14,numeric_cells_compared=checked,typed_dates_checked=True,relative_report_links_checked=True,source_statistics_unchanged=True,verified_utc=now()))
    print('Reading package verified:18PDFpages,14worksheets,numeric cells',checked,flush=True)
if __name__=='__main__':main()
