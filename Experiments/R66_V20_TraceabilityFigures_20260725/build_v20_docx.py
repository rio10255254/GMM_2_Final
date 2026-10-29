from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
PROJECT_ROOT = Path(__file__).resolve().parents[2]

SOURCE = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v19_第六章更新版.docx")
TARGET = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v20_量測位置補圖版.docx")
FIGDIR = Path(str(PROJECT_ROOT / 'Experiments/R66_V20_TraceabilityFigures_20260725/figures'))

FIGURES = {
    "N01": ("01_static_50x_pillar_identity.png",
            "50倍靜態CD驗證影像之中心九顆pillar編號。編號已依拍攝方向與SEM資料逐顆對應；本圖只用於辨識結構，不另外疊加圓心標記。影像以原生0–255灰階顯示，未逐圖拉伸對比。", 13.0),
    "N02": ("02_50x_p29_radial_profile_location.png",
            "50倍P29之徑向剖面取樣方式。白心藍點為由Raw影像估計一次後固定的局部取樣原點，藍色虛線為初始參考半徑而非擬合邊界；淡藍線示意多方向取樣，橘線只標出其中一個方向。圖6-28的曲線由360個等角方向彙整而成，不是橘線所代表的單一剖面。", 11.5),
    "N03": ("03_20x_p29_radial_profile_location.png",
            "20倍P29之徑向剖面取樣方式。白心藍點為由Raw影像估計一次後固定的局部取樣原點，藍色虛線為初始參考半徑而非擬合邊界；橘線只示意一個方向，圖6-33則彙整360個等角方向。Raw、Wiener與MOSR共用相同原點及方向。", 11.5),
    "N04": ("04_50x_p29_cd_sampling_location.png",
            "50倍P29之水平CD取樣位置。白心藍點為由Raw影像估計一次後固定的局部取樣原點；橘色帶涵蓋中心列及上下各三列，七列灰階先平均成單一水平剖面，再由原點向左右尋找最近的相對閾值交點。三種方法共用相同原點與取樣範圍。", 11.5),
    "N05": ("05_20x_p29_cd_sampling_location.png",
            "20倍P29之水平CD取樣位置。白心藍點為Raw影像固定的局部取樣原點，橘色帶表示先平均成單一剖面的七個像素列；因20倍取樣較粗，搜尋範圍會涵蓋相鄰結構，但量測僅採用離該原點最近的一對左右交點。", 11.5),
    "N06": ("06_displacement_measurement_location.png",
            "位移量測位置與相鄰影格追蹤方式。（a）同一pillar在−0.55R、−0.275R、0、0.275R與0.55R的五條水平弦；（b）淡藍虛線區為依上一影格位置與預期步距建立的左右搜尋範圍，橘點為最大梯度邊界，藍色短線只表示兩側邊界中點在該水平弦上的位置，不代表二維圓心。平台回授僅作為量測後的獨立位移參考，不參與邊界搜尋。", 14.5),
    "N07": ("07_50x_via_identity.png",
            "50倍via序列之中心九顆結構編號。V1至V9依SEM資料的實體排列順序定義；本圖只用於辨識結構，不疊加圓心標記。", 13.0),
    "N08": ("08_20x_via_identity.png",
            "20倍via序列之中心九顆結構編號。圖中排列已依20倍拍攝方向完成上下翻轉，並與50倍對應至相同實體via。", 13.0),
    "N09": ("09_line_measurement_regions.png",
            "50倍線條影像的量測區域與剖面方向。藍框為垂直線條區域，橘框為水平線條區域；框內直線表示形成中位數灰階剖面的方向。", 14.5),
    "N10": ("10_ablation_p29_matched_comparison.png",
            "消融實驗之P29同位置影像比較。由左至右為Raw、完整模型與移除原始影像位置基準的模型輸出；三張影像均使用原生0–255灰階與相同裁切範圍，未逐圖調整對比。", 15.0),
    "N11": ("11_multimodel_50x_pillar_matched_comparison.png",
            "50倍pillar之Raw、NAFNet、EDSR-x1與MOSR同位置影像比較。四種方法使用相同影格與裁切範圍，並直接以原生0–255灰階顯示，未對各模型輸出個別拉伸對比。", 15.5),
    "N12": ("12_multimodel_50x_via_matched_comparison.png",
            "50倍via之Raw、NAFNet、EDSR-x1與MOSR同位置影像比較。四種方法使用相同影格與裁切範圍，並直接以原生0–255灰階顯示，圖中對比差異來自模型輸出本身。", 15.5),
    "N13": ("13_50x_sequence_pillar_identity.png",
            "50倍位移序列與多模型比較所使用的九顆pillar編號。排列依該序列的實際拍攝方向與SEM資料對應；本圖只用於辨識結構，不疊加圓心標記。", 13.0),
    "N14": ("14_20x_sequence_pillar_identity.png",
            "20倍位移序列與多模型比較所使用的九顆pillar編號。排列已依20倍拍攝方向與SEM資料逐顆對應；本圖只用於辨識結構，不疊加圓心標記。", 13.0),
}


def paragraph_by_text(doc, prefix, style=None):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise RuntimeError(f"Paragraph not found: {prefix}")


def image_template_before(caption):
    node = caption._p.getprevious()
    while node is not None:
        if node.tag.endswith("}p"):
            for p in caption._parent.paragraphs:
                if p._p is node:
                    return p
        node = node.getprevious()
    raise RuntimeError("No image paragraph template")


def make_figure(doc, key, image_template, caption_template):
    filename, caption, width_cm = FIGURES[key]
    image_path = FIGDIR / filename
    if not image_path.is_file():
        raise FileNotFoundError(image_path)

    ip = doc.add_paragraph()
    ip._p.get_or_add_pPr().getparent().remove(ip._p.pPr)
    ip._p.insert(0, deepcopy(image_template._p.pPr))
    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.add_run().add_picture(str(image_path), width=Cm(width_cm))

    cp = doc.add_paragraph()
    cp._p.get_or_add_pPr().getparent().remove(cp._p.pPr)
    cp._p.insert(0, deepcopy(caption_template._p.pPr))
    cp.style = caption_template.style
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run(f"圖 6-{key} {caption}")
    return ip, cp


def insert_before(paragraph, elements):
    for element in elements:
        paragraph._p.addprevious(element._p)


def insert_after(paragraph, elements):
    anchor = paragraph._p
    for element in elements:
        anchor.addnext(element._p)
        anchor = element._p


def append_reference(paragraph, text):
    if paragraph.text and not paragraph.text.endswith(("。", "；")):
        paragraph.add_run("。")
    paragraph.add_run(text)


def update_runs(paragraph, transform):
    old = paragraph.text
    new = transform(old)
    if new == old:
        return
    if paragraph.runs:
        paragraph.runs[0].text = new
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new)


def set_paragraph_text(paragraph, text):
    update_runs(paragraph, lambda _: text)


def main():
    doc = Document(str(SOURCE))
    cap_template = paragraph_by_text(doc, "圖 6-24", "Figure Caption")
    img_template = image_template_before(cap_template)

    # Make the shared measurement definitions match the implemented procedures.
    set_paragraph_text(
        paragraph_by_text(doc, "量測分兩種模式。"),
        "量測分兩種模式。各結構先由預先指定的九宮格位置建立局部搜尋區，再從Raw影像估計一次局部取樣原點；同一影格的Wiener與MOSR均沿用該原點，不重新定位。水平模式先將原點所在列及上下各三列共七列灰階取平均，形成單一水平剖面，再由原點向左右尋找最近的一對閾值交點；徑向模式則由同一原點沿等角方向取樣，先彙整成徑向灰階剖面，再由原點向外尋找交點。第6.1節的PSF與邊緣剖面使用360個方向；第6.2.4節及多模型統一CD比較使用720個方向，以降低單一方向像素取樣造成的波動。"
    )
    set_paragraph_text(
        paragraph_by_text(doc, "每顆 pillar 以固定中心"),
        "每顆pillar以固定中心與參考半徑建立搜尋區域，並在相對中心高度−0.55R、−0.275R、0、0.275R與0.55R處取得五條水平弦。每條弦先依圓形參考半徑估計左右邊界位置，再於局部搜尋範圍內尋找符合結構明暗極性的最大梯度；峰值以三點二次內插取得亞像素位置[82,83]，左右邊界的中點即為該弦的中心估計。"
    )
    set_paragraph_text(
        paragraph_by_text(doc, "式（4-18）的意思是"),
        "式（4-18）的意思是：第j顆pillar在第k條弦上的中心，取該弦左右邊界的中點；該顆pillar的中心則取有效弦中心的中位數。量測至少須保留三條有效弦，且左右寬度、中心偏移、邊界強度及弦間中心離散度須通過共同門檻；未通過者直接排除，不以其他位置替代。"
    )
    set_paragraph_text(
        paragraph_by_text(doc, "取成對邊界的中點"),
        "取成對邊界的中點，可降低單側模糊或局部亮環對中心估計的干擾；取多條弦的中位數，則可抑制單一剖面受雜訊或局部缺陷影響而產生的離群值。相鄰影格的搜尋位置由上一影格已接受的中心加上預期步距建立，平台實際回授只在影像量測完成後作為獨立位移參考，不參與邊界搜尋。取得各影格的中心後，相鄰影格之間的位移即為中心差值乘上物體空間取樣尺度："
    )

    set_paragraph_text(
        paragraph_by_text(doc, "CD量測使用Raw影像鎖定的中心"),
        "CD量測先由預先指定的九宮格位置在Raw影像中估計每顆pillar的局部取樣原點與參考半徑，Wiener與MOSR沿用同一原點，不各自重新定位。SEM量測資料提供的是半徑，因此先乘以2轉為直徑；物體空間取樣為69.0 nm/pixel。相對閾值由35%至65%逐級掃描，並以水平七列平均剖面與720個等角方向平均徑向剖面交叉檢查。每顆pillar、每種方法共得到14個CD，逐顆圖只呈現最小值與最大值，整體統計則使用全部閾值、兩種剖面與九顆pillar的誤差分布。"
    )
    set_paragraph_text(
        paragraph_by_text(doc, "兩張閾值敏感度圖說明排名翻轉的原因"),
        "兩張閾值敏感度圖說明排名翻轉的原因。Wiener曲線在50%附近形成尖銳低谷，閾值稍微提高或降低後誤差便快速上升；MOSR雖不在每一個閾值都最低，但跨越兩種剖面方式時整體誤差較低，也沒有依賴單一交點才能勝出。這項分析不否定Wiener能使影像變銳利，而是把『視覺銳化』與『可穩定量測的CD修復』分開。水平七列平均剖面與720個等角方向平均徑向剖面的閾值敏感度分別如圖6-34與圖6-35所示。"
    )
    set_paragraph_text(
        paragraph_by_text(doc, "圖 6-34 水平七列剖面", "Figure Caption"),
        "圖 6-34 水平七列平均剖面的CD閾值敏感度。中心列及上下各三列先平均成單一剖面，再比較Raw、Wiener與MOSR在共同pillar身分、中心與SEM參考下的平均CD絕對誤差。"
    )
    set_paragraph_text(
        paragraph_by_text(doc, "圖 6-35 360度平均徑向剖面", "Figure Caption"),
        "圖 6-35 720個等角方向平均徑向剖面的CD閾值敏感度。比較Raw、Wiener與MOSR在共同pillar身分、中心與SEM參考下的平均CD絕對誤差。"
    )

    # Keep the old figure numbers before inserting new captions.
    old_caption_number = {}
    in_ch6 = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("6.1"):
            in_ch6 = True
        if t == "結論與未來工作":
            in_ch6 = False
        if in_ch6 and p.style.name == "Figure Caption":
            m = re.match(r"圖\s*6-(\d+)", t)
            if m:
                old_caption_number[p._p] = int(m.group(1))

    # Edge-profile location figures.
    cap27 = paragraph_by_text(doc, "圖 6-27", "Figure Caption")
    set_paragraph_text(
        cap27,
        "圖 6-27 50倍Raw、Wiener與MOSR的徑向邊緣剖面及過渡寬度。剖面由前圖所示局部原點出發，彙整360個等角方向後形成，並非前圖橘線所代表的單一方向或單一水平剖面；曲線呈現中心pillar的方向中位數，點與誤差範圍則彙整九顆pillar的10–90%過渡寬度。"
    )
    insert_before(image_template_before(cap27), make_figure(doc, "N02", img_template, cap_template))
    cap31 = paragraph_by_text(doc, "圖 6-31", "Figure Caption")
    set_paragraph_text(
        cap31,
        "圖 6-31 20倍Raw、Wiener與MOSR的徑向邊緣剖面及過渡寬度。剖面由前圖所示局部原點出發，彙整360個等角方向後形成，不對應單一橫剖線；曲線呈現中心pillar的方向中位數，點與誤差範圍則彙整九顆pillar的10–90%過渡寬度。"
    )
    insert_before(image_template_before(cap31), make_figure(doc, "N03", img_template, cap_template))

    # CD identity and sampling-location figures.
    p_cd50 = paragraph_by_text(doc, "為避免以單一有利閾值")
    set_paragraph_text(
        p_cd50,
        "為避免以單一有利閾值評估影像重建方法，本圖對Raw、Wiener與MOSR都不保留50%代表值，也不以任何中央值取代。對每顆pillar、每種方法分別掃描35–65%七個相對閾值，並使用水平七列平均剖面與720個等角方向平均徑向剖面，共形成14個量測值；圖中只畫出其中最小與最大直徑。SEM圓點為已知尺寸參考，不參與上下界計算。九顆pillar在35%至65%閾值與兩種剖面方法下的CD上下界整理於表6-6與圖6-32；pillar編號與SEM對應如圖6-N01，水平七列平均範圍如圖6-N04。"
    )
    insert_after(p_cd50, [*make_figure(doc, "N01", img_template, cap_template),
                          *make_figure(doc, "N04", img_template, cap_template)])
    p_cd20 = paragraph_by_text(doc, "為確認20倍結果可延伸至非PSF校正影格")
    set_paragraph_text(
        p_cd20,
        "為確認20倍結果可延伸至非PSF校正影格，本節選用1.0 µm位移序列的最後一張影像進行跨閾值CD檢查。該影像未用於PSF擬合，九顆pillar身分與SEM直徑維持不變；中心只由Raw影像估計一次共同網格位移，Wiener與MOSR不重新定位。每種方法同樣使用35%至65%七個閾值、水平七列平均剖面與720個等角方向平均徑向剖面，共得到126個尺寸誤差。20倍非PSF校正影格的跨閾值結果整理於圖6-36與表6-8，實際水平取樣範圍如圖6-N05。"
    )
    insert_after(p_cd20, make_figure(doc, "N05", img_template, cap_template))

    # Displacement measurement definition.
    p_disp = paragraph_by_text(doc, "本節只比較Raw、固定PSF之Wiener反卷積與MOSR")
    set_paragraph_text(
        p_disp,
        "本節只比較Raw、固定PSF之Wiener反卷積與MOSR，不加入消融條件或其他神經網路。四組序列均使用相同的五弦最大梯度邊界定位方法；每組30個步次由Raw影像的可量測性與運動品質事先固定，三種方法共用相同影格、九顆pillar、由首張Raw影像估計的局部取樣原點與預期步距。相鄰影格的搜尋位置由上一影格已接受的水平中心加上預期步距建立，PI平台實際回授只在追蹤完成後作為獨立位移參考，不參與邊界搜尋。五弦取樣、搜尋範圍與左右邊界定位方式如圖6-N06所示。"
    )
    insert_after(p_disp, make_figure(doc, "N06", img_template, cap_template))

    # Ablation and model-comparison visual evidence.
    p_ablation = paragraph_by_text(doc, "消融實驗與前節位移改善結果分開呈現")
    set_paragraph_text(
        p_ablation,
        "消融實驗與前節位移改善結果分開呈現。各消融條件固定起始權重、資料切分、PSF、訓練預算與評估影格，只改變待檢查的模型設計；圖中的範圍來自兩次固定初始化，用來判斷效果方向是否可重現。完整模型與移除原始影像位置基準的P29同位置輸出如圖6-N10所示。"
    )
    insert_after(p_ablation, make_figure(doc, "N10", img_template, cap_template))
    p_models = paragraph_by_text(doc, "本節比較Raw、NAFNet、EDSR-x1與MOSR")
    set_paragraph_text(
        p_models,
        "本節比較Raw、NAFNet、EDSR-x1與MOSR，不納入Wiener與消融條件。Pillar與via皆先由Raw影像估計一次九顆結構的局部取樣原點，其他模型沿用相同原點；位移沿用五弦最大梯度方法，CD則掃描35%至65%七個相對門檻，並同時使用水平七列平均剖面與720個等角方向平均徑向剖面。四種方法因此在相同結構身分、物體空間尺度與SEM參考下接受比較。影像比較直接使用各模型輸出的原生0–255灰階，不做逐模型對比拉伸；因此NAFNet或EDSR-x1較高的對比屬於模型輸出本身，而非作圖強化。CD採局部相對門檻、位移採梯度峰值，單純的線性對比變化不會自動形成量測優勢，但模型造成的非線性邊界、振鈴與局部飽和仍會完整反映在結果中。同一50倍pillar與via區域的影像比較分別如圖6-N11與圖6-N12所示。"
    )
    insert_after(p_models, [*make_figure(doc, "N11", img_template, cap_template),
                            *make_figure(doc, "N12", img_template, cap_template)])

    # Identity maps used by per-structure CD plots.
    p_unified = paragraph_by_text(doc, "每一倍率與形貌均量測九顆結構")
    set_paragraph_text(
        p_unified,
        "每一倍率與形貌均量測九顆結構；四種方法共用由Raw影像估計一次的局部取樣原點，並將35%至65%七個門檻、水平七列平均剖面與720個等角方向平均徑向剖面全部納入。每個倍率、形貌與方法共有126筆尺寸量測，主要統計量為相對逐顆SEM直徑的中位絕對誤差。20倍與50倍影像依各自拍攝方向對應至相同的實體pillar編號；via的SEM半徑則先乘以2轉為直徑。50倍與20倍pillar編號分別如圖6-N13與圖6-N14；via編號分別如圖6-N07與圖6-N08。"
    )
    insert_after(p_unified, [*make_figure(doc, "N13", img_template, cap_template),
                             *make_figure(doc, "N14", img_template, cap_template),
                             *make_figure(doc, "N07", img_template, cap_template),
                             *make_figure(doc, "N08", img_template, cap_template)])

    # Line-pattern measurement location.
    cap64 = paragraph_by_text(doc, "圖 6-64", "Figure Caption")
    insert_before(image_template_before(cap64), make_figure(doc, "N09", img_template, cap_template))

    # Add explicit in-text links to the new evidence.
    append_reference(paragraph_by_text(doc, "由三組影像可觀察到"), "徑向剖面的實際中心與取樣方向如圖6-N02所示。")
    append_reference(paragraph_by_text(doc, "20×影像的物體空間取樣較粗"), "20倍P29的共同取樣位置如圖6-N03所示。")
    append_reference(paragraph_by_text(doc, "線條區域以Raw固定相同邊界位置"), "實際量測區域與剖面方向如圖6-N09所示。")

    # Determine new sequential numbering for every Chapter 6 caption.
    existing_map = {}
    new_map = {}
    sequence = []
    in_ch6 = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("6.1"):
            in_ch6 = True
        if t == "結論與未來工作":
            in_ch6 = False
        if in_ch6 and p.style.name == "Figure Caption":
            sequence.append(p)

    for new_num, p in enumerate(sequence, start=1):
        old_num = old_caption_number.get(p._p)
        if old_num is not None:
            existing_map[old_num] = new_num
        marker = re.match(r"圖\s*6-(N\d+)", p.text.strip())
        if marker:
            new_map[marker.group(1)] = new_num

    # First update references in Chapter 6 and Chapter 7 using collision-safe tokens.
    def ref_transform(text):
        for old, new in existing_map.items():
            text = re.sub(rf"圖\s*6-{old}(?!\d)", f"@@OLD{old}@@", text)
        for marker, new in new_map.items():
            text = text.replace(f"圖6-{marker}", f"@@{marker}@@")
            text = text.replace(f"圖 6-{marker}", f"@@{marker}@@")
        for old, new in existing_map.items():
            text = text.replace(f"@@OLD{old}@@", f"圖6-{new}")
        for marker, new in new_map.items():
            text = text.replace(f"@@{marker}@@", f"圖6-{new}")
        return text

    in_scope = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("6.1"):
            in_scope = True
        if t.startswith("參考文獻"):
            in_scope = False
        if in_scope and p.style.name != "Figure Caption":
            update_runs(p, ref_transform)

    # Finally set every caption number in document order.
    for new_num, p in enumerate(sequence, start=1):
        update_runs(p, lambda t, n=new_num: re.sub(r"^圖\s*6-(?:\d+|N\d+)", f"圖 6-{n}", t))

    # Normalize new captions and image paragraphs to the existing thesis layout.
    for p in doc.paragraphs:
        if p.style.name == "Figure Caption" and re.match(r"圖\s*6-\d+", p.text.strip()):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = cap_template.paragraph_format.space_before
            p.paragraph_format.space_after = cap_template.paragraph_format.space_after
            p.paragraph_format.line_spacing = cap_template.paragraph_format.line_spacing

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(TARGET))

    qa = TARGET.with_suffix(".figure_map.txt")
    with qa.open("w", encoding="utf-8") as f:
        f.write("Existing figure map (v19 -> v20)\n")
        for old in sorted(existing_map):
            f.write(f"6-{old} -> 6-{existing_map[old]}\n")
        f.write("\nNew figure map\n")
        for marker in sorted(new_map):
            f.write(f"{marker} -> 6-{new_map[marker]}\n")
    print(TARGET)
    print(f"Chapter 6 figures: {len(sequence)}")


if __name__ == "__main__":
    main()
