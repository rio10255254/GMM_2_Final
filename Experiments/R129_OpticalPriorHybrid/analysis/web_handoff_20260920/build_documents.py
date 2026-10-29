"""Create four self-contained thesis handoff documents and numerical attachments."""
from pathlib import Path
import json,hashlib,shutil,re,zipfile,csv,sys
import numpy as np
import pandas as pd
import subprocess
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
PROJECT_ROOT = Path(__file__).resolve().parents[4]

sys.stdout.reconfigure(encoding='utf-8')
REPO=Path(str(PROJECT_ROOT))
SRC=REPO/'Deliverables/R129_Thesis_Rewrite_Ready_20260920'
HAND=SRC/'04_Results/R129_Handoff'
DATA=HAND/'01_Raw_R129_Comparison'
OUT=REPO/'Deliverables/R129_Web_Thesis_Documents_20260920'
WORK=REPO/'Experiments/R129_OpticalPriorHybrid/results/web_handoff_20260920'
OUT.mkdir(exist_ok=True);WORK.mkdir(parents=True,exist_ok=True)
(OUT/'Editable_Word').mkdir(exist_ok=True);(WORK/'figures').mkdir(exist_ok=True)
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
for item in load(SRC/'MANIFEST.json'):assert sha(SRC/item['path'])==item['sha256'],item['path']
def textformat(value):
 if isinstance(value,(int,np.integer)):return str(value)
 if isinstance(value,(float,np.floating)):return '' if pd.isna(value) else f'{value:.3f}'
 return str(value)

class Doc:
 def __init__(self,name,title,scope):
  self.name=name;self.d=Document();s=self.d.sections[0]
  s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.72);s.right_margin=Inches(.72)
  for style in ['Normal','Title','Heading 1','Heading 2','Caption']:
   st=self.d.styles[style];st.font.name='Calibri';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei');st.font.color.rgb=RGBColor(0,0,0)
  self.d.styles['Normal'].font.size=Pt(10.5);self.d.styles['Normal'].paragraph_format.line_spacing=1.15;self.d.styles['Normal'].paragraph_format.space_after=Pt(6)
  self.d.styles['Title'].font.size=Pt(22);self.d.styles['Heading 1'].font.size=Pt(15);self.d.styles['Heading 2'].font.size=Pt(12)
  for n in ['Heading 1','Heading 2']:self.d.styles[n].paragraph_format.keep_with_next=True
  self.d.styles['Caption'].font.size=Pt(9);self.d.styles['Caption'].font.italic=False
  h=s.header.paragraphs[0];h.text='MOSR 論文改寫資料  2026年9月20日';h.style='Caption'
  f=s.footer.paragraphs[0];f.alignment=2;f.add_run(name+'  |  ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
  self.d.add_paragraph(title,'Title');self.p(scope)
  self.p('共同約定：本文只採固定R129作為MOSR方法；保留ACM輪廓與量測結果；本次移除消融與多模型排行榜。所有結果均依修正後同九顆、30次掃描、20倍實拍解讀。這是可供網頁版獨立閱讀的備稿文件，不是已改完的原論文。')
 def p(self,t):return self.d.add_paragraph(t)
 def h(self,t,level=1):return self.d.add_paragraph(t,'Heading '+str(level))
 def page(self,t):
  p=self.h(t);p.paragraph_format.page_break_before=True
 def bullets(self,items):
  for t in items:self.p('• '+t)
 def table(self,headers,rows,widths=None,font=9):
  t=self.d.add_table(rows=1, cols=len(headers));t.autofit=False
  if widths is None:widths=[7.0/len(headers)]*len(headers)
  for i,w in enumerate(widths):t.columns[i].width=Inches(w)
  for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
  for row in rows:
   cells=t.add_row().cells
   for i,v in enumerate(row):cells[i].text=textformat(v)
  for j,row in enumerate(t.rows):
   for i,c in enumerate(row.cells):
    c.width=Inches(widths[i]);tcPr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     x=OxmlElement('w:'+edge);x.set(qn('w:w'),'60' if edge in ['top','bottom'] else '80');x.set(qn('w:type'),'dxa');m.append(x)
    tcPr.append(m);bd=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');bd.append(x)
    tcPr.append(bd);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5E9ED' if j==0 else ('F5F7F8' if j%2==0 else 'FFFFFF'));tcPr.append(shade)
    for p in c.paragraphs:
     p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
     for run in p.runs:run.font.size=Pt(font);run.bold=(j==0)
   pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
  rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
  self.d.add_paragraph().paragraph_format.space_after=Pt(2)
  return t
 def figure(self,path,caption,maxh=6.5,maxw=6.95):
  path=Path(path)
  with Image.open(path) as im:w,h=im.size
  wi=min(maxw,maxh*w/h)
  p=self.d.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Inches(wi))
  self.d.add_paragraph(caption,'Caption')
 def end(self,source):
  self.h('資料來源與本文件範圍',2);self.p(source);self.p('數值完整精度與逐筆資料見05_完整數據.xlsx；更底層的64點邊界與720方向資料另附完整數據壓縮包。這些附件是資料本身，無須存取作者電腦。')
  for root in [self.d._element,self.d.styles.element]:
   for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
  if len(sys.argv)>1 and not self.name.startswith(sys.argv[1]):return
  path=OUT/'Editable_Word'/f'{self.name}.docx';self.d.save(path);print('Created',path.name,flush=True)

# Split the training examples into two readable embedded pages using the
# existing scientific-plot runtime (matplotlib is absent from bundled Python).
plot_code = r'''
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
SRC=Path(r'Deliverables/R129_Thesis_Rewrite_Ready_20260920')
WORK=Path(r'Experiments/R129_OpticalPriorHybrid/results/web_handoff_20260920')
cats=[('irregular_curved','Irregular boundaries'),('corners_lines_junctions','Lines and junctions'),('dense_close_structures','Dense structures'),('textured_edges','Textured edges'),('clean_identity','Clean control'),('noisy_identity','Noisy control'),('flat_negative','Flat control')]
for no,sel in [(1,cats[:4]),(2,cats[4:])]:
 fig,axs=plt.subplots(len(sel),3,figsize=(9,2.6*len(sel)),layout='constrained')
 for row,(cat,label) in enumerate(sel):
  with np.load(SRC/f'03_Dataset/examples/{cat}.npz') as z:im=z['epoch0_observed_uint8'];target=z['target']
  arrays=[im[0,128:256,128:256]/255,target[0,128:256,128:256],im[1,128:256,128:256]/255]
  for col,a in enumerate(arrays):
   axs[row,col].imshow(a,cmap='gray',vmin=0,vmax=1,interpolation='nearest');axs[row,col].set_xticks([]);axs[row,col].set_yticks([])
   if row==0:axs[row,col].set_title(['Observed A','Clean target A','Observed shifted B'][col],fontsize=10)
  axs[row,0].set_ylabel(label,fontsize=10)
 fig.savefig(WORK/f'figures/training_examples_{no}.png',dpi=160,bbox_inches='tight');plt.close(fig)
'''
subprocess.run(['python','-c',plot_code],check=True)

# DOCUMENT 1
d=Doc('01_完整改寫指南','MOSR論文改寫指南','本文件給將接手改寫的網頁版對話使用。它包含已決定的研究範圍、七章改寫安排、圖表配置、文字取捨與可直接使用的交接指令。請先閱讀本文件，再讀方法與結果文件。')
d.h('研究主線與已定決策')
d.p('論文主線是以實測光學資訊輔助顯微影像重建，再使用一致的ACM輪廓與圓心量測流程，評估正常拍攝下微位移、尺寸及輪廓的重拍穩定性。正文方法名稱沿用MOSR，實作身分是R129，不再介紹多個候選模型。')
d.bullets(['以原v34完整標註版為底稿，保留七章及主要光學、設備與量測內容。','V42作為現用模型的架構、資料、訓練、圖片與結果須完整抽換；不只替換文字名稱。','ACM與Contour是教授要求的正文重點，必須完整保留方法與實驗結果。','本次刪除消融及所有消融回指，不改名成模組檢查加回；未來再重新設計。','不安排R128／R61／V42排行，不以擊敗舊模型作改稿前提。','主要改善先寫；必要限制簡短且數據完整，不另開失敗專章，也不把未改善結果改寫成改善。'])
d.h('上傳與閱讀方式',2)
d.p('網頁版先讀本文件與原始v34論文；接著讀02模型與訓練、03量測與實驗、04結果與寫作。05試算表保存所有表格明細。若需要重算原邊界，再上傳06底層數據壓縮包。四份文件的圖與必要表格已內嵌，不需要讀資料夾或本機路徑。')
d.page('七章改寫安排')
d.table(['部分','保留與改寫'],[
('摘要與第1章','保留量測導向問題，以MOSR自己的實拍結果重寫貢獻；刪舊排名與消融成果。'),('第2章文獻','保留光學、重建、量測與平移一致性背景；僅更新與現用架構不符的銜接，不扩張為大型XAI綜述。'),('第3章理論','光學、PSF、BADS與量測一致性大部分保留；3.6改為共享多次修正的神經重建背景。'),('第4章方法','4.4架構、4.5資料、4.6訓練整段換新；4.7保留ACM粗輪廓、最後取邊與LSC／尺寸／形狀計算。'),('第5章實作','保留設備；換實際資料及訓練設定、固定九顆和30次掃描配置；刪多模型訓練與消融配置。'),('第6章結果','以Raw對MOSR呈現位移、Contour重複性及SEM；效率簡述，刪舊排行與消融。'),('第7章結論','聚焦已測條件的穩定性與流程可追查性；限制與未來工作簡短交代。')],[1.5,5.5])
d.p('原题目「結合點擴散函數先驗之量測導向顯微影像重建：關鍵尺寸與微位移量測」可先保留。尺寸仍有實質分析，無須因不是所有數值改善就刪掉尺寸主題。不要另加XAI到題目、宣稱完全白箱或任意影像都物理正確。')
d.page('模型資料與訓練的替換位置')
d.table(['原節','替換內容'],[
('4.4.1','同尺寸輸出與共享四次修正；內部4倍表示不等於物理解析度4倍。'),('4.4.2','實測平均PSF與候選融合在輸入端運作，不是逐像素空間變異校正。'),('4.4.3','邊緣線索與改善上取樣；刪未接入的獨立方向感知邊界分支。'),('4.4.4與4.4.5','原圖參照、受限細節輸出、品質調幅及外框保留；品質調幅不是可信概率。'),('4.5','512／128／128母場景、七類來源、生成配對、角色與分組；展示真實訓練輸入範例。'),('4.6','實際影像、邊緣、配對殘差三項訓練目標與接續學習來源。'),('5.4','固定最終權重、16循環、8192更新、訓練環境、資料與評估配置。')],[1.35,5.65])
d.p('初始化及固定光學資料有真實歷史來源。完整抽換V42現用模型，不代表把沿用的校正資料改稱新量測，也不代表隱去初始化來源。正文可用「前期訓練後接續學習」敘述，不必展開內部版本競賽。')
d.h('刪除清單',2)
d.p('原5.4.7及6.2.6消融章節與專屬圖表刪除；同步移除摘要、研究貢獻、討論、結論、目錄及交叉引用中的消融說法。原5.4.5、5.4.6及6.2.7的多模型配置／排名刪除，必要的拍攝配置改併入新資料說明。')
d.page('結果章的完整配置')
d.table(['建議位置','正文內容','主要材料'],[
('6.1','光學校正及原有理論對應；核對其中舊MOSR圖是否V42','原論文保留的校正資料'),('6.2.1','實際接續訓練狀態','訓練與開發曲線'),('6.2.2','修正後九顆與共同評估規則','全視野與九顆位置圖'),('6.2.3','ACM及最後邊界量測方法','原圖／重建取邊示意'),('6.2.4','輪廓、尺寸重複性與SEM對照','30次起點／終點輪廓、尺寸及SEM圖'),('6.2.5','短步與距起點位移','平均曲線與±1SD、主要結果表'),('原6.2.6','刪消融，不留空白待做章節','無'),('效率與討論','單模型時間、條件及必要限制','不做速度排名')],[1.0,3.45,2.55])
d.p('建議主圖順序：架構 → 訓練資料例子 → 訓練曲線 → 九顆位置 → ACM取邊 → 位移 → 起末輪廓 → 尺寸／SEM。完整明細由附件保留，不必全部貼到正文。Word圖號依實際章節更新。')
d.page('如何寫主張與限制')
d.table(['可以寫','不要延伸成'],[
('短步PI差距及重拍波動降低','所有位置、距離、樣品、日期都更準'),('30次輪廓更集中且九顆皆改善','整條邊界或真實尺寸已完全恢復'),('去平均半徑後與SEM形狀更接近','保留尺寸的完整輪廓也必然更準'),('模型與資料來源可以追查','每一步都有已證明的物理意義或完整XAI'),('同一模型內有光學與上取樣設計','各模組的獨立效益已由消融證明'),('本次單次重拍波動及平均值分量已估','4nm絕對準確度或95%完整總不確定度')],[3.5,3.5])
d.p('有效範圍是同片樣品、同一工作階段、20倍、九顆及本次量法。不要把9顆×31位置×30次或720角度當成大量獨立樣品。結果與建議要區分；待做工作不能寫成已完成。')
d.h('可直接用的主題段落',2)
d.p('本研究以正常拍攝條件下的顯微影像量測為目的，提出結合實測光學資訊的MOSR重建流程。重建後採用與原圖相同的ACM輪廓及圓心量測方式，並透過平台回授、30次重拍及SEM參照，評估微位移與輪廓讀取的穩定性。本文重點是已測條件下的實際量測改善，並將重拍一致性與完整尺寸準確度分開解讀。')
d.page('交給網頁版的完整指令')
d.p('請依我上傳的原v34論文與四份MOSR備稿文件進行改寫，在副本保留原七章架構。本文只採固定R129作為MOSR方法，完整抽換V42作為現用模型的架構、資料、訓練、圖片與結果；保留正確的光學理論、設備、尺度及量測方法。')
d.p('ACM主動輪廓與Contour結果是正文重點。方法須分清ACM粗輪廓、沿法線最大梯度的最後取邊、圓心／尺寸／形狀計算。結果用修正後同九顆的Raw與R129資料，呈現位移、輪廓及尺寸重拍穩定性，以及適當的SEM對照。')
d.p('本次不寫消融，不加入其他模型排行榜，不要求新拍攝或新訓練。主要改善先呈現，必要限制簡短且數據完整；不要把形狀改善寫成所有尺寸準確度改善，不把標準差或平均值分量稱為完整總不確定度。未提供的數字或已知不確定的SEM方向不得自行補成確定。')
d.p('先完成第4、5章方法與訓練，再寫第6章結果，最後更新第7章、緒論貢獻與中英文摘要。核對公式符號、方法名称、數字、圖表、目錄、交叉引用及參考文獻。完整數據在05試算表；四份文件內的圖表與敘述已可獨立使用，不要要求存取作者電腦路徑。')
d.h('交付前檢查',2)
d.bullets(['全文搜尋舊模型參數、epoch、速度名次與消融結論，檢查語意而非只搜V42字串。','未使用的舊模組公式與圖不留在方法章；沿用的真實校正來源不改寫。','正文與摘要都使用4.125nm等新九顆數字，不混入舊4.092nm或其他模型。','SEM方向推定與資料內不確定度範圍留在對應圖說，避免離開限制單獨摘錄。','所有圖片內嵌，表格數值與附件一致；改稿完成後檢查每頁排版。'])
d.end('來源：使用者逐次決策、原v34完整標註版的131個章節標題、2026年9月20日R129修正九顆交接成果，以及已核對的模型與訓練記錄。本文件是改寫方案，不是新實驗。')

# DOCUMENT 2
d=Doc('02_模型與訓練完整說明','MOSR模型與訓練完整說明','本文件完整說明本文固定模型的組成、輸入與輸出、數學運算、訓練資料生成、分組、實際學習目標及訓練履歷，可獨立用來重寫方法與實作章。')
d.h('固定模型身分')
d.table(['項目','內容'],[('本文名稱','MOSR；內部固定版本R129 r129_long_01 step08192'),('權重識別','086608339c6418b9dfd650f6809b1f7b45e65b6eb11452afe7df45c5b24ff1fd'),('實拍输出','與原始影像同尺寸，量測使用浮點重建'),('外框設定','全視野距邊界80像素內保留Raw，之後16像素線性過渡'),('模型輸入','觀測灰階影像；使用固定繼承光學資料，不輸入該影像SEM、PI答案或真實幾何'),('學習履歷','先前已訓練初始化，再完成本次8192次接續更新；不是從零訓練')],[1.4,5.6])
d.p('架構由輸入端光學融合、共享四次重建和最後原圖參照輸出組成。固定量測與SEM／PI比較在模型外，不反向修改重建。研究主張是整體流程的實測效果，不因存在某模組就宣稱其独立效益已證實。')
d.page('整體架構與資料流')
d.figure(SRC/'05_Paper_Figures/F01_model_and_measurement.png','圖2-1 本文方法流程。最初Raw影像同時提供最後輸出的參照；經光學處理的影像則作為四次修正的固定基底。',maxh=4)
d.p('兩個「輸入參照」要分清：四次更新內部用的是光學融合後的影像；最後受限細節加入的參照則是原始Raw。四次更新共用同一套神經參數，每步有自己的修正強度。最終輸出沒有增加原影像像素數。')
d.table(['部分','已實作內容','不等同於'],[('光學融合','固定平均PSF及4個候選反推結果融合','逐像素空間變異PSF校正'),('共享重建','64通道，四次共用程序','四個獨立部署模型'),('改善上取樣','雙線性插值、固定二項濾波與卷積','已證明四倍物理解析度'),('最後輸出','Raw參照、受限高頻與平均亮度修正','保證所有圓心、尺寸不改變'),('品質調節','影像統計調整細節強度','已校準的可信概率')],[1.3,3.1,2.6])
d.page('實測光學融合的運算')
d.p('固定光學資料源於既有校正的25個101×101核之平均，以及4個候選核。各核正規化使總和為1。核的光學原點以強度質心處理，在頻域校正相位；影像先作反射邊界擴展，再執行反推。正則值λ＝0.0565沿用既有校正設定。這些是固定校正資料，不是每個合成場景的真實模糊參數。')
d.p('對核k的反推濾波可記為 Gₖ＝(1＋λ)Hₖ*／(|Hₖ|²＋λ)，其中Hₖ是已校正原點的頻域核。係數(1＋λ)使零頻增益在正規化核條件下為1。平均核的反推影像记为W₀(x)，四候選的加權反推记为ΣwₖWₖ(x)。')
d.p('光學輸出 o＝clip[x＋α(W₀(x)−x)＋β(ΣwₖWₖ(x)−x),0,1]。α可學習並限制於0.05至0.35；β限制於0.005至0.1。這是有效值範圍，不是最後權重一定等於初值。')
d.p('四個權重由整張影像的平均、標準差、水平及垂直相鄰灰階差的平均絕對值，經4→8→4的小型網路與softmax產生，softmax前除以2.5。它是全圖權重，不是每個像素各有校正核。候選核相近，額外自適應效益尚未證實；論文可說使用實測光學融合，不寫成候選選擇已獨立提升精度。')
d.h('校正資料與神經訓練的界線',2)
d.p('只繼承固定光學資料，不繼承舊V42神經權重。光學資料有歷史SEM校正背景；神經訓練的生成目標分離，不等於整套系統從未用過SEM。未新增20倍／50倍分別獨立校正，不能聲稱固定核與所有真實條件完全吻合。')
d.page('四次共享修正與上取樣')
d.table(['項目','設定'],[('固定基底','光學融合影像o'),('內部狀態','將o以重複像素提升至4倍網格，作為初始z₀'),('網路輸入','7張在2倍網格上的特徵圖'),('特徵來源','目前狀態縮至2倍、基底提升至2倍、基底與目前原尺寸結果之差、基底的水平梯度、垂直梯度、拉普拉斯及5×5局部平均'),('主幹','7→64的3×3卷積，4個殘差區塊，擴張率依序1、2、4、2'),('各區塊','兩個3×3卷積，SiLU中間啟動，0.1倍殘差更新'),('上取樣','雙線性2倍放大，反射填補，逐通道固定3×3二項核，再64→1卷積及tanh'),('共享次數','4；參數共用，每一步強度可學習')],[1.35,5.65])
d.p('二項濾波核為(1/16)×[[1,2,1],[2,4,2],[1,2,1]]。每一步可記為zᵢ＝clip[zᵢ₋₁＋0.25·sigmoid(sᵢ)·Pθ(Fᵢ),0,1]。Pθ為共享重建網路，Fᵢ是上述7張影像線索。步長受限與值域截斷不構成真實幾何完全不變的保證。')
d.p('目前原尺寸影像由基底o加上狀態相對初始狀態之差的4×4平均得到；最終神經結果也以同樣方式回到原尺寸。輸入端光學只做一次，不能把共享更新描述成每一步都完整執行光學再成像或已證明的最佳化解。')
d.page('最後輸出與實際訓練目標')
d.p('最後以原圖x為參照。以長度13、σ＝2的可分離高斯核從神經影像n取出高頻部分h＝n−G*n。影像品質量q由Raw的梯度強度與其相對5×5局部平均的差異計算；強度為0.35×(0.78＋0.3q)。先作y＝x＋強度×h，再修正平均亮度，使之與x平均相同，最後截斷至0–1。')
d.p('q＝signal／(signal＋3·noise＋10⁻⁸)，signal為水平與垂直灰階差絕對值平均的平均，noise為Raw減局部平均之絕對值平均。它是修正幅度控制，不是預測區間或信心分數。全視野外框保留由推論端另作一次，不能將內部分塊邊緣各自混回Raw。')
d.h('本次接續訓練的三項要求',2)
d.p('令e＝預測−生成目標。在中央監督區域，L＝L影像＋0.2L梯度＋0.2L配對。影像項衡量e；梯度項為水平／垂直相鄰殘差的平均；配對項比較第二張殘差與依已知位移平移後的第一張殘差。無可觀測移動的控制場景不啟用配對項。')
d.p('使用平滑穩健灰階誤差ρ(e)＝255·mean[e²／(sqrt(e²＋ε²)＋ε)]，ε＝0.5/255。梯度項取两方向ρ的平均。配對不是要求輸出位移追隨真實PI命令，而是利用生成時已知的平移檢查重建誤差一致性。')
d.p('本階段不使用內部4倍目標或逐步成像一致性項，不使用未被讀入的SDF或材料mask作幾何損失；資料檔中有欄位不代表訓練實際用到。此點必須與原稿四組舊損失整段抽換。')
d.page('資料生成分組與來源角色')
cc=pd.read_csv(SRC/'03_Dataset/category_counts.csv')
cn={'clean_identity':'清楚影像控制','corners_lines_junctions':'線條角點交會','dense_close_structures':'密集鄰近結構','flat_negative':'平坦區域控制','irregular_curved':'不規則曲線','noisy_identity':'雜訊控制','textured_edges':'紋理邊界'}
d.table(['類別','訓練','開發','既有生成測試'],[(cn[x.category],x.train,x.dev,x.test) for x in cc.itertuples()]+[('合計',512,128,128)],[3.1,1.3,1.3,1.3])
d.p('每個來源產生兩張有已知平移關係的觀測及乾淨目標。形成順序為連續幾何與本徵紋理移動、相機固定照明、同對影像固定模糊、感測器面積積分、獨立雜訊、8位元量化。幾何、亮暗、紋理、模糊與位移分開控制。先以不同母場景種子分train／dev／test，再产生衍生影像。')
d.p('原生成視野384×384，獨立8倍網格成像，保存4倍內部參照與原尺寸目標。實際訓練讀取原生區間32:352，形成320×320輸入；其中96:224為128×128監督核心。真實上下文保留，不以補零或縮放冒充。')
d.p('每個epoch按已固定順序走完512來源，雜訊依來源與epoch重新抽樣。資料包中的原source與每次實際訓練新雜訊輸入可能不同；本文件例圖已按第一次實際雜訊重建，並對上訓練輸入紀錄。768個來源陣列及metadata雜湊核對一致，三分組幾何雜湊交集為0。此檢查的範圍是這份來源包，不能單憑一個重疊數字推論所有歷史資料均獨立。')
d.page('訓練資料範例一')
d.figure(WORK/'figures/training_examples_1.png','圖2-2 每列為一類實際訓練來源，依序為觀測A、乾淨目標A、位移後觀測B。中央128×128顯示區，全部使用相同0–1灰階。各類選編號最小的訓練來源，不按效果選圖。',maxh=7.1)
d.page('訓練資料範例二')
d.figure(WORK/'figures/training_examples_2.png','圖2-3 清楚影像、雜訊與平坦區域控制。已知乾淨目標用於生成訓練監督；它們不是待驗證實拍或SEM答案。',maxh=6.5)
d.p('對真實影像的有效性由實拍結果另外檢驗。生成資料覆蓋多種形狀，不等於對任何材料、亮度、尺度或本來清楚的邊界都不產生改動。')
d.page('訓練設定與履歷')
d.table(['項目','實際設定'],[('本次更新','8192次，16個循環，每次一個來源的兩張影像'),('優化器','AdamW；最高學習率10⁻⁴，最低10⁻⁵；128次線性暖身後餘弦下降'),('正則與穩定','weight decay＝10⁻⁵，梯度上限5'),('亂數設定','模型種子1299001，順序1299002，雜訊1299003'),('保存與檢查','每128次保存；0及每512次對固定128開發来源評估，共17次'),('固定實拍面板','0、512、2048、4096、8192；屬既有研究資料，不是新未知樣品'),('最終生成測試','固定8192後評估既有128來源；不改稱新未見實拍確認集'),('選用模型','本次論文使用固定最後8192；保留中間紀錄，不事後把最好中間點改稱預定終點')],[1.5,5.5])
d.p('模型主幹承接先前合成資料預訓練權重，再作輸出頭轉換與24次短程接入，形成此次接續學習初值。更早的P型預訓練包含成像一致性監督，但R129接續階段不啟用該項。這是同一最終模型的來源說明；不在正文展示歷史四模型選拔表。')
d.p('保存環境記錄為Windows 11、Python 3.12.7、PyTorch 2.3.0+cu121、NumPy 1.26.4、SciPy 1.13.1、scikit-image 0.24、OpenCV 4.12。使用FP32、關閉TF32及確定性設定。這是原訓練記錄，不是宣稱任何設備都位元完全相同。')
d.page('訓練曲線與可引用結論')
d.figure(SRC/'05_Paper_Figures/F03_training_and_development.png','圖2-4 接續訓練曲線。左為每128更新平均訓練目標；右為固定128開發母場景的目標。橫軸不含先前預訓練。',maxh=3.7)
dev=pd.read_csv(SRC/'02_Method_and_Training/development_curve.csv')
d.table(['更新','開發平均目標'],dev.values.tolist(),[2,5],font=9)
d.p('曲線支持此訓練設定下目標下降且後段變化較小；不以生成誤差下降百分比代替真實位移或SEM改善。最終既有生成測試平均目標1.97272與開發群體不同，不拿它和初始開發分數計算提升比例。')
d.page('前期訓練到最終模型的銜接')
d.p('以下只交代同一最終模型的訓練來源，不是要求論文增加模型排行榜或消融對照。正文可壓縮為一段預訓練與接續學習說明，實作表保留必要設定。')
d.table(['階段','已完成內容'],[('合成預訓練','使用R128 P_s1285002固定step8192初始化來源；512母場景循環16次。AdamW最高10⁻³降至10⁻⁵，暖身128，weight decay 10⁻⁵，梯度上限5。'),('前期監督','影像1.0、4倍內部目標0.25、梯度0.2、配對參考殘差0.2；另成像一致性項權重0.02。生成核、噪聲與已知位移只供前向後的監督，不作影像重建輸入。'),('前期成像項','從內部表示以生成核再成像，與觀測核對；中間灰階用含量化變異1/12的高斯近似，飽和0／255用尾端概率，雜訊尺度下限0.5灰階。'),('接入與適配','承接共享主幹；四相位輸出頭取平均轉成單輸出頭，加入實測光學、改善上取樣與原圖參照。此轉換不保證原前向相同。以相同來源作24次短程適配，AdamW學習率10⁻⁵。'),('本文接續長訓練','由適配後模型開始，重建優化器，8192次新增更新；影像1、梯度0.2、配對0.2。前期的內部目標及成像項在此階段沒有啟用。')],[1.35,5.65])
d.p('前期神經權重與固定光學校正資料是兩條不同來源：前期神經模型來自合成資料學習，光學資料來自既有實測校正。來源清單中的r128字樣是資料及初始化身分，不表示正文要介紹另一個部署模型。最終權重仍以本文件第一頁SHA256唯一識別。')
d.end('來源：固定R129訓練套件的model、data、objectives、training.json、run_plan、8192筆更新及17次開發紀錄；初期來源方法与校正紀錄。本文件依實際運作撰寫，未新增消融或訓練。')

# DOCUMENT 3
d=Doc('03_ACM與實驗方法','ACM輪廓與實拍驗證完整方法','本文件獨立說明930張實拍的取樣、Raw追蹤、ACM粗輪廓、最後取邊、圓心位移、尺寸與輪廓統計、SEM對應及不確定度分量。它可用來重寫共同量測方法與實驗配置。')
d.h('拍攝與資料單位')
d.table(['項目','實際內容'],[('拍攝','同片樣品、同一工作階段，30次掃描，每次31位置'),('影像','每方法930張；同一批原圖分別直接量測或經固定MOSR重建'),('物件','十字交界附近固定3×3九顆；每方法8370條輪廓'),('相機及倍率','Basler acA2040-120um，20倍，1544×2064；此批記錄曝光3000 μs'),('尺度','名義0.1725 μm/px，即172.5 nm/px；非完整獨立校正不確定度'),('移動','相鄰約1 μm，距起點約1–30 μm，PI x/y為實際記錄回授，單位μm'),('兩類位移','每方法900個相鄰短步＋900個距起點讀值；不是1800次獨立實驗')],[1.5,5.5])
d.p('目前9顆來自修正後取樣區，不是舊中央九顆。部分原CSV欄仍叫center9，只是沿用名稱。資料以影像座標與固定物件身份判定，不能混入舊4.092nm等結果。30次是同時段重拍，不是30個樣品或跨日研究。')
d.page('固定取樣位置與追蹤')
d.figure(SRC/'05_Paper_Figures/F05_selected_nine.png','圖3-1 全視野與實際量測九顆。首張位置約x＝1592、1621、1650，y＝726、755、784；九顆局部編號固定，未按模型成績挑選。',maxh=4.6)
d.p('Raw閾值連通區提供物件位置，按連續影像作最近且唯一的對應。重建圖不參與選顆或追蹤，PI答案也不用來決定哪些輪廓要留下。930張追蹤均有結果。局部位置與全域Raw移動預測之最大差小於0.636像素，屬此次追蹤核對結果。')
d.p('SEM對應集合為22、23、24、28、29、30、34、35、36。左右鏡像與90度方向皆可對應同集合；逐顆排列仍推定。平均尺寸與重拍統計不受這兩種排列互換影響。此處不把SEM方向未定誤寫成整組集合不存在。')
d.page('ACM粗輪廓到最後取邊')
d.table(['步驟','實際操作'],[('局部視窗','以Raw追蹤中心裁半窗14像素，名義初始半徑8.330159px'),('ACM初始化','幾何圓初始化Morphological Chan–Vese，13次更新，smoothing＝1，兩區權重均1'),('物件選擇','選距種子最近的連通區並填補小孔，取得封閉粗輪廓'),('粗輪廓重取樣','沿弧長均分64點，計算切線及朝外法線'),('法線剖面','每點沿法線±2.4px，每0.1px取樣；三次插值及reflect邊界'),('剖面處理','以σ＝0.65px平滑；對暗柱使用正向灰階梯度峰值'),('細定位','峰值鄰近三點拋物線微調，限制在相鄰取樣範圍'),('最後輸出','64個最大梯度邊界點與其強度；供共同圓與形狀量測')],[1.45,5.55])
d.p('ACM首先穩定地找出物件大致範圍；最後量測用的是最大梯度位置，不是直接把二值遮罩外緣當精密邊界。實作每次測量前重置形態曲率操作狀態，避免上一次物件造成跨次狀態差異。相同設定作用於Raw及MOSR。')
d.page('實拍取邊示意')
d.figure(SRC/'05_Paper_Figures/F04_ACM_and_final_boundary.png','圖3-2 同一首掃描首張第一顆物件。上排是原圖与MOSR，下排黃虛線為ACM粗輪廓、藍線為最後量測邊界；共用灰階範圍及視窗，不分別拉對比。',maxh=7)
d.page('圓心位移與尺度換算')
d.p('LSC使用全部64個最後邊界點，尋找圓心(cₓ,cᵧ)與半徑R，使Σ(||pₖ−c||−R)²最小。直徑為2R。圓擬合提供量測定義，不代表假設物體應該完美正圓；輪廓形狀仍由原取邊點保留。')
d.p('相鄰位移直接以第j張圓心減第j−1張；距起點直接以第j張減第0張，不將各步中位數加總。每顆先求x／y位移，再對九顆取分量中位數，得到合併移動向量。單顆結果另存，不與合併主分數混用。')
d.p('沿移動方向的單位向量只由前6次Raw與PI求得，所有方法共用。主表採名義172.5nm/px，沿方向投影後與PI記錄位移相減，得到e＝影像位移−PI位移。另保存後24次及PI校準尺度的敏感度視圖；正文不能把不同尺度或群體的數字拼在一起。')
d.p('MZC、MIC、MCC如在原稿保留，須按其定義說明：MZC是最小環帶寬度，MIC／MCC提供內接／外接尺寸；它們與LSC的任務不同。本次主要新位移使用LSC，不把歷史其他模型的LSC／MZC對照改名為此模型結果。')
d.h('固定與變動因素',2)
d.p('固定模型、來源影像、追蹤、量測視窗、尺度、取邊設定及浮點精度。相同位置序號對應同一移動程序，PI實際位置有細小差異，不強迫每次完全相等。所有完整掃描保留，沒有剔除最差結果。')
d.page('位移尺寸與輪廓統計定義')
d.table(['量','計算與意義'],[('位移平均絕對PI差距','對指定群體取mean(|e|)，單位nm；不是絕對儀器準確度'),('位移平均偏差','mean(e)，保留正負；不能與MAE互換'),('位移重拍波動','每固定步／位置 across 30次計樣本SD，再對30個位置平均'),('直徑重拍波動','每固定位置×物件的30次樣本變異數，跨31×9平均後開根號'),('輪廓重拍波動','各次只扣圓心，以固定720方向表示半徑；跨30次計變異數，再跨位置／顆粒／方向平均後開根號'),('形狀重拍波動','先扣每條輪廓的平均半徑，再按相同方式計重拍波動'),('最大直徑偏離','相對同一位置同一物件30次平均直徑的最大差，不是對SEM的最差誤差'),('觀察最大PI差距','本次觀測到的最大|e|，不是未來保證界線')],[1.7,5.3])
d.p('64邊界點插值成720方向只為共同角度比較，不增加實際解析度。取中心後不旋轉、不縮放，不利用自由變形使輪廓重合。扣掉平均半徑的形狀指標會移除整體大小差，因此必須與保留尺寸的輪廓和直徑指標分清。')
d.p('範圍估計以完整掃描成對重抽，保留每次內部物件與位置的相關性；另以連續3次掃描成組檢查。不能把900步、9顆或720角度当900個、9倍或720倍獨立實驗。')
d.page('SEM對照與不確定度分量')
d.p('歷史SEM有明確36顆排序，使用者確認與新實拍為同片。九顆集合匹配已建立；工作方向採左右鏡像，但同十字區鏡像／90度候選相關0.978951／0.978681接近，未唯一確認。兩候選逐顆結果均保留，不挑更好分數決定方向。SEM ID28局部邊界不規則保留，尚未認證為實物缺陷。')
d.p('SEM集合平均直徑差是光學九顆平均直徑減SEM九顆平均直徑，不是逐顆平均絕對誤差。完整徑向差保留尺寸；shape差扣平均半徑。以上都是固定操作定義下的參照比較，不能自動等同唯一材料真邊界。')
d.h('資料內不確定度的實際方法',2)
d.p('60個固定步／位置各有30次誤差。單次重拍以原始誤差的樣本SD表示，不先去漂移再只報較小波動。30次平均值另用誤差＝截距＋中心化經過分鐘＋中心化前6／後24區塊指示量的模型。截距等於此時間設計下平均誤差。')
d.p('平均值變異使用Newey–West Bartlett時間相關估計，主要延遲3次，有限樣本修正30/27，約11分鐘拍攝中斷的兩區塊之間不作延遲乘積；延遲2及5次作敏感度。uA是平均值的資料內分量，2uA只是k＝2顯示，不保證95%完整涵蓋率。')
d.p('完整系統u與U仍未知，未合入PI校正、倍率／方向校正、跨日／同步等未提供項目。單次SD不能直接除√30當單张精度；均值分量小也不抹去平均偏差。沒有套用偏差修正。')
d.end('來源：修正後九顆固定取樣設定、原始64點邊界、R136 LSC量測核心及2026年9月20日重新核對結果、SEM方向報告、不確定度method.json。本文不引入新拍攝或新量測。')

# DOCUMENT 4
d=Doc('04_完整結果與圖表','MOSR完整結果圖表與寫作材料','本文件以Raw及固定MOSR兩組完整呈現位移、Contour、尺寸、SEM與資料內不確定度。必要方法與圖說已內嵌；主要表格保留不利結果，讓後續改寫可聚焦改善而不誤述證據。')
d.h('資料及讀法摘要')
d.p('同片樣品、同時段，30次掃描×31位置＝930影像／方法；每影像九顆，8370輪廓／方法。使用修正後十字附近九顆，Raw追蹤與同一ACM粗輪廓→最大梯度取邊→LSC。主位移為九顆中心差的分量中位數；每方法900短步及900距起點讀值。名義172.5nm/px，主表取全部30次。')
d.p('MOSR身分：R129長訓練step08192，SHA256為086608339c6418b9dfd650f6809b1f7b45e65b6eb11452afe7df45c5b24ff1fd，原80/16外框。不是W16，不使用舊中央九顆4.092nm成績。')
main=pd.read_csv(DATA/'main_comparison.csv');names={'mae_nm':'平均絕對PI差','repeatability_nm':'重拍波動','maximum_error_nm':'觀察最大PI差','diameter_repeat_sd_nm':'直徑重拍波動','contour_repeat_sd_nm':'輪廓重拍波動','shape_repeat_sd_nm':'去平均半徑形狀波動','diameter_repeat_absmax_nm':'最大直徑偏離'}
d.table(['項目','Raw nm','MOSR nm','降低 %'],[(('短步 ' if x.scope=='adjacent' else '距起點 ' if x.scope=='start' else '')+names[x.metric],x.Raw_nm,x.R129_nm,x.reduction_percent) for x in main.itertuples()],[3.25,1.25,1.25,1.25])
d.p('位移重拍採固定位置SD的平均；尺寸／輪廓採固定位置與顆粒變異數平均開根號。輪廓為半徑方向波動，直徑則為直徑波動，不能視為同一量。')
d.page('位移曲線及重拍波動')
d.figure(SRC/'05_Paper_Figures/F06_displacement_mean_sd.png','圖4-1 灰為Raw，藍為MOSR。曲線為各固定位置30次平均PI差，陰影為±1樣本SD；左為相鄰約1μm，右為距起點約1–30μm。陰影不是95%總不確定度。',maxh=4)
d.p('短步平均絕對差由8.167降至4.125nm，約降低49.5%；重拍波動由9.725降至4.251nm，約降低56.3%。30次掃描短步平均差皆改善。全路徑平均差由19.081降至17.985nm，20次改善、10次較差。全路徑平均與終點偏差分開解讀。')
pair=pd.read_csv(DATA/'paired_displacement_evidence.csv').query("metric == 'mae_nm'")
d.table(['範圍','MOSR−Raw nm','95%下界','95%上界','改善次數'],[(x.kind,x.candidate_minus_reference_mean_nm,x.ci95_low,x.ci95_high,x.improved_scans) for x in pair.itertuples()],[1.4,1.65,1.35,1.35,1.25])
d.p('範圍由完整掃描配對重抽；連續3次成組的敏感度也在附件。它支持本批差異，不代表跨日、跨材料或所有量測距離。')
d.page('終點與單顆讀值')
ep=pd.read_csv(DATA/'endpoint_30um.csv')
d.table(['方法','約30μm偏差 nm','重拍SD nm','觀察最小 nm','觀察最大 nm'],[(x.model,x.bias_nm,x.repeatability_sd_nm,x.observed_min_nm,x.observed_max_nm) for x in ep.itertuples()],[1,1.7,1.4,1.45,1.45])
d.p('約30μm終點平均偏差由+40.892變成+42.933nm，重拍SD由16.124降至14.299nm。正文可在相應表格及一兩句討論交代，不將全路徑平均改善誤寫成終點偏差改善。')
ind=pd.read_csv(DATA/'individual_circle_displacement_summary.csv')
g=ind[ind.kind=='adjacent'];piv=g.pivot(index='pillar_id',columns='model',values=['mae_nm','repeatability_nm'])
d.table(['顆粒','Raw MAE','MOSR MAE','Raw重拍','MOSR重拍'],[(i,*[piv.loc[i,(a,m)] for a,m in [('mae_nm','Raw'),('mae_nm','R129'),('repeatability_nm','Raw'),('repeatability_nm','R129')]]) for i in piv.index],[.8,1.55,1.55,1.55,1.55])
d.p('此表是單顆短步，單位nm，與九顆合併主結果不同。完整單顆短步與距起點共32400筆明細在試算表，不把4.125nm寫成每顆皆有此絕對準確度。')
d.page('輪廓重拍的主要結果')
d.p('Contour結果是本文正文重點。同位置同物件的30次比較中，輪廓重拍波動63.758→18.437nm，降低71.1%；去平均半徑的形狀波動61.667→15.789nm，降低74.4%。九顆的輪廓重拍波動全部降低，顯示共同取邊方法在重建後得到更一致的讀值。')
per=pd.read_csv(DATA/'nine_pillar_size_shape_comparison.csv');pv=per.pivot(index='local_pillar',columns='model',values=['diameter_repeat_sd_nm','contour_repeat_sd_nm','shape_repeat_sd_nm'])
d.table(['顆粒','Raw尺寸','MOSR尺寸','Raw輪廓','MOSR輪廓','MOSR形狀'],[(i,*[pv.loc[i,(a,m)] for a,m in [('diameter_repeat_sd_nm','Raw'),('diameter_repeat_sd_nm','R129'),('contour_repeat_sd_nm','Raw'),('contour_repeat_sd_nm','R129'),('shape_repeat_sd_nm','R129')]]) for i in pv.index],[.55,1.29,1.29,1.29,1.29,1.29])
d.p('單位nm；尺寸是直徑波動，輪廓／形狀是徑向波動。形狀先扣每條輪廓平均半徑，不能以此表宣稱尺寸真值改善。圈線更集中與更圓也是不同概念；樣品不必原本正圓。')
pc=pd.read_csv(DATA/'paired_shape_size_evidence.csv')
d.table(['指標','MOSR−Raw nm','95%下界','95%上界'],[(names[x.metric],x.difference_nm,x.ci_low_nm,x.ci_high_nm) for x in pc.itertuples()],[2.6,1.6,1.4,1.4])
d.page('起始位置的完整輪廓')
d.figure(SRC/'05_Paper_Figures/F07_contours_start.png','圖4-2 位置0，九顆各30次取邊。只對齊各次LSC圓心，不旋轉、不縮放。使用原始64點連線，灰為Raw、藍為MOSR。',maxh=7.5)
d.page('末端位置的完整輪廓')
d.figure(SRC/'05_Paper_Figures/F08_contours_end.png','圖4-3 位置30，與起點圖相同顆粒、尺度、30次及畫法。起點和終點為固定展示位置，不按結果挑選。',maxh=7.5)
d.page('尺寸平均與重拍穩定性')
d.figure(SRC/'05_Paper_Figures/F09_diameter_and_SEM.png','圖4-4 左為平均直徑與SEM九顆集合參照；右為各位置重拍波動。更穩定與更接近SEM尺寸需分別判讀。',maxh=3.8)
sg=pd.read_csv(DATA/'SEM_group_mean_size.csv')
d.table(['方法','平均直徑 nm','SEM平均 nm','集合有號差 nm'],[(x.model,x.mean_diameter_nm,x.SEM_group_mean_diameter_nm,x.group_signed_diameter_difference_nm) for x in sg.itertuples()],[1,2,2,2])
d.p('MOSR直徑重拍波動降低約41.3%，但集合平均尺寸偏大的差距由185.182增至224.472nm。這項集合平均不受兩種顆粒排列互換影響。必要限制留在此處即可，無須反覆在其他段落擴寫，也不能刪掉相關數據。')
d.page('與SEM的形狀對照')
sd=pd.read_csv(DATA/'SEM_direction_sensitivity.csv')
d.table(['方法','方向','直徑差 nm','完整徑向差 nm','去尺寸形狀差 nm'],[(x.model,x.mapping_candidate,x.diameter_mae_nm,x.radial_mae_nm,x.shape_mae_nm) for x in sd.itertuples()],[.8,.65,1.65,1.9,2])
d.p('方向1為左右鏡像工作推定，方向2為保留的90度候選。方向1下去平均半徑形狀差56.348→32.636nm，降低約42.1%；方向2下也改善，幅度不同。完整徑向差保留了尺寸差，不能用去尺寸形狀差替代。')
d.p('集合ID為22、23、24、28、29、30、34、35、36。方向1排列22/23/24、28/29/30、34/35/36；另一排列22/28/34、23/29/35、24/30/36。原始實拍與SEM同片，逐顆方向不是唯一確認。ID28局部不規則不預先判為實物缺陷。')
d.p('可寫：在目前兩種合理對應下，去除平均尺寸差後的形狀更接近SEM，重拍輪廓也更一致。不要寫：所有局部缺口都已恢復，或完整邊界絕對誤差必然下降。')
d.page('SEM輪廓圖')
d.figure(SRC/'05_Paper_Figures/F10_SEM_contours_inferred.png','圖4-5 黑為SEM、灰為Raw、藍為MOSR。圖中採方向1工作推定，30次輪廓在起點對齊圓心；尺度與方向限制仍適用。',maxh=7.5)
d.page('不確定度分量與寫法')
u=pd.read_csv(HAND/'02_R129_Uncertainty/r129_uncertainty_by_position.csv')
d.p('本表包含60個固定步／位置、每組30次的資料內分量。單次s包含此工作階段的重拍波動；平均值uA考慮時間漂移、前6／後24區塊與時間相關。k＝2列只是2uA，不是完整95%總不確定度。完整系統u與U仍未知，沒有將缺少的校正項目設為零。')
sel=u[((u.kind=='adjacent')&(u.position_index==1))|((u.kind=='start')&(u.position_index.isin([10,20,30])))]
d.table(['類別／位置','偏差 nm','單次s nm','平均uA nm','2uA nm'],[(f'{x.kind} {x.position_index}',x.bias_nm,x.single_session_repeatability_s_nm,x.mean_uA_time_model_nm,x.mean_UA_k2_time_model_nm) for x in sel.itertuples()],[1.8,1.3,1.3,1.3,1.3])
d.p('終點平均PI差42.932539nm，單次s＝14.299373nm，平均uA＝3.672449nm，2uA＝7.344899nm。平均值分量較小不會消除42.933nm的偏差，也不代表單張達7.345nm總不確定度。')
d.p('主要時間模型是截距、中心化經過分鐘、中心化前6／後24指示量；以Bartlett權重、延遲3、有限修正30/27估平均值變異，跨約11分鐘中斷不作延遲乘積，延遲2、5另列敏感度。完整公式與27欄全60列在試算表及數據附件。')
d.page('所有固定位置的位移統計')
ps=pd.read_csv(DATA/'fixed_position_statistics_120.csv')
for kind,label in [('adjacent','相鄰短步'),('start','距起點')]:
 if kind=='start':d.page('距起點的所有位置統計')
 else:d.h(label,2)
 pv=ps[ps.kind==kind].pivot(index='position_index',columns='model',values=['bias_nm','repeatability_sd_nm'])
 d.table(['位置','Raw偏差','MOSR偏差','Raw SD','MOSR SD'],[(i,*[pv.loc[i,(a,m)] for a,m in [('bias_nm','Raw'),('bias_nm','R129'),('repeatability_sd_nm','Raw'),('repeatability_sd_nm','R129')]]) for i in pv.index],[.7,1.575,1.575,1.575,1.575],font=8)
 d.p('全部單位nm，每列每方法30次。位置序號不是額外獨立樣品；完整最小、最大、MAE、橫向波動及PI距離在試算表。')
d.page('可供正文使用的結果敘述')
d.p('在固定20倍拍攝條件下，本文MOSR將相鄰短步相對PI回授的平均絕對差距由8.167nm降至4.125nm，並將固定位置重拍標準差由9.725nm降至4.251nm。30次掃描的平均短步差距均降低，支持此量測流程於已測條件下的實用改善。')
d.p('透過相同ACM粗輪廓與最大梯度取邊程序，MOSR使同位置30次的輪廓波動由63.758nm降至18.437nm，九顆皆有改善。扣除平均半徑後的形狀重拍波動及SEM形狀差亦降低，顯示重建可提升輪廓讀取的一致性。')
d.p('本文將位移、輪廓重複性與SEM尺寸差分開呈現。集合尺寸偏大及長距離終點偏差仍存在，因此不以重拍波動降低概括所有絕對準確度。主要結論限定本片樣品、本次工作階段與固定量測方法。')
d.h('核對範圍',2)
d.p('來源交接記錄包含16740個圓心獨立重算、16740條720方向輪廓再算、4張影像重播取邊、120組位移均值／SD與60個不確定度分量核對。本次另核对交接包284個清單檔，並从保存明細重算主要摘要；未新增模型推論或把原紀錄改成另一個版本。')
d.end('來源：2026年9月20日R129修正九顆交接包的Raw／R129主CSV、逐顆與逐位置表、SEM兩方向敏感度、不確定度方法與分量。圖表均來自實際資料；没有消融、重訓或其他模型比較。')

# Copy the original Word as the actual reference to be uploaded with these documents.
shutil.copy2(SRC/'01_Original_Thesis/thesis_v34_original_unchanged.docx',OUT/'00_原論文_v34_未修改.docx')
assert sha(OUT/'00_原論文_v34_未修改.docx')=='d85a7987bbea8f5f58585334104f9d0425f3a42a92e51bbbb187cca2c5a9fc12'

# Prepare typed matrices for the artifact-tool workbook. Every source CSV retained.
mapping=[
('主要結果',DATA/'main_comparison.csv'),('掃描配對差距',DATA/'paired_displacement_evidence.csv'),('重拍配對差距',DATA/'paired_repeatability_evidence.csv'),('形狀尺寸配對',DATA/'paired_shape_size_evidence.csv'),('終點結果',DATA/'endpoint_30um.csv'),('九顆尺寸形狀',DATA/'nine_pillar_size_shape_comparison.csv'),('SEM集合尺寸',DATA/'SEM_group_mean_size.csv'),('SEM兩方向',DATA/'SEM_direction_sensitivity.csv'),('120固定位置',DATA/'fixed_position_statistics_120.csv'),('120繪圖數據',DATA/'overlay_mean_sd_plot_data.csv'),('62位置尺寸',DATA/'diameter_by_position_62.csv'),('1440角度統計',DATA/'angular_repeatability_1440.csv'),('逐掃描位移',DATA/'per_scan_displacement.csv'),('前6後24',DATA/'time_block_displacement.csv'),('單顆位移摘要',DATA/'individual_circle_displacement_summary.csv'),('最差案例對照',DATA/'paired_worst_cases.csv'),('3600合併位移',DATA/'all_displacements_3600.csv'),('32400單顆位移',DATA/'individual_circle_displacements_32400.csv'),('16740輪廓',DATA/'all_contours_16740.csv'),('60不確定度分量',HAND/'02_R129_Uncertainty/r129_uncertainty_by_position.csv'),('訓練分類',SRC/'03_Dataset/category_counts.csv'),('768資料來源',SRC/'03_Dataset/all_768_parent_sources.csv'),('8192訓練更新',SRC/'02_Method_and_Training/training_curve_all_steps.csv'),('17開發檢查',SRC/'02_Method_and_Training/development_curve.csv'),('131章節處理',SRC/'00_Plan/section_action_map.csv')]
inputs=[]
for name,p in mapping:
 df=pd.read_csv(p);df=df.astype(object).where(pd.notna(df),None)
 # Local paths are provenance strings, never needed to read this workbook.
 inputs.append(dict(name=name,source=p.name,sha256=sha(p),columns=list(df.columns),rows=df.values.tolist()))
(WORK/'workbook_inputs.json').write_text(json.dumps(inputs,ensure_ascii=False,allow_nan=False),encoding='utf-8')
# Full raw numerical substrate: CSV tables, NPZ boundaries/profiles, metadata and definitions.
with zipfile.ZipFile(OUT/'06_完整底層數據.zip','w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for p in sorted(HAND.rglob('*')):
  if p.is_file() and p.suffix.lower() in ['.csv','.npz','.json','.md','.py']:
   z.write(p,'R129_results/'+str(p.relative_to(HAND)).replace('\\','/'))
 for p in sorted((SRC/'03_Dataset').rglob('*')):
  if p.is_file():z.write(p,'training_materials/'+str(p.relative_to(SRC/'03_Dataset')).replace('\\','/'))
 for p in sorted((SRC/'02_Method_and_Training/records').glob('*')):
  if p.is_file():z.write(p,'training_records/'+p.name)
 for p in sorted((SRC/'06_Internal_Sources/model').glob('*.py')):z.write(p,'model_definition/'+p.name)
 for p in sorted((SRC/'06_Internal_Sources/training').glob('*.py')):z.write(p,'training_definition/'+p.name)
 z.writestr('READ_ME.txt','本包包含數據本身，無須讀取本機路徑。CSV可直接讀；NPZ可用numpy.load讀。Raw/R129半徑資料shape為30×31×9×720，單位nm；邊界點為每掃描31×9×64×2像素座標。文件03說明算法，05試算表含主要及逐筆CSV。原來源記錄的绝對路徑只是追溯文字。此包不是獨立訓練環境；768組完整訓練陣列和930張完整影像未包含，所有撰稿統計與必要範例均已包含。')
summary=dict(documents=4,source_word_unchanged=True,workbook_sheets=len(inputs)+1,source_csv_rows={x['name']:len(x['rows']) for x in inputs},embedded_figures=True,no_word_thesis_rewrite=True)
(WORK/'build_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)
