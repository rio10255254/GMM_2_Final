"""Independent numerical and deliverable audit for the separate 1 um poster."""
from pathlib import Path
from collections import defaultdict
from statistics import median,stdev,mean
import json,csv,hashlib,io,zipfile,math
from lxml import etree
import openpyxl,pymupdf
PROJECT_ROOT = Path(__file__).resolve().parents[3]
ROOT=Path(__file__).resolve().parents[1];DATA=ROOT/'results/separate_conditions';OUT=ROOT/'poster_final'
PPTX=OUT/'IMPACT2026_MOSR_1um_Separate_v2.pptx';PDF=PPTX.with_suffix('.pdf')
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def rows(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
data=load(DATA/'poster_data.json');a=load(DATA/'20x_complete.json');b=load(DATA/'50x_complete.json')
assert a['passed'] and b['passed'] and a['frames']==930 and b['images']==86
assert a['excluded_steps']==b['excluded_pairs']==0
assert b['max_fresh_reconstruction_difference']<1e-7 and b['learned_parameters_unchanged']
assert b['full_vs_tiled_max_difference']<1e-6 and b['single_psf_reinstallation_exact']
axis=load(Path(str(PROJECT_ROOT / 'Experiments/R152_R148v2_SafeAdaptivePSF/evaluation/calibration.json')))['axis_unit_vector']
ci=defaultdict(dict)
for r in rows(DATA/'20x_all_centers.csv'):
    ci[(r['run'],r['method'],int(r['index']))][int(r['pillar_id'])]=(float(r['center_x_px']),float(r['center_y_px']))
step_groups=defaultdict(list);max_step_err=0
for r in rows(DATA/'20x_steps.csv'):
    j=int(r['step']);start=ci[(r['run'],r['method'],j-1)];end=ci[(r['run'],r['method'],j)]
    assert set(start)==set(end)==set(range(1,10))
    dx=median([end[i][0]-start[i][0] for i in start]);dy=median([end[i][1]-start[i][1] for i in start])
    expected=(dx*axis[0]+dy*axis[1])*172.5
    max_step_err=max(max_step_err,abs(expected-float(r['displacement_nm'])))
    step_groups[(r['method'],int(r['run_index']))].append(float(r['displacement_nm']))
assert len(step_groups)==90 and max_step_err<1e-8
values20={m:[] for m in ['Raw','Wiener','MOSR']}
for (m,run),v in sorted(step_groups.items(),key=lambda kv:kv[0][1]):
    assert len(v)==30
    sd=stdev(v);assert abs(sd-data['twenty']['scan_sd_nm'][m][run-1])<1e-8
    values20[m].append(sd)
for m,v in values20.items():
    pooled=math.sqrt(mean([x*x for x in v]));assert abs(pooled-data['twenty']['pooled_within_scan_sd_nm'][m])<1e-8
fi=defaultdict(dict)
for r in rows(DATA/'50x_centers.csv'):fi[(r['sequence'],r['method'],int(r['frame_index']))][int(r['pillar_id'])]=float(r['center_x_px'])
fifty_groups=defaultdict(list)
locked=defaultdict(set)
for r in rows(ROOT/'sources/original_locked_pairs.csv'):
    if r['magnification']=='50x':locked[r['sequence']].add((int(r['start_frame_index']),int(r['end_frame_index'])))
actual=defaultdict(set)
for r in rows(DATA/'50x_steps.csv'):
    seq,m=r['sequence'],r['method'];i,j=int(r['start']),int(r['end'])
    c,d=fi[(seq,m,i)],fi[(seq,m,j)];assert len(c)==len(d)==9 and set(c)==set(d)
    expected=abs(median([d[k]-c[k] for k in c]))*69
    assert abs(expected-float(r['step_nm']))<1e-8
    fifty_groups[(seq,m)].append(float(r['step_nm']));actual[(seq,m)].add((i,j))
for (seq,m),v in fifty_groups.items():
    assert len(v)==30 and actual[(seq,m)]==locked[seq]
    if seq=='50x_x_1um':assert abs(stdev(v)-data['fifty']['step_sd_nm'][m])<1e-8
for relative,expected in data['source_hashes'].items():assert sha(ROOT/relative)==expected
with zipfile.ZipFile(PPTX) as z:
    size=etree.fromstring(z.read('ppt/presentation.xml')).find('p:sldSz',NS)
    assert (size.get('cx'),size.get('cy'))==('32400000','43200000')
    slide=etree.fromstring(z.read('ppt/slides/slide1.xml'));visible='\n'.join(slide.xpath('//a:t/text()',namespaces=NS))
    for forbidden in ['R152','0.5','19.5%','60.8%','120 pairs','four conditions','Mean row']:
        assert forbidden not in visible,forbidden
    for expected in ['52.3%','23.5%','36.9–64.6%','30 repeated scans','30 locked pairs']:
        assert expected in visible,expected
    charts=[n for n in z.namelist() if '/charts/' in n and n.endswith('.xml') and '/_rels/' not in n];assert len(charts)==2
    for n in charts:
        c=etree.fromstring(z.read(n));is_line=c.find('.//c:lineChart',NS) is not None
        series=c.findall('.//c:ser',NS);assert len(series)==3
        for item,m in zip(series,['Raw','Wiener','MOSR']):
            assert item.find('.//c:tx//c:v',NS).text==m
            nums=[float(v.text) for v in item.findall('.//c:val//c:pt/c:v',NS)]
            target=data['twenty']['scan_sd_nm'][m] if is_line else [data['fifty']['step_sd_nm'][m]]
            assert len(nums)==len(target) and max(abs(x-y) for x,y in zip(nums,target))<5.1e-7
    books=[n for n in z.namelist() if n.endswith('.xlsx')];assert len(books)==2
    numeric=[]
    for n in books:
        w=openpyxl.load_workbook(io.BytesIO(z.read(n)),data_only=True)
        is_line=w.active.max_row==31
        for col,m in zip([2,4,6],['Raw','Wiener','MOSR']):
            assert w.active.cell(1,col).value==m
            vals=[w.active.cell(r,col).value for r in range(2,w.active.max_row+1)]
            target=data['twenty']['scan_sd_nm'][m] if is_line else [data['fifty']['step_sd_nm'][m]]
            assert len(vals)==len(target) and max(abs(x-y) for x,y in zip(vals,target))<5.1e-7
            numeric.extend(vals)
    assert len(numeric)==93
    assert len(slide.findall('.//a:tbl',NS))==1
    media=[z.read(n) for n in z.namelist() if '/media/' in n]
    for image in data['images'].values():assert Path(image['path']).read_bytes() in media
with zipfile.ZipFile(ROOT/'sources/impact2026_poster_template.pptx') as z:assert z.read('ppt/media/image2.jpeg') in media
fit=load(ROOT/'poster_build/separate_text_fit_v2.json');assert not any(x['overflow'] for x in fit)
validation=load(ROOT/'poster_build/mosr_separate_validation_v2.json');assert validation['presentationLayout']['warning_count']==0
doc=pymupdf.open(PDF);assert len(doc)==1
page=doc[0];cm=[page.rect.width/72*2.54,page.rect.height/72*2.54];assert abs(cm[0]-90)<.01 and abs(cm[1]-120)<.01
text=page.get_text()
for v in ['52.3%','23.5%','10.200','4.865','5.938','4.542','Fig. 1.','Fig. 2.','Fig. 3.']:assert v in text,v
assert 'R152' not in text and '0.5' not in text
page.get_pixmap(matrix=pymupdf.Matrix(.65,.65)).save(ROOT/'poster_build/separate_pdf_review.png')
result=dict(passed=True,twenty_x_scans=30,twenty_x_frames=930,twenty_x_steps_per_method=900,twenty_x_all_three_methods_remeasured=True,
            fifty_x_fresh_reconstruction_frames=86,fifty_x_one_um_main_pairs=30,fifty_x_half_um_supplement_pairs=30,
            no_cross_condition_average=True,twenty_x_max_center_to_step_replay_error_nm=max_step_err,
            fifty_x_fresh_reconstruction_max_difference=b['max_fresh_reconstruction_difference'],native_charts=2,native_chart_points=93,embedded_workbooks=2,
            editable_table_count=1,source_image_bytes_preserved=True,template_logo_unchanged=True,text_overflows=0,print_size_cm=cm,
            pptx_sha256=sha(PPTX),pdf_sha256=sha(PDF),paper_code=None)
(ROOT/'poster_build/separate_final_verification.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,indent=2))
