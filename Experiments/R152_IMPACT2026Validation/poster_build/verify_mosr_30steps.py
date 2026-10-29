"""Validate every selected step, curve point and output in the 30-step poster."""
from pathlib import Path
from collections import defaultdict
from statistics import median,stdev
import json,csv,hashlib,io,zipfile
from lxml import etree
import openpyxl,pymupdf
PROJECT_ROOT = Path(__file__).resolve().parents[3]
ROOT=Path(__file__).resolve().parents[1];DATA=ROOT/'results/separate_conditions'
PPTX=ROOT/'poster_final/IMPACT2026_MOSR_1um_30Steps_v2.pptx';PDF=PPTX.with_suffix('.pdf')
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def rows(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
d=load(DATA/'poster_data_30steps.json');lock=load(DATA/'single_scan_presentation_lock.json')
src=Path(str(PROJECT_ROOT / 'Experiments/R152_R148v2_SafeAdaptivePSF'))
plan=load(src/'evaluation/capture_plan.json');cal=load(src/'evaluation/calibration.json')
first=next(r for r in plan['runs'] if r['run_index']>cal['raw_calibration_runs'])
assert first['run']==d['twenty']['run']==lock['20x_run'] and first['run_index']==7
assert load(DATA/'20x_complete.json')['passed'] and load(DATA/'50x_complete.json')['passed']
assert load(DATA/'50x_complete.json')['max_fresh_reconstruction_difference']==0
ci=defaultdict(dict)
for r in rows(DATA/'20x_all_centers.csv'):
    if r['run']==first['run']:ci[(r['method'],int(r['index']))][int(r['pillar_id'])]=(float(r['center_x_px']),float(r['center_y_px']))
steps20=rows(DATA/'main_20x_30steps.csv');assert len(steps20)==90 and all(r['run']==first['run'] for r in steps20)
axis=cal['axis_unit_vector'];maxerr=0
for r in steps20:
    m,j=r['method'],int(r['step']);a=ci[(m,j-1)];b=ci[(m,j)];assert len(a)==len(b)==9
    delta=[median([b[i][k]-a[i][k] for i in a]) for k in [0,1]]
    val=(delta[0]*axis[0]+delta[1]*axis[1])*172.5
    maxerr=max(maxerr,abs(val-float(r['displacement_nm'])));assert abs(val-d['twenty']['steps_nm'][m][j-1])<1e-8
assert maxerr<1e-8
steps50=rows(DATA/'main_50x_30pairs.csv');assert len(steps50)==90
for m in ['Raw','Wiener','MOSR']:
    q=sorted([r for r in steps50 if r['method']==m],key=lambda r:int(r['start']))
    assert len(q)==30 and all(r['sequence']=='50x_x_1um' for r in q)
    assert all(abs(float(r['step_nm'])-v)<1e-8 for r,v in zip(q,d['fifty']['steps_nm'][m]))
    for cohort in ['twenty','fifty']:
        s=stdev(d[cohort]['steps_nm'][m]);assert abs(s-d[cohort]['step_sd_nm'][m])<1e-8
for c in ['twenty','fifty']:
    gain=100*(1-d[c]['step_sd_nm']['MOSR']/d[c]['step_sd_nm']['Raw']);assert abs(gain-d[c]['relative_reduction_pct'])<1e-10
for p,expected in d['source_hashes'].items():assert sha(ROOT/p)==expected
with zipfile.ZipFile(PPTX) as z:
    size=etree.fromstring(z.read('ppt/presentation.xml')).find('p:sldSz',NS);assert (size.get('cx'),size.get('cy'))==('32400000','43200000')
    slide=etree.fromstring(z.read('ppt/slides/slide1.xml'));visible='\n'.join(slide.xpath('//a:t/text()',namespaces=NS))
    for v in ['R152','52.3%','30 repeated scans','pooled','0.5 µm','19.5%']:assert v not in visible,v
    for v in ['57.2%','23.5%','9.506','4.068','5.938','4.542','30 single steps','0.1-pixel']:assert v in visible,v
    charts=sorted(n for n in z.namelist() if '/charts/' in n and n.endswith('.xml') and '/_rels/' not in n);assert len(charts)==2
    for name,cohort in zip(charts,['twenty','fifty']):
        c=etree.fromstring(z.read(name));assert c.find('.//c:lineChart',NS) is not None
        for ser,m in zip(c.findall('.//c:ser',NS),['Raw','Wiener','MOSR']):
            assert ser.find('.//c:tx//c:v',NS).text==m
            v=[float(p.text) for p in ser.findall('.//c:val//c:pt/c:v',NS)]
            expected=[x-1000 for x in d[cohort]['steps_nm'][m]];assert len(v)==30
            assert max(abs(a-b) for a,b in zip(v,expected))<5.1e-7
        assert float(c.find('.//c:valAx/c:scaling/c:min',NS).get('val'))==-30
        assert float(c.find('.//c:valAx/c:scaling/c:max',NS).get('val'))==30
    books=sorted(n for n in z.namelist() if n.endswith('.xlsx'));assert len(books)==2
    count=0
    for n,cohort in zip(books,['twenty','fifty']):
        w=openpyxl.load_workbook(io.BytesIO(z.read(n)),data_only=True);assert w.active.max_row==31
        for col,m in zip([2,4,6],['Raw','Wiener','MOSR']):
            assert w.active.cell(1,col).value==m
            v=[w.active.cell(r,col).value for r in range(2,32)];expected=[x-1000 for x in d[cohort]['steps_nm'][m]]
            assert max(abs(a-b) for a,b in zip(v,expected))<5.1e-7;count+=len(v)
    assert count==180
    media=[z.read(n) for n in z.namelist() if '/media/' in n]
    for asset in d['images'].values():assert Path(asset['path']).read_bytes() in media
with zipfile.ZipFile(ROOT/'sources/impact2026_poster_template.pptx') as z:assert z.read('ppt/media/image2.jpeg') in media
assert not any(r['overflow'] for r in load(ROOT/'poster_build/thirty_text_fit_v2.json'))
assert load(ROOT/'poster_build/mosr_30steps_validation_v2.json')['presentationLayout']['warning_count']==0
doc=pymupdf.open(PDF);assert len(doc)==1
page=doc[0];cm=[page.rect.width/72*2.54,page.rect.height/72*2.54];assert abs(cm[0]-90)<.01 and abs(cm[1]-120)<.01
text=page.get_text()
for v in ['57.2%','23.5%','9.506','4.068','5.938','4.542','Fig. 1.','Fig. 2.','Fig. 3.']:assert v in text,v
assert 'R152' not in text and 'pooled' not in text and '52.3%' not in text
page.get_pixmap(matrix=pymupdf.Matrix(.65,.65)).save(ROOT/'poster_build/thirty_pdf_review.png')
result=dict(passed=True,main_conditions=2,steps_per_condition_per_method=30,methods=['Raw','Wiener','MOSR'],twenty_x_run=first['run'],twenty_x_run_index=7,selection_rule=lock['selection_rule'],no_cross_scan_or_condition_pooling=True,curves_are_observed_step_minus_1000_nm=True,identical_curve_limits_nm=[-30,30],twenty_x_center_to_step_max_error_nm=maxerr,native_charts=2,native_points=180,embedded_workbooks=2,source_images_preserved=True,template_logo_unchanged=True,text_overflow=0,print_size_cm=cm,pptx_sha256=sha(PPTX),pdf_sha256=sha(PDF),paper_code=None)
(ROOT/'poster_build/thirty_final_verification.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,indent=2))
