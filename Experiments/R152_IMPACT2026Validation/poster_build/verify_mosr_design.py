"""Check publication data, genuine geometry, native editability structure and disclosure scope."""
from pathlib import Path
from statistics import mean,stdev
import json,hashlib,zipfile,io,re
from lxml import etree
import openpyxl,pymupdf
PROJECT_ROOT = Path(__file__).resolve().parents[3]
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'poster_final'
PPTX=OUT/'IMPACT2026_MOSR_Design_TW0243_v3.pptx';PDF=PPTX.with_suffix('.pdf')
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
load=lambda p:json.loads(p.read_text(encoding='utf-8-sig'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
d=load(ROOT/'results/separate_conditions/poster_data_30steps.json');b=load(ROOT/'poster_build/design_assets/measured_boundary.json')
rec=load(ROOT/'poster_build/mosr_design_receipt_v3.json')
assert rec['paperCode']=='TW0243' and not rec['fullArchitectureOrTrainingRecipeDisclosed'] and not rec['extraThirtyScanResultsDisclosed']
assert sha(Path(PROJECT_ROOT / b['source_npz']))==b['source_sha256'] and sha(Path(PROJECT_ROOT / b['source_image']))==b['source_image_sha256']
assert sha(Path(b['image_path']))==b['image_sha256']
for rel,h in d['source_hashes'].items():assert sha(ROOT/rel)==h
for cohort in ['twenty','fifty']:
    c=d[cohort]
    for m in ['Raw','Wiener','MOSR']:assert len(c['steps_nm'][m])==30 and abs(stdev(c['steps_nm'][m])-c['step_sd_nm'][m])<1e-8
    assert abs(100*(1-c['step_sd_nm']['MOSR']/c['step_sd_nm']['Raw'])-c['relative_reduction_pct'])<1e-9
delta_mean=abs(mean(d['twenty']['steps_nm']['Raw'])-mean(d['twenty']['steps_nm']['MOSR']))
delta_sd=d['twenty']['step_sd_nm']['Raw']-d['twenty']['step_sd_nm']['MOSR']
assert abs(delta_mean-rec['meanStepChange20xNm'])<1e-9 and abs(delta_sd-rec['stepSdDecrease20xNm'])<1e-9
with zipfile.ZipFile(PPTX) as z:
    p=etree.fromstring(z.read('ppt/presentation.xml'));size=p.find('p:sldSz',NS)
    assert (size.get('cx'),size.get('cy'))==('32400000','43200000')
    slide=etree.fromstring(z.read('ppt/slides/slide1.xml'));text='\n'.join(slide.xpath('//a:t/text()',namespaces=NS))
    for v in ['TW0243','57.2%','23.5%','0.10 nm','5.44 nm','Calibrated optical guidance','LSC centers','Median shift','Original-image reference']:
        assert v in text,v
    for forbidden in ['52.3%','30 repeated scans','pooled','0.5 µm','19.5%','R152','R148','8192','1,207,105','AdamW','teacher','safe_scale']:
        assert forbidden not in text,forbidden
    # Public notes and properties must not leak the internal dossier or local source paths.
    for name in z.namelist():
        if name.endswith('.xml'):
            xml=z.read(name).decode('utf-8')
            for token in ['R152','R148v2','SafeAdaptive','003412_XY_0','Thesis_MOSR_Evidence','GMM_2_Final','0.013800','8192']:
                assert token not in xml,(name,token)
    shapes={s.find('p:nvSpPr/p:cNvPr',NS).get('name'):s for s in slide.findall('p:cSld/p:spTree/p:sp',NS)}
    for n in ['model_input','model_context','model_native','model_optical','model_output','model_bounded_addition','readout_centers','readout_shift','readout_sd']:
        assert n in shapes,n
    connectors=slide.findall('p:cSld/p:spTree/p:cxnSp',NS);assert len(connectors)>=10
    for c in connectors:
        head=c.find('p:spPr/a:ln/a:headEnd',NS)
        assert head is None or head.get('type')=='none','An arrow points toward its source'
    points=[n for n in shapes if n.startswith('measured_boundary_point_')];assert len(points)==64
    for i,(x,y) in enumerate(b['boundary_xy_px'],1):
        s=shapes[f'measured_boundary_point_{i}'];off=s.find('p:spPr/a:xfrm/a:off',NS)
        expected_x=(140+(x+.5)*240/29-1.9)*9375;expected_y=(3675+(y+.5)*240/29-1.9)*9375
        assert abs(float(off.get('x'))-expected_x)<2 and abs(float(off.get('y'))-expected_y)<2
    tables=slide.findall('.//a:tbl',NS);assert len(tables)==3
    for table,cohort in zip(tables[1:],['twenty','fifty']):
        rows=table.findall('a:tr',NS);assert len(rows)==2
        vals=[''.join(c.xpath('.//a:t/text()',namespaces=NS)) for c in rows[1].findall('a:tc',NS)]
        assert vals[1:]==[f"{d[cohort]['step_sd_nm'][m]:.3f}" for m in ['Raw','Wiener','MOSR']]
    for table in tables:
        for pr in table.findall('.//a:tcPr',NS):
            assert pr.find('a:lnL/a:noFill',NS) is not None and pr.find('a:lnR/a:noFill',NS) is not None
    charts=sorted(n for n in z.namelist() if '/charts/' in n and n.endswith('.xml') and '/_rels/' not in n);assert len(charts)==2
    for name,cohort in zip(charts,['twenty','fifty']):
        c=etree.fromstring(z.read(name));assert c.find('.//c:lineChart',NS) is not None
        for ser,m in zip(c.findall('.//c:ser',NS),['Raw','Wiener','MOSR']):
            assert ser.find('.//c:tx//c:v',NS).text==m
            nums=[float(v.text) for v in ser.findall('.//c:val//c:pt/c:v',NS)]
            assert len(nums)==30 and max(abs(v-(w-1000)) for v,w in zip(nums,d[cohort]['steps_nm'][m]))<5.1e-7
        assert float(c.find('.//c:valAx/c:scaling/c:min',NS).get('val'))==-30
        assert float(c.find('.//c:valAx/c:scaling/c:max',NS).get('val'))==30
    books=sorted(n for n in z.namelist() if n.endswith('.xlsx'));assert len(books)==2
    for name,cohort in zip(books,['twenty','fifty']):
        w=openpyxl.load_workbook(io.BytesIO(z.read(name)),data_only=True);assert w.active.max_row==31
        for col,m in zip([2,4,6],['Raw','Wiener','MOSR']):
            assert w.active.cell(1,col).value==m
            vals=[w.active.cell(i+2,col).value for i in range(30)]
            assert max(abs(v-(x-1000)) for v,x in zip(vals,d[cohort]['steps_nm'][m]))<5.1e-7
    media=[z.read(n) for n in z.namelist() if '/media/' in n]
    for asset in d['images'].values():assert Path(asset['path']).read_bytes() in media
    assert Path(b['image_path']).read_bytes() in media
with zipfile.ZipFile(ROOT/'sources/impact2026_poster_template.pptx') as z:assert z.read('ppt/media/image2.jpeg') in media
fit=load(ROOT/'poster_build/design_text_fit_v3.json');assert not any(r['overflow'] for r in fit)
validation=load(ROOT/'poster_build/mosr_design_validation_v3.json');assert validation['presentationLayout']['warning_count']==0
doc=pymupdf.open(PDF);assert len(doc)==1
page=doc[0];cm=[page.rect.width/72*2.54,page.rect.height/72*2.54];assert abs(cm[0]-90)<.01 and abs(cm[1]-120)<.01
pdftext=page.get_text()
for v in ['TW0243','57.2%','23.5%','9.506','5.794','4.068','5.938','4.722','4.542','0.10','5.44','Fig. 1.','Fig. 2.','Fig. 3.','Fig. 4.','Fig. 5.']:assert v in pdftext,v
page.get_pixmap(matrix=pymupdf.Matrix(.65,.65)).save(ROOT/'poster_build/design_pdf_review.png')
result=dict(passed=True,source_snapshot_unchanged=sha(Path(rec['source']))==rec['sourceSha256'],paper_code='TW0243',native_charts=2,native_tables=3,curve_points=180,embedded_workbooks=2,editable_diagram_connectors=len(connectors),measured_boundary_points=64,boundary_geometry_matches_measurement=True,source_images_preserved=True,logo_preserved=True,no_private_dossier_strings_in_pptx_xml=True,no_full_training_recipe_or_extra_thirty_scan_results=True,mean_step_change20x_nm=delta_mean,step_sd_reduction20x_nm=delta_sd,text_overflow=0,pdf_pages=1,size_cm=cm,pptx_sha256=sha(PPTX),pdf_sha256=sha(PDF))
(ROOT/'poster_build/design_final_verification.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,indent=2))
