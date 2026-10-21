from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "v18_source_copy.docx"
OUTPUT = ROOT / "論文完整版_20260716_v19_第六章更新版.docx"
AUDIT = ROOT / "v19_build_audit.json"

FIGURE_ROOT = Path(
    r"D:\GMM_2_Final\Thesis_MOSR_Evidence\Final_Data_Figures"
    r"\07_V19_OLD_METHOD_FIGURE_PACKAGE_20260724"
)
DISPLACEMENT = FIGURE_ROOT / "01_displacement_improvement"
ABLATION = FIGURE_ROOT / "02_ablation"
MODEL = FIGURE_ROOT / "03_model_comparison"

FIGURE_WIDTH = Cm(14.6)


def find_paragraph(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def replace_paragraph_text(paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.add_run(text)


def append_before(anchor, element) -> None:
    anchor._p.addprevious(element)


def add_paragraph_before(
    document: Document,
    anchor,
    text: str = "",
    style: str = "Normal",
    alignment=None,
    keep_with_next: bool | None = None,
):
    paragraph = document.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    if alignment is not None:
        paragraph.alignment = alignment
    if keep_with_next is not None:
        paragraph.paragraph_format.keep_with_next = keep_with_next
    append_before(anchor, paragraph._p)
    return paragraph


def add_figure_before(
    document: Document,
    anchor,
    image_path: Path,
    number: int,
    caption: str,
) -> None:
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=FIGURE_WIDTH)
    append_before(anchor, paragraph._p)
    add_paragraph_before(
        document,
        anchor,
        f"圖 6-{number} {caption}",
        style="Figure Caption",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=False,
    )


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        element = tc_mar.find(qn(f"w:{margin}"))
        if element is None:
            element = OxmlElement(f"w:{margin}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_table_before(
    document: Document,
    anchor,
    number: int,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    widths_cm: list[float],
    page_break_before: bool = False,
) -> None:
    caption_paragraph = add_paragraph_before(
        document,
        anchor,
        f"表 6-{number} {caption}",
        style="Table Caption",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True,
    )
    caption_paragraph.paragraph_format.page_break_before = page_break_before
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col_index, header in enumerate(headers):
        cell = table.rows[0].cells[col_index]
        cell.text = header
    for values in rows:
        row = table.add_row()
        for col_index, value in enumerate(values):
            row.cells[col_index].text = str(value)
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
    append_before(anchor, table._tbl)


def remove_between(start_paragraph, end_paragraph) -> None:
    current = start_paragraph._p
    end = end_paragraph._p
    parent = current.getparent()
    while current is not None and current is not end:
        next_element = current.getnext()
        parent.remove(current)
        current = next_element


def clone_previous_figure(caption_paragraph):
    current = caption_paragraph._p.getprevious()
    while current is not None:
        if current.xpath(
            ".//*[local-name()='drawing' or local-name()='pict']"
        ):
            return deepcopy(current)
        current = current.getprevious()
    raise ValueError(f"No drawing found before: {caption_paragraph.text}")


def insert_cloned_figure_before(
    document: Document,
    anchor,
    drawing_element,
    number: int,
    caption: str,
) -> None:
    clean_drawing_element = deepcopy(drawing_element)
    for text_node in clean_drawing_element.xpath(".//*[local-name()='t']"):
        text_node.text = ""
    append_before(anchor, clean_drawing_element)
    add_paragraph_before(
        document,
        anchor,
        f"圖 6-{number} {caption}",
        style="Figure Caption",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=False,
    )


def replace_section_body(
    document: Document, heading_prefix: str, new_paragraphs: list[str]
) -> None:
    heading = find_paragraph(document, heading_prefix)
    current = heading._p.getnext()
    anchor_element = None
    while current is not None:
        if current.tag == qn("w:p"):
            style_nodes = current.xpath("./w:pPr/w:pStyle")
            style_id = style_nodes[0].get(qn("w:val")) if style_nodes else ""
            if style_id in {"1", "21", "2", "Heading1", "Heading2"}:
                anchor_element = current
                break
        current = current.getnext()
    if anchor_element is None:
        raise ValueError(f"Next heading not found after {heading_prefix}")
    current = heading._p.getnext()
    parent = heading._p.getparent()
    while current is not None and current is not anchor_element:
        next_element = current.getnext()
        parent.remove(current)
        current = next_element
    anchor_paragraph = None
    for paragraph in document.paragraphs:
        if paragraph._p is anchor_element:
            anchor_paragraph = paragraph
            break
    if anchor_paragraph is None:
        raise ValueError(f"Anchor paragraph unavailable after {heading_prefix}")
    for text in new_paragraphs:
        add_paragraph_before(document, anchor_paragraph, text, style="Normal")


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build() -> None:
    document = Document(SOURCE)

    replace_paragraph_text(
        find_paragraph(document, "第六章呈現"),
        "第六章呈現PSF校正與反卷積結果，接著以彼此分開的位移改善、配對式消融與多模型比較三組實驗，檢查Raw、Wiener與MOSR的尺寸與位移量測表現，並加入via、線條、EDSR-x1、NAFNet及全解析度影像推論效能。",
    )
    replace_paragraph_text(
        find_paragraph(document, "本節評估MOSR在HS-100MG樣品上的重建與量測表現"),
        "本節評估MOSR在HS-100MG樣品上的重建與量測表現。後續結果分為三組彼此獨立的實驗：第6.2.5節比較Raw、Wiener與MOSR的位移改善；第6.2.6節以配對式消融檢查抗混疊上取樣與以輸入影像為基準的殘差輸出；第6.2.7節則比較Raw、NAFNet、EDSR-x1與MOSR在pillar、via與線條上的位移、CD、邊緣及運算效能。三組實驗各自回答不同問題，圖表與統計不混合計算。",
    )
    replace_paragraph_text(
        find_paragraph(document, "第6.2.3至第6.2.5節使用完成訓練後固定的MOSR模型"),
        "第6.2.3至第6.2.5節使用完成訓練後固定的MOSR模型，評估期間不更新權重。第6.2.6節固定起始權重、資料切分、光學先驗與訓練預算建立配對式消融比較；第6.2.7節的EDSR-x1與NAFNet則依共同資料與相同輸出尺度完成訓練。主要位移結果、消融結果與多模型結果只在各自的共同條件內比較。",
    )
    replace_paragraph_text(
        find_paragraph(document, "跨越35%至65%閾值與兩種剖面方式後"),
        "跨越35%至65%閾值與兩種剖面方式後，MOSR具有最窄的CD上下界與最低的絕對誤差中位數，顯示其表觀邊界較不受量測規則改變影響。第6.2.7節沿用同一套九顆結構、多門檻、水平與徑向剖面規則，比較Raw、NAFNet、EDSR-x1與MOSR在pillar與via上的SEM對照CD，使第三組實驗與前述CD驗證具有相同的判讀基礎。",
    )

    old_line_figure = clone_previous_figure(find_paragraph(document, "圖 6-54"))
    section_start = find_paragraph(document, "6.2.5 位移量測結果")
    section_end = find_paragraph(document, "6.2.8 綜合討論")
    remove_between(section_start, section_end)
    anchor = section_end

    # 6.2.5 位移改善實驗
    add_paragraph_before(
        document, anchor, "6.2.5 位移改善實驗", style="Heading 3"
    )
    add_paragraph_before(
        document,
        anchor,
        "本節只比較Raw、固定PSF之Wiener反卷積與MOSR，不加入消融條件或其他神經網路。四組序列均沿用原有的最大梯度邊界定位方法；每組30個步次由Raw影像的可量測性與運動品質事先固定，三種方法共用相同影格、九顆pillar、中心位置與PI平台實際回授位移。",
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.5.1 四組序列之逐次位移",
        style="Heading 4",
    )
    add_figure_before(
        document,
        anchor,
        DISPLACEMENT / "01_50x_x_0p5um_step_trace.png",
        37,
        "50倍、0.5 µm位移序列之逐次量測結果。虛線為PI平台實際回授位移，Raw、Wiener與MOSR使用相同的30個相鄰影格。",
    )
    add_figure_before(
        document,
        anchor,
        DISPLACEMENT / "02_50x_x_1um_step_trace.png",
        38,
        "50倍、1.0 µm位移序列之逐次量測結果。各點為中心九顆pillar位移的中位數。",
    )
    add_figure_before(
        document,
        anchor,
        DISPLACEMENT / "03_20x_0p5um_step_trace.png",
        39,
        "20倍、0.5 µm位移序列之逐次量測結果。橫軸為實際納入分析的30個步次。",
    )
    add_figure_before(
        document,
        anchor,
        DISPLACEMENT / "04_20x_1um_step_trace.png",
        40,
        "20倍、1.0 µm位移序列之逐次量測結果。三種方法共用PI平台回授與固定影格名單。",
    )
    add_paragraph_before(
        document,
        anchor,
        "四組序列中，MOSR相對PI平台實際回授位移的平均絕對偏差依序為3.07、3.36、11.74與8.81 nm，均低於對應Raw結果。50倍兩組中，MOSR同時得到最低的平台偏差與九顆pillar位移離散度；20倍兩組則仍維持相對Raw的整體改善。",
    )
    add_table_before(
        document,
        anchor,
        9,
        "四組位移改善實驗相對PI平台實際回授位移之平均絕對偏差。單位為nm，每組均為30個共同步次。",
        ["條件", "Raw", "Wiener", "MOSR"],
        [
            ["50倍、0.5 µm", "5.25", "4.07", "3.07"],
            ["50倍、1.0 µm", "6.98", "3.89", "3.36"],
            ["20倍、0.5 µm", "13.13", "12.78", "11.74"],
            ["20倍、1.0 µm", "12.82", "9.13", "8.81"],
        ],
        [4.6, 3.1, 3.1, 3.1],
    )
    add_table_before(
        document,
        anchor,
        10,
        "四組位移改善實驗之九顆pillar位移離散度。單位為nm，數值越低表示同一影格內的量測越一致。",
        ["條件", "Raw", "Wiener", "MOSR"],
        [
            ["50倍、0.5 µm", "13.03", "6.86", "4.48"],
            ["50倍、1.0 µm", "12.06", "7.07", "3.73"],
            ["20倍、0.5 µm", "25.14", "8.70", "9.13"],
            ["20倍、1.0 µm", "20.28", "9.03", "8.53"],
        ],
        [4.6, 3.1, 3.1, 3.1],
        page_break_before=True,
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.5.2 跨序列整體結果",
        style="Heading 4",
    )
    add_figure_before(
        document,
        anchor,
        DISPLACEMENT / "05_sequence_displacement_mae.png",
        41,
        "四組位移序列相對PI平台實際回授位移的平均絕對偏差。數值越低表示影像量測與平台回授越一致。",
    )
    add_figure_before(
        document,
        anchor,
        DISPLACEMENT / "06_sequence_step_sd.png",
        42,
        "四組位移序列的逐次位移量標準差。數值越低表示相鄰步次間的量測波動越小。",
    )
    add_figure_before(
        document,
        anchor,
        DISPLACEMENT / "07_sequence_within_grid.png",
        43,
        "四組位移序列中九顆pillar的位移離散度。數值越低表示相同實體平移在不同pillar上的估測越一致。",
    )
    add_paragraph_before(
        document,
        anchor,
        "四組序列合併後，Raw、Wiener與MOSR相對PI平台實際回授位移的平均絕對偏差為9.54、7.46與6.74 nm；九顆pillar位移離散度則為17.63、7.92與6.47 nm。MOSR相較Raw分別降低29.36%與63.29%，且改善不是由改變整段位移比例所造成。",
    )

    # 6.2.6 消融實驗
    add_paragraph_before(
        document, anchor, "6.2.6 配對式消融實驗", style="Heading 3"
    )
    add_paragraph_before(
        document,
        anchor,
        "消融實驗與前節位移改善結果分開呈現。各消融條件固定起始權重、資料切分、PSF、訓練預算與評估影格，只改變待檢查的模型設計；圖中的範圍來自兩次固定初始化，用來判斷效果方向是否可重現。",
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.6.1 抗混疊上取樣與位移一致性",
        style="Heading 4",
    )
    add_figure_before(
        document,
        anchor,
        ABLATION / "01_shift_safe_within_grid_effect.png",
        44,
        "一般上取樣相對抗混疊上取樣所增加的九顆pillar位移離散度。正值表示移除抗混疊設計後一致性下降。",
    )
    add_figure_before(
        document,
        anchor,
        ABLATION / "02_shift_safe_mae_effect.png",
        45,
        "一般上取樣相對抗混疊上取樣所增加的位移平均絕對偏差。正值表示量測結果更偏離PI平台回授。",
    )
    add_figure_before(
        document,
        anchor,
        ABLATION / "03_50x_x_0p5um_upsampling_step_trace.png",
        46,
        "50倍、0.5 µm序列中抗混疊上取樣與一般上取樣的逐次位移比較。陰影表示兩次固定初始化結果的範圍。",
    )
    add_figure_before(
        document,
        anchor,
        ABLATION / "04_50x_x_1um_upsampling_step_trace.png",
        47,
        "50倍、1.0 µm序列中抗混疊上取樣與一般上取樣的逐次位移比較。",
    )
    add_figure_before(
        document,
        anchor,
        ABLATION / "05_20x_0p5um_upsampling_step_trace.png",
        48,
        "20倍、0.5 µm序列中抗混疊上取樣與一般上取樣的逐次位移比較。",
    )
    add_figure_before(
        document,
        anchor,
        ABLATION / "06_20x_1um_upsampling_step_trace.png",
        49,
        "20倍、1.0 µm序列中抗混疊上取樣與一般上取樣的逐次位移比較。",
    )
    add_paragraph_before(
        document,
        anchor,
        "四組序列合併後，一般上取樣使九顆pillar位移離散度增加0.977 nm，即14.31%，並使相對PI平台回授的平均絕對偏差增加0.424 nm，即6.43%。兩次固定初始化的方向一致，支持抗混疊上取樣主要改善跨pillar的相對位置一致性。",
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.6.2 以原始影像為基準的殘差輸出與邊界保護",
        style="Heading 4",
    )
    add_figure_before(
        document,
        anchor,
        ABLATION / "07_raw_position_anchor_20x_cd_effect.png",
        50,
        "移除以原始影像為基準的殘差輸出後，20倍pillar CD中位絕對誤差的增加量。正值表示保留原始位置基準時誤差較低。",
    )
    add_figure_before(
        document,
        anchor,
        ABLATION / "08_raw_position_anchor_20x_per_pillar_cd.png",
        51,
        "完整模型與移除原始影像位置基準後，各pillar之CD絕對誤差比較。誤差棒表示兩次固定初始化的結果範圍。",
    )
    add_paragraph_before(
        document,
        anchor,
        "移除以原始影像為基準的殘差輸出後，20倍pillar的CD中位絕對誤差平均增加88.94 nm，即73.97%，且兩次固定初始化皆呈現誤差增加。此結果說明該設計的主要作用不是單純提高對比，而是限制重建輪廓偏離原始位置基準。",
    )
    add_table_before(
        document,
        anchor,
        11,
        "兩項主要模型設計的配對式消融結果。正值表示移除或替換該設計後量測變差。",
        ["模型設計", "主要量測", "變化", "判讀"],
        [
            [
                "抗混疊上取樣",
                "九顆pillar位移離散度",
                "+0.977 nm（+14.31%）",
                "改善跨位置一致性",
            ],
            [
                "抗混疊上取樣",
                "位移平均絕對偏差",
                "+0.424 nm（+6.43%）",
                "降低平台回授偏差",
            ],
            [
                "原始影像位置基準",
                "20倍pillar CD中位絕對誤差",
                "+88.94 nm（+73.97%）",
                "限制輪廓偏移",
            ],
        ],
        [3.3, 4.2, 3.5, 3.6],
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.6.3 消融結果小結",
        style="Heading 4",
    )
    add_paragraph_before(
        document,
        anchor,
        "兩項主要設計對應不同層級的量測問題：抗混疊上取樣降低不同pillar之間的位移分歧；以原始影像為基準的殘差輸出則維持最終輪廓位置與尺寸。消融結果因此把模型設計分別連結至位移一致性與邊界重建精度，而不是只以整體流程解釋結果。",
    )

    # 6.2.7 多模型與跨形貌比較
    add_paragraph_before(
        document,
        anchor,
        "6.2.7 多模型與跨形貌比較",
        style="Heading 3",
    )
    add_paragraph_before(
        document,
        anchor,
        "本節比較Raw、NAFNet、EDSR-x1與MOSR，不納入Wiener與消融條件。Pillar與via皆使用Raw固定的九顆結構中心；位移沿用原有最大梯度方法，CD則與前述實驗相同，掃描35%至65%七個相對門檻並同時使用水平與徑向剖面。四種方法因此在相同結構身分、物體空間尺度與SEM參考下接受比較。",
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.7.1 Pillar與via位移比較",
        style="Heading 4",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "01_pillar_50x_x_0p5um_model_step_trace.png",
        52,
        "50倍、0.5 µm pillar序列之Raw、NAFNet、EDSR-x1與MOSR逐次位移比較。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "03_pillar_20x_0p5um_model_step_trace.png",
        53,
        "20倍、0.5 µm pillar序列之Raw、NAFNet、EDSR-x1與MOSR逐次位移比較。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "05_via_50x_x_0p5um_model_step_trace.png",
        54,
        "50倍、0.5 µm via序列之Raw、NAFNet、EDSR-x1與MOSR逐次位移比較。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "07_via_20x_0p5um_model_step_trace.png",
        55,
        "20倍、0.5 µm via序列之Raw、NAFNet、EDSR-x1與MOSR逐次位移比較。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "09_model_displacement_mae.png",
        56,
        "不同倍率與形貌下，各方法相對PI平台實際回授位移的平均絕對偏差。每個柱狀值合併0.5 µm與1.0 µm兩組序列。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "10_model_within_grid.png",
        57,
        "不同倍率與形貌下，各方法對九顆結構量得之位移離散度。數值越低表示跨結構位置的一致性越高。",
    )
    add_paragraph_before(
        document,
        anchor,
        "在50倍pillar與via中，MOSR的位移平均絕對偏差分別為3.21與3.12 nm，均明顯低於Raw，並接近EDSR-x1的2.87與3.07 nm。20倍pillar中，MOSR的九宮格位移離散度為8.83 nm，為四種方法最低；20倍via則以EDSR-x1的位移指標較低。這組結果顯示不同網路在位移指標上的排序會隨倍率與形貌改變，因此後續再以SEM對照CD檢查邊界位置是否同步改善。",
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.7.2 統一規則之SEM對照CD比較",
        style="Heading 4",
    )
    add_paragraph_before(
        document,
        anchor,
        "每一倍率與形貌均量測九顆結構；四種方法共用由Raw影像固定的中心，並將35%至65%七個門檻與水平、徑向兩種剖面全部納入。每個倍率、形貌與方法共有126筆尺寸量測，主要統計量為相對逐顆SEM直徑的中位絕對誤差。20倍與50倍影像依各自拍攝方向對應至相同的實體pillar編號；via的SEM半徑則先乘以2轉為直徑。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "15_model_cd_median_abs_error.png",
        58,
        "50倍與20倍之pillar、via在統一量測規則下的CD中位絕對誤差。每種方法均合併九顆結構、七個門檻與兩種剖面。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "17_50x_pillar_cd_range.png",
        59,
        "50倍九顆pillar的SEM參考直徑與各方法CD量測範圍。符號為全部門檻與剖面結果的中位數，誤差棒為最小值至最大值。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "18_20x_pillar_cd_range.png",
        60,
        "20倍九顆pillar的SEM參考直徑與各方法CD量測範圍。結構編號已依20倍拍攝方向與SEM資料正確對應。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "19_50x_via_cd_range.png",
        61,
        "50倍九顆via的SEM參考直徑與各方法CD量測範圍。",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "20_20x_via_cd_range.png",
        62,
        "20倍九顆via的SEM參考直徑與各方法CD量測範圍。",
    )
    add_table_before(
        document,
        anchor,
        12,
        "多模型比較之CD中位絕對誤差。單位為nm，每格均包含九顆結構、七個門檻與兩種剖面。",
        ["條件", "Raw", "NAFNet", "EDSR-x1", "MOSR"],
        [
            ["50倍pillar", "107.57", "51.42", "40.23", "53.24"],
            ["20倍pillar", "169.44", "94.67", "384.39", "75.39"],
            ["50倍via", "246.77", "178.16", "154.76", "181.20"],
            ["20倍via", "497.53", "446.97", "420.54", "329.01"],
        ],
        [3.7, 2.7, 2.7, 2.7, 2.7],
        page_break_before=True,
    )
    add_paragraph_before(
        document,
        anchor,
        "MOSR在四個條件中均低於Raw：50倍pillar、20倍pillar、50倍via與20倍via的誤差分別降低50.51%、55.51%、26.57%與33.87%。在取樣較粗的20倍pillar與20倍via中，MOSR為四種方法最低；50倍兩組則由EDSR-x1取得最低誤差。結果顯示MOSR的CD優勢主要出現在較粗取樣條件，而不是依單一形貌或單一門檻建立。",
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.7.3 邊緣過渡與線條週期",
        style="Heading 4",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "11_model_edge_width.png",
        63,
        "不同倍率與圓形結構下，各方法的平均10–90%邊緣過渡寬度。數值越低表示亮暗轉換集中於較窄範圍。",
    )
    insert_cloned_figure_before(
        document,
        anchor,
        old_line_figure,
        64,
        "Raw與MOSR於水平及垂直線條上的10–90%邊緣過渡寬度。重建前後維持相同影像取樣尺寸。",
    )
    add_paragraph_before(
        document,
        anchor,
        "線條區域以Raw固定相同邊界位置。水平與垂直線條的邊緣寬度由683.8與706.1 nm降至587.1與579.2 nm，分別改善14.2%與18.0%；表觀週期相對Raw僅變化1.95與0.72 nm。這表示線條邊緣變集中時，週期性的幾何關係仍被保留。",
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.7.4 全解析度影像推論效能",
        style="Heading 4",
    )
    add_figure_before(
        document,
        anchor,
        MODEL / "12_model_full_frame_runtime.png",
        65,
        "NAFNet、EDSR-x1與MOSR處理單張2048×1536影像的完整推論時間。數值為相同顯示卡上八次正式推論的中位數。",
    )
    add_table_before(
        document,
        anchor,
        13,
        "相同顯示卡下的全解析度影像推論效能與模型規模。",
        ["方法", "參數量（M）", "中位時間（s）", "峰值記憶體（GB）", "吞吐量（MP/s）"],
        [
            ["NAFNet", "4.182", "0.662", "1.32", "4.75"],
            ["EDSR-x1", "4.071", "3.781", "2.31", "0.83"],
            ["MOSR", "4.024", "0.408", "5.34", "7.71"],
        ],
        [3.2, 2.7, 3.0, 3.2, 2.7],
    )
    add_paragraph_before(
        document,
        anchor,
        "在各模型的正式完整影像推論路徑下，MOSR的中位時間為0.408 s，低於NAFNet的0.662 s與EDSR-x1的3.781 s；三者參數量均約為四百萬。MOSR的峰值GPU記憶體較高，但仍能在RTX A6000上直接完成整張影像推論。",
    )
    add_paragraph_before(
        document,
        anchor,
        "6.2.7.5 多模型與跨形貌結果小結",
        style="Heading 4",
    )
    add_paragraph_before(
        document,
        anchor,
        "多模型比較顯示，EDSR-x1在多數位移指標與50倍CD中表現突出；MOSR則在兩組20倍CD取得最低誤差，並在四種CD條件中皆相對Raw改善。配合pillar主實驗、via與線條結果，可將MOSR的主要效果歸納為較粗取樣下的邊界位置修復、跨pillar位移一致性，以及不改變輸出尺寸的完整影像重建。",
    )

    # Rewrite the comprehensive discussion.
    discussion_paragraphs = []
    current = section_end._p.getnext()
    while current is not None and current.tag == qn("w:p"):
        style_nodes = current.xpath("./w:pPr/w:pStyle")
        style_id = style_nodes[0].get(qn("w:val")) if style_nodes else ""
        if style_id in {"1", "Heading1"}:
            break
        discussion_paragraphs.append(current)
        current = current.getnext()
    for element in discussion_paragraphs:
        element.getparent().remove(element)
    chapter_seven_heading = find_paragraph(document, "結論與未來工作")
    for text in [
        "第6.1節確認局部擬合PSF能使Wiener反卷積在20倍與50倍影像中縮窄邊緣過渡，因此擬合結果與實際模糊方向大致一致；同時，Wiener的振鈴與門檻敏感度說明影像變銳利不必然等同於尺寸量測更準確。",
        "主要位移改善實驗使用Raw固定的四組30步名單。MOSR在120個共同步次中取得最低的平台回授位移平均絕對偏差與九顆pillar位移離散度。配對式消融進一步顯示，抗混疊上取樣主要降低跨pillar位移分歧，以原始影像為基準的殘差輸出則限制20倍CD邊界偏移。",
        "多模型比較採相同位移規則與統一CD量測方式。MOSR的CD誤差在四個倍率與形貌條件中均低於Raw，並在20倍pillar與20倍via最低；EDSR-x1則在50倍CD與多數位移指標較低。這些結果把MOSR的優勢範圍定位於較粗取樣下的尺寸修復與主要實驗中的跨位置穩定性，而非單純追求最窄邊緣。",
    ]:
        add_paragraph_before(document, chapter_seven_heading, text, style="Normal")

    replace_paragraph_text(
        find_paragraph(document, "本研究以量測導向超解析為主軸，完成光學先驗建立"),
        "本研究以量測導向超解析為主軸，完成PSF校正、單張等尺寸影像重建、SEM對照CD量測、PI閉迴路位移量測、配對式消融、多模型與跨形貌比較，以及完整影像推論效能評估。本章依研究問題整理主要結論與量化結果，並提出後續工作。",
    )

    replace_section_body(
        document,
        "7.1 研究結論",
        [
            "本研究的第一個問題，是影像重建後的邊界是否更容易定位，且量得尺寸是否更接近獨立SEM參考。固定PSF的Wiener反卷積在20倍與50倍均縮窄邊緣過渡，但同時提高振鈴與門檻敏感度；MOSR在主要pillar比較中取得較低的跨門檻CD誤差，顯示尺寸修復不能只以視覺銳利度判斷。",
            "第二個問題，是重建後的微位移是否仍保持穩定。四組位移序列共120個共同步次中，MOSR相對PI平台實際回授位移的平均絕對偏差與九顆pillar位移離散度均低於Raw與Wiener。改善主要來自降低步與步、pillar與pillar之間的量測波動，而不是改變整段位移比例。",
            "第三個問題，是模型設計能否對應至具體量測效果。配對式消融顯示，抗混疊上取樣降低跨pillar位移分歧；以原始影像為基準的殘差輸出則降低20倍CD誤差。兩項設計分別處理解碼階段的空間一致性與輸出階段的邊界位置保護。",
            "跨模型比較進一步顯示，MOSR在50倍與20倍的pillar、via CD均優於Raw，並在兩組20倍條件取得最低誤差；EDSR-x1則在50倍CD與多數位移指標表現較佳。這項結果說明MOSR的主要優勢集中於較粗取樣條件下的邊界與尺寸修復。",
        ],
    )
    replace_section_body(
        document,
        "7.2 主要方法貢獻",
        [
            "本研究的第一項主要貢獻，是把解碼上取樣造成的位置敏感性轉化為可由九顆結構共同位移直接檢查的問題，並以抗混疊上取樣降低不同空間位置的量測分歧。第二項主要貢獻，是以原始影像作為低頻輪廓與中心位置基準，只允許模型加入受限制的高頻修正，使影像變清楚時仍保留可量測的邊界位置。",
            "本研究並建立同一套由影像重建連接至後端量測的驗證方式：CD使用Raw固定中心、逐顆SEM尺寸、35%至65%門檻及水平與徑向剖面；位移則使用PI平台實際回授位置、固定30步名單與九顆結構共同量測。模型設計因此可由尺寸與位移結果直接判讀，而不只依靠視覺示例。",
        ],
    )
    replace_section_body(
        document,
        "7.3 主要量化結果",
        [
            "在主要50倍pillar CD比較中，Raw、Wiener與MOSR跨七個門檻與兩種剖面的絕對誤差中位數分別為95.3、123.7與76.1 nm；MOSR相較Raw降低20.16%，相較Wiener降低38.48%。九顆pillar的CD上下界寬度中位數則由Raw的354.5 nm與Wiener的429.6 nm降至MOSR的143.3 nm。",
            "在四組位移序列共120個共同步次中，Raw、Wiener與MOSR相對PI平台實際回授位移的平均絕對偏差為9.54、7.46與6.74 nm，九顆pillar位移離散度為17.63、7.92與6.47 nm。MOSR相較Raw分別降低29.36%與63.29%。",
            "消融實驗中，一般上取樣使九顆pillar位移離散度增加0.977 nm，即14.31%；移除原始影像位置基準後，20倍pillar CD中位絕對誤差增加88.94 nm，即73.97%。兩項結果分別支持抗混疊上取樣的跨位置一致性作用，以及原始影像位置基準的邊界保護作用。",
            "在多模型CD比較中，MOSR相對Raw於50倍pillar、20倍pillar、50倍via與20倍via分別改善50.51%、55.51%、26.57%與33.87%。MOSR在兩組20倍條件最低，EDSR-x1在兩組50倍條件最低。完整2048×1536影像的中位推論時間為0.408 s。",
        ],
    )
    replace_section_body(
        document,
        "7.4 驗證範圍與量測可信度",
        [
            "本研究在同一受控顯微平台中，以HS-100MG樣品的pillar、via、水平與垂直線條完成驗證，涵蓋20倍與50倍物鏡，以及0.5 µm與1.0 µm位移步長。SEM提供九顆pillar與九顆via的逐顆尺寸，PI平台實際回授位置提供共同位移參考。",
            "所有方法共用影像、結構身分、Raw固定中心、物體空間尺度、門檻與統計規則。第三組多模型CD亦重新採用與主要CD實驗相同的九顆結構、七個門檻及兩種剖面，因此不同模型的結果可在同一量測基礎上比較。",
        ],
    )
    replace_section_body(
        document,
        "7.5 未來工作",
        [
            "後續可擴充不同批次樣品、更多pillar與via尺寸、方形結構及混合線寬版圖，再於不同視野位置、焦平面、曝光與照明條件下建立分層測試[50,54,55,56]。若更換相機、物鏡或照明模組，則重新建立物體空間尺度與光學校正資料，再判斷模型需完整訓練或少量校正[75,77,90]。",
            "後續亦可使用雷射干涉儀與PI平台實際回授位置同步量測，進一步建立外部量測追溯關係，並將重複性、像素尺度、SEM參考、平台回授與模型變動整合為完整不確定度預算[84,85]。模型端則可針對20倍via與不同線寬資料增加訓練覆蓋，延伸較粗取樣條件下的尺寸修復能力。",
        ],
    )
    replace_section_body(
        document,
        "7.6 小結",
        [
            "本研究證明，顯微影像變清楚後仍須以尺寸與位移檢查其量測價值。MOSR以抗混疊上取樣維持跨位置關係，並以原始影像為基準的殘差輸出限制最終輪廓偏移；主要位移實驗、配對式消融與統一CD多模型比較分別提供整體效果、模組作用與跨形貌證據。結果顯示MOSR在主要pillar量測與較粗取樣的20倍CD中具有穩定改善，並能以單張完整影像完成等尺寸重建。",
        ],
    )

    set_update_fields(document)
    document.core_properties.title = "量測導向超解析顯微影像重建與微位移量測"
    document.core_properties.comments = (
        "v19：更新第六章位移改善、消融與多模型比較圖表，"
        "並同步調整第七章結論與未來工作。"
    )
    document.save(OUTPUT)

    audit = {
        "source": str(SOURCE),
        "output": str(OUTPUT),
        "figure_numbers_added": list(range(37, 66)),
        "table_numbers_added": list(range(9, 14)),
        "figure_count_added": 29,
        "table_count_added": 5,
        "chapter_six_structure": [
            "6.2.5 位移改善實驗",
            "6.2.6 配對式消融實驗",
            "6.2.7 多模型與跨形貌比較",
            "6.2.8 綜合討論",
        ],
        "chapter_seven_sections": [
            "7.1 研究結論",
            "7.2 主要方法貢獻",
            "7.3 主要量化結果",
            "7.4 驗證範圍與量測可信度",
            "7.5 未來工作",
            "7.6 小結",
        ],
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    build()
