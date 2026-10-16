from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "論文完整版_20260716_v19_第六章更新版.docx"
REPORT = ROOT / "v19_structure_qa.json"


def section_text(document: Document, start: str, end: str) -> str:
    collecting = False
    result: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""
        if text == start and style_name == "Heading 1":
            collecting = True
        if collecting:
            result.append(text)
        if collecting and text == end and style_name == "Heading 1":
            break
    return "\n".join(result)


document = Document(DOCX)
chapter_six = section_text(document, "實驗結果與分析", "結論與未來工作")
chapter_seven = section_text(document, "結論與未來工作", "參考文獻")
all_table_text = "\n".join(
    cell.text
    for table in document.tables
    for row in table.rows
    for cell in row.cells
)

headings = [
    paragraph.text.strip()
    for paragraph in document.paragraphs
    if paragraph.style
    and paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    and (
        paragraph.text.strip() in {"實驗結果與分析", "結論與未來工作"}
        or re.match(r"^[67]\.", paragraph.text.strip())
    )
]

figure_numbers = [
    int(match.group(1))
    for paragraph in document.paragraphs
    if paragraph.style and paragraph.style.name == "Figure Caption"
    for match in [re.match(r"^圖 6-(\d+)\s", paragraph.text.strip())]
    if match
]
table_numbers = [
    int(match.group(1))
    for paragraph in document.paragraphs
    if paragraph.style and paragraph.style.name == "Table Caption"
    for match in [re.match(r"^表 6-(\d+)\s", paragraph.text.strip())]
    if match
]

checks = {
    "figure_numbers_are_1_to_65": figure_numbers == list(range(1, 66)),
    "table_numbers_are_1_to_13": table_numbers == list(range(1, 14)),
    "three_experiment_sections_present": all(
        heading in chapter_six
        for heading in (
            "6.2.5 位移改善實驗",
            "6.2.6 配對式消融實驗",
            "6.2.7 多模型與跨形貌比較",
        )
    ),
    "old_midpoint_wording_removed": "CD區間中點" not in chapter_six + chapter_seven,
    "base_unet_removed": all(
        term not in chapter_six + chapter_seven
        for term in ("Base U-Net", "基礎U-Net", "基礎 U-Net")
    ),
    "chapter_seven_sections_present": all(
        heading in chapter_seven
        for heading in (
            "7.1 研究結論",
            "7.2 主要方法貢獻",
            "7.3 主要量化結果",
            "7.4 驗證範圍與量測可信度",
            "7.5 未來工作",
            "7.6 小結",
        )
    ),
    "latest_runtime_values_present": all(
        value in all_table_text
        for value in ("0.408", "0.662", "3.781")
    ),
}

report = {
    "docx": str(DOCX),
    "paragraph_count": len(document.paragraphs),
    "table_count": len(document.tables),
    "inline_shape_count": len(document.inline_shapes),
    "headings": headings,
    "figure_numbers": figure_numbers,
    "table_numbers": table_numbers,
    "checks": checks,
    "all_checks_passed": all(checks.values()),
}
REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

print(json.dumps(report, ensure_ascii=False, indent=2))
if not report["all_checks_passed"]:
    raise SystemExit(1)
