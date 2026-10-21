from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path

from PIL import Image
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
SOURCE = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用"
    r"\論文完整版_20260716_v19_第六章更新版.docx"
)
V18 = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用"
    r"\論文完整版_20260716_v18_圖表整理版.docx"
)
OUTPUT = ROOT / "論文完整版_20260716_v19_第六章更新版.docx"
BACKUP = ROOT / "v19_before_format_alignment.docx"
AUDIT = ROOT / "v19_format_alignment_audit.json"
OLD_IMAGES = ROOT / "old_figure_set"
CROPPED_IMAGES = ROOT / "old_figure_set_cropped"
CROPPED_IMAGES.mkdir(exist_ok=True)

FIGURE_NUMBERS_WITH_SCALE_LABEL = [24, 25, 26, 28, 29, 30]


def find_paragraph(document: Document, prefix: str, style_name: str | None = None):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix) and (
            style_name is None or paragraph.style.name == style_name
        ):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def paragraph_index(document: Document, target: Paragraph) -> int:
    return next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is target._p
    )


def copy_paragraph_properties(
    source: Paragraph,
    target: Paragraph,
    preserve_page_break: bool = False,
) -> None:
    had_page_break = (
        target._p.pPr is not None
        and target._p.pPr.find(qn("w:pageBreakBefore")) is not None
    )
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, deepcopy(source._p.pPr))
    if preserve_page_break and had_page_break:
        ppr = target._p.get_or_add_pPr()
        if ppr.find(qn("w:pageBreakBefore")) is None:
            ppr.append(OxmlElement("w:pageBreakBefore"))


def copy_run_properties(source_run, target_run) -> None:
    if target_run._r.rPr is not None:
        target_run._r.remove(target_run._r.rPr)
    if source_run._r.rPr is not None:
        target_run._r.insert(0, deepcopy(source_run._r.rPr))


def format_paragraph(
    paragraph: Paragraph,
    paragraph_reference: Paragraph,
    run_reference: Paragraph | None = None,
    preserve_page_break: bool = False,
) -> None:
    copy_paragraph_properties(
        paragraph_reference,
        paragraph,
        preserve_page_break=preserve_page_break,
    )
    source_paragraph = run_reference or paragraph_reference
    source_run = next((run for run in source_paragraph.runs if run.text), None)
    if source_run is None:
        return
    for run in paragraph.runs:
        copy_run_properties(source_run, run)


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def tables_by_caption(document: Document) -> dict[int, Table]:
    result: dict[int, Table] = {}
    last_caption_number = None
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            match = re.match(r"^表 6-(\d+)\s", block.text.strip())
            last_caption_number = int(match.group(1)) if match else None
        elif last_caption_number is not None:
            result[last_caption_number] = block
            last_caption_number = None
    return result


def copy_cell_visual_properties(source_cell, target_cell, copy_shading: bool) -> None:
    source_tcpr = source_cell._tc.get_or_add_tcPr()
    target_tcpr = target_cell._tc.get_or_add_tcPr()

    target_margins = target_tcpr.find(qn("w:tcMar"))
    if target_margins is not None:
        target_tcpr.remove(target_margins)

    for tag in ("w:shd", "w:vAlign"):
        existing = target_tcpr.find(qn(tag))
        if existing is not None:
            target_tcpr.remove(existing)
        source_element = source_tcpr.find(qn(tag))
        if source_element is not None and (tag != "w:shd" or copy_shading):
            target_tcpr.append(deepcopy(source_element))


def align_new_tables(document: Document) -> int:
    tables = tables_by_caption(document)
    reference = tables[8]
    header_paragraph = reference.cell(0, 0).paragraphs[0]
    header_run = next(run for run in header_paragraph.runs if run.text)
    body_paragraph = reference.cell(1, 0).paragraphs[0]
    body_run = next(run for run in body_paragraph.runs if run.text)

    changed = 0
    for number in range(9, 14):
        table = tables[number]
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                reference_cell = reference.cell(0 if row_index == 0 else 1, 0)
                copy_cell_visual_properties(
                    reference_cell,
                    cell,
                    copy_shading=row_index == 0,
                )
                for paragraph in cell.paragraphs:
                    source_paragraph = (
                        header_paragraph if row_index == 0 else body_paragraph
                    )
                    source_run = header_run if row_index == 0 else body_run
                    copy_paragraph_properties(source_paragraph, paragraph)
                    for run in paragraph.runs:
                        copy_run_properties(source_run, run)
                changed += 1
    return changed


def crop_old_images() -> dict[int, Path]:
    cropped: dict[int, Path] = {}
    for number in FIGURE_NUMBERS_WITH_SCALE_LABEL:
        source = OLD_IMAGES / f"old_figure_6_{number}.png"
        output = CROPPED_IMAGES / f"figure_6_{number}_with_5um_label.png"
        image = Image.open(source).convert("RGB")
        # The old figure title occupies the top 52 pixels.  Keep the original
        # microscopy crop and its in-image 5 µm label, but remove the redundant
        # English title because the Word caption already carries that role.
        cropped_image = image.crop((0, 52, image.width, image.height))
        cropped_image.save(output, quality=95)
        cropped[number] = output
    return cropped


def replace_figure_image(
    document: Document,
    number: int,
    image_path: Path,
) -> None:
    caption = find_paragraph(document, f"圖 6-{number} ", "Figure Caption")
    current = caption._p.getprevious()
    image_paragraph = None
    width_emu = None
    while current is not None:
        drawings = current.xpath(".//*[local-name()='drawing']")
        if drawings:
            extents = current.xpath(".//*[local-name()='extent']")
            width_emu = int(extents[0].get("cx")) if extents else None
            image_paragraph = Paragraph(current, document)
            break
        current = current.getprevious()
    if image_paragraph is None or width_emu is None:
        raise RuntimeError(f"Figure paragraph not found for figure 6-{number}")

    for run in list(image_paragraph.runs):
        image_paragraph._p.remove(run._r)
    image_paragraph.add_run().add_picture(str(image_path), width=Emu(width_emu))


def apply() -> None:
    BACKUP.write_bytes(SOURCE.read_bytes())
    document = Document(SOURCE)
    v18 = Document(V18)

    body_reference = find_paragraph(
        document, "為避免以單一有利閾值評估影像重建方法"
    )
    heading3_reference = find_paragraph(
        document, "6.2.4 SEM 對照之 CD 修正與閾值穩健性"
    )
    heading4_reference = find_paragraph(
        document, "6.2.4.1 50× 靜態影像之跨閾值比較"
    )
    figure_caption_reference = find_paragraph(
        document, "圖 6-24 ", "Figure Caption"
    )
    table_caption_reference = find_paragraph(document, "表 6-8 ", "Table Caption")

    v18_chapter_seven = find_paragraph(v18, "結論與未來工作", "Heading 1")
    v18_chapter_seven_index = paragraph_index(v18, v18_chapter_seven)
    chapter7_intro_reference = next(
        paragraph
        for paragraph in v18.paragraphs[v18_chapter_seven_index + 1 :]
        if paragraph.style.name == "Normal" and paragraph.text.strip()
    )
    chapter7_body_reference = find_paragraph(v18, "第一個研究問題關注")

    modified_run_prefixes = [
        "第六章呈現PSF校正與反卷積結果",
        "本節評估MOSR在HS-100MG樣品上的重建與量測表現",
        "第6.2.3至第6.2.5節使用完成訓練後固定的MOSR模型",
        "跨越35%至65%閾值與兩種剖面方式後",
    ]
    run_reference = next(run for run in body_reference.runs if run.text)
    modified_run_count = 0
    for prefix in modified_run_prefixes:
        paragraph = find_paragraph(document, prefix)
        for run in paragraph.runs:
            copy_run_properties(run_reference, run)
            modified_run_count += 1

    start = paragraph_index(
        document,
        find_paragraph(document, "6.2.5 位移改善實驗")
    )
    chapter_seven_heading = find_paragraph(document, "結論與未來工作", "Heading 1")
    chapter_seven_index = paragraph_index(document, chapter_seven_heading)

    counts = {
        "heading3": 0,
        "heading4": 0,
        "body_chapter6": 0,
        "figure_caption": 0,
        "table_caption": 0,
        "chapter7_intro": 0,
        "chapter7_body": 0,
        "modified_runs_before_section": modified_run_count,
    }

    for paragraph in document.paragraphs[start:chapter_seven_index]:
        text = paragraph.text.strip()
        style_name = paragraph.style.name
        if (
            style_name == "Heading 3"
            and text.startswith(("6.2.5 ", "6.2.6 ", "6.2.7 "))
        ):
            format_paragraph(paragraph, heading3_reference)
            counts["heading3"] += 1
        elif style_name == "Heading 4" and text.startswith("6.2."):
            format_paragraph(paragraph, heading4_reference)
            counts["heading4"] += 1
        elif style_name == "Figure Caption":
            match = re.match(r"^圖 6-(\d+)\s", text)
            if match and 37 <= int(match.group(1)) <= 65:
                format_paragraph(paragraph, figure_caption_reference)
                counts["figure_caption"] += 1
        elif style_name == "Table Caption":
            match = re.match(r"^表 6-(\d+)\s", text)
            if match and 9 <= int(match.group(1)) <= 13:
                format_paragraph(
                    paragraph,
                    table_caption_reference,
                    preserve_page_break=True,
                )
                counts["table_caption"] += 1
        elif (
            style_name == "Normal"
            and text
            and not paragraph._p.xpath(".//*[local-name()='drawing']")
        ):
            format_paragraph(paragraph, body_reference)
            counts["body_chapter6"] += 1

    references_heading = find_paragraph(document, "參考文獻", "Heading 1")
    references_index = paragraph_index(document, references_heading)
    chapter7_normal_seen = 0
    for paragraph in document.paragraphs[chapter_seven_index + 1 : references_index]:
        if paragraph.style.name != "Normal" or not paragraph.text.strip():
            continue
        if chapter7_normal_seen == 0:
            format_paragraph(paragraph, chapter7_intro_reference)
            counts["chapter7_intro"] += 1
        else:
            format_paragraph(paragraph, chapter7_body_reference)
            counts["chapter7_body"] += 1
        chapter7_normal_seen += 1

    counts["table_cells"] = align_new_tables(document)

    cropped_images = crop_old_images()
    for number, image_path in cropped_images.items():
        replace_figure_image(document, number, image_path)
    counts["restored_scaled_images"] = len(cropped_images)

    document.save(OUTPUT)
    AUDIT.write_text(
        json.dumps(
            {
                "source": str(SOURCE),
                "backup": str(BACKUP),
                "output": str(OUTPUT),
                "counts": counts,
                "restored_figures": FIGURE_NUMBERS_WITH_SCALE_LABEL,
                "style_references": {
                    "body": body_reference.text,
                    "heading3": heading3_reference.text,
                    "heading4": heading4_reference.text,
                    "figure_caption": figure_caption_reference.text,
                    "table_caption": table_caption_reference.text,
                    "chapter7_intro": chapter7_intro_reference.text,
                    "chapter7_body": chapter7_body_reference.text,
                },
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(OUTPUT)
    print(json.dumps(counts, ensure_ascii=False))


if __name__ == "__main__":
    apply()
