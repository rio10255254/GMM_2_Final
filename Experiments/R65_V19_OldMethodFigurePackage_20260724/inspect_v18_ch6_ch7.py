from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path(__file__).with_name("v18_source_copy.docx")
OUTPUT = Path(__file__).with_name("v18_ch6_ch7_structure.json")


def iter_blocks(document: Document):
    body = document.element.body
    for index, child in enumerate(body.iterchildren()):
        if child.tag == qn("w:p"):
            yield index, "paragraph", Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield index, "table", Table(child, document)


def drawing_count(paragraph: Paragraph) -> int:
    return len(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def table_preview(table: Table) -> str:
    rows = []
    for row in table.rows[:3]:
        rows.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
    return " || ".join(rows)


def is_chapter_marker(text: str, chapter_number: int) -> bool:
    chinese = {6: "六", 7: "七"}[chapter_number]
    compact = re.sub(r"\s+", "", text)
    return (
        compact.startswith(f"第{chinese}章")
        or re.match(rf"^{chapter_number}(?:章|[.、．])", compact) is not None
    )


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    document = Document(SOURCE)
    blocks = []
    for body_index, kind, block in iter_blocks(document):
        if kind == "paragraph":
            text = block.text.strip()
            blocks.append(
                {
                    "body_index": body_index,
                    "kind": kind,
                    "style": block.style.name if block.style else "",
                    "text": text,
                    "drawings": drawing_count(block),
                    "keep_with_next": block.paragraph_format.keep_with_next,
                    "page_break_before": block.paragraph_format.page_break_before,
                }
            )
        else:
            blocks.append(
                {
                    "body_index": body_index,
                    "kind": kind,
                    "rows": len(block.rows),
                    "cols": len(block.columns),
                    "text": table_preview(block),
                }
            )

    chapter_six = next(
        i
        for i, block in enumerate(blocks)
        if block["kind"] == "paragraph"
        and block["style"].startswith("Heading 1")
        and "實驗結果" in block["text"]
    )
    chapter_seven = next(
        i
        for i, block in enumerate(blocks)
        if i > chapter_six
        and block["kind"] == "paragraph"
        and block["style"].startswith("Heading 1")
        and "結論與未來工作" in block["text"]
    )
    selected = blocks[max(0, chapter_six - 5) : min(len(blocks), chapter_seven + 160)]

    output = {
        "source": str(SOURCE),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "chapter_six_block_list_index": chapter_six,
        "chapter_seven_block_list_index": chapter_seven,
        "selected_blocks": selected,
    }
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")

    for block in selected:
        if block["kind"] == "paragraph":
            print(
                f'{block["body_index"]:04d} P '
                f'[{block["style"]}] d={block["drawings"]} '
                f'{block["text"]}'
            )
        else:
            print(
                f'{block["body_index"]:04d} T '
                f'{block["rows"]}x{block["cols"]} {block["text"]}'
            )
    print(f"\nWrote {OUTPUT}")


if __name__ == "__main__":
    main()
