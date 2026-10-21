from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "論文完整版_20260716_v19_第六章更新版.docx"
CROPPED = ROOT / "old_figure_set_cropped"
REPORT = ROOT / "v19_format_alignment_qa.json"


def find_paragraph(document, prefix, style_name=None):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix) and (
            style_name is None or paragraph.style.name == style_name
        ):
            return paragraph
    raise ValueError(prefix)


def paragraph_index(document, target):
    return next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is target._p
    )


def ppr_signature(paragraph):
    ppr = paragraph._p.pPr
    if ppr is None:
        return None
    clone = ppr.__deepcopy__({})
    for tag in ("w:rsidR", "w:rsidRDefault", "w:rsidP", "w:pageBreakBefore"):
        for node in clone.findall(qn(tag)):
            clone.remove(node)
    return clone.xml


def rpr_signature(run):
    rpr = run._r.rPr
    if rpr is None:
        return None
    clone = rpr.__deepcopy__({})
    for tag in ("w:rsidR", "w:rsidDel"):
        for node in clone.findall(qn(tag)):
            clone.remove(node)
    return clone.xml


def first_nonempty_run(paragraph):
    return next(run for run in paragraph.runs if run.text.strip())


def has_page_break_before(paragraph):
    return (
        paragraph._p.pPr is not None
        and paragraph._p.pPr.find(qn("w:pageBreakBefore")) is not None
    )


def drawing_blob_before_caption(document, caption):
    index = paragraph_index(document, caption)
    for paragraph in reversed(document.paragraphs[:index]):
        blips = paragraph._p.xpath(".//a:blip")
        if not blips:
            continue
        relationship_id = blips[-1].get(qn("r:embed"))
        return document.part.related_parts[relationship_id].blob
    raise ValueError(f"No drawing found before {caption.text}")


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def main():
    document = Document(DOCX)

    h3_reference = find_paragraph(document, "6.2.4 SEM", "Heading 3")
    h4_reference = find_paragraph(document, "6.2.4.1", "Heading 4")
    body_reference = find_paragraph(document, "為避免以單一有利閾值")
    figure_caption_reference = find_paragraph(document, "圖 6-24", "Figure Caption")
    table_caption_reference = find_paragraph(document, "表 6-8", "Table Caption")

    h3_targets = [
        find_paragraph(document, prefix, "Heading 3")
        for prefix in ("6.2.5", "6.2.6", "6.2.7")
    ]
    h4_targets = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 4"
        and re.match(r"6\.2\.[567]\.", paragraph.text.strip())
    ]
    figure_caption_targets = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Figure Caption"
        and (
            match := re.match(r"圖\s*6-(\d+)", paragraph.text.strip())
        )
        and 37 <= int(match.group(1)) <= 65
    ]
    table_caption_targets = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Table Caption"
        and (
            match := re.match(r"表\s*6-(\d+)", paragraph.text.strip())
        )
        and 9 <= int(match.group(1)) <= 13
    ]

    checks = {}
    checks["heading3_paragraph_format"] = all(
        ppr_signature(target) == ppr_signature(h3_reference) for target in h3_targets
    )
    checks["heading3_run_format"] = all(
        rpr_signature(first_nonempty_run(target))
        == rpr_signature(first_nonempty_run(h3_reference))
        for target in h3_targets
    )
    checks["heading4_paragraph_format"] = all(
        ppr_signature(target) == ppr_signature(h4_reference) for target in h4_targets
    )
    checks["heading4_run_format"] = all(
        rpr_signature(first_nonempty_run(target))
        == rpr_signature(first_nonempty_run(h4_reference))
        for target in h4_targets
    )
    checks["figure_caption_paragraph_format"] = all(
        ppr_signature(target) == ppr_signature(figure_caption_reference)
        for target in figure_caption_targets
    )
    checks["figure_caption_run_format"] = all(
        rpr_signature(first_nonempty_run(target))
        == rpr_signature(first_nonempty_run(figure_caption_reference))
        for target in figure_caption_targets
    )
    checks["table_caption_paragraph_format"] = all(
        ppr_signature(target) == ppr_signature(table_caption_reference)
        for target in table_caption_targets
    )
    checks["table_caption_run_format"] = all(
        rpr_signature(first_nonempty_run(target))
        == rpr_signature(first_nonempty_run(table_caption_reference))
        for target in table_caption_targets
    )
    checks["table_caption_page_breaks_preserved"] = all(
        has_page_break_before(
            find_paragraph(document, f"表 6-{number}", "Table Caption")
        )
        for number in (10, 12)
    )

    start = paragraph_index(document, find_paragraph(document, "6.2.5 位移改善實驗"))
    end = paragraph_index(document, find_paragraph(document, "結論與未來工作", "Heading 1"))
    chapter6_new_body = [
        paragraph
        for paragraph in document.paragraphs[start:end]
        if paragraph.style.name == "Normal" and paragraph.text.strip()
    ]
    checks["chapter6_body_paragraph_format"] = all(
        ppr_signature(paragraph) == ppr_signature(body_reference)
        for paragraph in chapter6_new_body
    )
    checks["chapter6_body_run_format"] = all(
        rpr_signature(first_nonempty_run(paragraph))
        == rpr_signature(first_nonempty_run(body_reference))
        for paragraph in chapter6_new_body
    )

    new_tables = document.tables[-5:]
    reference_table = document.tables[-6]
    checks["new_table_header_fill"] = all(
        (shading := table.cell(0, 0)._tc.tcPr.find(qn("w:shd"))) is not None
        and shading.get(qn("w:fill")) == "E7E6E6"
        for table in new_tables
    )
    checks["new_table_header_run_format"] = all(
        rpr_signature(first_nonempty_run(table.cell(0, 0).paragraphs[0]))
        == rpr_signature(first_nonempty_run(reference_table.cell(0, 0).paragraphs[0]))
        for table in new_tables
    )
    checks["new_table_body_run_format"] = all(
        rpr_signature(first_nonempty_run(table.cell(1, 0).paragraphs[0]))
        == rpr_signature(first_nonempty_run(reference_table.cell(1, 0).paragraphs[0]))
        for table in new_tables
    )

    restored = {}
    for figure_number in (24, 25, 26, 28, 29, 30):
        caption = find_paragraph(document, f"圖 6-{figure_number}", "Figure Caption")
        embedded_hash = sha256(drawing_blob_before_caption(document, caption))
        source_file = next(CROPPED.glob(f"figure_6_{figure_number:02d}_*.png"))
        source_hash = sha256(source_file.read_bytes())
        restored[str(figure_number)] = {
            "embedded_sha256": embedded_hash,
            "source_sha256": source_hash,
            "matches": embedded_hash == source_hash,
        }
    checks["six_scaled_images_embedded"] = all(
        item["matches"] for item in restored.values()
    )

    report = {
        "docx": str(DOCX),
        "counts": {
            "heading3_targets": len(h3_targets),
            "heading4_targets": len(h4_targets),
            "chapter6_body_targets": len(chapter6_new_body),
            "figure_caption_targets": len(figure_caption_targets),
            "table_caption_targets": len(table_caption_targets),
            "new_tables": len(new_tables),
            "restored_images": len(restored),
        },
        "checks": checks,
        "restored_images": restored,
        "all_checks_passed": all(checks.values()),
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if not report["all_checks_passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
