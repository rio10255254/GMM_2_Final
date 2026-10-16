from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用"
    r"\論文完整版_20260716_v9.docx"
)
WORK_DIR = Path(
    r"D:\GMM_2_Final\Experiments\R63_CausalBoundaryMotionAblation"
    r"\results\thesis_v10_citation_logic"
)
WORKING = WORK_DIR / "thesis_v10_citation_working.docx"
FORMAL = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用"
    r"\論文完整版_20260716_v10.docx"
)
AUDIT = WORK_DIR / "thesis_v10_citation_audit.json"
CHANGELOG = WORK_DIR / "thesis_v10_citation_changes.json"


REPLACEMENTS: dict[str, str] = {
    "隨著半導體製程的特徵尺寸持續縮小": (
        "隨著半導體製程的特徵尺寸持續縮小，關鍵尺寸、結構位置與製程穩定度的量測要求也隨之提高。"
        "散射場顯微等光學方法可在非接觸條件下取得大面積資料，並已用於半導體結構與關鍵尺寸量測[78,79]；"
        "深度學習也逐步被導入光學量測與CD重建[37,89]，並用於晶片顯微影像的SEM-like轉換[90]。"
        "因此，顯微成像仍是半導體與精密製造的重要量測方式。"
    ),
    "當待測結構逐漸接近光學系統的解析極限時": (
        "當待測結構逐漸接近光學系統的解析極限時，解析度、景深與視野之間的取捨會更為明顯。"
        "有限孔徑與繞射共同限制成像解析能力[13,17,18,20]；提高物鏡數值孔徑雖可改善橫向解析能力，"
        "卻也會縮小景深並提高系統對失焦、機械位移、熱漂移與環境雜訊的敏感度。對量測工作而言，"
        "這些擾動最終會反映在邊界座標、結構尺寸與位移軌跡，而不只是影像外觀。"
    ),
    "近年來，深度學習已廣泛應用於影像重建與超解析": (
        "近年來，深度學習已廣泛應用於顯微影像重建，包括跨模態與電子顯微解析度提升[71,72]，"
        "以及單張復原、通用模型與像差補償[35,74,75,77]。這些方法可改善視覺清晰度，但像素或感知品質與幾何失真並非同一問題[30,31]；"
        "重建模型也可能產生不穩定或不存在的細節[63,64]。若影像用於CD或其他定量量測，空間與定量資訊是否保留仍須另行驗證[89,93]。"
        "本研究因此將評估重點由主觀清晰度轉向尺寸與位移是否能被穩定量測。"
    ),
    "在光學成像系統中，點擴散函數": (
        "在光學成像系統中，點擴散函數（Point Spread Function, PSF）描述系統對空間資訊的模糊與擴散[13,50]。"
        "影像反卷積與重建屬於逆問題，其穩定性取決於光學模型是否合理，以及雜訊與其他擾動是否受到控制[46,52]。"
        "若模型只以像素誤差或感知相似度為目標，即使影像數值看似改善，也可能在幾何定位上留下不易察覺的偏差[63]。"
    ),
    "本研究將影像重建視為量測流程的一部分": (
        "本研究將影像重建視為量測流程的一部分，並以重建後的定位精度與量測穩定性作為主要判斷依據。"
        "像素相似度或視覺品質不能單獨代表量測可靠性[30,31]，而重建不穩定或人工細節也可能改變幾何位置[63,64]。"
        "基於此，本研究提出量測導向超解析（Measurement-Oriented Super-Resolution, MOSR），不以PSNR、SSIM或主觀銳利度作為唯一目標，"
        "而以結構尺寸、邊界位置、位移一致性及其量測不確定度作為判斷依據[84,85]。"
    ),
    "本研究針對兩個容易影響量測的模型行為進行設計": (
        "本研究針對兩個容易影響量測的模型行為進行設計。首先，下採樣與上採樣可能使模型對取樣相位敏感[39,58,61,91]，"
        "邊界補值也可能使網路利用絕對位置[59]；因此，本研究在特徵放大時加入抗混疊處理。其次，網路即使輸出更銳利的影像，"
        "仍可能改變低頻輪廓、主要亮暗關係與中心位置，因此在輸出端保留原始影像的位置基準。訓練時再以影像、邊界、運動與輸出保守性四類損失限制模型行為；"
        "頻域與梯度約束的相關基礎見文獻[67,68,69,70]，平移一致性的成像重建作法則見文獻[62,92]。"
    ),
    "本研究的貢獻依模型設計、訓練目標與實驗證據分為四項": (
        "本研究的貢獻依模型設計、訓練目標與實驗證據分為四項。既有研究已將深度學習超解析用於CD量測[89]與晶片顯微影像轉換[90]，"
        "也已分別探討上採樣穩定性與影像重建的平移等變性[91,92]，以及定量資訊保存[93]。因此，本研究不把這些個別原理，"
        "或光學校正、資料建立、影像重建與後端量測的流程串接宣稱為首創，而將學術貢獻限定在針對CD與微位移量測所形成的模組設計、"
        "失效對應損失，以及可直接檢查其作用的配對量測證據。"
    ),
    "貢獻一：建立面向微位移量測的抗混疊上採樣模組": (
        "貢獻一：建立面向微位移量測的抗混疊上採樣模組。抗混疊與平移穩定網路已有明確基礎[39,58,61]，"
        "上採樣過程的頻譜穩定性及影像重建的平移等變性亦已有直接研究[60,91,92]。本研究的貢獻在於把這些原理具體轉化為顯微重建解碼器中的空間對應保護，"
        "並以同一影格九顆pillar的真實位移一致性作為直接驗證。配對消融顯示，相較於一般上採樣，該模組使四組真實位移序列合併後的九顆pillar不一致量降低約14.3%，"
        "其可被支持的主要作用為改善跨空間位置的相對幾何一致性。"
    ),
    "貢獻二：建立原始影像位置保留機制": (
        "貢獻二：建立原始影像位置保留機制。殘差學習與輸入跳接本身已有成熟研究基礎[42,43,44]；本研究的差異不在於首次使用殘差，"
        "而在於將輸出限制為以原始影像提供低頻位置與大尺度亮暗基準，再加入幅度受限的高頻邊界修正。移除該模組後，20倍CD中位絕對誤差與兩倍率振鈴明顯惡化，"
        "顯示其主要價值是避免邊界與亮暗拓樸脫離輸入。由於50倍CD的單一誤差方向不同，本研究不將其延伸為所有倍率均必然降低CD的通則。"
    ),
    "貢獻三：建立依量測失敗形式分組的功能性損失": (
        "貢獻三：建立依量測失敗形式分組的功能性損失。頻域損失與頻譜偏差已有相關研究[67,68]，梯度引導與任務特定損失亦有既有基礎[69,70]；"
        "平移等變約束則可用於成像重建[62]。本研究的貢獻是把影像內容、邊界幾何、運動一致性與輸出保守性四組目標，分別對應外觀失真、輪廓偏移、影格依賴與非必要修改等量測風險。"
        "消融結果未支持將完整改善獨立歸因於任一單項損失，因此論文只主張這套失效對應設計的整體功能，不誇大個別損失的因果效果。"
    ),
    "貢獻四：建立可區分整體效果與單項作用的配對量測證據": (
        "貢獻四：建立可區分整體效果與單項作用的配對量測證據。原始影像、Wiener反卷積與MOSR共用影像名單、pillar身分、中心、物體空間尺度、門檻與統計規則，"
        "並以亞像素配準[80,81]、邊界定位[82,83]及量測不確定度原則[84,85]檢查逐顆SEM尺寸與真實位移序列。另以兩個固定訓練種子、匹配訓練預算與配對重抽樣[87]比較模組及損失移除，"
        "使『完整方法有效』與『哪一個設計負責何種改善』不再混為同一結論。"
    ),
    "從繞射光學觀點出發": (
        "從繞射光學觀點出發，成像系統的解析能力根本上受限於有限孔徑所造成的繞射效應[13,29]。當點光源經由圓形孔徑成像時，"
        "其在像平面上形成具主峰與旁瓣的擴散圖樣；Abbe與Rayleigh的研究進一步建立了解析能力與可分辨條件的基礎[17,18,20]。"
        "此一概念構成後續PSF建模與解析度分析的基礎。"
    ),
    "系統的光學響應與成像特性可由點擴散函數": (
        "系統的光學響應與成像特性可由點擴散函數（Point Spread Function, PSF）表徵。實際PSF會受到像差、離焦、光瞳與視野位置影響[51,54,55,56]，"
        "因此理論模型或簡化高斯核不一定能完整代表真實系統。若重建所用PSF與實際模糊不一致，後續反卷積與超解析便可能產生偽影或量測偏差[50,53]。"
        "因此，許多研究會在實驗條件下校正PSF，再將其用於重建或分析。"
    ),
    "維納反卷積可視為「頻率域濾波器」": (
        "維納反卷積可視為頻率域中的受控反濾波器；當PSF已知且需要快速處理大量影像時，具有計算效率與雜訊控制上的優勢[46]。"
        "Richardson–Lucy方法則以迭代方式逐步修正影像，並可加入正則化以抑制雜訊放大[47,48,49]。兩類傳統反卷積都對PSF準確性與退化假設敏感，"
        "當實際模糊具有空間變異或核失配時，復原品質容易不穩定[52,53,55,56]。因此，後續研究逐漸轉向以學習式模型補充或取代固定反卷積流程。"
    ),
    "相較於傳統反卷積依賴明確的成像模型": (
        "相較於傳統反卷積依賴明確的成像模型與PSF／雜訊假設，卷積神經網路可由資料學習低品質影像至高品質影像的非線性映射[40,43,44]。"
        "此類方法已用於內容感知、跨模態及電子顯微影像復原[35,71,72]，也延伸至單張顯微重建、物理資訊擴散與像差補償[74,76,77]；"
        "近期亦開始重視跨條件的通用復原能力[75]。其優點是推論快速且可涵蓋複雜退化，限制則是結果高度依賴訓練資料與實際成像分布是否一致。"
        "本節依代表性模型演進回顧CNN超解析與重建的設計思路。"
    ),
    "整體而言，CNN超解析由早期淺層映射": (
        "整體而言，CNN超解析由早期淺層映射逐步發展至殘差學習、多尺度特徵融合與注意力機制[40,43,44,45]。"
        "顯微領域則已涵蓋內容感知復原、跨模態、電子顯微及定位式超解析[35,71,72,73,74]，並進一步發展通用模型、物理資訊重建與像差補償[75,76,77]。"
        "這類方法具有快速推論與非線性表達能力，但效果高度依賴訓練資料與退化模型。為提高真實成像條件下的可靠度，本研究進一步將PSF、頻域與幾何限制納入模型設計及驗證。"
    ),
    "雖然卷積神經網路已能在視覺品質上顯著改善超解析影像": (
        "雖然卷積神經網路已能在視覺品質上顯著改善超解析影像，多數評估仍以像素誤差、峰值訊噪比或結構相似度為主；"
        "但像素失真與感知品質之間存在取捨[30,31]。顯微超解析的公平比較還必須檢查定位精度與分析軟體的行為[32,33]，"
        "而深度重建本身也可能出現不穩定或人工結構[63,64]。因此，若輸出將用於量測，評估應進一步延伸至下游任務與量測誤差[70,89]。"
    ),
    "對結構化幾何目標與精密位移量測而言": (
        "對結構化幾何目標與精密位移量測而言，關鍵問題不是影像內容的分類，而是邊界能否在亞像素尺度被穩定定位[82,83]，"
        "以及同一結構的微小位移能否被一致追蹤[80,81]。PSNR或SSIM無法直接區分『銳利但位置不穩定』與『視覺改善較小但量測更一致』的重建結果。"
        "因而，本研究將影像重建視為量測鏈中的中間步驟，並把尺寸誤差、位移重複性及多位置一致性作為核心判斷依據。"
    ),
    "本研究提出的MOSR即針對此一應用情境": (
        "本研究提出的MOSR即針對此一應用情境，將邊緣過渡、跨門檻CD、位移重複性及多pillar空間一致性納入訓練設計與驗證流程，"
        "使影像重建由『看得清楚』轉向『量得穩定』。後端量測以亞像素配準[80,81]與邊界定位方法[82,83]取得物體空間幾何量，"
        "再依量測不確定度原則解讀結果[84,85]。"
    ),
    "以超解析影像支援尺寸量測並非本研究首次提出": (
        "以超解析影像支援尺寸量測並非本研究首次提出。Kim與Pahk已將低解析光學影像轉換為較高解析條件，並以有機層的解析度與CD量測評估結果；"
        "近期晶片顯微研究也將光學超解析影像轉換為近似SEM的影像[89,90]。顯微重建領域另已有跨模態與電子顯微解析度提升[71,72]，"
        "以及單張超解析、通用顯微復原與像差補償等方法[74,75,77]。這些研究說明，以深度學習改善顯微影像並服務後續定量分析已有前例。"
    ),
    "為使深度學習模型在有限資料下保持合理的物理行為": (
        "為使深度學習模型在有限資料下保持合理的物理行為，既有方法會以方程殘差、成像正向模型或資料一致性限制解空間[34,65,66]。"
        "在顯微影像中，內容感知復原、定位式超解析與物理資訊擴散模型也分別利用資料與成像條件建立重建約束[35,36,76]；"
        "等變成像則進一步以已知轉換關係提供學習訊號[62]。本研究沿用此一思路，將重建影像應滿足的幾何與頻域條件轉寫為可計算的限制。"
    ),
    "其中一項重要條件為平移等變性": (
        "其中一項重要條件為平移等變性。理想情況下，輸入影像平移後，模型輸出也應產生相同幅度的平移[38,62]。"
        "下採樣與非線性可能使模型對取樣相位敏感[39,58,61]，邊界補值也可能洩漏絕對位置[59]；上採樣則可能引入額外的頻譜與空間不穩定[60,91]。"
        "影像重建中的平移等變性因此需要在架構與訓練兩端共同檢查[92]。本研究利用已知位移影像對，在訓練與驗證階段直接檢查跨影格幾何關係。"
    ),
    "除平移關係外，本研究亦限制高頻能量與邊界微分行為": (
        "除平移關係外，本研究亦限制高頻能量與邊界微分行為。神經網路具有頻譜偏差，而頻域損失可直接調整不同頻率的學習比重[67,68]；"
        "梯度引導與任務特定損失則可讓訓練目標更接近結構與下游需求[69,70]。這些限制用來降低過衝、振鈴及不合理邊界變化，"
        "但不保證輸出必然等同於真實物體；深度重建仍須防範不穩定與人工細節[63,64]。"
    ),
    "綜合上述代表性研究，本研究不將抗混疊": (
        "綜合上述代表性研究，本研究不將抗混疊、殘差連接、頻域損失或平移等變性本身列為首創。殘差式重建已有成熟基礎[42,43]；"
        "抗混疊與穩定上採樣亦已有相關方法[58,61,91]；任務特定損失及重建等變性也已有前例[70,92]。本研究的差異在於，"
        "以單張同尺寸顯微重建為部署形式，將上採樣的空間對應保護與輸入參考的受限輸出修正，直接對應到CD與微位移兩類量測風險，"
        "再以亞像素位移、既有CD研究及定量資訊保存原則界定可被支持的作用[80,89,93]。因此，論文的創新主張是量測問題下的具體設計與證據，"
        "而不是宣稱個別基礎元件前所未見。"
    ),
    "下採樣層（down-sampling layer）": (
        "下採樣層（down-sampling layer）的目的在於降低特徵圖尺寸，以減少後續運算量並擴張有效感受野。"
        "常見操作包含最大池化與平均池化，其方法與特性已有專門回顧[24]；以Fast R-CNN為例，RoI pooling會將不同大小的候選區域轉為固定尺寸特徵[23]。池化可保留局部代表值並抑制部分小尺度擾動，"
        "但降解析度同時會改變取樣相位，因此其位置穩定性仍須另外檢查。"
    ),
    "影像重建、分割與定位不只需要辨識影像內容": (
        "影像重建、分割與定位不只需要辨識影像內容，也必須保留物體的空間位置。因此，輸入影像整體移動時，模型輸出也應朝相同方向移動相同距離。"
        "這項要求可由群組等變網路與等變成像的理論描述[38,62]，在影像重建任務中則有直接的平移等變設計[92]。"
    ),
    "這種輸入移動後，輸出隨之產生相同位移的性質稱為平移等變性": (
        "這種輸入移動後，輸出隨之產生相同位移的性質稱為平移等變性[38,62]。對像素級重建與定位而言，輸出本身包含空間座標，因此這項性質特別重要。"
    ),
    "下採樣前加入低通濾波": (
        "下採樣前加入低通濾波，例如抗混疊池化（blur pooling），可降低模型對取樣相位的敏感性[39]。但這類方法不代表所有整數與亞像素位移下都能嚴格等變[58,61]。"
    ),
    "平移一致性可同時由架構與訓練改善": (
        "平移一致性可同時由架構與訓練改善。架構上可採用抗混疊取樣與對位移較穩定的網路設計[39,58,61]，並另外處理上採樣穩定性[91]與邊界補值造成的位置偏差[59]；"
        "訓練時則可要求平移前後的輸出維持一致[62,92]。"
    ),
    "實際卷積神經網路通常只能近似滿足平移等變性": (
        "實際卷積神經網路通常只能近似滿足平移等變性，因此需要以數值指標比較兩條處理路徑的差異[38,62,92]。"
    ),
    "本研究希望重建後的顯微影像仍能可靠地用於量測": (
        "本研究希望重建後的顯微影像仍能可靠地用於量測，而不只是在視覺上更銳利。像素或感知品質不能單獨保證幾何正確[30,31]，"
        "深度重建也可能出現不穩定或人工細節[63,64]。因此，評估內容包含結構邊界能否穩定定位、尺寸是否接近參考量測，以及連續影像中的微小位移能否被一致追蹤；"
        "此一方向與CD量測及定量資訊保存的需求一致[85,89,93]。"
    ),
    "多位置PSF不要求整個視野共用相同的成像核": (
        "多位置PSF不要求整個視野共用相同的成像核，而是以數個代表性光學響應描述空間變異模糊[50,54,55,56]。"
        "模型再以依輸入調整的動態融合方式組合這些候選響應[57]，降低只根據影像紋理任意銳化的風險。其融合方式於第4.2.4節說明，實際網路配置則列於第五章。"
    ),
    "本研究的模組貢獻不在於首次提出抗混疊": (
        "本研究的模組貢獻不在於首次提出抗混疊、殘差連接或平移等變原理。抗混疊與位移穩定網路已有相關研究[39,58,61]，"
        "殘差式重建已有成熟基礎[42,43]，上採樣穩定性與重建等變性也已有直接方法[91,92]。本研究則針對兩個可被量測的失敗來源建立對應設計："
        "抗混疊上採樣用於降低解碼器對像素相位與空間位置的敏感性；原始影像位置保留則限制輸出邊界脫離輸入的風險。"
        "其他模組負責提供光學資訊與限制修正幅度，其獨立作用由第六章消融結果另行界定。"
    ),
    "一般上採樣可能在特徵放大時引入取樣相位差異": (
        "解碼器中的取樣與非線性可能使特徵對像素相位敏感[39,58,61]，一般上採樣又可能引入額外的頻譜偽影[60,91]；邊界補值亦可能使不同位置得到不同響應[59]。"
        "這類誤差會直接影響亞像素邊界與位移量測。本研究的抗混疊上採樣依序進行雙線性插值、反射填補、逐通道低通濾波與卷積，使尺寸恢復與特徵整合分開處理。"
        "模型另保留PSF殘差修正，使影像重建能參考固定的光學先驗。各層通道數、卷積核與連接位置列於第5.4.2節。"
    ),
    "本研究不直接將邊界修正結果作為最終輸出": (
        "殘差式重建通常只學習輸入與目標之間的差異[42,43,44]，而定量影像處理還必須避免破壞原有空間與強度資訊[93]。"
        "基於此，本研究不直接將邊界修正結果作為最終輸出，而以原始影像作為低頻輪廓、主要亮暗與中心位置的基準。模型預測只提供局部高頻細節，再以受限比例加入原始影像。"
    ),
    "此輸出方式不讓網路的大尺度低頻預測完整取代輸入": (
        "此輸出方式不讓網路的大尺度低頻預測完整取代輸入。即使主幹網路對背景亮度或輪廓中心產生偏移，最終影像仍以原始影像為位置基準，使模型的主要修正集中在pillar邊界附近。"
    ),
    "本研究的訓練目標不是單純降低輸出影像": (
        "本研究的訓練目標不是單純降低輸出影像與參考影像之間的像素差異，因為像素失真與感知品質並非完全一致[30,31]。"
        "頻域損失可補充不同空間頻率的學習[67,68]，梯度與任務特定損失可強化結構及下游需求[69,70]，等變約束則用於維持轉換前後的幾何關係[62,92]。"
        "本研究據此要求模型在改善邊界可辨識度的同時，維持微結構原有的尺寸、中心位置及相鄰影格間的位移關係。"
    ),
    "影像重建模型可能對結構所在的像素位置敏感": (
        "影像重建模型可能對結構所在的像素位置敏感。下採樣與非線性會造成取樣相位依賴[39,58,61]，上採樣也可能引入額外的位置差異[91]，邊界補值則可能使模型利用絕對位置[59]。"
        "因此，即使輸入內容相同，整張影像向左或向右移動少量像素後，輸出輪廓仍可能產生不同變化。這種位置依賴性對一般辨識任務未必明顯，但會直接影響亞像素位移量測。"
    ),
    "每顆 pillar 以固定中心與參考半徑建立搜尋區域": (
        "每顆pillar以固定中心與參考半徑建立搜尋區域，從中心附近多條對稱剖面尋找左右或上下邊界。邊界峰值以局部內插取得亞像素位置[82,83]，"
        "再以相對邊界的中點形成各剖面的中心估計；此結果可進一步用於影格間的亞像素位移追蹤[80,81]。"
    ),
    "同一次平台移動理論上會使九宮格內九顆 pillar": (
        "同一次平台移動理論上會使九宮格內九顆pillar朝相近方向移動。因此，每一步先以亞像素配準與位移估計取得九顆局部結果[80,81]，"
        "再以中位數形成代表位移，並以經尺度校正的中位數絕對偏差描述九宮格內不一致量[88]。"
    ),
    "CD 誤差同時以平均絕對誤差": (
        "CD誤差同時以平均絕對誤差、絕對誤差中位數、均方根誤差、平均帶號偏差與最大絕對誤差描述。"
        "多個指標並用，可避免正負誤差互相抵銷，並同時呈現典型誤差、系統偏差與極端值；此解讀遵循量測不確定度與方法一致性評估的基本原則[84,85,86]。"
    ),
    "位移結果則報告代表位移相對名義指令的差距": (
        "位移結果則報告代表位移相對名義指令的差距、帶號位移序列的步間波動、九宮格內不一致量與跨軸位移；亞像素位移的估計基礎見文獻[80,81]。"
        "相同影格對與相同彙整規則適用於所有方法，並依共同量測與不確定度原則解讀[84,85]，避免因不同的資料篩選造成表面改善。"
    ),
    "本章建立以物理量測為核心的單張顯微影像重建方法": (
        "本章建立以物理量測為核心的單張顯微影像重建方法，並把研究主張分層。光學校正、多來源資料與共同後端尺規提供可追溯的研究架構；"
        "抗混疊上採樣與原始影像位置保留為本研究主要檢驗的模型設計，其貢獻在於面向CD與微位移的具體實作及配對證據；"
        "功能性損失則把量測失敗形式轉為可訓練的目標。"
    ),
    "每個條件與隨機初始化均完成相同的靜態CD": (
        "每個條件與隨機初始化均完成相同的靜態CD、真實位移與合成位移評估，共產生18,272張候選重建影像。"
        "結果同時檢查兩組初始化的方向、相對差異是否達10%、95%配對區間是否高於零，以及主要限制指標是否惡化超過15%。"
        "配對區間以4,000次重抽樣估計[87]；靜態CD以pillar、真實位移以步驟、合成位移以案例為抽樣單位。"
    ),
    "目前局部PSF與多位置PSF提供了描述成像模糊的兩種尺度": (
        "目前局部PSF與多位置PSF提供了描述成像模糊的兩種尺度；局部與空間變異模糊的相關基礎見文獻[50,54,55,56]。"
        "未來可在不同視野位置、焦平面、曝光與照明條件下建立更密集的校正資料，並以重複拍攝估計每一位置PSF參數的變異範圍；"
        "光瞳函數的相位校正亦可作為補充[51]。如此可區分真正的空間像差、短期對焦變動與擬合不確定度，避免將所有差異都歸因於單一模糊核。"
    ),
    "現有位移結果已以平台名義指令完成步間散佈": (
        "現有位移結果已以平台名義指令完成步間散佈與九顆pillar一致性比較，其影像位移估計採用亞像素定位原則[80,81]。"
        "未來加入雷射干涉儀、外部電容式感測器或具校正證書的編碼器，可同步取得每一步的獨立回授，進一步分離平台、影像定位與重建三類變動來源，"
        "並建立更完整的量測追溯鏈[84,85]。"
    ),
    "結合獨立回授後，可建立包含重複性": (
        "結合獨立回授後，可建立包含重複性、系統偏差、像素尺度、SEM參考與模型變動的量測不確定度預算[84,85]，"
        "並以方法一致性分析檢查系統偏差[86]。不同倍率與位置的區間可由重抽樣估計[87]，而重尾或離群條件下的散布可搭配穩健尺度統計[88]，"
        "使位移結果具有明確的追溯鏈與可信範圍。"
    ),
    "目前消融結果支持抗混疊上採樣": (
        "目前消融結果支持抗混疊上採樣對多顆pillar位移一致性的改善，以及原始影像位置保留對邊界的保護作用，但實驗只包含兩組隨機初始化，單一損失的效果也尚未明確。"
        "未來可增加獨立訓練次數，並在不同樣品與倍率下重複相同比較，以確認兩項設計是否維持一致效果。"
    ),
    "量測導向損失函數可先依影像內容": (
        "量測導向損失函數可先依影像內容、邊界幾何、運動一致性與輸出保守性四組進行比較，再於效果明確的組別中分析個別子項。"
        "這樣可降低多重比較造成的偶然結果，也有助於判斷單一損失是否被其他機制補償。"
    ),
    "後續模型選擇應同時考量邊緣寬度": (
        "後續模型選擇應同時考量邊緣寬度、跨門檻CD、振鈴、位移重複性與多顆pillar一致性，而不能只依訓練損失或主觀銳利度決定。"
        "系統也可依量測不確定度原則輸出可信範圍[85]，並檢查空間與定量資訊是否被保留[93]。可信度不足時，則保留Raw結果、降低修正幅度或重新取像。"
    ),
}


def body_before_references(document: Document) -> list[Paragraph]:
    paragraphs = document.paragraphs
    reference_index = next(
        i for i, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "參考文獻"
    )
    return paragraphs[:reference_index]


def find_one(document: Document, prefix: str) -> Paragraph:
    matches = [
        p for p in body_before_references(document) if p.text.strip().startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one body paragraph starting with {prefix!r}, found {len(matches)}"
        )
    return matches[0]


def replace_visible_text(paragraph: Paragraph, text: str) -> None:
    xml = paragraph._p.xml
    if "oMath" in xml or "w:drawing" in xml or "w:fldChar" in xml:
        raise RuntimeError(f"Unsafe whole-paragraph replacement: {paragraph.text[:80]!r}")
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = deepcopy(run._r.rPr)
            break
    element = paragraph._p
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_rpr)


def w_text_nodes(paragraph: Paragraph):
    return paragraph._p.xpath(".//w:t")


def replace_w_text_span(
    paragraph: Paragraph, start: int, end: int, replacement: str
) -> None:
    nodes = w_text_nodes(paragraph)
    texts = [node.text or "" for node in nodes]
    total = sum(len(text) for text in texts)
    if not (0 <= start < end <= total):
        raise ValueError((start, end, total, paragraph.text))
    cumulative = 0
    start_node = start_offset = end_node = end_offset = None
    for index, value in enumerate(texts):
        next_cumulative = cumulative + len(value)
        if start_node is None and start < next_cumulative:
            start_node = index
            start_offset = start - cumulative
        if end <= next_cumulative:
            end_node = index
            end_offset = end - cumulative
            break
        cumulative = next_cumulative
    if start_node is None or end_node is None:
        raise RuntimeError("Could not map visible text span to Word text nodes")
    if start_node == end_node:
        value = texts[start_node]
        nodes[start_node].text = (
            value[:start_offset] + replacement + value[end_offset:]
        )
        return
    nodes[start_node].text = texts[start_node][:start_offset] + replacement
    for index in range(start_node + 1, end_node):
        nodes[index].text = ""
    nodes[end_node].text = texts[end_node][end_offset:]


def visible_w_text(paragraph: Paragraph) -> str:
    return "".join(node.text or "" for node in w_text_nodes(paragraph))


def replace_fragment(paragraph: Paragraph, old: str, new: str) -> None:
    text = visible_w_text(paragraph)
    occurrences = text.count(old)
    if occurrences != 1:
        raise RuntimeError(
            f"Expected one occurrence of {old!r}, found {occurrences} in {text!r}"
        )
    start = text.index(old)
    replace_w_text_span(paragraph, start, start + len(old), new)


def citation_numbers(text: str) -> list[int]:
    values: list[int] = []
    for match in re.finditer(r"\[([0-9][^\]]*)\]", text):
        content = match.group(1)
        if re.fullmatch(r"\d+(?:,\d+)*", content):
            values.extend(int(value) for value in content.split(","))
    return values


def audit_document(path: Path) -> dict:
    document = Document(str(path))
    body = body_before_references(document)
    bibliography = [
        p.text.strip()
        for p in document.paragraphs
        if p.style is not None and p.style.name == "EndNote Bibliography"
    ]
    bibliography_numbers = [
        int(match.group(1))
        for text in bibliography
        if (match := re.match(r"\[(\d+)\]", text))
    ]
    malformed = []
    groups: list[list[int]] = []
    duplicate_groups = []
    citation_paragraphs = 0
    for index, paragraph in enumerate(body):
        matches = list(re.finditer(r"\[([0-9][^\]]*)\]", paragraph.text))
        if matches:
            citation_paragraphs += 1
        seen: set[tuple[int, ...]] = set()
        for match in matches:
            content = match.group(1)
            if not re.fullmatch(r"\d+(?:,\d+)*", content):
                malformed.append(
                    {
                        "paragraph": index,
                        "citation": match.group(0),
                        "text": paragraph.text,
                    }
                )
                continue
            group = [int(value) for value in content.split(",")]
            groups.append(group)
            key = tuple(group)
            if key in seen:
                duplicate_groups.append(
                    {"paragraph": index, "citation": match.group(0), "text": paragraph.text}
                )
            seen.add(key)
    flattened = [value for group in groups for value in group]
    expected = set(range(1, 94))
    used = set(flattened)
    distribution = Counter(len(group) for group in groups)
    return {
        "document": str(path),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "bibliography_count": len(bibliography),
        "bibliography_numbers_ok": bibliography_numbers == list(range(1, 94)),
        "citation_paragraphs": citation_paragraphs,
        "citation_groups": len(groups),
        "citation_occurrences": len(flattened),
        "group_size_distribution": dict(sorted(distribution.items())),
        "maximum_group_size": max((len(group) for group in groups), default=0),
        "groups_over_five": sum(len(group) > 5 for group in groups),
        "unused_references": sorted(expected - used),
        "out_of_range_citations": sorted(used - expected),
        "malformed_citations": malformed,
        "duplicate_groups_within_paragraph": duplicate_groups,
        "citation_spacing_violations": sum(
            len(re.findall(r"\[\d+,\s+\d+", paragraph.text)) for paragraph in body
        ),
        "separated_citation_groups": sum(
            len(re.findall(r"\[\d+\]\s*,\s*\[\d+\]", paragraph.text))
            for paragraph in body
        ),
        "per_reference_occurrences": {
            str(number): flattened.count(number) for number in range(1, 94)
        },
    }


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    document = Document(str(SOURCE))
    changes = []
    for prefix, new_text in REPLACEMENTS.items():
        paragraph = find_one(document, prefix)
        old_text = paragraph.text
        replace_visible_text(paragraph, new_text)
        changes.append({"prefix": prefix, "before": old_text, "after": new_text})

    # These paragraphs contain inline Word equations. Only their visible text is
    # adjusted so the equations and their formatting remain intact.
    equation_safe_edits = [
        (
            "因此，一個模型即使對整數像素平移表現良好",
            "[39,58,61]",
            "[58,61]",
        ),
        (
            "現代 CNN 通常不只包含步幅為",
            "為了擴大感受野或恢復影像解析度，網路經常使用下採樣與上採樣；為了維持影像大小，卷積層也經常在邊界補值。這些操作都可能破壞原本的平移關係。[39,58,59,60,61,91,92]",
            "下採樣可能改變取樣相位與高頻資訊[39,58,61]；上採樣也可能引入新的頻譜與空間不穩定[60,91]。此外，為維持影像大小而使用的邊界補值，可能讓模型利用絕對位置[59]。因此，三類操作都需要分別檢查。",
        ),
        (
            "步幅為 的卷積、最大池化",
            "[39,58,60]",
            "[39,61]",
        ),
        (
            "為限制此問題，本研究加入平移一致性損失",
            "[38,39,58,61,62,92]",
            "[38,62,92]",
        ),
    ]
    for prefix, old, new in equation_safe_edits:
        paragraph = find_one(document, prefix)
        before = paragraph.text
        replace_fragment(paragraph, old, new)
        changes.append({"prefix": prefix, "before": before, "after": paragraph.text})

    simple_fragment_edits = [
        (
            "一般而言，成像過程可近似為物體",
            "[46,47,48,50,52]",
            "[46,50]",
        ),
        (
            "轉置卷積則使用可學習的卷積核進行上採樣",
            "[27,91]",
            "[27,60,91]",
        ),
        (
            "當物體位於影像中央時，它周圍主要是真實像素",
            "模型看到的局部內容也已經不同，因此輸出不再只是簡單平移。",
            "模型看到的局部內容也已經不同，因此輸出不再只是簡單平移[59]。",
        ),
    ]
    for prefix, old, new in simple_fragment_edits:
        paragraph = find_one(document, prefix)
        before = paragraph.text
        replace_fragment(paragraph, old, new)
        changes.append({"prefix": prefix, "before": before, "after": paragraph.text})

    document.save(str(WORKING))
    audit = audit_document(WORKING)
    failures = []
    if audit["bibliography_count"] != 93 or not audit["bibliography_numbers_ok"]:
        failures.append("bibliography numbering or count")
    if audit["unused_references"]:
        failures.append(f"unused references: {audit['unused_references']}")
    if audit["out_of_range_citations"]:
        failures.append(f"out-of-range citations: {audit['out_of_range_citations']}")
    if audit["malformed_citations"]:
        failures.append(f"malformed citations: {len(audit['malformed_citations'])}")
    if audit["duplicate_groups_within_paragraph"]:
        failures.append(
            "duplicate groups: "
            f"{len(audit['duplicate_groups_within_paragraph'])}"
        )
    if audit["citation_spacing_violations"]:
        failures.append(
            f"citation spacing: {audit['citation_spacing_violations']}"
        )
    if audit["separated_citation_groups"]:
        failures.append(
            f"separated groups: {audit['separated_citation_groups']}"
        )
    if audit["groups_over_five"]:
        failures.append(f"groups over five: {audit['groups_over_five']}")

    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    CHANGELOG.write_text(
        json.dumps(changes, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if failures:
        raise RuntimeError("; ".join(failures))

    shutil.copy2(WORKING, FORMAL)
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    print(f"Wrote working copy: {WORKING}")
    print(f"Wrote formal copy: {FORMAL}")


if __name__ == "__main__":
    main()
