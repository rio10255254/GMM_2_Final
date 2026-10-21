from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCUMENT = Path(
    r"D:\GMM_2_Final\Experiments\R63_CausalBoundaryMotionAblation"
    r"\results\thesis_v9_citations\thesis_v9_citation_working.docx"
)
AUDIT_JSON = DOCUMENT.with_name("thesis_v9_citation_audit.json")


NEW_REFERENCES = [
    '[46] N. Wiener, Extrapolation, Interpolation, and Smoothing of Stationary Time Series. Cambridge, MA, USA: MIT Press, 1949.',
    '[47] W. H. Richardson, "Bayesian-based iterative method of image restoration," Journal of the Optical Society of America, vol. 62, no. 1, pp. 55–59, 1972.',
    '[48] L. B. Lucy, "An iterative technique for the rectification of observed distributions," The Astronomical Journal, vol. 79, pp. 745–754, 1974.',
    '[49] N. Dey, L. Blanc-Feraud, C. Zimmer, P. Roux, Z. Kam, J.-C. Olivo-Marin, and J. Zerubia, "Richardson–Lucy algorithm with total variation regularization for 3D confocal microscope deconvolution," Microscopy Research and Technique, vol. 69, no. 4, pp. 260–266, 2006.',
    '[50] J. G. Nagy and D. P. O’Leary, "Restoring images degraded by spatially variant blur," SIAM Journal on Scientific Computing, vol. 19, no. 4, pp. 1063–1082, 1998.',
    '[51] B. M. Hanser, M. G. L. Gustafsson, D. A. Agard, and J. W. Sedat, "Phase-retrieved pupil functions in wide-field fluorescence microscopy," Journal of Microscopy, vol. 216, no. 1, pp. 32–48, 2004.',
    '[52] A. Levin, Y. Weiss, F. Durand, and W. T. Freeman, "Understanding and evaluating blind deconvolution algorithms," in Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, pp. 1964–1971, 2009.',
    '[53] N. Efrat, D. Glasner, A. Apartsin, B. Nadler, and A. Levin, "Accurate blur models vs. image priors in single image super-resolution," in Proceedings of the IEEE International Conference on Computer Vision, pp. 2832–2839, 2013.',
    '[54] A. Shajkofci and M. Liebling, "Semi-blind spatially-variant deconvolution in optical microscopy with local point spread function estimation by use of convolutional neural networks," in Proceedings of the IEEE International Conference on Image Processing, pp. 3818–3822, 2018.',
    '[55] A. Shajkofci and M. Liebling, "Spatially-variant CNN-based point spread function estimation for blind deconvolution and depth estimation in optical microscopy," IEEE Transactions on Image Processing, vol. 29, pp. 5848–5861, 2020.',
    '[56] B. Toader, J. Boulanger, Y. Korolev, M. O. Lenz, J. Manton, C.-B. Schönlieb, and L. Mureşan, "Image reconstruction in light-sheet microscopy: spatially varying deconvolution and mixed noise," Journal of Mathematical Imaging and Vision, vol. 64, pp. 968–992, 2022.',
    '[57] X. Jia, B. De Brabandere, T. Tuytelaars, and L. Van Gool, "Dynamic filter networks," in Advances in Neural Information Processing Systems, vol. 29, 2016.',
    '[58] A. Chaman and I. Dokmanic, "Truly shift-invariant convolutional neural networks," in Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition, pp. 3773–3783, 2021.',
    '[59] O. S. Kayhan and J. C. van Gemert, "On translation invariance in CNNs: convolutional layers can exploit absolute spatial location," in Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition, pp. 14274–14285, 2020.',
    '[60] T. Karras, M. Aittala, S. Laine, E. Härkönen, J. Hellsten, J. Lehtinen, and T. Aila, "Alias-free generative adversarial networks," in Advances in Neural Information Processing Systems, vol. 34, 2021.',
    '[61] H. Michaeli, T. Michaeli, and D. Soudry, "Alias-free convnets: fractional shift invariance via polynomial activations," in Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition, pp. 16333–16342, 2023.',
    '[62] D. Chen, J. Tachella, and M. E. Davies, "Equivariant imaging: learning beyond the range space," in Proceedings of the IEEE/CVF International Conference on Computer Vision, pp. 4379–4388, 2021.',
    '[63] V. Antun, F. Renna, C. Poon, B. Adcock, and A. C. Hansen, "On instabilities of deep learning in image reconstruction and the potential costs of AI," Proceedings of the National Academy of Sciences, vol. 117, no. 48, pp. 30088–30095, 2020.',
    '[64] S. Bhadra, V. A. Kelkar, F. J. Brooks, and M. A. Anastasio, "On hallucinations in tomographic image reconstruction," IEEE Transactions on Medical Imaging, vol. 40, no. 11, pp. 3249–3260, 2021.',
    '[65] H. K. Aggarwal, M. P. Mani, and M. Jacob, "MoDL: model-based deep learning architecture for inverse problems," IEEE Transactions on Medical Imaging, vol. 38, no. 2, pp. 394–405, 2019.',
    '[66] H. Gupta, K. H. Jin, H. Q. Nguyen, M. T. McCann, and M. Unser, "CNN-based projected gradient descent for consistent CT image reconstruction," IEEE Transactions on Medical Imaging, vol. 37, no. 6, pp. 1440–1453, 2018.',
    '[67] L. Jiang, B. Dai, W. Wu, and C. C. Loy, "Focal frequency loss for image reconstruction and synthesis," in Proceedings of the IEEE/CVF International Conference on Computer Vision, pp. 13919–13929, 2021.',
    '[68] N. Rahaman et al., "On the spectral bias of neural networks," in Proceedings of the International Conference on Machine Learning, pp. 5301–5310, 2019.',
    '[69] C. Ma, Y. Rao, Y. Cheng, C. Chen, J. Lu, and J. Zhou, "Structure-preserving super resolution with gradient guidance," in Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition, pp. 7769–7778, 2020.',
    '[70] A. Mustafa, A. Mikhailiuk, D. A. Iliescu, V. Babbar, and R. K. Mantiuk, "Training a task-specific image reconstruction loss," in Proceedings of the IEEE/CVF Winter Conference on Applications of Computer Vision, pp. 2319–2328, 2022.',
    '[71] H. Wang, Y. Rivenson, Y. Jin, Z. Wei, R. Gao, H. Günaydın, L. A. Bentolila, C. Kural, and A. Ozcan, "Deep learning enables cross-modality super-resolution in fluorescence microscopy," Nature Methods, vol. 16, no. 1, pp. 103–110, 2019.',
    '[72] K. de Haan, Y. Rivenson, Y. Wu, and A. Ozcan, "Resolution enhancement in scanning electron microscopy using deep learning," Scientific Reports, vol. 9, art. no. 12050, 2019.',
    '[73] W. Ouyang, A. Aristov, M. Lelek, X. Hao, and C. Zimmer, "Deep learning massively accelerates super-resolution localization microscopy," Nature Biotechnology, vol. 36, no. 5, pp. 460–468, 2018.',
    '[74] R. Chen et al., "Single-frame deep-learning super-resolution microscopy for intracellular dynamics imaging," Nature Communications, vol. 14, art. no. 2854, 2023.',
    '[75] C. Ma, W. Tan, R. He, and B. Yan, "Pretraining a foundation model for generalizable fluorescence microscopy-based image restoration," Nature Methods, vol. 21, no. 8, pp. 1558–1567, 2024.',
    '[76] R. Li et al., "Microscopy image reconstruction with physics-informed denoising diffusion probabilistic model," Communications Engineering, vol. 3, art. no. 186, 2024.',
    '[77] M. Guo et al., "Deep learning-based aberration compensation improves contrast and resolution in fluorescence microscopy," Nature Communications, vol. 16, art. no. 313, 2025.',
    '[78] R. M. Silver, B. M. Barnes, R. Attota, J. Jun, M. Stocker, E. Marx, and H. J. Patrick, "Scatterfield microscopy for extending the limits of image-based optical metrology," Applied Optics, vol. 46, no. 20, pp. 4248–4257, 2007.',
    '[79] B. M. Barnes, R. Attota, R. Quintanilha, Y.-J. Sohn, and R. M. Silver, "Characterizing a scatterfield optical platform for semiconductor metrology," Measurement Science and Technology, vol. 22, no. 2, art. no. 024003, 2011.',
    '[80] M. Guizar-Sicairos, S. T. Thurman, and J. R. Fienup, "Efficient subpixel image registration algorithms," Optics Letters, vol. 33, no. 2, pp. 156–158, 2008.',
    '[81] B. Pan, Z. Wang, and H. Xie, "Generalized spatial-gradient-based digital image correlation for displacement and shape measurement with subpixel accuracy," The Journal of Strain Analysis for Engineering Design, vol. 44, no. 8, pp. 659–669, 2009.',
    '[82] A. Trujillo-Pino, K. Krissian, M. Alemán-Flores, and D. Santana-Cedrés, "Accurate subpixel edge location based on partial area effect," Image and Vision Computing, vol. 31, no. 1, pp. 72–90, 2013.',
    '[83] P. Chen, F. Chen, Y. Han, and Z. Zhang, "Sub-pixel dimensional measurement with logistic edge model," Optik, vol. 125, no. 9, pp. 2076–2080, 2014.',
    '[84] N. F. Zhang, R. M. Silver, H. Zhou, and B. M. Barnes, "Improving optical measurement uncertainty with combined multitool metrology using a Bayesian approach," Applied Optics, vol. 51, no. 25, pp. 6196–6206, 2012.',
    '[85] Joint Committee for Guides in Metrology, Evaluation of Measurement Data—Guide to the Expression of Uncertainty in Measurement, JCGM 100:2008(E), 2008.',
    '[86] J. M. Bland and D. G. Altman, "Statistical methods for assessing agreement between two methods of clinical measurement," The Lancet, vol. 1, no. 8476, pp. 307–310, 1986.',
    '[87] B. Efron, "Bootstrap methods: another look at the jackknife," The Annals of Statistics, vol. 7, no. 1, pp. 1–26, 1979.',
    '[88] P. J. Rousseeuw and C. Croux, "Alternatives to the median absolute deviation," Journal of the American Statistical Association, vol. 88, no. 424, pp. 1273–1283, 1993.',
    '[89] S. Kim and H. J. Pahk, "A method for improving resolution and critical dimension measurement of an organic layer using deep learning superresolution," Current Optics and Photonics, vol. 2, no. 2, pp. 153–164, 2018.',
    '[90] H. Sun et al., "Deep learning-driven conversion of scanning superlens microscopy to high depth-of-field SEM-like imaging," Microsystems & Nanoengineering, vol. 11, art. no. 252, 2025.',
    '[91] S. Agnihotri, J. Grabinski, and M. Keuper, "Improving feature stability during upsampling—spectral artifacts and the importance of spatial context," in Computer Vision—ECCV 2024, pp. 357–376, 2024.',
    '[92] J. Hu, Z. Yao, L. Jin, H. He, and Y. Lu, "Enhancing image restoration transformer via adaptive translation equivariance," in Proceedings of the IEEE/CVF International Conference on Computer Vision, pp. 16047–16057, 2025.',
    '[93] J. Y. You et al., "Preserving spatial and quantitative information in unpaired biomedical image-to-image translation," Cell Reports Methods, vol. 5, no. 6, art. no. 101074, 2025.',
]


CONTRIBUTION_OVERVIEW = (
    "本研究的貢獻依模型設計、訓練目標與實驗證據分為四項。既有研究已將深度學習超解析用於"
    "顯微與關鍵尺寸量測，也已分別探討抗混疊、穩定上採樣、平移等變性及定量資訊保存"
    "[89,90,91,92,93]。因此，本研究不把這些個別原理或光學校正、資料建立、影像重建與後端"
    "量測的流程串接宣稱為首創，而將學術貢獻限定在針對CD與微位移量測所形成的模組設計、"
    "失效對應損失，以及可直接檢查其作用的配對量測證據。"
)

CONTRIBUTION_1 = (
    "貢獻一：建立面向微位移量測的抗混疊上採樣模組。抗混疊與平移等變網路已有明確基礎，"
    "近年亦有研究直接討論上採樣過程的特徵穩定性與影像重建的平移等變性"
    "[39,58,59,60,61,91,92]。本研究的貢獻在於把這些原理具體轉化為顯微重建解碼器中的空間"
    "對應保護，並以同一影格九顆pillar的真實位移一致性作為直接驗證。配對消融顯示，相較於"
    "一般上採樣，該模組使四組真實位移序列合併後的九顆pillar不一致量降低約14.3%，其可被"
    "支持的主要作用為改善跨空間位置的相對幾何一致性。"
)

CONTRIBUTION_2 = (
    "貢獻二：建立原始影像位置保留機制。殘差學習與輸入跳接本身已有成熟研究基礎"
    "[42,43,44]；本研究的差異不在於首次使用殘差，而在於將輸出限制為以原始影像提供低頻"
    "位置與大尺度亮暗基準，再加入幅度受限的高頻邊界修正。移除該模組後，20倍CD中位絕對"
    "誤差與兩倍率振鈴明顯惡化，顯示其主要價值是避免邊界與亮暗拓樸脫離輸入。由於50倍CD"
    "的單一誤差方向不同，本研究不將其延伸為所有倍率均必然降低CD的通則。"
)

CONTRIBUTION_3 = (
    "貢獻三：建立依量測失敗形式分組的功能性損失。頻域、梯度、任務特定及等變性損失皆有"
    "相關研究基礎[62,67,68,69,70]；本研究的貢獻是把影像內容、邊界幾何、運動一致性與輸出"
    "保守性四組目標，分別對應外觀失真、輪廓偏移、影格依賴與非必要修改等量測風險。消融"
    "結果未支持將完整改善獨立歸因於任一單項損失，因此論文只主張這套失效對應設計的整體"
    "功能，不誇大個別損失的因果效果。"
)

CONTRIBUTION_4 = (
    "貢獻四：建立可區分整體效果與單項作用的配對量測證據。原始影像、Wiener反卷積與MOSR"
    "共用影像名單、pillar身分、中心、物體空間尺度、門檻與統計規則，並以逐顆SEM尺寸、"
    "亞像素邊界與真實位移序列檢查重建結果[80,81,82,83,84,85]；另以兩個固定訓練種子、"
    "匹配訓練預算與配對重抽樣比較模組及損失移除，使『完整方法有效』與『哪一個設計負責"
    "何種改善』不再混為同一結論。"
)

CLOSEST_PRIOR = (
    "以超解析影像支援尺寸量測並非本研究首次提出。Kim與Pahk已將低解析光學影像轉換為較高"
    "解析條件，並以有機層的解析度與CD量測評估結果；近期晶片顯微研究也將光學超解析影像"
    "轉換為近似SEM的影像[89,90]。顯微重建領域另已有跨模態、電子顯微、單張超解析、通用"
    "顯微復原與像差補償等方法[71,72,74,75,77]。這些研究說明，以深度學習改善顯微影像並"
    "服務後續定量分析已有前例。"
)

NOVELTY_BOUNDARY = (
    "綜合上述代表性研究，本研究不將抗混疊、殘差連接、頻域損失或平移等變性本身列為首創。"
    "本研究的差異在於：以單張同尺寸顯微重建為部署形式，將上採樣的空間對應保護與輸入參考"
    "的受限輸出修正，直接對應到CD與微位移兩類量測風險；再以逐顆SEM參考、真實步進序列、"
    "共同後端尺規與配對消融界定各模組能被支持的作用[42,43,58,61,70,80,89,91,92,93]。"
    "因此，論文的創新主張是量測問題下的具體設計與證據，而不是宣稱個別基礎元件前所未見。"
)

METHOD_INNOVATION = (
    "本研究的模組貢獻不在於首次提出抗混疊、殘差連接或平移等變原理"
    "[39,42,43,58,61,91,92]，而是針對兩個可被量測的失敗來源建立對應設計：抗混疊上採樣"
    "用於降低解碼器對像素相位與空間位置的敏感性；原始影像位置保留則限制輸出邊界脫離輸入的"
    "風險。其他模組負責提供光學資訊與限制修正幅度，其獨立作用由第六章消融結果另行界定。"
)

CHAPTER4_SUMMARY = (
    "本章建立以物理量測為核心的單張顯微影像重建方法，並把研究主張分層。光學校正、多來源"
    "資料與共同後端尺規提供可追溯的研究架構；抗混疊上採樣與原始影像位置保留為本研究主要"
    "檢驗的模型設計，其貢獻在於面向CD與微位移的具體實作及配對證據，而不是將抗混疊、"
    "殘差或平移等變原理本身宣稱為首創[39,42,43,58,61,91,92]；功能性損失則把量測失敗"
    "形式轉為可訓練的目標。"
)


def replace_visible_text(paragraph: Paragraph, text: str) -> None:
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = deepcopy(run._r.rPr)
            break
    element = paragraph._p
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_rpr)


def find_one(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = paragraph.style
    run = new_paragraph.add_run(text)
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(paragraph.runs[0]._r.rPr))
    return new_paragraph


def append_citation(document: Document, prefix: str, numbers: list[int]) -> None:
    paragraph = find_one(document, prefix)
    citation = "[" + ",".join(str(number) for number in numbers) + "]"
    if citation not in paragraph.text:
        paragraph.add_run(citation)


def w_text_nodes(paragraph: Paragraph):
    return paragraph._p.xpath(".//w:t")


def replace_w_text_span(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    nodes = w_text_nodes(paragraph)
    texts = [node.text or "" for node in nodes]
    total = sum(len(text) for text in texts)
    if not (0 <= start < end <= total):
        raise ValueError((start, end, total, paragraph.text))

    cumulative = 0
    start_node = start_offset = end_node = end_offset = None
    for index, text in enumerate(texts):
        next_cumulative = cumulative + len(text)
        if start_node is None and start < next_cumulative:
            start_node = index
            start_offset = start - cumulative
        if end <= next_cumulative:
            end_node = index
            end_offset = end - cumulative
            break
        cumulative = next_cumulative
    if start_node is None or end_node is None:
        raise RuntimeError("Could not map text span to nodes")

    if start_node == end_node:
        text = texts[start_node]
        nodes[start_node].text = text[:start_offset] + replacement + text[end_offset:]
        return

    prefix = texts[start_node][:start_offset]
    suffix = texts[end_node][end_offset:]
    nodes[start_node].text = prefix + replacement
    for index in range(start_node + 1, end_node):
        nodes[index].text = ""
    nodes[end_node].text = suffix


def normalize_citations(paragraph: Paragraph) -> int:
    changes = 0
    while True:
        text = "".join(node.text or "" for node in w_text_nodes(paragraph))
        match = re.search(r"\[(\d+)\s*[–—-]\s*(\d+)\]", text)
        if not match:
            break
        first, last = (int(match.group(1)), int(match.group(2)))
        if last < first or last - first > 30:
            break
        replacement = "[" + ",".join(str(number) for number in range(first, last + 1)) + "]"
        replace_w_text_span(paragraph, match.start(), match.end(), replacement)
        changes += 1

    while True:
        text = "".join(node.text or "" for node in w_text_nodes(paragraph))
        match = re.search(r"\[\d+\](?:\s*,\s*\[\d+\])+", text)
        if not match:
            break
        replacement = "[" + ",".join(re.findall(r"\d+", match.group(0))) + "]"
        replace_w_text_span(paragraph, match.start(), match.end(), replacement)
        changes += 1

    while True:
        text = "".join(node.text or "" for node in w_text_nodes(paragraph))
        match = re.search(r"\[\d+(?:\s*,\s*\d+)+\]", text)
        if not match:
            break
        replacement = "[" + ",".join(re.findall(r"\d+", match.group(0))) + "]"
        if replacement == match.group(0):
            break
        replace_w_text_span(paragraph, match.start(), match.end(), replacement)
        changes += 1
    return changes


def all_body_paragraphs(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def rebuild_bibliography(document: Document) -> list[str]:
    old = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style is not None and paragraph.style.name == "EndNote Bibliography"
    ]
    if len(old) != 45:
        raise RuntimeError(f"Expected 45 existing references, found {len(old)}")
    old_texts = [paragraph.text.strip() for paragraph in old]
    old_numbers = [int(re.match(r"\[(\d+)\]", text).group(1)) for text in old_texts]
    if old_numbers != list(range(1, 46)):
        raise RuntimeError(f"Unexpected existing bibliography numbering: {old_numbers}")

    template_ppr = deepcopy(old[0]._p.pPr)
    for paragraph in old:
        paragraph._p.getparent().remove(paragraph._p)

    references = old_texts + NEW_REFERENCES
    heading = find_one(document, "參考文獻")
    cursor = heading._p
    for text in references:
        new_p = OxmlElement("w:p")
        if template_ppr is not None:
            new_p.append(deepcopy(template_ppr))
        cursor.addnext(new_p)
        paragraph = Paragraph(new_p, heading._parent)
        paragraph.add_run(text)
        cursor = new_p
    return references


def body_before_references(document: Document) -> list[Paragraph]:
    paragraphs = document.paragraphs
    reference_index = next(
        index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "參考文獻"
    )
    return paragraphs[:reference_index]


def audit_document(path: Path) -> dict:
    document = Document(str(path))
    paragraphs = body_before_references(document)
    citation_numbers: list[int] = []
    malformed: list[dict] = []
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text
        for match in re.finditer(r"\[([0-9][^\]]*)\]", text):
            content = match.group(1)
            if not re.fullmatch(r"\d+(?:,\d+)*", content):
                malformed.append({"paragraph": index, "citation": match.group(0), "text": text})
                continue
            citation_numbers.extend(int(value) for value in content.split(","))

    bibliography = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style is not None and paragraph.style.name == "EndNote Bibliography"
    ]
    ref_numbers = [
        int(match.group(1))
        for text in bibliography
        if (match := re.match(r"\[(\d+)\]", text))
    ]
    cited = sorted(set(citation_numbers))
    expected = set(range(1, 94))
    return {
        "document": str(path),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "bibliography_count": len(bibliography),
        "bibliography_numbers_ok": ref_numbers == list(range(1, 94)),
        "citation_occurrences": len(citation_numbers),
        "unique_citations": cited,
        "unused_references": sorted(expected - set(cited)),
        "out_of_range_citations": sorted(set(cited) - expected),
        "malformed_citations": malformed,
        "citation_spacing_violations": sum(
            len(re.findall(r"\[\d+,\s+\d+", paragraph.text)) for paragraph in paragraphs
        ),
        "separated_citation_groups": sum(
            len(re.findall(r"\[\d+\]\s*,\s*\[\d+\]", paragraph.text)) for paragraph in paragraphs
        ),
    }


def main() -> None:
    if not DOCUMENT.exists():
        raise FileNotFoundError(DOCUMENT)

    document = Document(str(DOCUMENT))

    replace_visible_text(find_one(document, "本研究的主要貢獻可分為"), CONTRIBUTION_OVERVIEW)
    replace_visible_text(find_one(document, "貢獻一："), CONTRIBUTION_1)
    replace_visible_text(find_one(document, "貢獻二："), CONTRIBUTION_2)
    replace_visible_text(find_one(document, "貢獻三："), CONTRIBUTION_3)
    replace_visible_text(find_one(document, "貢獻四："), CONTRIBUTION_4)
    replace_visible_text(find_one(document, "抗混疊上採樣處理解碼器"), METHOD_INNOVATION)
    replace_visible_text(find_one(document, "本章建立以物理量測為目標的單張顯微影像重建方法"), CHAPTER4_SUMMARY)

    mosr_paragraph = find_one(document, "本研究提出的MOSR")
    replace_visible_text(
        mosr_paragraph,
        "本研究提出的MOSR即針對此一應用情境，將邊緣過渡、跨門檻CD、位移重複性及"
        "多pillar空間一致性納入訓練設計與驗證流程，使影像重建由『看得清楚』轉向『量得穩定』。"
        "其最終評估基準為可換算至物體空間的幾何量，並以既有的亞像素配準、邊界定位與量測"
        "不確定度方法作為後端量測基礎[80,81,82,83,84,85]。",
    )
    insert_after(mosr_paragraph, CLOSEST_PRIOR)
    insert_after(find_one(document, "本研究所稱的物理約束"), NOVELTY_BOUNDARY)

    citation_map = [
        ("隨著半導體製程的特徵尺寸持續縮小", [37,78,79,89,90]),
        ("當待測結構逐漸接近光學系統", [13,17,18,20,78,79]),
        ("近年來，深度學習已廣泛應用", [30,31,63,64,71,72,74,75,77,89,90,93]),
        ("在光學成像系統中，點擴散函數", [13,46,50,52,63]),
        ("本研究將影像重建視為量測流程", [30,31,63,64,70,84,85]),
        ("本研究針對兩個容易影響量測", [39,58,59,61,62,67,69,70,91,92,93]),
        ("從繞射光學觀點出發", [13,17,18,20,29]),
        ("系統的光學響應與成像特性可由點擴散函數", [50,51,53,54,55,56]),
        ("綜合上述研究可知，PSF 校正不僅是", [50,53,55,56]),
        ("一般而言，成像過程可近似為物體", [46,47,48,50,52]),
        ("維納反卷積的前提是", [46]),
        ("其中 與 分別為雜訊與原始訊號", [46,53]),
        ("理查德森–露西反卷積是經典", [47,48]),
        ("其中分式為逐點相除", [47,48,49]),
        ("關於半盲/盲延伸", [52,54,55]),
        ("維納反卷積可視為", [50,52,53,55,56]),
        ("相較於傳統反卷積依賴明確", [35,40,43,44,45,71,72,74,75,76,77]),
        ("整體而言，CNN超解析", [35,71,72,73,74,75,76,77]),
        ("雖然卷積神經網路已能在視覺品質", [30,31,32,33,63,64,70,84,85,86]),
        ("對結構化幾何目標與精密位移", [80,81,82,83]),
        ("為使深度學習模型在有限資料下", [34,62,65,66,76]),
        ("其中一項重要條件為平移等變性", [38,39,58,59,60,61,62,91,92]),
        ("除平移關係外，本研究亦限制高頻", [63,64,67,68,69,70]),
        ("影像重建、分割與定位不只需要", [38,39,58,59,61,62,92]),
        ("這種輸入移動後", [38,58,62]),
        ("因此，一個模型即使對整數像素", [39,58,61]),
        ("現代 CNN 通常不只包含", [39,58,59,60,61,91,92]),
        ("步幅為 的卷積、最大池化", [39,58,60]),
        ("下採樣前加入低通濾波", [39,58,61]),
        ("轉置卷積則使用可學習", [27,91]),
        ("平移一致性可同時由架構與訓練改善", [39,58,61,62,91,92]),
        ("實際卷積神經網路通常只能近似", [38,58,62]),
        ("整數像素平移可以直接移動影像", [80]),
        ("等變性殘差也可以直接作為訓練損失", [62,92]),
        ("平移一致性損失不能單獨使用", [62,65,66]),
        ("本研究希望重建後的顯微影像", [30,31,63,64,85,89,93]),
        ("式（4-1）所採用的空間不變假設", [50,54,55,56]),
        ("多位置PSF不要求", [50,54,55,56,57]),
        ("局部 PSF 適合說明特定九宮格", [50,54,55,56]),
        ("自適應融合分支使用數個代表性候選", [57]),
        ("一般上採樣可能在特徵放大時", [39,58,59,60,61,91,92]),
        ("本研究不直接將邊界修正結果", [42,43,44,69,93]),
        ("此輸出方式不讓網路的大尺度低頻預測", [42,43,44,63,64,93]),
        ("本研究的訓練目標不是單純降低", [30,31,62,67,68,69,70]),
        ("為限制此問題，本研究加入平移一致性損失", [38,39,58,61,62,92]),
        ("邊緣銳利度以正規化灰階剖面", [82,83]),
        ("每顆 pillar 以固定中心與參考半徑", [80,81,82,83]),
        ("同一次平台移動理論上會使九宮格", [80,81,88]),
        ("CD 誤差同時以平均絕對誤差", [84,85,86,87,88]),
        ("位移結果則報告代表位移相對", [80,81,84,85]),
        ("目前局部PSF與多位置PSF提供", [50,51,54,55,56]),
        ("若模型進一步應用於不同相機", [75,77,90]),
        ("現有位移結果已以平台名義指令", [80,81,84,85]),
        ("結合獨立回授後", [84,85,86,87,88]),
        ("目前消融結果支持抗混疊上採樣", [87]),
        ("量測導向損失函數可先依影像內容", [87,88]),
        ("後續模型選擇應同時考量", [85,93]),
    ]
    for prefix, numbers in citation_map:
        append_citation(document, prefix, numbers)

    normalization_changes = 0
    for paragraph in all_body_paragraphs(document):
        if paragraph.style is not None and paragraph.style.name == "EndNote Bibliography":
            continue
        normalization_changes += normalize_citations(paragraph)

    references = rebuild_bibliography(document)
    if len(references) != 93:
        raise RuntimeError(f"Expected 93 references, got {len(references)}")

    document.save(str(DOCUMENT))
    audit = audit_document(DOCUMENT)
    audit["normalization_changes"] = normalization_changes
    AUDIT_JSON.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")

    failures = []
    if audit["bibliography_count"] != 93:
        failures.append(f"bibliography_count={audit['bibliography_count']}")
    if not audit["bibliography_numbers_ok"]:
        failures.append("bibliography numbering")
    if audit["unused_references"]:
        failures.append(f"unused references={audit['unused_references']}")
    if audit["out_of_range_citations"]:
        failures.append(f"out-of-range citations={audit['out_of_range_citations']}")
    if audit["malformed_citations"]:
        failures.append(f"malformed citations={len(audit['malformed_citations'])}")
    if audit["citation_spacing_violations"]:
        failures.append(f"citation spacing={audit['citation_spacing_violations']}")
    if audit["separated_citation_groups"]:
        failures.append(f"separated groups={audit['separated_citation_groups']}")
    if failures:
        raise RuntimeError("; ".join(failures))

    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
