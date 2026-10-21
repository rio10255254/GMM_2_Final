from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v7.docx")
DESTINATION = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v8.docx")
FIGURE_DIR = Path(r"D:\GMM_2_Final\Experiments\R63_CausalBoundaryMotionAblation\figures")


def replace_paragraph_text(paragraph, text: str) -> None:
    """Replace visible text while retaining paragraph properties and run formatting."""
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = deepcopy(run._r.rPr)
            break

    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)

    run = paragraph.add_run(text)
    if first_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_rpr)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_in: float) -> None:
    cell.width = Inches(width_in)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def format_table(table, widths: list[float]) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")

    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[column_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index == 0:
                set_cell_shading(cell, "D9D9D9")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.bold = row_index == 0


class Inserter:
    def __init__(self, document: Document, reference_paragraph):
        self.document = document
        self.reference = reference_paragraph

    def paragraph(self, text: str, style: str = "Normal", keep_with_next: bool = False):
        paragraph = self.reference.insert_paragraph_before(text, style=style)
        paragraph.paragraph_format.keep_with_next = keep_with_next
        return paragraph

    def page_break(self):
        paragraph = self.reference.insert_paragraph_before(style="Normal")
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        return paragraph

    def figure(self, path: Path, caption: str, width_in: float = 5.75):
        if not path.exists():
            raise FileNotFoundError(path)
        picture_paragraph = self.reference.insert_paragraph_before(style="Normal")
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.keep_with_next = True
        picture_paragraph.paragraph_format.space_before = Pt(4)
        picture_paragraph.paragraph_format.space_after = Pt(0)
        picture_paragraph.add_run().add_picture(str(path), width=Inches(width_in))

        caption_paragraph = self.reference.insert_paragraph_before(caption, style="Figure Caption")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return picture_paragraph, caption_paragraph

    def table(self, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]):
        caption_paragraph = self.reference.insert_paragraph_before(caption, style="Table Caption")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.keep_with_next = True

        table = self.document.add_table(rows=1, cols=len(headers))
        for index, text in enumerate(headers):
            table.rows[0].cells[index].text = text
        for row_data in rows:
            cells = table.add_row().cells
            for index, text in enumerate(row_data):
                cells[index].text = text
        format_table(table, widths)
        self.reference._p.addprevious(table._tbl)
        return caption_paragraph, table


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not DESTINATION.exists():
        raise FileNotFoundError(DESTINATION)

    document = Document(str(DESTINATION))
    paragraphs = document.paragraphs
    if len(paragraphs) != 1195:
        raise RuntimeError(f"Unexpected paragraph count: {len(paragraphs)}")

    replacements: dict[int, tuple[str, str]] = {
        20: (
            "本研究提出量測導向超解析",
            "本研究提出量測導向超解析（Measurement-Oriented Super-Resolution, MOSR）作為精密顯微量測的單張影像重建方法。其學術創新不以光學校正、影像重建與後端量測的流程整合本身作為主張，而是針對卷積神經網路可能因上採樣與輸出修正改變幾何位置的問題，設計位移安全上採樣（ShiftSafe Upsampling）與原圖位置錨定（Raw-position Anchor）兩項核心模組；訓練端再依影像內容、邊界幾何、運動一致性與輸出保守性建立功能性損失。模型實際使用時只輸入單張影像，輸出尺寸與輸入相同；多影格與位移資訊僅用於訓練與驗證。局部與多位置點擴散函數則作為傳統反卷積基準與固定光學先驗。",
        ),
        22: (
            "結果顯示",
            "結果顯示，Wiener反卷積可提高對比與邊緣銳利度，MOSR則在跨門檻CD與多pillar位移一致性上較穩定。配對消融顯示，一般上採樣使四組序列合併的九顆pillar不一致量增加14.31%；移除原圖位置錨定使20倍CD中位絕對誤差增加73.97%，並使兩倍率振鈴明顯惡化。其餘單項損失未形成可獨立歸因的廣泛效果。因此，主要學術貢獻界定為具直接證據的模組設計，以及依量測失敗形式建立的功能性損失設計，而非整合流程本身。",
        ),
        23: (
            "關鍵字：",
            "關鍵字：量測導向超解析、位移安全上採樣、原圖位置錨定、功能性損失、關鍵尺寸、平移定位。",
        ),
        26: (
            "This thesis proposes",
            "This thesis proposes Measurement-Oriented Super-Resolution (MOSR) as a single-image reconstruction method for precision microscopy. Its academic contribution is not claimed to arise from workflow integration itself. Instead, it addresses geometry changes introduced by convolutional-network upsampling and output correction through two core modules: ShiftSafe Upsampling and a Raw-position Anchor. Training objectives are further organized by four measurement-related failure modes: image content, boundary geometry, motion consistency, and output conservatism. At inference, the model accepts one image and produces an output of the same size; multi-frame and displacement information is used only for training and evaluation. Local and multi-position PSFs serve as the conventional deconvolution baseline and a fixed optical prior, respectively.",
        ),
        28: (
            "The results show",
            "The results show that Wiener deconvolution effectively increases contrast and edge sharpness, whereas MOSR provides more stable overall performance for cross-threshold CD and multi-pillar displacement consistency. Matched ablation further showed that replacing ShiftSafe Upsampling with ordinary upsampling increased the pooled nine-pillar disagreement by 14.31%. Removing the Raw-position Anchor increased the 20× median absolute CD error by 73.97% and increased outer ringing by 246.16% and 392.37% at 50× and 20×, respectively. Other individual loss removals did not provide sufficiently broad evidence for independent causal attribution. The principal academic contributions are therefore assigned to the evidenced module designs and the measurement-failure-oriented loss formulation, rather than to workflow integration itself.",
        ),
        29: (
            "Keywords:",
            "Keywords: Measurement-Oriented Super-Resolution; ShiftSafe Upsampling; Raw-position Anchor; Functional Loss Design; Critical Dimension; Translational Localization; Measurement Stability.",
        ),
        129: (
            "在量測導向的框架之下",
            "在量測導向的框架之下，本研究聚焦於兩個可由模型設計直接處理的問題。第一，一般解碼器的上採樣可能使相同結構在不同像素相位或不同視野位置產生不一致的邊界位移，因此需要在特徵放大時保留穩定的空間對應。第二，即使輸出影像更銳利，網路仍可能改寫低頻輪廓、亮暗拓樸與原始中心位置，因此需要在輸出端保留可追溯的原圖位置基準。訓練目標則不能只以像素相似度表示，而應依影像外觀、邊界幾何、影格間運動與不必要輸出修改等失敗形式分別建立約束，使模組設計與損失設計共同服務於尺寸及位移量測。",
        ),
        131: (
            "建立以量測可靠性",
            "設計位移安全上採樣模組，降低解碼上採樣造成的像素相位差異，使同一影格內不同位置的pillar維持較一致的相對位移。",
        ),
        132: (
            "結合局部與多位置",
            "設計原圖位置錨定與量測失效對應的功能性損失，使模型在增加可用邊界細節時，仍保留原始低頻位置、亮暗拓樸與可檢查的輸出範圍。",
        ),
        133: (
            "以20倍與50倍",
            "以20倍與50倍的邊緣、SEM對照CD及四組位移序列建立共同量測與配對消融，分別檢查完整方法效益、模組邊際作用與損失設計的可歸因範圍。",
        ),
        135: (
            "本研究的貢獻可歸納",
            "本研究的貢獻依模組設計、損失設計與實驗證據分為四項。MOSR提供研究問題與評估語境，光學校正、資料建立、重建及後端量測則構成實作與驗證架構；上述流程本身不被列為獨立的學術創新。",
        ),
        136: (
            "貢獻一：",
            "貢獻一：提出位移安全上採樣模組。該模組針對解碼器放大特徵時可能產生的像素相位差異，以對齊後的特徵重排與受控融合維持輸入、特徵與輸出之間的空間對應。配對消融顯示，相較於一般上採樣，該模組使四組真實位移序列合併後的九顆pillar不一致量降低約14.3%，其主要作用可被限縮為改善跨空間位置的相對幾何一致性。",
        ),
        137: (
            "為使不同方法",
            "貢獻二：提出原圖位置錨定模組。模型不直接以自由生成的完整影像取代輸入，而是由原始影像提供低頻位置與大尺度亮暗基準，只加入受限的高頻邊界修正。移除該模組後，20倍CD中位絕對誤差與兩倍率振鈴明顯惡化，證實其主要價值是防止邊界和亮暗拓樸脫離輸入；由於50倍CD的單一誤差方向不同，本研究不將其寫成所有倍率均必然降低CD的通則。",
        ),
        138: (
            "貢獻二：",
            "貢獻三：建立依量測失敗形式分組的功能性損失。影像內容、邊界幾何、運動一致性與輸出保守性四組目標，分別對應外觀失真、輪廓偏移、影格依賴與非必要修改。此處的創新在於把一般影像相似目標改寫為可對應後端量測風險的目標系統；消融結果未支持將完整改善獨立歸因於任一單項損失，因此論文只主張其整體設計功能，不誇大個別損失的因果效果。",
        ),
        139: (
            "貢獻三：",
            "貢獻四：建立可區分整體效果與單項作用的配對量測證據。原始影像、Wiener反卷積與MOSR共用影像名單、pillar身分、中心、物體空間尺度、門檻與統計規則；另以兩個固定訓練種子、匹配訓練預算與配對重抽樣比較模組及損失移除，使『完整方法有效』與『哪一個設計負責何種改善』不再混為同一結論。",
        ),
        140: (
            "本研究以同一顯微系統",
            "上述結果支持位移安全上採樣的跨pillar一致性作用，以及原圖位置錨定的邊界安全作用；功能性損失則保留為量測導向的共同訓練設計。核心驗證仍限於同一顯微系統、HS-100MG樣品及20倍與50倍成像資料域，更廣泛的跨樣品、跨系統及更多訓練重複於第七章列為後續工作。",
        ),
        143: (
            "第一章為緒論",
            "第一章為緒論，闡述研究背景、研究動機、量測導向超解析重建（MOSR）問題框架，以及位移安全上採樣、原圖位置錨定、功能性損失與配對量測驗證四項主要貢獻。",
        ),
        146: (
            "第四章說明",
            "第四章說明量測導向重建的研究方法，依序建立成像與物體空間座標、局部與多位置PSF策略、位移安全上採樣、原圖位置錨定、量測失效對應的功能性損失，以及Raw、Wiener與所提模型共用的尺寸與位移量測規則。",
        ),
        147: (
            "第五章描述",
            "第五章描述顯微光學系統、樣品與SEM參考量測、軟硬體環境、模型架構、訓練資料分組、實際微調設定及模組與損失函數的匹配消融配置。",
        ),
        148: (
            "第六章呈現",
            "第六章呈現PSF校正與反卷積驗證，比較原始影像、Wiener反卷積及MOSR在邊緣、CD與位移量測上的結果，並以配對消融界定位移安全上採樣、原圖位置錨定及功能性損失可被支持的作用範圍。",
        ),
        566: (
            "本研究的核心創新",
            "本研究的完整方法鏈用來確保光學資訊、影像重建與後端量測可在同一座標與規則下接受檢查，但流程串接本身不是本研究的主要學術創新。核心創新集中於兩項模型模組與一項訓練目標設計：位移安全上採樣降低不同像素相位與空間位置之間的幾何差異；原圖位置錨定限制網路任意改寫低頻輪廓與亮暗拓樸；功能性損失則把影像內容、邊界幾何、運動一致性與輸出保守性分別對應到具體量測失敗形式。PSF、多來源資料與共同後端尺規負責提供光學依據、監督來源與驗證條件。",
        ),
        568: (
            "4.1.4",
            "4.1.4 研究架構與模組、損失函數創新主軸",
        ),
        569: (
            "本研究的方法鏈",
            "本研究的方法鏈依序包含光學校正、監督資料建立、模型學習、單張影像重建與後端量測。此架構的功能是固定座標、資料角色與驗證順序，使後續實驗可追溯；它是承載創新設計與實驗判定的研究架構，而不是單獨被宣稱為演算法創新。",
        ),
        570: (
            "光學校正階段",
            "第一項核心模組創新為位移安全上採樣。一般解碼器在放大特徵圖時可能因插值、重排或邊界處理而產生位置相依的輸出；本研究以保持空間對應為優先的上採樣模組，使同一結構在不同像素相位與不同視野位置下維持較一致的邊界移動。",
        ),
        571: (
            "資料建立階段",
            "第二項核心模組創新為原圖位置錨定。模型以原始影像保留低頻位置、大尺度亮暗與主要輪廓，只將受限的模型高頻成分加入原圖，避免整張網路預測直接取代輸入影像。此設計允許邊界修復，同時保留可追溯的原始座標基準。",
        ),
        572: (
            "模型學習階段",
            "第三項創新為量測失效對應的功能性損失設計。訓練目標不只比較像素，而是分成影像內容、邊界幾何、運動一致性與輸出保守性四組，分別處理外觀失真、輪廓偏移、影格依賴及非必要修改，使每一組目標都能對應後端CD或位移量測可能發生的失敗。",
        ),
        573: (
            "模型部署時",
            "局部與多位置PSF、功能分離的資料來源、受限邊界殘差及品質門控均保留為重要支援設計。它們分別提供光學依據、監督訊息與輸出保護，但本研究不因其被納入完整流程，就將每一項都宣稱為已獨立證明的主要創新。",
        ),
        574: (
            "因此，本研究的創新主軸",
            "因此，本研究的主張層級為：位移安全上採樣與原圖位置錨定是可由消融直接檢查的核心模組；功能性損失是由量測失敗形式推導的訓練設計；MOSR方法鏈、PSF先驗、多來源監督與共同量測尺規則負責使上述設計能被公平實作與驗證。",
        ),
        630: (
            "因此，本研究不是讓神經網路",
            "因此，本研究不是讓神經網路完全自由地重新生成影像，而是在特徵放大與最終輸出兩個最容易改變幾何位置的階段加入專門限制。位移安全上採樣維持解碼過程的空間對應；原圖位置錨定保留輸入影像的低頻位置與亮暗拓樸；光學PSF、殘差上限及輸入品質評估則提供額外的物理與輸出保護。",
        ),
        631: (
            "整體模型包含四項",
            "整體模型以兩項核心模組為主，並搭配三項支援機制。核心模組為位移安全上採樣與原圖位置錨定；前者作用於解碼器的特徵放大，後者作用於最終影像形成。支援機制包括PSF殘差修正、方向感知且幅度受限的邊界修正，以及依輸入品質調節修正量的門控。輸入與輸出維持相同尺寸，使所有處理沿用同一像素座標與物體空間尺度。",
        ),
        632: (
            "這些機制共同構成",
            "本研究的模組創新重點不是把既有元件集中到同一網路，而是針對兩個明確失敗來源提出對應設計：以位移安全上採樣處理解碼器的位置相依性，以原圖位置錨定處理輸出邊界脫離輸入的風險。其他模組負責提供光學資訊與限制修正幅度，其獨立作用由第六章消融結果另行界定。",
        ),
        633: (
            "4.3.1",
            "4.3.1 同尺寸編碼器—解碼器與位移安全上採樣",
        ),
        636: (
            "一般上採樣方法",
            "一般上採樣方法可能在放大特徵圖時引入像素相位偏移、棋盤狀紋理，或使相同結構在不同影像位置產生不同輸出。這類位置誤差對一般視覺任務可能不明顯，但對亞像素邊界定位會直接轉化為CD或位移誤差。本研究因此使用位移安全上採樣（ShiftSafe Upsampling）：先以固定幾何規則恢復空間尺寸，再將同尺度跳接特徵對齊後融合，避免可學習上採樣同時承擔內容生成與座標重排。模型內部另保留Residual PSF Correction，使網路在學習影像內容時參考固定光學先驗。各層通道數、卷積核與連接位置於第5.4.2節說明。",
        ),
        654: (
            "4.3.3",
            "4.3.3 原圖位置錨定與受限高頻細節注入",
        ),
        656: (
            "因此，本研究不直接將",
            "因此，本研究不直接將邊界修正結果作為最終輸出，而是建立原圖位置錨定（Raw-position Anchor）：原始影像作為低頻位置、大尺度亮暗與主要輪廓基準，模型預測只提供局部高頻細節，再以受限比例加入原圖。",
        ),
        661: (
            "式（4-11）可理解",
            "式（4-11）可理解為『原始影像加上受限的模型高頻細節』。其中，原始影像提供低頻位置、主要背景與大尺度亮暗；模型輸出則提供局部邊界修正。係數0.35表示只加入35%的高頻成分，而不是對整張影像的對比乘上1.35。",
        ),
        662: (
            "這項設計與直接輸出",
            "原圖位置錨定與直接輸出完整網路預測的差異，在於模型的大尺度低頻結果不會完整取代輸入。即使主幹網路對背景亮度或輪廓中心產生偏移，最終影像仍以原圖作為位置基底，使模型較大的自由度集中在與pillar邊界相關的局部快速變化。",
        ),
        674: (
            "綜合而言，本研究利用三層",
            "綜合而言，原圖位置錨定是最終輸出的主要位置保護；受限邊界殘差與輸入品質門控則分別限制局部單點修正及整體修正幅度。三者作用層級不同：錨定保留低頻座標與亮暗拓樸，殘差限制局部增量，品質門控依訊號與噪聲調整修正強度。",
        ),
        675: (
            "三項機制共同限制",
            "三項機制共同限制模型的輸出自由度，但是否每一項都能單獨造成CD或位移改善，不能只由結構動機判定，仍須以匹配重訓的消融條件及共同量測規則檢查。",
        ),
        676: (
            "本研究模型的主要創新",
            "本研究模型的主要模組創新因而界定為位移安全上採樣與原圖位置錨定：前者降低解碼上採樣造成的跨位置相位差異，後者降低輸出邊界脫離原始位置與亮暗拓樸的風險。方向感知邊界殘差、品質門控與PSF修正保留為支援設計；下一節則由損失函數處理無法只靠架構硬限制表達的影像、幾何與運動關係。",
        ),
        707: (
            "綜合而言，本研究在訓練資料上的主要創新",
            "綜合而言，多來源資料的功能分工是功能性損失得以成立的監督基礎，而不是單獨被主張為主要學術創新。真實顯微影像提供外觀，合成幾何與SEM參考提供輪廓與尺寸，已知位移及連續影格提供位置關係，穩定性資料則限制模型對非幾何因素的過度反應。這些資料經由分組切分後，分別餵給與量測失敗形式對應的損失項。",
        ),
        708: (
            "4.5",
            "4.5 量測失效對應的功能性損失與分階段微調",
        ),
        711: (
            "因此，本研究將訓練目標",
            "因此，本研究不是依常見損失名稱任意堆疊項目，而是先列出量測可能失敗的方式，再反推所需的訓練約束。總目標分為影像內容、邊界幾何、運動一致性及輸出保守性四組，分別回答輸出是否仍像真實顯微影像、邊界是否位於合理位置、相同結構在不同影格中的位置關係是否被保留，以及模型是否對非量測區域進行不必要修改。",
        ),
        725: (
            "需要區分的是",
            "需要區分的是，第4.3節中的殘差截斷、原圖位置錨定及品質門控屬於模型架構上的限制；輸出保守性損失則屬於訓練時的軟性懲罰。前者直接限制模型能如何形成輸出，後者在模型嘗試產生不合理修改時增加訓練代價，兩者不是同一機制。",
        ),
        727: (
            "本研究在訓練設計上的主要特色",
            "本研究在損失設計上的創新，是將一般影像重建中的『影像相似』目標，改寫為與實際量測失敗一一對應的功能性約束。這項設計的學術內容是問題分解與目標定義，而不是宣稱損失項數量越多越好。各組損失在完整模型中共同作用；其能否被單獨歸因，則以第六章消融結果為準。",
        ),
        728: (
            "4.5.2",
            "4.5.2 平移一致性損失",
        ),
        735: (
            "平移一致性只能限制",
            "平移一致性只能限制相對關係，不能單獨決定絕對幾何。若模型在所有位置都把pillar直徑固定縮小5%，兩條平移路徑仍可能保持一致，但絕對CD依然錯誤。因此，式（4-14）需與SEM參考幾何、合成形狀、邊界距離及原圖位置錨定共同使用。",
        ),
        736: (
            "換言之，平移一致性回答",
            "換言之，平移一致性回答『相同物體移動後，模型是否仍以相同方式處理它』，而SEM與幾何監督回答『模型處理後的尺寸是否合理』。兩者在設計上互補，但此互補關係不等同於每一項都已被獨立證明；其邊際作用由匹配消融另行判讀。",
        ),
        793: (
            "本章建立一套",
            "本章建立以物理量測為核心的單張顯微影像重建方法，並把研究主張明確分層。光學校正、多來源資料與共同後端尺規提供可追溯的研究架構；位移安全上採樣與原圖位置錨定構成核心模組創新；功能性損失則把量測失敗形式轉為可訓練的目標。",
        ),
        794: (
            "光學部分以中心九宮格",
            "光學部分以中心九宮格的逐顆參考幾何擬合局部PSF，並以多位置PSF建立空間變異融合模組。這些PSF在完整模型與後續消融條件中均保持固定，用來提供共同光學先驗，而不是作為本研究消融歸因的對象。",
        ),
        795: (
            "模型部分採用單張同尺寸輸出",
            "模型部分採用單張同尺寸輸出。位移安全上採樣降低解碼器造成的跨位置相位差異，原圖位置錨定保留低頻位置與亮暗拓樸；方向感知邊界殘差、品質門控及PSF修正則限制或引導局部修正。訓練目標依影像、邊界、運動與保守性四種量測失敗形式分組。",
        ),
        796: (
            "最後，Raw、Wiener",
            "最後，Raw、Wiener與所提模型均使用相同座標、中心、邊界與統計規則。第六章除檢查完整模型是否改善邊界、CD與位移，也以匹配重訓消融區分核心模組、功能性損失與支援機制的可歸因範圍，使創新主張建立在具體證據上。",
        ),
        901: (
            "本研究的重建網路",
            "本研究的重建網路以Encoder–Decoder與跳接結構為主體，核心模組為ShiftSafe Upsampling與Raw-position Anchor，並結合Boundary Correction及空間變異PSF融合模組。圖5-19呈現各模組在特徵提取、解析度恢復、光學修正、邊界補償與最終輸出形成中的相對位置。",
        ),
        903: (
            "解碼器主要負責",
            "解碼器主要負責將深層特徵逐步還原至原始解析度，並透過跳接融合淺層邊界資訊。為避免一般上採樣在特徵放大時造成位置誤差或棋盤狀偽影，本研究使用ShiftSafe Upsampling，將尺寸恢復與內容融合分開處理，使上採樣後的特徵維持較穩定的空間對應。後續命名中帶有Decoder者表示解碼器路徑，帶有ShiftSafe Upsampling者表示位移安全上採樣；相關層級參數列於表5-15。",
        ),
        912: (
            "模型輸出前另設有Boundary Correction",
            "模型輸出前設有Boundary Correction分支，針對局部邊界預測幅度受限的殘差；其後的Raw-position Anchor再以原始影像保留低頻位置與大尺度亮暗，只把受限高頻修正加入原圖。品質門控則依輸入的邊界訊號與平坦區噪聲調整整體修正量。三者依序對應局部邊界補償、原始位置保護與輸入品質調節。",
        ),
        913: (
            "整體而言",
            "整體而言，網路以Encoder–Decoder與ShiftSafe Upsampling完成同尺寸多尺度重建，以空間變異PSF融合模組導入固定光學先驗，再由Boundary Correction、Raw-position Anchor與品質門控形成受控輸出。核心模組與支援模組的邊際作用於第6.2.6節以匹配消融檢查。",
        ),
        1006: (
            "本節評估量測導向超解析",
            "本節評估量測導向超解析（MOSR）模型在100MG樣品影像上的實際重建效果。前半部以同一批Raw影像、相同九顆pillar、相同120個位移步與相同量測規則，比較Raw、Wiener反卷積及採用模型；後半部再以匹配重訓的消融控制組，分離位移安全上採樣、原圖位置錨定、功能性損失與其他支援模組的邊際作用。如此可將『完整方法是否有效』與『哪一項設計負責何種改善』分開回答。",
        ),
        1007: (
            "本節所稱 MOSR",
            "第6.2.3至第6.2.5節所稱MOSR，為完成訓練後固定的採用模型；評估期間不更新權重，輸入與輸出尺寸相同。第6.2.6節則使用相同起始權重、資料切分、光學先驗與訓練預算建立兩個種子的完整控制模型及消融模型。消融數值用於配對比較，不取代前述採用模型的正式Raw、Wiener與MOSR結果。",
        ),
        1011: (
            "訓練紀錄中可可靠對應",
            "訓練紀錄中可可靠對應的三項損失皆由第1回合下降至第120回合：重建損失由0.01938降至0.01640，下降15.37%；梯度幅值損失由0.15801降至0.13115，下降17.00%；邊緣方向損失由0.004871降至0.001837，下降62.28%。圖中呈現每一回合的原始數值，沒有平滑或挑選局部區間。這些曲線只支持最佳化已收斂至較低誤差，不能單獨證明任一損失造成CD或位移改善；後者由量測結果與匹配消融判定。",
        ),
        1122: (
            "6.2.6",
            "6.2.7 綜合討論與研究限制",
        ),
        1123: (
            "綜合影像與量測結果",
            "綜合完整模型比較與配對消融，可整理為四項主要發現。第一，Wiener與MOSR都能使邊界比Raw集中，但Wiener伴隨較明顯的門檻敏感度；採用MOSR在跨門檻CD與120步九顆pillar一致性上較穩定。第二，兩種子完整控制模型相對Raw同時降低兩倍率CD誤差、四組序列步間波動、指令偏差與九宮格不一致量，確認MOSR的量測效益不是單一採用權重的偶然結果。第三，位移安全上採樣相較一般上採樣使四序列合併的九宮格不一致量降低14.31%，直接支持其跨位置相對幾何作用。第四，原圖位置錨定在20倍CD及兩倍率振鈴上提供明顯保護，但CD收益具有倍率依賴性。",
        ),
        1124: (
            "本節結果的適用範圍",
            "本節結果的適用範圍仍為目前成像系統、樣品與資料域，不延伸主張跨樣品或跨系統泛化。平台只有名義指令而沒有獨立位置回授，因此位移結果支持重複性、指令一致性與九宮格內一致性，而非絕對位移準確度。消融只包含兩個訓練種子，配對區間主要反映pillar、步驟與方向的條件變異；多數單項損失移除未達預定效果門檻，不能被選擇性解讀為獨立貢獻。完整控制模型在50倍的外圍振鈴比Raw增加86.14%，亦顯示邊界集中與假影代價必須同時報告。",
        ),
        1125: (
            "因此，本節最有力且較保守的結論",
            "因此，本研究最有力的學術結論不再是『整合一套流程即可改善影像』，而是兩項可被直接驗證的模組作用：位移安全上採樣降低不同pillar之間的位置分歧，原圖位置錨定防止邊界與亮暗拓樸脫離輸入。功能性損失把量測失敗形式轉為訓練目標，是方法設計的重要部分；但現有消融只支持其共同角色，尚不足以把完整MOSR的全部改善歸因於任一單項損失。",
        ),
        1078: (
            "跨全部量測條件後",
            "綜合全部量測後，中位絕對CD誤差由Raw的192.1 nm降至Wiener的161.2 nm與MOSR的117.0 nm，相對降幅為16.1%與39.1%。這說明Wiener可有限改善20倍CD一致性，而MOSR在相同跨門檻規則下的整體誤差較低。",
        ),
        1128: (
            "本研究已建立",
            "本研究已完成光學校正、影像重建、尺寸與位移量測及匹配消融，並將主要貢獻收斂為位移安全上採樣、原圖位置錨定與量測失效對應的功能性損失。後續工作應擴大模組與損失的跨條件因果驗證，再逐步延伸至更多光學條件、樣品類型與可追溯回授。",
        ),
        1138: (
            "7.4",
            "7.4 模組與損失函數的擴大消融",
        ),
        1139: (
            "後續模型可由",
            "目前消融已支持位移安全上採樣對跨pillar一致性的作用，以及原圖位置錨定對邊界安全的作用，但只有兩個訓練種子，且功能性損失多數未形成可獨立歸因的差異。後續可增加獨立訓練種子、由隨機初始化進行完整重訓，並在不同樣品與倍率重複相同配對比較，以確認兩項模組的效果是否跨資料域維持。",
        ),
        1140: (
            "此外，可讓模型",
            "對功能性損失則可採分層因子設計，先比較影像內容、邊界幾何、運動一致性與輸出保守性四大群組，再於確定有效的群組內拆解子項。如此可降低多重比較造成的偶然結果，也能區分『某項損失可被其他機制補償』與『該項損失本身沒有作用』兩種不同情況。",
        ),
        1141: (
            "7.5",
            "7.5 量測指標整合與實際部署",
        ),
        1142: (
            "為提升實務使用效率",
            "後續模型選擇應同時考量邊緣寬度、跨門檻CD穩健度、振鈴、位移重複性與多顆空間一致性，而不是只依訓練損失或主觀銳利度決定。亦可讓模型輸出量測可信度，標示低對比、邊界不完整、超出訓練範圍或可能振鈴的區域；當可信度不足時，系統保留Raw、降低修正強度或要求重新取像。",
        ),
        1143: (
            "在運算層面",
            "在部署層面，可將樣品方向判定、pillar與via分區、中心九宮格定位、影像品質檢查、重建及後端量測自動化，並固定保存灰階範圍、量測中心、門檻設定、排除影格與排除原因。後續再評估單張推論時間、記憶體需求與連續取像吞吐量，使位移安全上採樣與原圖位置錨定可在不犧牲可稽核性的前提下用於實際顯微量測。",
        ),
        1145: (
            "未來研究的核心方向",
            "未來研究的核心方向，是把目前已被支持的模組作用擴展為跨位置、跨樣品與跨系統均可重複驗證的能力，並以更大規模消融釐清功能性損失的個別與互補作用。隨著獨立位移真值、完整不確定度預算及自動品質控管加入，量測導向超解析可進一步發展為具有可追溯性與明確適用界線的顯微量測工具。",
        ),
    }

    for index, (expected_prefix, new_text) in replacements.items():
        actual = paragraphs[index].text.strip()
        if not actual.startswith(expected_prefix):
            raise RuntimeError(
                f"Paragraph {index} mismatch. Expected prefix {expected_prefix!r}, got {actual[:80]!r}"
            )
        replace_paragraph_text(paragraphs[index], new_text)

    # Add the ablation configuration to Chapter 5 before the chapter summary.
    chapter5_summary = next(p for p in document.paragraphs if p.text.strip() == "5.5 小結")
    insert5 = Inserter(document, chapter5_summary)
    insert5.paragraph("5.4.5 模組與損失函數消融配置", style="Heading 3")
    insert5.paragraph(
        "為區分完整模型效果與單一設計的邊際作用，本研究建立匹配重訓的消融實驗。全部條件保留相同的起始權重、固定資料切分、批次設定、更新次數、檢查點選擇規則與PSF光學先驗；PSF路徑在所有條件中保持啟用，不列為本輪消融變因。"
    )
    insert5.paragraph(
        "實驗共包含16個條件與2個固定訓練種子，合計32次有效訓練。比較內容包括幾何與運動損失的2×2組合、七項細分損失移除、三項輸出模組移除，以及位移安全上採樣與一般上採樣的參數匹配比較。所有條件依相同收斂規則訓練至60回合，並在第1至第60回合中以共同規則選擇檢查點。"
    )
    insert5.paragraph(
        "每一條件與種子均完成相同的靜態CD、真實位移與合成位移評估，共產生18,272張候選重建影像。主要差異以兩個種子方向是否一致、相對效果是否達10%、95%配對區間是否高於零，以及主要護欄是否惡化超過15%共同判定。配對區間以4,000次重抽樣估計；靜態CD以pillar、真實位移以步驟、合成位移以案例為抽樣單位。"
    )
    insert5.paragraph(
        "第6.2.3至第6.2.5節使用的採用模型仍固定為第120回合權重；消融實驗則以共同起點和共同60回合預算估計模組或損失的邊際作用。兩者目的不同，消融控制組的絕對數值不取代採用模型的正式三方法結果，論文只以成對差異進行因果解讀。"
    )

    # Add the complete Chapter 6 ablation section before the existing discussion.
    discussion = next(p for p in document.paragraphs if p.text.strip() == "6.2.7 綜合討論與研究限制")
    insert6 = Inserter(document, discussion)
    insert6.page_break()
    insert6.paragraph("6.2.6 模組與損失函數消融實驗", style="Heading 3")
    insert6.paragraph(
        "前述Raw、Wiener與採用MOSR模型的比較，回答完整重建方法是否有助於後端量測；但完整模型同時包含多個模組與損失，不能由整體結果直接推論任何單一設計必然是改善來源。因此，本節以匹配重訓消融固定除目標變因外的其餘條件，分別檢查位移安全上採樣、原圖位置錨定、功能性損失與其他輸出保護的邊際作用。"
    )

    insert6.paragraph("6.2.6.1 實驗設計與公平比較條件", style="Heading 4")
    insert6.paragraph(
        "消融共建立16個比較條件與2個固定種子，合計32次有效訓練。每一條件沿用相同起始權重、固定資料切分、PSF、批次、更新次數與檢查點選擇方式；每一條件、種子均完成571張候選影像重建。表6-13整理各比較層級，PSF在全部條件中維持固定，不因本節結果被移除或重新選擇。"
    )
    insert6.table(
        "表 6-13 模組與損失函數消融之匹配比較條件",
        ["比較層級", "條件數", "唯一主要變動", "共同固定項目"],
        [
            ["完整控制與2×2因子", "4", "幾何損失、運動損失之保留或移除", "起始權重、PSF、資料切分、訓練預算"],
            ["細分損失移除", "7", "每次移除一項幾何或運動子損失", "其餘損失與全部模型模組"],
            ["輸出模組移除", "3", "邊界殘差、原圖位置錨定或品質門控", "相同資料、預算與光學先驗"],
            ["上採樣匹配", "2", "位移安全上採樣或一般上採樣", "可訓練參數規模、初始權重與其餘模組"],
        ],
        [1.30, 0.60, 2.10, 2.15],
    )
    insert6.paragraph(
        "所有條件依共同收斂規則訓練至60回合，並在第1至第60回合中以相同規則選擇檢查點；32次訓練中有29次選到第60回合、2次選到第56回合、1次選到第24回合。主要判定要求兩個種子方向一致、相對效果至少10%、95%配對區間完全高於零，且主要護欄不得惡化超過15%。小於5%視為可忽略，5%至10%或只有單一倍率一致者只列為趨勢。"
    )
    insert6.paragraph(
        "為維持與前節主實驗的分工，本節中的Full MOSR指相同起點下重新訓練的完整控制模型，而不是第6.2.3至第6.2.5節採用的第120回合模型。兩組完整模型的絕對數值可能略有差異；本節只使用同一消融套件內的成對差異，避免把訓練預算或檢查點差異誤認為模組效果。"
    )

    insert6.paragraph("6.2.6.2 完整控制模型相對Raw之量測效果", style="Heading 4")
    insert6.paragraph(
        "在拆解單項機制之前，先檢查完整控制模型是否重現主實驗的整體方向。兩個種子平均後，50倍CD中位絕對誤差由Raw的96.34 nm降為78.54 nm，下降18.48%；20倍由190.39 nm降為120.25 nm，下降36.84%。圖6-47中的兩個種子標記方向一致，表示改善不只出現在單一訓練重複。"
    )
    insert6.figure(
        FIGURE_DIR / "15_full_vs_raw_cd_error.png",
        "圖 6-47 完整控制模型相對Raw之兩倍率SEM對照CD誤差",
    )
    insert6.paragraph(
        "跨門檻CD上下界寬度在50倍由354.46 nm降為115.47 nm，縮減67.42%；20倍由792.65 nm降為674.84 nm，縮減14.86%。10–90%邊界過渡寬度在50倍與20倍分別縮減28.85%與23.55%。因此，完整控制模型的收益同時表現在典型CD誤差、門檻敏感度與邊界集中程度，而不是只依賴單一50%交點。"
    )
    insert6.paragraph(
        "四組真實位移序列各固定使用由Raw先行排序選出的30步，候選模型不參與影格挑選。圖6-48顯示完整控制模型在四組序列均降低逐步位移波動，其中20倍、0.5 µm由16.25 nm降為14.67 nm，改善9.73%，略低於10%實質門檻；其餘三組改善47.82%、56.38%與44.71%。四組合併後，逐步位移波動由11.64 nm降為8.49 nm，下降27.03%。"
    )
    insert6.figure(
        FIGURE_DIR / "16_full_vs_raw_step_stability.png",
        "圖 6-48 完整控制模型相對Raw之四組真實序列步間穩定性",
    )
    insert6.paragraph(
        "位移主軸中最明顯的整體改善出現在九顆pillar的空間一致性。圖6-49顯示四組序列均向較低不一致量移動；合併120步後，九宮格不一致量由Raw的17.63 nm降為6.96 nm，下降60.52%，命令位移絕對誤差亦由9.54 nm降為6.70 nm，下降29.84%。這組完整控制結果確認主實驗的改善方向可在兩個匹配訓練種子中重現。"
    )
    insert6.figure(
        FIGURE_DIR / "17_full_vs_raw_grid_agreement.png",
        "圖 6-49 完整控制模型相對Raw之九顆pillar位移一致性",
    )

    insert6.paragraph("6.2.6.3 原圖位置錨定之邊界保護作用", style="Heading 4")
    insert6.paragraph(
        "原圖位置錨定以Raw保留低頻位置、大尺度亮暗與主要輪廓，只將受限的模型高頻成分加入原圖。消融時移除該模組並以相同預算重新訓練，使比較回答『保留原圖位置基準是否必要』，而不是在推論時臨時關閉一個未重新適應的分支。"
    )
    insert6.figure(
        FIGURE_DIR / "18_raw_position_anchor_cd_error.png",
        "圖 6-50 移除原圖位置錨定對兩倍率SEM對照CD誤差之影響",
    )
    insert6.paragraph(
        "20倍移除錨定後，CD中位絕對誤差增加88.94 nm，即增加73.97%；兩個種子的代價分別為78.44與99.44 nm，95%配對區間為49.29至125.29 nm，符合方向、幅度與區間要求。50倍則出現相反方向，移除後誤差降低15.09 nm。故原圖位置錨定不能被寫成所有倍率都必然使單一CD誤差下降，其CD收益具有倍率依賴性。"
    )
    insert6.paragraph(
        "然而，邊界安全護欄在兩倍率具有一致方向：移除錨定後，50倍與20倍外圍振鈴分別增加246.16%與392.37%。若只在推論時直接關閉錨定，兩組靜態CD輸出亦因徑向對比不符合既定量測條件而判為無效，且較大合成位移出現明顯惡化。綜合CD、振鈴與失敗檢查，最適當的結論是原圖位置錨定防止邊界與亮暗拓樸脫離輸入；其保護作用在較粗取樣的20倍CD上最明顯。"
    )

    insert6.paragraph("6.2.6.4 位移安全上採樣之空間一致性", style="Heading 4")
    insert6.paragraph(
        "位移安全上採樣與一般上採樣的比較維持相同可訓練參數規模、相同初始權重、資料、訓練預算與光學先驗，只更換特徵放大方式。圖6-51中一般上採樣的九宮格不一致量在四組序列均高於位移安全上採樣，表示其差異主要反映不同視野位置之間的相對幾何，而不是整體對比或影格名單改變。"
    )
    insert6.figure(
        FIGURE_DIR / "19_upsampling_grid_agreement.png",
        "圖 6-51 位移安全上採樣與一般上採樣之九顆pillar位移一致性",
    )
    insert6.paragraph(
        "四序列合併後，一般上採樣的九宮格不一致量比位移安全上採樣高0.977 nm，即增加14.31%，95%配對區間為0.213至1.714 nm；50倍合併增加0.610 nm或14.47%，區間為0.142至1.091 nm。20倍合併亦增加14.24%，但區間跨越零，因此只列為同方向趨勢。這項結果直接支持位移安全上採樣降低跨pillar空間分歧。"
    )
    insert6.paragraph(
        "相對地，四序列合併的逐步位移波動只增加4.68%，95%區間跨越零；合成次像素位移向量誤差亦未在兩個種子得到一致方向。因此，本研究不把位移安全上採樣寫成全面提高所有次像素位移精度，而將其作用限縮為維持真實序列中不同空間位置的相對幾何一致性。"
    )

    insert6.paragraph("6.2.6.5 功能性損失與其他輸出模組", style="Heading 4")
    insert6.paragraph(
        "除兩項核心模組外，本研究亦移除廣義幾何錨定、廣義運動一致性、七項細分損失、邊界殘差分支與品質門控。表6-14將主要結果與可支持的主張並列。大多數單項移除的效果小於10%、95%區間跨越零，或兩倍率與兩個種子方向不一致，故不以零散的有利數字宣稱獨立因果作用。"
    )
    insert6.table(
        "表 6-14 核心模組、功能性損失與支援模組之消融判讀",
        ["比較項目", "主要量測", "移除或替換代價", "判讀"],
        [
            ["原圖位置錨定", "20倍CD中位絕對誤差", "+88.94 nm（+73.97%）", "支持20倍CD與邊界安全；CD具倍率依賴"],
            ["位移安全上採樣", "四序列九宮格不一致量", "+0.977 nm（+14.31%）", "支持跨pillar空間一致性"],
            ["廣義幾何錨定損失", "50倍CD誤差", "−1.19 nm（−1.52%）", "方向不利且未達門檻，不支持獨立歸因"],
            ["廣義運動一致性損失", "四序列步間波動", "+0.061 nm（+0.72%）", "差異可忽略，區間跨零"],
            ["相鄰影格一致性損失", "四序列步間波動", "+0.025 nm（+0.30%）", "差異可忽略，區間跨零"],
            ["邊界殘差分支", "兩倍率CD與護欄", "方向不一致", "不支持單獨歸因"],
            ["品質門控", "CD與位移", "效果小且方向不一致", "保留為支援機制，不作獨立效果主張"],
        ],
        [1.45, 1.55, 1.45, 1.70],
    )
    insert6.paragraph(
        "上述結果不等同於功能性損失沒有作用。所有候選皆由相同穩定起點進行有限預算的匹配加強訓練，移除一項後，其餘損失、原圖位置錨定、位移安全上採樣與輸出保護仍可在重新訓練時部分補償。現有證據支持功能性損失作為依量測失敗形式建立的共同目標系統，但不足以把完整MOSR改善歸因於任一單項損失。"
    )
    insert6.paragraph(
        "完整控制模型亦保留銳利度與假影之間的取捨。圖6-52顯示，50倍外圍振鈴能量由Raw的0.0000718增至0.0001336，增加86.14%；20倍則由0.0001731降為0.0001273，下降26.47%。因此，50倍的CD與邊界改善伴隨較強外圍振鈴，不能將MOSR描述為在所有倍率都同時消除假影。"
    )
    insert6.figure(
        FIGURE_DIR / "20_full_vs_raw_ringing_tradeoff.png",
        "圖 6-52 完整控制模型相對Raw之外圍振鈴取捨",
    )

    insert6.paragraph("6.2.6.6 消融結果對創新主軸之界定", style="Heading 4")
    insert6.paragraph(
        "消融結果使本論文的創新主軸由『整合光學、網路與量測流程』進一步收斂為可被檢查的模型與訓練設計。第一，位移安全上採樣具有跨pillar位置一致性的直接證據；第二，原圖位置錨定具有邊界與亮暗拓樸安全的直接證據，並在20倍CD上呈現明顯收益；第三，功能性損失將量測失敗形式轉為訓練目標，但其作用應以整體設計描述，不拆成每一項均已獨立證明。"
    )
    insert6.paragraph(
        "PSF融合、多來源監督、方向感知邊界殘差、品質門控及共同量測流程仍為完整MOSR不可缺少的支援內容，但它們在本輪不是主要因果主張。此分層使學術創新與工程實作各自清楚：模組與損失設計回答模型如何降低幾何破壞，整體流程則確保這些設計能在相同物理座標與後端尺規下被公平驗證。"
    )

    # Ensure Word updates field results and list entries when the file is opened.
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(str(DESTINATION))
    print(f"Saved: {DESTINATION}")
    print(f"Paragraphs: {len(document.paragraphs)}")
    print(f"Tables: {len(document.tables)}")
    print(f"Inline shapes: {len(document.inline_shapes)}")


if __name__ == "__main__":
    main()
