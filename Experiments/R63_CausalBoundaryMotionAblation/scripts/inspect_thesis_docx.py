from __future__ import annotations

import argparse
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SECTION_RE = re.compile(r"^(?:第[一二三四五六七八九十]+章|\d+(?:\.\d+){0,3})(?:\s|　|$)")
CAPTION_RE = re.compile(r"^(?:圖|表|Figure|Table)\s*[0-9０-９一二三四五六七八九十]+[-－.]?")


def paragraph_outline_level(paragraph) -> str | None:
    ppr = paragraph._p.pPr
    if ppr is None:
        return None
    outline = ppr.find(qn("w:outlineLvl"))
    return outline.get(qn("w:val")) if outline is not None else None


def run_summary(paragraph) -> dict:
    fonts = Counter()
    sizes = Counter()
    bold = 0
    italic = 0
    for run in paragraph.runs:
        if run.font.name:
            fonts[run.font.name] += len(run.text) or 1
        if run.font.size:
            sizes[round(run.font.size.pt, 2)] += len(run.text) or 1
        if run.bold:
            bold += len(run.text) or 1
        if run.italic:
            italic += len(run.text) or 1
    return {
        "fonts": dict(fonts.most_common()),
        "sizes_pt": {str(key): value for key, value in sizes.most_common()},
        "bold_chars": bold,
        "italic_chars": italic,
    }


def inspect(input_path: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    doc = Document(str(input_path))
    paragraphs = []
    heading_candidates = []
    captions = []
    style_counts = Counter()
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style is not None else ""
        style_counts[style] += 1
        record = {
            "index": index,
            "style": style,
            "outline_level": paragraph_outline_level(paragraph),
            "text": text,
            "run_summary": run_summary(paragraph),
        }
        paragraphs.append(record)
        lower_style = style.lower()
        if text and ("heading" in lower_style or "標題" in style or SECTION_RE.match(text)):
            heading_candidates.append(record)
        if text and CAPTION_RE.match(text):
            captions.append(record)

    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row_index, row in enumerate(table.rows):
            rows.append({
                "row": row_index,
                "cells": ["\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip()) for cell in row.cells],
            })
        tables.append({
            "index": table_index,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "preview": rows[:5],
        })

    sections = []
    for index, section in enumerate(doc.sections):
        sections.append({
            "index": index,
            "start_type": str(section.start_type),
            "page_width_in": round(section.page_width.inches, 4),
            "page_height_in": round(section.page_height.inches, 4),
            "left_margin_in": round(section.left_margin.inches, 4),
            "right_margin_in": round(section.right_margin.inches, 4),
            "top_margin_in": round(section.top_margin.inches, 4),
            "bottom_margin_in": round(section.bottom_margin.inches, 4),
        })

    with zipfile.ZipFile(input_path) as archive:
        names = archive.namelist()
        media = [name for name in names if name.startswith("word/media/")]
        comments_present = "word/comments.xml" in names
        footnotes_present = "word/footnotes.xml" in names
        endnotes_present = "word/endnotes.xml" in names
        document_xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        tracked_insertions = document_xml.count("<w:ins")
        tracked_deletions = document_xml.count("<w:del")

    inventory = {
        "input": str(input_path),
        "size_bytes": input_path.stat().st_size,
        "paragraph_count": len(paragraphs),
        "nonempty_paragraph_count": sum(bool(p["text"]) for p in paragraphs),
        "table_count": len(tables),
        "section_count": len(sections),
        "inline_shape_count": len(doc.inline_shapes),
        "media_part_count": len(media),
        "comments_present": comments_present,
        "footnotes_present": footnotes_present,
        "endnotes_present": endnotes_present,
        "tracked_insertions": tracked_insertions,
        "tracked_deletions": tracked_deletions,
        "style_counts": dict(style_counts.most_common()),
        "sections": sections,
        "heading_candidates": heading_candidates,
        "captions": captions,
        "tables": tables,
        "media_parts": media,
    }

    (output_dir / "v7_inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with (output_dir / "v7_paragraphs.tsv").open("w", encoding="utf-8", newline="") as stream:
        stream.write("index\tstyle\toutline_level\ttext\n")
        for paragraph in paragraphs:
            safe = paragraph["text"].replace("\t", " ").replace("\r", " ").replace("\n", "\\n")
            stream.write(f"{paragraph['index']}\t{paragraph['style']}\t{paragraph['outline_level'] or ''}\t{safe}\n")
    with (output_dir / "v7_fulltext.txt").open("w", encoding="utf-8") as stream:
        for paragraph in paragraphs:
            if paragraph["text"]:
                stream.write(f"[{paragraph['index']:05d}] [{paragraph['style']}] {paragraph['text']}\n")

    print(json.dumps({
        "paragraphs": len(paragraphs),
        "headings": len(heading_candidates),
        "captions": len(captions),
        "tables": len(tables),
        "sections": len(sections),
        "inline_shapes": len(doc.inline_shapes),
        "media_parts": len(media),
        "output_dir": str(output_dir),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    inspect(args.input, args.output_dir)
