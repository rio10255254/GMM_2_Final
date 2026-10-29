from __future__ import annotations

import os
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
PROJECT_ROOT = Path(__file__).resolve().parents[3]


DOCUMENT = Path(
    str(PROJECT_ROOT / 'Experiments/R63_CausalBoundaryMotionAblation/results/thesis_v8_language_cleanup/thesis_v8_language_working.docx')
)
ARCHITECTURE_DIR = DOCUMENT.parent / "architecture_figures"
ABLATION_FIGURE_DIR = Path(
    str(PROJECT_ROOT / 'Experiments/R63_CausalBoundaryMotionAblation/figures')
)


def first_run_properties(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def make_run(text: str, rpr=None):
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(deepcopy(rpr))
    node = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    return run


def rewrite_endnote_paragraph(paragraph, before: str, after: str) -> None:
    p = paragraph._p
    rpr = first_run_properties(paragraph)
    children = list(p)

    begin_index = None
    end_index = None
    for index, child in enumerate(children):
        for fld in child.xpath(".//w:fldChar"):
            kind = fld.get(qn("w:fldCharType"))
            if kind == "begin" and begin_index is None:
                begin_index = index
            elif kind == "end":
                end_index = index

    if begin_index is None or end_index is None:
        raise RuntimeError("Expected EndNote citation field was not found")

    for child in children[1:begin_index]:
        p.remove(child)
    field_begin = list(p)[1]
    field_begin.addprevious(make_run(before, rpr))

    children = list(p)
    actual_end = None
    for index, child in enumerate(children):
        if any(f.get(qn("w:fldCharType")) == "end" for f in child.xpath(".//w:fldChar")):
            actual_end = index
            break
    if actual_end is None:
        raise RuntimeError("EndNote field end disappeared")
    for child in children[actual_end + 1 :]:
        p.remove(child)
    p.append(make_run(after, rpr))


def replace_plain_paragraph(paragraph, expected: str, text: str) -> None:
    if paragraph.text.strip() != expected:
        raise RuntimeError(f"Unexpected paragraph text: {paragraph.text!r}")
    rpr = first_run_properties(paragraph)
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    p.append(make_run(text, rpr))


def replace_cell_text(cell, expected: str, text: str) -> None:
    if cell.text.strip() != expected:
        raise RuntimeError(f"Unexpected cell text: {cell.text!r}")
    replace_plain_paragraph(cell.paragraphs[0], expected, text)


def replace_media(document_path: Path) -> None:
    replacements = {
        "word/media/image41.png": ARCHITECTURE_DIR / "figure_5_19_main.png",
        "word/media/image42.png": ARCHITECTURE_DIR / "figure_5_19_legend.png",
        "word/media/image43.png": ARCHITECTURE_DIR / "figure_5_20_psf_fusion.png",
        "word/media/image92.png": ABLATION_FIGURE_DIR / "15_full_vs_raw_cd_error.png",
        "word/media/image93.png": ABLATION_FIGURE_DIR / "16_full_vs_raw_step_stability.png",
        "word/media/image94.png": ABLATION_FIGURE_DIR / "17_full_vs_raw_grid_agreement.png",
        "word/media/image95.png": ABLATION_FIGURE_DIR / "18_raw_position_anchor_cd_error.png",
        "word/media/image96.png": ABLATION_FIGURE_DIR / "19_upsampling_grid_agreement.png",
        "word/media/image97.png": ABLATION_FIGURE_DIR / "20_full_vs_raw_ringing_tradeoff.png",
    }
    for source in replacements.values():
        if not source.exists():
            raise FileNotFoundError(source)

    temporary = document_path.with_suffix(".media-rebuild.docx")
    with zipfile.ZipFile(document_path, "r") as source_zip, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED
    ) as output_zip:
        names = set(source_zip.namelist())
        missing = set(replacements) - names
        if missing:
            raise RuntimeError(f"Missing media entries: {sorted(missing)}")
        for item in source_zip.infolist():
            if item.filename in replacements:
                output_zip.writestr(item, replacements[item.filename].read_bytes())
            else:
                output_zip.writestr(item, source_zip.read(item.filename))
    os.replace(temporary, document_path)


def main() -> None:
    if not DOCUMENT.exists():
        raise FileNotFoundError(DOCUMENT)

    document = Document(str(DOCUMENT))
    if len(document.paragraphs) != 1247:
        raise RuntimeError(f"Unexpected paragraph count: {len(document.paragraphs)}")

    rewrite_endnote_paragraph(
        document.paragraphs[190],
        "PSF校正除了要取得成像核，也必須評估標靶、平台與模型假設帶來的誤差。Li等人指出，在三維單分子定位顯微鏡中，離焦會同時改變PSF形狀與尺度，平台讀取的深度不宜直接視為真實深度；其研究因此以多顆微球的量測PSF與理論PSF擬合，建立平台讀值與推定深度之間的校正關係",
        "。該流程可降低系統性偏差，但若PSF仍有模型失配或量測雜訊，結果的散布不一定同步下降。這說明準確度與精密度必須分開評估。",
    )

    caption_updates = {
        1138: (
            "表 6-13 模組與損失函數消融之匹配比較條件",
            "表 6-13 模組與損失函數消融實驗之比較條件",
        ),
        1144: (
            "圖 6-47 完整控制模型相對Raw之兩倍率SEM對照CD誤差",
            "圖 6-47 完整MOSR模型相對Raw之兩倍率SEM對照CD誤差",
        ),
        1148: (
            "圖 6-48 完整控制模型相對Raw之四組真實序列步間穩定性",
            "圖 6-48 完整MOSR模型相對Raw之四組真實序列步間穩定性",
        ),
        1151: (
            "圖 6-49 完整控制模型相對Raw之九顆pillar位移一致性",
            "圖 6-49 完整MOSR模型相對Raw之九顆pillar位移一致性",
        ),
        1155: (
            "圖 6-50 移除原始影像位置保留對兩倍率SEM對照CD誤差之影響",
            "圖 6-50 原始影像位置保留對兩倍率SEM對照CD誤差之影響",
        ),
        1161: (
            "圖 6-51 抗混疊上採樣與一般上採樣之九顆pillar位移一致性",
            "圖 6-51 抗混疊與一般上採樣之九顆pillar位移一致性",
        ),
        1166: (
            "表 6-14 核心模組、量測導向損失函數與支援模組之消融判讀",
            "表 6-14 主要模組、量測導向損失函數與輔助模組之消融結果",
        ),
        1170: (
            "圖 6-52 完整控制模型相對Raw之外圍振鈴取捨",
            "圖 6-52 完整MOSR模型相對Raw之外圍振鈴比較",
        ),
    }
    for index, (expected, text) in caption_updates.items():
        replace_plain_paragraph(document.paragraphs[index], expected, text)

    table109 = document.tables[109]
    replace_cell_text(table109.cell(1, 0), "完整控制與2×2因子", "完整模型與2×2因子設計")
    replace_cell_text(table109.cell(2, 0), "細分損失移除", "單一損失移除")
    replace_cell_text(table109.cell(4, 0), "上採樣匹配", "上採樣方式比較")

    table110 = document.tables[110]
    replace_cell_text(table110.cell(3, 0), "廣義幾何錨定損失", "幾何約束損失")
    replace_cell_text(table110.cell(4, 0), "廣義運動一致性損失", "運動一致性損失")

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(str(DOCUMENT))
    replace_media(DOCUMENT)

    print(f"Saved: {DOCUMENT}")
    print("Updated EndNote paragraph, captions, table labels, and 9 embedded figures")


if __name__ == "__main__":
    main()
