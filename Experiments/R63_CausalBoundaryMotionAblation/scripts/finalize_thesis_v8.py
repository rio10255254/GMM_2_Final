from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
PROJECT_ROOT = Path(__file__).resolve().parents[3]


DOCUMENT = Path(
    str(PROJECT_ROOT / 'Experiments/R63_CausalBoundaryMotionAblation/results/thesis_v8_work/thesis_v8_working.docx')
)

RESULTS_TEXT = (
    "結果顯示，Wiener反卷積可提高對比與邊緣銳利度，MOSR則在跨門檻CD與多pillar位移一致性上較穩定。"
    "配對消融顯示，一般上採樣使四組序列合併的九顆pillar不一致量增加14.31%；移除原圖位置錨定使20倍CD"
    "中位絕對誤差增加73.97%，並使兩倍率振鈴明顯惡化。其餘單項損失未形成可獨立歸因的廣泛效果。"
    "因此，主要學術貢獻界定為具直接證據的模組設計，以及依量測失敗形式建立的功能性損失設計，而非整合流程本身。"
)

KEYWORDS_TEXT = (
    "關鍵字：量測導向超解析、位移安全上採樣、原圖位置錨定、功能性損失、關鍵尺寸、平移定位。"
)

CD_AUDIT_TEXT = (
    "全體量測之中位絕對CD誤差由Raw的192.1 nm降至Wiener的161.2 nm與MOSR的117.0 nm，"
    "降幅分別為16.1%與39.1%，其中MOSR最低。"
)

INNOVATION_SCOPE_TEXT = (
    "消融結果將本論文的學術創新明確界定為可被檢查的模型與訓練設計，而光學校正、影像重建與後端量測僅作為驗證架構。"
    "第一，位移安全上採樣具有跨pillar位置一致性的直接證據；第二，原圖位置錨定具有邊界與亮暗拓樸安全的直接證據，"
    "並在20倍CD上呈現明顯收益；第三，功能性損失將量測失敗形式轉為訓練目標，但其作用應以整體設計描述，"
    "不拆成每一項均已獨立證明。"
)


def replace_visible_text(paragraph, text: str) -> None:
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = deepcopy(run._r.rPr)
            break

    element = paragraph._p
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)

    run = paragraph.add_run(text)
    if first_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_rpr)


def find_one(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def find_one_any(document: Document, prefixes: tuple[str, ...]):
    matches = [
        p
        for p in document.paragraphs
        if any(p.text.strip().startswith(prefix) for prefix in prefixes)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefixes!r}, found {len(matches)}")
    return matches[0]


def main() -> None:
    if not DOCUMENT.exists():
        raise FileNotFoundError(DOCUMENT)

    document = Document(str(DOCUMENT))
    replace_visible_text(find_one(document, "結果顯示，Wiener反卷積"), RESULTS_TEXT)
    replace_visible_text(find_one(document, "關鍵字：量測導向超解析"), KEYWORDS_TEXT)
    replace_visible_text(
        find_one_any(
            document,
            ("跨全部量測條件後", "綜合全部量測後", "全體量測之中位絕對CD誤差"),
        ),
        CD_AUDIT_TEXT,
    )
    replace_visible_text(
        find_one_any(
            document,
            ("消融結果使本論文的創新主軸", "消融結果將本論文的學術創新明確界定"),
        ),
        INNOVATION_SCOPE_TEXT,
    )
    document.save(str(DOCUMENT))
    print(f"Saved: {DOCUMENT}")
    print(f"Paragraphs: {len(document.paragraphs)}")
    print(f"Tables: {len(document.tables)}")
    print(f"Inline shapes: {len(document.inline_shapes)}")


if __name__ == "__main__":
    main()
