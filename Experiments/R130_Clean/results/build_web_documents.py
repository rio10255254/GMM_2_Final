from pathlib import Path
import json,hashlib,shutil,zipfile,re,collections,math
import numpy as np,pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
PROJECT_ROOT = Path(__file__).resolve().parents[3]
R=Path(str(PROJECT_ROOT / 'Experiments/R130_Clean'))
H=Path(str(PROJECT_ROOT / 'Deliverables/Thesis_Handoffs/R130_Clean_Corrected_Corner9_Validation_Handoff_20260920'))
P=Path(str(PROJECT_ROOT / 'Deliverables/Thesis_Handoffs/R130_Clean_Web_Thesis_Documents_20260920'))
assert not (P/"MANIFEST.json").exists(), "Do not overwrite sealed delivery"
P.mkdir(exist_ok=True)
(P/'完整數據').mkdir(exist_ok=True);(P/'圖表').mkdir(exist_ok=True);(P/'閱讀文字版').mkdir(exist_ok=True);(P/'製作與檢查').mkdir(exist_ok=True)
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
# Validate the sealed source package before reusing any report.
for x in load(H/'MANIFEST.json'):assert sha(H/x['path'])==x['sha256'],x['path']
S=R/'results/comparison'
tables={p.stem:pd.read_csv(p) for p in S.glob('*.csv')}
meta=[load(R/'data'/x['metadata']) for x in load(R/'data/manifest.json')['entries']]
mdf=pd.DataFrame(meta)
counts=pd.crosstab(mdf.category,mdf.split).reset_index().rename(columns={'category':'資料類型'})
noise=mdf.groupby('split').noise_sigma_gray.agg(['min','max']).reset_index()
# Preserve every table without rounding in the data companion.
for src in sorted(S.glob('*.csv')):shutil.copy2(src,P/'完整數據'/src.name)
for name in ['individual_circle_displacements','nine_pillar_size_shape','geometry_by_position','angular_repeatability']:
 shutil.copy2(H/'02_R129_R130_Comparison_and_Runtime'/(name+'.csv'),P/'完整數據'/(name+'.csv'))
for name in ['time_block_sensitivity','mean_shape_amplitudes','independently_recomputed_displacement']:
 shutil.copy2(R/'results/independent_review'/(name+'.csv'),P/'完整數據'/(name+'.csv'))
shutil.copy2(R/'results/reconstruction_runtime.csv',P/'完整數據/reconstruction_runtime.csv')
shutil.copy2(H/'05_Verification_and_Sources/all_930_image_sources.csv',P/'完整數據/all_930_image_sources.csv')
shutil.copy2(R/'data/manifest.json',P/'完整數據/training_data_manifest.json')
counts.to_csv(P/'完整數據/training_category_counts.csv',index=False,encoding='utf-8-sig')
mdf.to_json(P/'完整數據/training_metadata.json',orient='records',force_ascii=False,indent=2)
for name in ['profile.json','profile_tiled_fresh.json','preflight.json','independent_final_training_audit.json','clean_control_summary.json','final_review_complete.json']:
 shutil.copy2(R/'results'/name,P/'完整數據'/name)
shutil.copy2(R/'config_r130_clean.yaml',P/'完整數據/config_r130_clean.yaml')
shutil.copy2(R/'runs/formal_8192/updates.jsonl',P/'完整數據/training_updates.jsonl')
# Preserve measured boundary arrays as optional originals, not millions of fabricated independent observations.
for src in (H/'05_Verification_and_Sources/measured_centers').rglob('*.npz'):
 dest=P/'完整數據/輪廓原始點'/src.relative_to(H/'05_Verification_and_Sources/measured_centers');dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src,dest)
for src in S.glob('*radial_profiles.npz'):shutil.copy2(src,P/'完整數據'/src.name)
figs={}
for name,src in {
 'training':R/'reports/figures/training_curves.png',
 'displacement':H/'02_R129_R130_Comparison_and_Runtime/figures/displacement_mean_sd.png',
 'repeatability':H/'02_R129_R130_Comparison_and_Runtime/figures/geometry_repeatability.png',
 'pillars':H/'02_R129_R130_Comparison_and_Runtime/figures/per_pillar_repeatability.png',
 'contour0':H/'01_Raw_R130_Comparison/figures/contours_position00.png',
 'contour30':H/'01_Raw_R130_Comparison/figures/contours_position30.png',
 'selected':H/'03_Sampling_and_SEM/inherited_sampling_reference/selected_nine.png',
 'example':R/'reports/figures/native_nine_example.png',
 'fullfield':R/'results/first_frame_visual_check.png',
 'SEMshape':R/'reports/figures/independent_mean_shapes.png',
 'worst':R/'reports/figures/worst_adjacent_observation.png',
}.items():
 dest=P/'圖表'/(name+'.png');shutil.copy2(src,dest);figs[name]=dest
class Report:
 def __init__(self,num,title,intro):
  self.stem=f'{num:02d}_{title}';self.doc=Document();sec=self.doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
  for n in ['Normal','Title','Heading 1','Heading 2','Caption']:
   st=self.doc.styles[n];st.font.name='Microsoft JhengHei';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei');st.font.color.rgb=RGBColor(0,0,0)
  
  for border in self.doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
  st=self.doc.styles['Normal'];st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.15
  self.doc.styles['Title'].font.size=Pt(23);self.doc.styles['Heading 1'].font.size=Pt(16);self.doc.styles['Heading 2'].font.size=Pt(12)
  foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('R130 Clean　2026 09 20　｜　');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
  self.md=[f'# {title}',intro];self.doc.add_paragraph(title,'Title');self.doc.add_paragraph('網頁版論文構思用獨立資料文件　2026年9月20日');self.doc.add_paragraph(intro)
 def h(self,text):self.doc.add_heading(text,1);self.md+=['## '+text]
 def p(self,text):
  for line in text.strip().split('\n\n'):self.doc.add_paragraph(line);self.md.append(line)
 def table(self,df,title):
  df=df.copy();self.p(title);self.doc.paragraphs[-1].paragraph_format.keep_with_next=True
  for c in df:
   if pd.api.types.is_float_dtype(df[c]):df[c]=df[c].map(lambda v:'' if pd.isna(v) else f'{v:.6f}')
  table=self.doc.add_table(rows=1, cols=len(df.columns));table.style='Light Shading Accent 1'
  for cell,col in zip(table.rows[0].cells,df.columns):cell.text=str(col)
  repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
  for row in df.itertuples(index=False,name=None):
   cells=table.add_row().cells
   for cell,val in zip(cells,row):cell.text=str(val)
  for row in table.rows:
   pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
   for cell in row.cells:
    for p in cell.paragraphs:
     p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
     for run in p.runs:run.font.size=Pt(8.5)
  self.md.append("\n".join(["|"+"|".join(map(str,df.columns))+"|", "|"+"|".join(["---"]*len(df.columns))+"|"]+["|"+"|".join(str(v).replace("|","/") for v in row)+"|" for row in df.itertuples(index=False,name=None)]))
 def fig(self,key,caption):
  from PIL import Image
  w,h=Image.open(figs[key]).size;width=min(6.2 if key=='selected' else 6.9,7.4*w/h)
  p=self.doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(figs[key]),width=Inches(width));self.doc.add_paragraph(caption,'Caption');self.md.append(f'圖 {key}：{caption}。圖像已嵌入同名 Word 文件；獨立圖片為 圖表/{key}.png。')
 def finish(self):
  self.doc.save(P/(self.stem+'.docx'));(P/'閱讀文字版'/(self.stem+'.md')).write_text('\n\n'.join(self.md),encoding='utf-8')
# 01
x=Report(1,'閱讀總覽與研究判斷','這組文件讓未接觸本機資料的讀者理解 R130 Clean 的模型、訓練、真實驗證與論文改寫方向。所有主要數字來自 2026年9月20日已完成的實驗。核心結論是：短步位移表現與 R129 接近，尺寸重拍更穩定、處理较快，但尺寸偏大仍存在。是否改用 R130 作為論文唯一模型，尚未由本次整理自動決定。')
x.h('文件閱讀順序');x.table(pd.DataFrame([
 ['01','研究全貌、關鍵結論、證據界線'],['02','實際模型結構、全部訓練設定、資料來源'],['03','930 張真實影像、ACM 與 LSC、SEM、統計定義'],['04','主要完整數據、逐位置和逐顆表、全部主要圖'],['05','對照原論文章節的改寫方案與可用敘述']],columns=['文件','內容']),'先讀本文件，再依序閱讀 02 至 05。每份均可獨立理解，不靠電腦上的絕對路徑補全文義。')
x.h('研究真正要回答的問題');x.p('在正常曝光、照明與相機設定可固定的顯微拍攝中，影像重建能否使輪廓定位及位移量測更穩定，同時保留可檢查的訓練來源與處理步驟？最終應用只需要一個模型。研究目的不是證明影像越銳利越好，也不是追求原樣本不具有的完美圓形。\n\n現有實驗分成三個問題：重拍是否一致、與 PI 位移是否接近、與 SEM 尺寸及形狀是否相符。三者不能混成一個「最好」分數。')
x.h('主要結果速讀');x.table(pd.DataFrame([
 ['相鄰位移平均絕對差 nm',8.166620,4.125379,4.121476],['起點位移平均絕對差 nm',19.080973,17.984507,17.948881],['直徑重拍標準差 nm',32.750251,19.230772,16.421786],['輪廓重拍標準差 nm',63.758024,18.436707,18.579026],['去尺寸形狀重拍標準差 nm',61.666738,15.788824,16.669974],['平均直徑高於 SEM nm',185.182117,224.472228,219.408680]],columns=['指標','Raw','R129','R130']),'Raw 為未經神經重建的原圖。完整未截短數值與定義見文件 03、04。')
x.p('相鄰位移 R130 與 R129 的差值只有 −0.003902 nm，配對 95% 區間包含零，因此不能宣稱位移顯著超越 R129。R130 直徑重拍標準差較 R129 約降低 14.6%；去尺寸形狀重拍標準差反而約高 5.6%。輪廓更穩定，並不等於尺寸更接近 SEM。')
x.h('哪個模型特色可以寫');x.p('R130 實際具備原尺寸運算、四次共用網路修正、以原圖為參考的有界修正、局部影像線索、從零開始且只用合成配對的訓練。R130 實測 PSF 輸入為關閉，也没有抗混疊上採樣模組，因為這一版沒有上採樣路徑。\n\n這些是可以查證的設計，不等於已證明每個模組都有獨立效益、每次修正都有物理意義，或已成立學術上的首次創新。未執行消融，也未做本次文獻新穎性檢索。')
x.h('如何交給下一個網頁對話');x.p('請將五份 Word 文件一起提供；需要精確查表時，再提供相應 CSV，或閱讀文字版。所有 Word 圖片已嵌入檔案；若接收端只讀到文字，可另上傳圖表資料夾中的圖片。完整數據壓縮包必須解開後選取所需檔案，不能假設接收端必定會自動讀完壓縮檔。\n\n建議交辦文字：請先讀完這五份文件，區分已完成證據、研究解釋和未完成事項。以單一最終模型、正常拍攝、LSC 位移與 ACM 輪廓為主軸，提出盡量保留七章架構的改寫方案。先指出採用 R130 會改變哪些原有創新，不要自行把 PSF 融合、抗混疊上採樣、消融或完整 XAI 寫回去。所有量化敘述必須對回文件中的實際方法與數值。')
x.h('本組文件的範圍');x.p('本組文件完整呈現本輪既有驗證，不代表已完成所有可能實驗。沒有新增獨立樣本、未見過的真實保留集、完整不確定度合成、重新公平訓練全部歷史模型或新的消融。大批量逐筆資料附在完整數據；文件 04 收錄適合直接閱讀的彙總與逐位置、逐顆結果。\n\nR119 舊完成宣稱及其後被推翻的收據不能用來補強本次模型。本次證據以 R130 的訓練、重建、量測及独立檢查記錄為準；來源與查核不等於把既有真實影像重新變成未使用過的資料。');x.finish()
# 02
x=Report(2,'模型結構與訓練資料完整規格','本文件描述已實際完成的 R130 Clean，供撰寫方法與實作章。不是未執行的設計提案。模型最後訓練 8192 次、300101 個參數；未載入舊神經權重，實測光學輸入關閉。')
x.h('完整資料流');x.p('單通道灰階原圖 x 先轉為 0 到 1，模型全程維持原影像尺寸。初始修正 r 為零，初始影像 z 等於 x。每一回合把目前 z、原图 x、z−x、原圖水平方向差、垂直方向差、拉普拉斯差分、5×5 局部均值七個通道交给同一套網路。\n\n網路先以 3×3 卷積將 7 通道變成 64 通道，再經四個殘差區塊，最後 3×3 卷積輸出一個通道的建議修正。四個區塊的膨脹間距為 1、2、4、2；每個區塊有兩層卷積，第一層之後使用 SiLU，區塊輸出乘 0.1 後加回。輸入層之後也有 SiLU。')
x.table(pd.DataFrame([['輸入','1 通道，FP32，0 至 1'],['局部線索','z、x、z−x、dx、dy、lap、5×5 mean'],['主幹','7→64；四個雙卷積區塊；64→1'],['四區塊間距','1、2、4、2'],['修正回合','4，主幹權重共用'],['每回合步幅','0.25×sigmoid(s)，初始 0.05，各回合獨立可學'],['累積修正範圍','−0.35 到 +0.35'],['輸出範圍','0 到 1'],['參數總數','300101'],['有效依賴半徑','82 原圖像素'],['分塊保留邊界','96 像素，核心 512']],columns=['項目','實際設定']),'實際模型設定，對應執行前固定的設定檔及程式。')
x.h('每次修正如何限制');x.p('主幹先輸出尚未限制的數值，僅經一次 tanh 將建議修正限制在 −1 到 1。當回合步幅為 0.25 乘上可學參數的 sigmoid；把步幅乘以建議修正後加到累積修正 r，再將 r 限制在 ±0.35。最後 z 等於原圖 x 加上 r，並限制到 0 到 1。四回合都以同一張原圖為參考。\n\n末層卷積的初始權重與偏置為零，所以訓練前輸出精確等於原圖。這僅是初始化的性質，不能寫成訓練後所有清楚影像都不變。固定上限也是設計選擇，沒有經本輪消融證明最優。')
x.h('局部性與原尺寸設計');x.p('所有卷積採鏡射邊界。沒有全圖均值或標準差作為條件，沒有頻域處理、全局品質閘門、末端高通注入、亮度重新置中，也没有 2 倍特徵網格或 4 倍隱藏影像狀態。這有利於檢查局部影響與分塊一致性，但並不單獨證明任意亞像素位移等變或細節都真實。\n\n完整影像與分塊最大差為 0.000000149011612，低於預設 0.000002 容許差。分塊只複製有效核心，不羽化混合。真實 930 張使用完整影像方式；分塊是另行檢查，不是另一套量測結果。')
x.h('合成資料的來源與構成');x.p('沿用 R129 獨立訓練包內的既有合成配對，母場景識別沿用 r128 前綴，不代表載入 R128 模型。共 768 個母場景，訓練 512、開發 128、既有生成測試 128。每個場景含一對影像及已知合成移動。\n\n本輪只原值提取 source、noiseless、target。source 是量化含噪輸入，noiseless 用來在每個訓練循環重抽噪聲，target 是合成乾淨目標。沒有將原存的合成 kernel4、隱藏幾何或光學係數輸入模型。原值相同核對不等於重新獨立建立資料。\n\n原資料生成記錄包含連續幾何與內在紋理位移、固定於相機的照明、合成模糊、像素面積積分、独立噪聲與 uint8 量化。原場景使用較高解析度生成，最後模型仍只吃原尺寸影像。重要區別：本輪不讀實測 PSF，並不代表合成訓練影像生成時沒有模擬模糊。合成答案是監督學習的目標，不是把真實待驗證 SEM 答案混進訓練。')
x.table(counts,'所有 768 個場景的類型與分組數量，由本地資料清單逐筆統計。')
x.table(noise.rename(columns={'split':'分組','min':'最小噪聲灰階','max':'最大噪聲灰階'}),'既有各組噪聲設定的範圍。訓練每循環重抽噪聲，開發與測試固定。')
x.p('各組母場景 ID 與幾何雜湊交集均為零；此項可排除清單所記錄的同一母場景重複，但不是所有統計相似性或真實世界泛化的證明。沒有實拍、SEM 尺寸、SEM 輪廓或 PI 位移作為本輪訓練監督。')
x.h('裁切與配對監督');x.p('資料原尺寸 384×384，訓練取座標 32 到 351 的 320×320 區塊。僅監督這個區塊的 96 到 223，即中間 128×128；對應原場景 128 到 255。這個外圍上下文用來隔離邊界對中央監督的影響，不是把整張圖都拿來算損失。\n\n同一對 A、B 共享場景與合成位移。配對項對齊的是「重建减乾淨目標」的殘差，而不是強迫兩張不同位置的影像相同。有可觀測位移時才啟用；非可觀測場景配對項為零。位移對齊採雙線性插值，中央監督位置遠離補零邊界。')
x.h('損失與訓練設定');x.p('總損失為影像項加 0.2 倍梯度項，再加 0.2 倍配對殘差項。穩健灰階誤差對殘差 e 計算 255 乘以平均值「平方根(e² 加 ε²) 減 ε」，ε 為 0.5/255；程式以等價的分式計算，以減少抵消。這不是 PSNR、不是 nm，也不是每像素絕對誤差的簡單算術平均。\n\n梯度項是中央監督殘差的相鄰水平差與相鄰垂直差各算穩健誤差後取平均。配對項是 B 殘差減去按已知合成位移對齊的 A 殘差，再算同一穩健誤差。')
x.table(pd.DataFrame([['初始化','全新，沒有舊權重、優化器、EMA 或步幅'],['正式更新','8192 次；16×512'],['優化器','AdamW；權重衰減 0.00001'],['學習率','最高 0.001，最低 0.00001'],['排程','前 128 次暖身，之後餘弦下降'],['梯度限制','5.0'],['模型 隨機種子','1309001'],['順序 隨機種子','1309002，與循環編號共同決定'],['噪聲 隨機種子','1309003，與循環及場景編號共同決定'],['數值設定','FP32，TF32 關閉，確定性設定'],['開發檢查','第 0 次及每 512 次，共 17 次'],['保留模型','固定最後第 8192 次，不按真實效果挑選'],['正式訓練時間 秒','3915.683392'],['訓練前試跑','另起 32 次的獨立模型，完成後丟棄']],columns=['項目','設定']),'所有設定對應本輪實際執行，沒有為得到較佳真實結果延長訓練。')
x.fig('training','正式訓練與固定開發檢查曲線；訓練曲線依 128 次更新區塊彙整。')
rows=[]
for path in sorted((R/'runs/formal_8192/evaluations').glob('dev_*.json')):
 v=load(path)
 # Keep the full original JSON separately if schema differs.
 shutil.copy2(path,P/'完整數據'/path.name)
x.table(pd.DataFrame([['初始開發',5.724764379,4.642031369,3.745995140,1.667669470],['最終開發',2.484037877,1.990103397,1.684617807,.785054771],['最終生成測試',1.761407281,1.395734119,1.207525678,.620840176]],columns=['組別','總損失','影像項','梯度項','配對項']),'初始與最終分項均取自實際評估；全部原始開發記錄附於完整數據。')
x.h('模型識別與核對範圍');x.p('最終模型 SHA256 為 559afeff214ff15e92ce910d7df7ead12994074b6173dabe9a7f21a8b40372cc。初始參數 SHA256 為 0d4a4b3ed24b224df65640a3e0ec243c24b270dde6204ac5de00c40707e3a4a3。核對包含全部 8192 次順序、種子與損失算術、優化器步數，以及抽選 48 個訓練輸入的獨立重建；不是把全部 8192 份輸入都重新生成過。\n\n此證據支持可追溯、可核查的執行，不能單憑它證明每個卷積特徵的語義或神經修正的物理因果性。');x.finish()
# 03
x=Report(3,'真實影像驗證方法與指標定義','本文件使接收者可以正確解讀 R130 的全部量測結果。採用同一批 30 趟重拍、每趟 31 個位置、同一組九顆，共 930 張與每方法 8370 次圓量測。Raw、R129、R130 的共同量測流程保持一致。')
x.h('資料和實驗單位');x.p('影像為正常曝光的真實光學顯微影像，20×，Basler acA2040-120um，1544×2064，曝光 3000 μs。相機與照明可固定；這批驗證不主張低光、任意曝光或跨儀器效果。每趟從位置 0 到 30，名義相鄰約 1 μm，以實際 PI 回授差為比較參考，不用名義 1 μm 當答案。\n\n採同一樣本與同一工作階段的重拍，早 6 趟和晚 24 趟之間存在時間間隔。30 趟是配對重抽樣的基本單位；同一影像的九顆、同一輪廓的 64 點和插值的 720 方向都不視為獨立重拍。這批真實資料曾被既有研究使用，不是未見過的保留資料。')
x.fig('selected','本輪實際測量的角落九顆；沿用更正後 R129 的選取，沒有為 R130 重新挑顆粒。')
x.h('相同輪廓量測步驟');cfg=load(R/'evaluation/measurement_config.json');x.p('先使用同一套原圖追蹤種子，逐顆取局部區域，執行 ACM 輪廓演化，再沿輪廓法線取得正確極性的最大梯度邊界點。每圈保留 64 個實際量測點，使用最小平方圓擬合求圓心與半徑。曲率運算在每個目標重設，避免上一顆的狀態污染下一顆。三方法不各自換參數或依 PI 調整輪廓。')
x.p('固定設定：局部半寬 14 像素，初始半徑約 8.330159 像素，64 個取樣點，ACM 演化 13 次，法線半寬 2.4 像素，剖面步距 0.1 像素，平滑 σ=0.65，暗物體極性。完整設定同時附於本文資料包。')
shutil.copy2(R/'evaluation/measurement_config.json',P/'完整數據/measurement_config.json')
for n in ['capture_plan.json','calibration.json','sem_mapping_candidates.json','sem_scale.json']:shutil.copy2(R/'evaluation'/n,P/'完整數據'/n)
x.h('位移如何計算');x.p('先對每顆的兩位置圓心做差，得到九組水平及垂直位移；水平方向九個值取中位數，垂直方向九個值取中位數，組合成一個位移向量。先取分量中位數再投影，不能改成先投影再取中位數。\n\n固定移動軸為 (−0.9998766245119538, 0.015707824629190823)，乘固定尺度 0.1725 μm/px，得到影像沿軸位移。影像位移減去 PI 的 y 回授差，再乘 1000，為本文以 nm 表示的有號差。軸來自前六趟 Raw；主表仍報全部 30 趟，不能把這全部稱為完全獨立的軸校準後驗證。估計尺度 0.17238058808029322 μm/px 未用於主表。\n\n相鄰法每個 j 使用 j−(j−1)；起點法每個 j 使用 j−0。每種方法每種位移各有 30×30=900 筆；三方法兩種位移共 5400 筆。相對起點直接算端點差，不再次扣位移，也不加總逐段中位數。')
x.table(pd.DataFrame([['mae_nm','全部 900 筆有號差取絕對值後平均'],['bias_nm','全部 900 筆有號差平均'],['repeatability_nm','每個位置用 30 趟求樣本標準差，再平均 30 個標準差'],['maximum_error_nm','本批 900 筆的最大絕對差，非保證上限'],['endpoint mean std','只取 j=30 相對起點的 30 筆，求平均和樣本標準差']],columns=['欄位','定義']),'所有差值以 PI 為參考，尚未合成 PI 自身及尺度的完整不確定度。')
x.h('尺寸及輪廓的統計方式');x.p('直徑是擬合半徑乘 2，再乘 172.5 nm/px。逐固定位置、固定顆粒在 30 趟間求樣本變異數；對 31×9 組變異數取平均再開根號，為直徑重複性。不是所有不同顆粒直徑混在一起求一次標準差。\n\n輪廓先以每次擬合圓心置中，將 64 個量測點的極角與半徑作週期插值到 720 個共同方向。各位置、顆粒、方向在 30 趟間的樣本變異數取平均再開根號，為輪廓重複性。扣掉每條輪廓自身的平均半徑後，同樣計算去尺寸形狀重複性。\n\n置中會消除平移，所以此輪廓指標不能代替圓心位移指標；去平均半徑會消除整體大小，所以去尺寸形狀更接近也不能代替直徑準確度。輪廓疊圖只扣平移，没有任意縮放、旋轉最佳化或自由變形以美化吻合。')
x.h('SEM 如何對照');mapping=load(R/'evaluation/sem_mapping_candidates.json');x.table(pd.DataFrame([[i+1,str(v['sem_ids_by_new_nine']),str(v['candidate'])] for i,v in enumerate(mapping)],columns=['候選','九顆 SEM 對應順序','既有方向設定']),'九顆集合相同，逐顆對應方向仍有兩個候選；沒有按模型表現挑有利方向。')
x.p('SEM 參考點使用既有尺度並重新置中。參考平均直徑為 2614.438543257 nm。平均直徑差為方法平均直徑減去這個參考。完整徑向差為同方向半徑差的平均絕對值；去尺寸形狀差為雙方各自扣平均半徑後再取平均絕對差。\n\nSEM 的完整輪廓不是擬合成完美圓後再比較。樣本原本不完全圓，本文不把「圓度變高」當作成功目標。真實形狀保留應以輪廓及 SEM 相符程度解釋，並保持方向對應不確定性的限制。')
x.h('配對區間及時間敏感度');x.p('相鄰與起點的平均絕對差比較，先每趟求平均，再以同一趟 R130 減參考方法，30 趟整趟配對抽樣，共 2000 次，種子 130136；取重抽樣分布的第 2.5 與 97.5 百分位。尺寸與輪廓標準差則同樣重抽完整趟次，再重算整體統計，不能把每個點單獨重抽。\n\n另對位移差做時間區塊敏感度：早 6 趟及晚 24 趟各自以長度 3 趟的循環區塊抽樣，2000 次，種子 13013601，不跨時間間隔。它檢查特定短程時間相依的敏感程度，不是已排除所有漂移。\n\n所附 95% 區間是方法差值的不確定範圍，不是完整量測系統的擴展不確定度。圖中平均值±1 標準差是重拍分散，不是平均值的 95% 信賴區間。')
x.h('驗證檢查與證據限制');x.p('既有独立查核驗證 930 張重建內容識別各不相同且吻合紀錄，重新擬合新模型 8370 圈，最大座標差約 0.000000321 像素；重新計算整體位移差約 1.78×10⁻¹⁵ nm。後續封裝又從全部圓心重算 5400 筆位移，結果相符。\n\n這些查核支持算術與來源可靠，不能由此推論全新樣本、其他倍率、其他工作階段、任意曝光或所有亞像素位移都有效。若全文仍使用「絕對精度」「完整 XAI」「未見過的保留集」或「每個模組皆有效」等說法，現有证據不足。');x.finish()
# 04
x=Report(4,'完整實測數據與圖表','本文件完整列出本輪三方法主要結果、逐位置及逐顆統計和圖表。主表與細表保留至小數六位供核對，這個顯示精度不是實際量測能力。所有未四捨五入的逐筆 CSV 與原始輪廓點另附完整數據。')
x.h('位移主要結果')
for kind,cn in [('adjacent','相鄰位移'),('start','相對起點位移')]:
 v=tables['displacement_summary'];v=v[v.kind==kind].drop(columns='kind').rename(columns={'model':'方法','mae_nm':'平均絕對差','bias_nm':'有號平均差','repeatability_nm':'重拍標準差','maximum_error_nm':'最大絕對差'});x.table(v,cn+'，單位 nm，每方法 900 筆。')
x.fig('displacement','平均有號差與±1 個重拍標準差。左為相鄰，右為相對起點；不是 95% 區間。')
x.table(tables['endpoint_30um'].rename(columns={'model':'方法','mean':'平均差','std':'標準差','min':'最小差','max':'最大差'}),'約 30 μm 終點的 30 趟結果，單位 nm。')
x.h('尺寸與輪廓主要結果')
v=tables['shape_size_summary'];x.table(v[['model','diameter_mean_nm','diameter_repeat_sd_nm','diameter_repeat_absmax_nm']].rename(columns={'model':'方法','diameter_mean_nm':'平均直徑','diameter_repeat_sd_nm':'直徑重拍標準差','diameter_repeat_absmax_nm':'最大重拍偏離'}),'單位 nm；最大重拍偏離以同一位置和顆粒的 30 趟平均為參考。')
x.table(v[['model','contour_repeat_sd_nm','shape_repeat_sd_nm']].rename(columns={'model':'方法','contour_repeat_sd_nm':'輪廓重拍標準差','shape_repeat_sd_nm':'去尺寸形狀標準差'}),'輪廓包含半徑變動；去尺寸形狀扣掉每圈平均半徑。')
x.fig('repeatability','固定位置與固定顆粒的重拍統計。小值代表較一致，不能直接推論更準確。')
x.h('SEM 對照完整數值')
for cand in [1,2]:
 v=tables['SEM_comparison'];v=v[v.mapping_candidate==cand][['model','mean_diameter_difference_nm','full_radial_difference_nm','de_size_shape_difference_nm']].rename(columns={'model':'方法','mean_diameter_difference_nm':'有號直徑差','full_radial_difference_nm':'完整徑向差','de_size_shape_difference_nm':'去尺寸形狀差'});x.table(v,f'候選方向 {cand}，單位 nm。共同 SEM 平均直徑 2614.438543257 nm。')
x.p('方向一的去尺寸形狀 R130 較 R129 小，方向二則較大；不能只挑方向一。兩方向均顯示重建直徑仍較 SEM 大。輪廓重拍穩定與 SEM 相符程度是兩個問題。')
x.fig('SEMshape','九顆的平均去尺寸形狀。黑線 SEM 使用候選方向一，仍屬推定對應；方向二完整數字已列上表。')
x.h('整趟配對比較範圍')
for ref in ['Raw','R129']:
 v=tables['paired_scan_intervals'];v=v[v.reference==ref].drop(columns='reference');v['metric']=v.metric.replace({'adjacent_mae_nm':'相鄰平均絕對差','start_mae_nm':'起點平均絕對差','diameter_repeat_sd_nm':'直徑重拍標準差','contour_repeat_sd_nm':'輪廓重拍標準差','shape_repeat_sd_nm':'去尺寸形狀標準差'});v=v.rename(columns={'metric':'指標','difference_nm':'R130減參考','ci95_low':'95%下界','ci95_high':'95%上界','improved_scans':'改善趟數'});x.table(v,'參考方法 '+ref+'。負值代表 R130 較小；空白改善趟數不是零。')
v=pd.read_csv(R/'results/independent_review/time_block_sensitivity.csv').rename(columns={'kind':'位移種類','reference':'參考方法','mean_difference_nm':'平均差值','block3_separate_6_24_low_nm':'95%下界','block3_separate_6_24_high_nm':'95%上界'});x.table(v,'時間區塊敏感度的完整原表，差值單位 nm；欄位定義同文件 03。')
x.h('真實影像及輪廓展示');x.fig('fullfield','第一趟第一個位置的完整視野與局部示例，用來檢查整體是否保留。')
x.fig('example','同一真實影像區域的 Raw、R129、R130；三者共用顯示灰階範圍，視覺锐利程度不等於計量準確度。')
x.fig('contour0','起點九顆各 30 趟輪廓，灰色 Raw、橘色 R130；每條線為實際 64 點，置中但不縮放。')
x.fig('contour30','終點九顆各 30 趟輪廓，僅扣平移。')
x.fig('worst','R130 本批最大相鄰位移差個案，004747_XY_0，第 18 位置，20.414296 nm；沒有因此剔除該影像。')
x.h('速度與記憶體')
x.table(pd.DataFrame([['R130 本輪 930 張純重建中位數',1.38404875],['R130 本輪 930 張含讀寫中位數',1.86689665],['R129 歷史含讀寫中位數',11.18238745],['R130 完整影像暖機後 約',1.382],['R130 分塊暖機後 約',2.577]],columns=['時間範圍','秒每張']),'硬體 RTX A6000。R129 數字為歷史記錄，沒有本輪匹配重測，不能與 R130 純重建直接算倍數。')
x.table(pd.DataFrame([['完整影像',10865539072/2**30,19163774976/2**30],['獨立新程序分塊',1757112320/2**30,4301258752/2**30]],columns=['模式','峰值已配置 GiB','峰值保留 GiB']),'分塊不受前次完整影像的快取影響；這些數值不包含所有驅動程式或其他程序，不保證 4 GB 顯示卡可用。')
x.p('ACM 與圓量測平均約 0.0890 秒／張，中位約 0.09625 秒／張；這個時間不包含載入或神經重建。模型參數少並不保證快：運算仍受影像尺寸、四回合、通道數與實作方式影響。本次支持實測時間差，未做逐模組效率因果消融。')
x.h('九顆個別尺寸和輪廓結果')
v=pd.read_csv(H/'02_R129_R130_Comparison_and_Runtime/nine_pillar_size_shape.csv')
for model in ['Raw','R129','R130-Clean']:
 z=v[v.model==model].drop(columns='model').rename(columns={'pillar_id':'顆粒','diameter_mean_nm':'平均直徑','diameter_repeat_sd_nm':'直徑標準差','contour_repeat_sd_nm':'輪廓標準差','shape_repeat_sd_nm':'形狀標準差'});x.table(z,model+'，單位 nm；各顆數字彙整全部 31 位置及 30 趟。')
x.fig('pillars','九顆個別重拍結果。每一點仍是同一工作階段資料，不能當九個獨立樣本。')
x.h('每個位置的位移統計')
v=tables['position_statistics']
for kind in ['adjacent','start']:
 for model in ['Raw','R129','R130-Clean']:
  z=v[(v.kind==kind)&(v.model==model)][['position','mean_error_nm','repeat_sd_nm','mean_PI_um']].rename(columns={'position':'位置','mean_error_nm':'平均差 nm','repeat_sd_nm':'重拍SD nm','mean_PI_um':'PI均值 μm'});x.table(z,model+' / '+kind+'；每列 30 趟。')
x.h('詳細數據清單和讀取方式');x.p('完整數據內 all_circles.csv 共 25110 筆，all_displacements.csv 共 5400 筆，individual_circle_displacements.csv 共 48600 筆。另有 180 組逐位置位移、180 組逐趟位移統計、27 組逐顆尺寸輪廓、93 組逐位置尺寸輪廓、2160 組逐方向重複性，以及 930 張處理時間和來源。逐位置幾何完整表另外附於閱讀文字版 06。\n\n實際邊界點保留於輪廓原始點，每方法 30 個檔，每檔 points 為 31×9×64×2。radial_profiles 是將其轉成 720 個方向後的資料，屬插值，不能當作額外獨立觀測。');x.finish()
# 05 original thesis outline is read-only and never modified.
thesis=Path(r'C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文_v34_SEM與不利論述_完整標註版.docx')
old=Document(thesis);heads=[p.text for p in old.paragraphs if p.style.name.startswith('Heading')]
x=Report(5,'論文改寫構思與章節對照','本文件依原 v34 論文實際章節標題，以及本輪 R130 的可核對結果，提出改寫方案。這是規劃文件，沒有改動原 Word 論文，也沒有替使用者決定 R130 已正式取代 R129。若最終採 R130，必須一併調整原有光學融合與上採樣創新敘述。')
x.h('論文主線建議');x.p('建議主題聚焦「可追溯合成訓練與局部迭代重建對顯微位移及輪廓量測的影響」。中文或英文題名須由作者再定，不把尚未證明的新穎性、物理可解釋性或超解析能力寫進題名。\n\n正文以原圖和最終採用的一個模型為主；R129 對照作為模型選擇的内部證據或必要時的附錄。無須以比贏 V42 當換模型前提，也不能用「V42 訓練可能有答案」替代本模型的有效性證據。對舊模型的資料來源爭議若要寫，必須另引用已確認的來源稽核，而本包沒有重新做 V42 全面公平比較。')
x.h('哪些內容保留及調整');x.table(pd.DataFrame([
 ['第1章 緒論','保留量測需求；研究貢獻改為本次可證實的來源、架構與量測結果'],
 ['第2章 文獻回顧','保留成像、重建、量測導向背景；弱化與實際模型無關的架構介紹'],
 ['第3章 理論背景','保留光學成像、取樣、位移與輪廓理論；U Net 理論改成殘差與共享迭代'],
 ['第4章 研究方法','完整替換 4.4 模型、4.5 資料、4.6 損失；4.7 保留 ACM 與 LSC 主線'],
 ['第5章 系統與實作','保留實際硬體；5.4 用本次設定與訓練紀錄替换'],
 ['第6章 實驗結果','採原圖對單一模型；位移、輪廓重拍、SEM、時間依序呈現'],
 ['第7章 結論','分開可用性、重複性、尺寸偏差及適用範圍；不主張全部指標皆優']],columns=['章節','改寫方案']),'以原有七章架構為主，避免讓論文變成內部版本沿革。')
x.h('需要完整替換的模型小節');x.p('原 4.4.1 是同尺寸編碼器解碼器與抗混疊上採樣，4.4.2 是空間變動 PSF 融合，4.4.3 是方向感知與局部強度調節殘差，4.4.4 包含高通細節注入，4.4.5 為輸入品質感知調節。若採 R130，不能只把 V42 字樣换成 R130：這些小節必須按實際架構重寫。\n\n建議新 4.4.1 描述原尺寸七通道局部線索；4.4.2 描述共享殘差主幹；4.4.3 描述四次有界修正；4.4.4 描述原圖參考與初始恆等；4.4.5 描述有限影響範圍與分塊驗證。圖中的箭頭、通道、回授、步幅與原圖相加應與文件 02 完全一致。\n\n原實測 PSF 校準若有自身完整結果，可作为獨立的成像特性研究保留，但要明確說本次重建模型未使用該 PSF。若它是論文原先最主要創新，採 R130 將實質改變論文方法主張；這一點需要作者與指導教授決定，不能靠文字包裝掩蓋。')
x.h('訓練資料章的重寫');x.p('把舊四層多來源監督、真實或 SEM 錨點及微調敘述，替換成此次 512／128／128 個合成母場景、一對已知位移、固定原始目標和每循環重抽噪聲。說清楚目標如何來自合成幾何和成像模擬，不把合成訓練來源寫成從真實 SEM 取得答案。\n\n母場景分組與幾何識別零重疊可以寫，但生成測試和真實評估都是既有資料，不能寫成全新獨立未見測試。資料合成使用模糊機制，模型本身不讀實測光學參數，兩者需分开。損失改成三項實際目標，不沿用舊四組量測導向損失的名稱。')
x.h('結果章的建議順序');x.table(pd.DataFrame([
 ['6.2.1','模型與訓練完成狀態','文件02的設定與訓練曲線'],['6.2.2','真實資料與共同量測','文件03的九顆位置與流程'],['6.2.3','LSC位移與重拍範圍','文件04相鄰及起點表、平均±SD圖'],['6.2.4','ACM輪廓與尺寸重拍','文件04的九顆疊圖與逐顆統計'],['6.2.5','SEM尺寸及形狀對照','文件04兩方向候選表，分開直徑與去尺寸形狀'],['6.2.6','處理時間與適用範圍','本輪R130時間；如提R129須標歷史記錄'],['6.2.7','綜合討論','改善了什麼、沒有解決什麼、可追溯不等於完整XAI']],columns=['建議位置','內容','現有證據']),'這是建議新安排；原稿同號小節不應直接當作已經寫好。')
x.h('刪除或暫緩的部分');x.p('依使用者意向，消融實驗先刪除，包含原抗混疊上採樣拆除對照及其因果性結論。原七種外部模型與八模型競賽，若本次資料与流程並不一致，正文不直接沿用。原 SEM 36 顆、LSC 與 MZC 比較、舊微調結果和其他獨立數字不能改模型名稱後留存；沒有本包支持的新结果，應删除或明確標為獨立歷史研究。\n\n全篇也需要檢查摘要、英文摘要、貢獻條列、架構圖、圖說、表頭、參數量、訓練曲線、訓練資料來源、硬體時間與結論。不只替换方法章的模型名字。')
x.h('可直接據以改寫的結果段落');x.p('在同一批正常拍攝條件下取得的 30 趟重拍資料中，本方法的相鄰位移平均絕對 PI 差為 4.121 nm，原圖為 8.167 nm；相鄰位移重拍標準差由 9.725 nm 降至 4.184 nm。此結果支持本方法在本次樣本與工作階段下改善短步位移穩定性。\n\n尺寸和輪廓的重拍變動也下降：直徑重拍標準差由 32.750 nm 降至 16.422 nm，轮廓重拍標準差由 63.758 nm 降至 18.579 nm。然而，重建後平均直徑仍較 SEM 參考大約 219.409 nm，顯示重複性的改善不能直接解讀為尺寸真實度的改善。\n\n本研究採合成配對資料從零訓練，保存資料分組、設定、模型與逐筆量測來源，並以共同量測方式對照原圖和重建。此安排使執行與結果可追溯，但尚不足以宣稱學到的所有特徵皆具物理意義。')
x.h('可寫與不可直接寫');x.table(pd.DataFrame([
 ['本批正常條件下改善短步位移及輪廓重拍','任意樣本與曝光皆更準'],['原尺寸局部迭代且步驟可檢查','完整 XAI 已獲證明'],['實測光學輸入關閉且合成配對訓練','具備實測 PSF 融合創新'],['不需要上採樣的實際架構','已驗證抗混疊上採樣帶來提升'],['樣本原形狀與尺寸分開評估','圓度提高即表示更正確'],['本輪處理速度及歷史時間範圍','在新匹配速度測試中全面勝過所有模型'],['保留尺寸偏大的限制','尺寸已恢復真值']],columns=['現有資料支持','現有資料不支持']),'這是寫作證據邊界，不是要求另加大量實驗。')
x.h('原稿章節標題供對照');x.p('以下由原 v34 Word 的標題樣式直接抽出與改寫最相關的小節，供網頁對話掌握原骨架；不是原文全文。若要逐段改稿，仍需把原 Word 一起上傳。原稿 SHA256：'+sha(thesis))
for h in heads:
 if re.match(r'^(4\.[3-8]|5\.4|6\.|7\.)',h):x.p(h)
x.finish()
# All aggregate data as readable text, with exact CSV values rather than lossy table rounding.
full=['# 完整細表及欄位說明','這份文字附錄包含全部彙總、逐趟、逐位置、逐顆及逐方向細表，適合按需上傳。全部逐筆大表另外提供原值 CSV，避免把 79110 筆混排成難以閱讀的 Word 頁面。','模型名稱 Raw 是原圖；R129 是既有比較；R130-Clean 是本輪。index 或 j 是 0 起算拍攝位置或位移終點；run 是趟次；pillar_id 是 1 到 9；kind 的 adjacent 是相鄰法、start 是相對起點法。所有 nm 欄位為奈米，um 為微米，px 為像素；差值符號與统计定義見文件03。']
large={'all_circles','all_displacements','individual_circle_displacements','reconstruction_runtime','all_930_image_sources'}
for q in sorted((P/'完整數據').glob('*.csv')):
 if q.stem in large:continue
 full+=['## '+q.name,'```csv\n'+q.read_text(encoding='utf-8-sig')+'```']
(P/'閱讀文字版/06_完整統計細表.md').write_text('\n\n'.join(full),encoding='utf-8')
(P/'00_上傳與閱讀說明.md').write_text('''# 網頁版論文構思資料

先上傳 01 至 05 五份 Word 文件。主要圖表均已嵌入，每張圖旁也有文字說明及相應數字。需要全文改寫時，請另上傳原 v34 論文。

若需要精確查某個位置或某顆的數字，可再提供「閱讀文字版/06_完整統計細表.md」；如要重算，提供完整數據內的相關 CSV。完整數據另有全部邊界點及訓練紀錄。不要假設只上傳 ZIP，接收端就一定讀取了全部內容。

建議交辦文字：請先完整讀取這五份文件，確認本次實際模型、資料、指標和結果，再以單一模型為主線提出保留七章骨架的論文改寫方案。重點是正常拍攝下的位移與 ACM 輪廓，不以圓度或視覺銳利當作唯一目標。不得把 R129 的 PSF 融合與上採樣特色誤植到 R130，也不要補寫尚未做的消融、獨立保留測試或完整 XAI 證據。請先說明採用 R130 會保留及改变哪些原論文貢獻，再規劃逐章修改。

文字版和 Word 內容對應；文字版的圖片僅有圖說，如接收端未取得 Word 圖像，可另上傳「圖表」中的圖片。完整 930 張大影像不在此包逐張嵌入，完整量測結果與來源索引已附；本輪未執行的實驗不會以舊版本數字冒充。
''',encoding='utf-8')
# Source definitions carried inside the package without requiring local folders.
for name in ['model_r130_clean.py','dataset_r130_clean.py','losses_r130_clean.py','compare_measurements.py']:
 shutil.copy2(R/name,P/'製作與檢查'/name)
shutil.copy2(__file__,P/'製作與檢查/build_web_documents.py')
print(json.dumps(dict(output=str(P),documents=[p.name for p in P.glob('*.docx')],data_files=len(list((P/'完整數據').rglob('*')))),ensure_ascii=True))

