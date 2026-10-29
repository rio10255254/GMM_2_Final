from __future__ import annotations

import csv
import json
from pathlib import Path

import cv2
import numpy as np
from openpyxl import load_workbook
PROJECT_ROOT = Path(__file__).resolve().parents[5]


EXPERIMENT = Path(__file__).resolve().parent
OUTPUT = EXPERIMENT / "evaluation" / "via_reference"
SEM_WORKBOOK = Path(
    str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/20260618_100MG/SEM_data/via_pilar.xlsx')
)
SOURCE_CONFIG = Path(
    str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/Final_Data_Figures/01_CURRENT_THESIS_RESULTS_V42_CENTER9/config/sequence_center9_50x20x.json')
)


def read_gray(path: Path) -> np.ndarray:
    image = cv2.imdecode(np.fromfile(path, dtype=np.uint8), cv2.IMREAD_GRAYSCALE)
    if image is None:
        raise RuntimeError(f"Could not read {path}")
    return image


def write_png(path: Path, image: np.ndarray) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    success, encoded = cv2.imencode(".png", image)
    if not success:
        raise RuntimeError(f"Could not encode {path}")
    encoded.tofile(str(path))


def refine_bright_center(
    image: np.ndarray,
    expected_x: float,
    expected_y: float,
    pitch_px: float,
) -> tuple[float, float]:
    search_radius = int(round(pitch_px * 0.43))
    cx = int(round(expected_x))
    cy = int(round(expected_y))
    patch = image[
        cy - search_radius : cy + search_radius + 1,
        cx - search_radius : cx + search_radius + 1,
    ].astype(np.float64)
    if patch.shape != (2 * search_radius + 1, 2 * search_radius + 1):
        raise RuntimeError("Via center search window reached the image boundary")
    smoothed = cv2.GaussianBlur(
        patch,
        (0, 0),
        max(0.8, pitch_px * 0.025),
    )
    yy, xx = np.indices(smoothed.shape)
    xx0 = xx - search_radius
    yy0 = yy - search_radius
    support = xx0 * xx0 + yy0 * yy0 <= (pitch_px * 0.38) ** 2
    low = float(np.percentile(smoothed[support], 25.0))
    high = float(np.percentile(smoothed[support], 90.0))
    threshold = low + 0.22 * (high - low)
    weights = np.maximum(smoothed - threshold, 0.0) ** 2
    weights *= support
    total = float(weights.sum())
    if total <= 0.0:
        raise RuntimeError("Via center refinement produced zero weight")
    refined_x = expected_x + float((weights * xx).sum() / total) - search_radius
    refined_y = expected_y + float((weights * yy).sum() / total) - search_radius
    return refined_x, refined_y


def load_sem_reference() -> list[list[dict[str, float | str]]]:
    workbook = load_workbook(SEM_WORKBOOK, data_only=True, read_only=True)
    worksheet = workbook["SEM data"]
    matrix: list[list[dict[str, float | str]]] = []
    for row in range(11, 14):
        row_records: list[dict[str, float | str]] = []
        for column in range(7, 10):
            cell = worksheet.cell(row, column)
            radius_nm = float(cell.value)
            row_records.append(
                {
                    "cell": cell.coordinate,
                    "radius_nm": radius_nm,
                    "diameter_um": 2.0 * radius_nm / 1000.0,
                }
            )
        matrix.append(row_records)
    workbook.close()
    return matrix


def expected_grid(magnification: str) -> tuple[list[float], list[float]]:
    if magnification == "50x":
        return [799.0, 871.5, 943.5], [878.0, 951.0, 1024.0]
    if magnification == "20x":
        return [1665.0, 1694.0, 1723.0], [682.0, 711.0, 740.0]
    raise ValueError(magnification)


def overlay(
    image: np.ndarray,
    centers: list[dict[str, float | str]],
    radius_px: float,
) -> tuple[np.ndarray, np.ndarray]:
    canvas = cv2.cvtColor(image, cv2.COLOR_GRAY2BGR)
    for center in centers:
        x = int(round(float(center["x"])))
        y = int(round(float(center["y"])))
        cv2.circle(canvas, (x, y), int(round(radius_px)), (0, 220, 0), 2, cv2.LINE_AA)
        cv2.drawMarker(canvas, (x, y), (0, 0, 255), cv2.MARKER_CROSS, 11, 1)
        cv2.putText(
            canvas,
            str(center["sem_cell"]),
            (x - 14, y - int(round(radius_px)) - 7),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.42,
            (255, 255, 255),
            1,
            cv2.LINE_AA,
        )
    xs = np.asarray([float(center["x"]) for center in centers])
    ys = np.asarray([float(center["y"]) for center in centers])
    margin = int(round(radius_px * 2.0))
    x0 = max(0, int(np.floor(xs.min())) - margin)
    x1 = min(canvas.shape[1], int(np.ceil(xs.max())) + margin + 1)
    y0 = max(0, int(np.floor(ys.min())) - margin)
    y1 = min(canvas.shape[0], int(np.ceil(ys.max())) + margin + 1)
    return canvas, canvas[y0:y1, x0:x1]


def main() -> None:
    sem_matrix = load_sem_reference()
    sem_diameters = [
        float(record["diameter_um"])
        for row in sem_matrix
        for record in row
    ]
    source = json.loads(SOURCE_CONFIG.read_text(encoding="utf-8"))
    sequence_outputs: dict[str, dict[str, object]] = {}

    for sequence_name, sequence in source["sequences"].items():
        magnification = str(sequence["magnification"])
        raw_dir = Path(PROJECT_ROOT / sequence["raw_dir"])
        frame_path = raw_dir / str(sequence["first_frame"])
        image = read_gray(frame_path)
        x_expected, y_expected = expected_grid(magnification)
        pitch_px = float(sequence["pitch_px"])
        um_per_px = float(sequence["object_um_per_px"])
        refined = [
            [refine_bright_center(image, x, y, pitch_px) for x in x_expected]
            for y in y_expected
        ]

        centers: list[dict[str, float | str]] = []
        for physical_row in range(3):
            optical_row = physical_row if magnification == "50x" else 2 - physical_row
            for physical_column in range(3):
                x, y = refined[optical_row][physical_column]
                sem = sem_matrix[physical_row][physical_column]
                centers.append(
                    {
                        "sem_cell": str(sem["cell"]),
                        "physical_row": physical_row + 1,
                        "physical_column": physical_column + 4,
                        "x": x,
                        "y": y,
                        "sem_radius_nm": float(sem["radius_nm"]),
                        "sem_diameter_um": float(sem["diameter_um"]),
                    }
                )

        radius_px = float(np.median(sem_diameters)) / (2.0 * um_per_px)
        full_overlay, crop_overlay = overlay(image, centers, radius_px)
        write_png(OUTPUT / f"{sequence_name}_via_center9_overlay_full.png", full_overlay)
        write_png(OUTPUT / f"{sequence_name}_via_center9_overlay_crop.png", crop_overlay)

        sequence_outputs[sequence_name] = {
            "dataset": sequence["dataset"],
            "group": sequence["group"],
            "magnification": magnification,
            "raw_dir": str(raw_dir),
            "mosr_dir": str(PROJECT_ROOT / sequence["v42_dir"]),
            "first_frame": sequence["first_frame"],
            "frame_count": sequence["frame_count"],
            "object_um_per_px": um_per_px,
            "nominal_step_um": sequence["nominal_step_um"],
            "expected_signed_dx_px": sequence["expected_signed_dx_px"],
            "capture_transform": sequence["capture_transform"],
            "polarity": "bright",
            "pitch_px": pitch_px,
            "radius_px": radius_px,
            "search_radius_px": sequence["search_radius_px"],
            "min_valid_features": 7,
            "centers": centers,
        }

    output = {
        "version": 1,
        "purpose": "SEM-referenced center-nine via evaluation for 50x and 20x",
        "sem_workbook": str(SEM_WORKBOOK),
        "sem_sheet": "SEM data",
        "sem_cells": "G11:I13",
        "sem_value_definition": "Workbook values are radii in nanometers; diameter equals two times radius",
        "physical_selection": "Upper-right 3x3 corner of the lower-left via region, nearest the sample center",
        "orientation_rule": (
            "50x follows the physical orientation; 20x is vertically flipped, so physical SEM row 1 "
            "maps to the lowest optical row"
        ),
        "sem_reference": sem_matrix,
        "sequences": sequence_outputs,
    }
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / "via_center9_reference.json").write_text(
        json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with (OUTPUT / "via_center9_sem_reference.csv").open(
        "w", newline="", encoding="utf-8-sig"
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "physical_row",
                "physical_column",
                "sem_cell",
                "sem_radius_nm",
                "sem_diameter_nm",
            ],
        )
        writer.writeheader()
        for row_index, row in enumerate(sem_matrix, start=1):
            for column_index, record in enumerate(row, start=4):
                writer.writerow(
                    {
                        "physical_row": row_index,
                        "physical_column": column_index,
                        "sem_cell": record["cell"],
                        "sem_radius_nm": record["radius_nm"],
                        "sem_diameter_nm": float(record["diameter_um"]) * 1000.0,
                    }
                )
    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
