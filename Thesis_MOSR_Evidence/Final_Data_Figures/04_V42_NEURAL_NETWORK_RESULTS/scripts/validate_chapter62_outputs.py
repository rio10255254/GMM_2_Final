from __future__ import annotations

import csv
import math
import statistics
from collections import Counter
from datetime import datetime
from pathlib import Path

from PIL import Image
from docx import Document
PROJECT_ROOT = Path(__file__).resolve().parents[4]


PACKAGE = Path(__file__).resolve().parents[1]
TABLES = PACKAGE / "tables"
FIGURES = PACKAGE / "figures"
QA = PACKAGE / "qa"
SOURCE_TABLES = PACKAGE / "source_validation" / "wiener_baseline_20260715" / "tables"
SUPP_20X_TABLES = (
    PACKAGE
    / "source_validation"
    / "wiener_baseline_20260715"
    / "supplementary_20x_cd"
    / "tables"
)
PSF_PACKAGE = Path(
    str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/Final_Data_Figures/03_PSF_EXPERIMENT_RESULTS/01_global_bads_fit')
)
PSF_TABLES = PSF_PACKAGE / "tables" / "center9_pillar_fit_20x_50x"
PSF_FIGURES = PSF_PACKAGE / "figures" / "center9_pillar_fit_20x_50x"
PSF_SHARPNESS = (
    PSF_PACKAGE / "source_data" / "center9_pillar_fit_20x_50x" / "sharpness_validation"
)
WORD_DOC = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節"
    r"\第六章 實驗結果與分析_含6.2_MOSR與反卷積CD上下界分析.docx"
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def number(row: dict[str, str], field: str) -> float:
    value = float(row[field])
    if not math.isfinite(value):
        raise AssertionError(f"Non-finite value in {field}")
    return value


def close(actual: float, expected: float, label: str, tolerance: float = 1e-8) -> None:
    if not math.isclose(actual, expected, rel_tol=tolerance, abs_tol=tolerance):
        raise AssertionError(f"{label}: {actual} != {expected}")


def mean(values: list[float]) -> float:
    return sum(values) / len(values)


def median(values: list[float]) -> float:
    return statistics.median(values)


def main() -> None:
    selected_nsr = read_csv(SOURCE_TABLES / "wiener_selected_nsr.csv")
    assert {row["magnification"] for row in selected_nsr} == {"20x", "50x"}
    for row in selected_nsr:
        close(number(row, "selected_nsr"), 0.05, f"{row['magnification']} selected NSR")

    psf_per_pillar = read_csv(PSF_TABLES / "cd_before_after_per_pillar.csv")
    assert len(psf_per_pillar) == 18
    psf_summary = {
        row["magnification"]: row
        for row in read_csv(PSF_TABLES / "cd_before_after_summary.csv")
    }
    psf_sharpness = {
        row["magnification"]: row
        for row in read_csv(PSF_SHARPNESS / "sharpness_validation_metrics.csv")
    }
    for magnification in ("20x", "50x"):
        rows = [row for row in psf_per_pillar if row["magnification"] == magnification]
        assert len(rows) == 9
        summary = psf_summary[magnification]
        close(number(summary, "calibrated_common_nsr"), 0.05, f"{magnification} PSF-validation NSR")
        raw_errors = [number(row, "raw_abs_error_nm") for row in rows]
        wiener_errors = [number(row, "deconvolved_abs_error_nm") for row in rows]
        close(mean(raw_errors), number(summary, "raw_mae_nm"), f"{magnification} PSF raw MAE")
        close(mean(wiener_errors), number(summary, "deconvolved_mae_nm"), f"{magnification} PSF Wiener MAE")
        raw_widths = [number(row, "raw_edge_width_px") for row in rows]
        wiener_widths = [number(row, "deconvolved_edge_width_px") for row in rows]
        assert all(new < raw for raw, new in zip(raw_widths, wiener_widths))
        assert int(float(psf_sharpness[magnification]["pillars_sharpened"])) == 9

    psf_threshold = read_csv(PSF_TABLES / "cd_threshold_sensitivity.csv")
    assert len(psf_threshold) == 6
    assert {(row["magnification"], row["contrast_fraction"]) for row in psf_threshold} == {
        (magnification, threshold)
        for magnification in ("20x", "50x")
        for threshold in ("0.4", "0.5", "0.6")
    }
    revised_psf_figures = [
        PSF_FIGURES / "comparison" / name
        for name in (
            "07_edge_profile_20x.png",
            "08_edge_profile_50x.png",
            "09_edge_width_summary.png",
            "10_cd_mae_summary.png",
            "11_threshold_sensitivity_20x.png",
            "12_threshold_sensitivity_50x.png",
            "13_ringing_summary.png",
        )
    ]
    for path in revised_psf_figures:
        assert path.is_file(), f"Missing revised PSF figure: {path}"
        with Image.open(path) as image:
            assert image.size[0] >= 900 and image.size[1] >= 700
    for magnification in ("20x", "50x"):
        assert (PSF_FIGURES / magnification / "09_fitted_psf_wiener_deconvolution_pillar_3x3_spectrum.png").is_file()

    assert WORD_DOC.is_file()
    document = Document(str(WORD_DOC))
    visible_parts = [paragraph.text for paragraph in document.paragraphs]
    visible_parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    visible_text = "\n".join(visible_parts)
    assert visible_text.count("6.1.2 點擴散函數驗證") == 1
    assert visible_text.count("6.2 神經網路重建效果與位移量測結果") == 1
    assert "0.0565" not in visible_text
    assert "V42" not in visible_text
    assert len(document.inline_shapes) >= 35

    per_step = read_csv(TABLES / "displacement_three_way_per_step.csv")
    assert len(per_step) == 120
    counts = Counter(row["sequence"] for row in per_step)
    assert counts == Counter({"50x_x_0p5um": 30, "50x_x_1um": 30, "20x_0p5um": 30, "20x_1um": 30})
    for row in per_step:
        for field in ("raw_boundary_valid_pillars", "wiener_boundary_valid_pillars", "v42_boundary_valid_pillars"):
            assert int(float(row[field])) == 9
        for field in (
            "raw_boundary_dx_um", "wiener_boundary_dx_um", "v42_boundary_dx_um",
            "raw_within_grid_nm", "wiener_within_grid_nm", "v42_within_grid_nm",
        ):
            number(row, field)

    detailed = {row["sequence"]: row for row in read_csv(TABLES / "displacement_three_way_summary.csv")}
    for sequence, count in counts.items():
        rows = [row for row in per_step if row["sequence"] == sequence]
        summary = detailed[sequence]
        nominal = number(summary, "nominal_step_um")
        assert int(float(summary["paired_count"])) == count
        for prefix, dx_field, within_field in (
            ("raw", "raw_boundary_dx_um", "raw_within_grid_nm"),
            ("wiener", "wiener_boundary_dx_um", "wiener_within_grid_nm"),
            ("v42", "v42_boundary_dx_um", "v42_within_grid_nm"),
        ):
            magnitudes = [abs(number(row, dx_field)) for row in rows]
            within = [number(row, within_field) for row in rows]
            close(statistics.stdev(magnitudes), number(summary, f"{prefix}_step_sd_nm") / 1000.0, f"{sequence} {prefix} SD")
            close(mean([abs(value - nominal) for value in magnitudes]) * 1000.0, number(summary, f"{prefix}_command_mae_nm"), f"{sequence} {prefix} command MAE")
            close(mean(within), number(summary, f"{prefix}_within_mean_nm"), f"{sequence} {prefix} within mean")

    nominal_cd = read_csv(TABLES / "cd_50pct_nominal_per_pillar.csv")
    assert len(nominal_cd) == 9
    for row in nominal_cd:
        assert row["raw_status"] == row["wiener_status"] == row["v42_e120_status"] == "ok"
    nominal_summary = {row["case"]: row for row in read_csv(TABLES / "cd_50pct_nominal_summary.csv")}
    for case, field in (("raw", "raw_abs_error_nm"), ("wiener", "wiener_abs_error_nm"), ("v42_e120", "v42_e120_abs_error_nm")):
        values = [number(row, field) for row in nominal_cd]
        close(mean(values), number(nominal_summary[case], "mean_abs_error_nm"), f"{case} nominal CD MAE")
        close(max(values), number(nominal_summary[case], "max_abs_error_nm"), f"{case} nominal CD max")

    sensitivity = read_csv(SOURCE_TABLES / "three_way_cd_sensitivity_per_measurement.csv")
    bounds = read_csv(TABLES / "cd_per_pillar_bounds.csv")
    assert len(bounds) == 9
    for row in bounds:
        assert int(float(row["observations_per_method"])) == 14
        sem_index = int(float(row["sem_index"]))
        for case, prefix in (("raw", "raw"), ("wiener", "wiener"), ("v42_e120", "mosr")):
            values = [number(item, "diameter_um") for item in sensitivity if item["case"] == case and int(float(item["sem_index"])) == sem_index]
            assert len(values) == 14
            close(min(values), number(row, f"{prefix}_lower_um"), f"{case} P{sem_index} lower bound")
            close(max(values), number(row, f"{prefix}_upper_um"), f"{case} P{sem_index} upper bound")
    robust_summary = {row["case"]: row for row in read_csv(TABLES / "cd_threshold_robust_summary.csv")}
    for case in ("raw", "wiener", "v42_e120"):
        values = [number(row, "abs_error_nm") for row in sensitivity if row["case"] == case]
        assert len(values) == 126
        close(median(values), number(robust_summary[case], "pooled_median_abs_error_nm"), f"{case} robust median")
        close(mean(values), number(robust_summary[case], "pooled_mean_abs_error_nm"), f"{case} robust mean")
        close(math.sqrt(mean([value * value for value in values])), number(robust_summary[case], "pooled_rmse_nm"), f"{case} robust RMSE")

    supplementary_20x = {
        row["case"]: row
        for row in read_csv(SUPP_20X_TABLES / "20x_frame49_cd_summary.csv")
    }
    assert set(supplementary_20x) == {"raw", "wiener", "mosr"}
    for case, row in supplementary_20x.items():
        assert int(float(row["observation_count"])) == 126
        assert int(float(row["pillar_count"])) == 9
    assert number(supplementary_20x["wiener"], "pooled_median_abs_error_nm") < number(
        supplementary_20x["raw"], "pooled_median_abs_error_nm"
    )
    assert number(supplementary_20x["mosr"], "pooled_median_abs_error_nm") < number(
        supplementary_20x["wiener"], "pooled_median_abs_error_nm"
    )

    selection = {row["population"]: row for row in read_csv(TABLES / "selection_robustness_20x_0p5um.csv")}
    assert int(float(selection["All usable"]["step_count"])) == 32
    assert int(float(selection["Locked top 30"]["step_count"])) == 30

    obsolete = (
        FIGURES / "08_legacy_raw_v42_50pct_cd_error_summary.png",
        FIGURES / "09_legacy_raw_v42_50pct_cd_improvement_per_pillar.png",
        TABLES / "legacy_raw_v42_50pct_cd_summary.csv",
        TABLES / "displacement_detailed_summary.csv",
        TABLES / "displacement_pooled_summary.csv",
    )
    assert not any(path.exists() for path in obsolete)

    png_paths = sorted(FIGURES.glob("*.png"))
    csv_paths = sorted(TABLES.glob("*.csv"))
    assert len(png_paths) == 31
    assert len(csv_paths) == 19
    for path in png_paths:
        with Image.open(path) as image:
            width, height = image.size
            assert width >= 900 and height >= 700, f"Unexpectedly small figure: {path.name} {image.size}"

    QA.mkdir(parents=True, exist_ok=True)
    report = QA / "final_validation_report.txt"
    report.write_text(
        "\n".join(
            [
                f"Validated: {datetime.now().astimezone().isoformat(timespec='seconds')}",
                "Fixed NSR: 0.05 at 20x and 50x",
                "PSF validation: 18 pillar measurements recomputed with the same fixed NSR",
                "PSF edge narrowing: 9/9 pillars at both magnifications",
                "PSF threshold audit: 6 magnification-threshold summaries",
                "Revised PSF figures: 7 English MATLAB figures plus refreshed absolute spectra",
                f"Word structure: one revised 6.1.2 section, one 6.2 section, {len(document.inline_shapes)} embedded figures",
                "Word terminology: no obsolete NSR 0.0565 or visible internal model label",
                "Primary displacement rows: 120 (30 per sequence)",
                "Primary valid pillars: 9 in every Raw, Wiener, and MOSR measurement",
                "Nominal CD crossings: 9/9 valid in all three methods",
                "Per-pillar CD chart: full lower-upper bounds from 14 observations; no central Raw, Wiener, or MOSR operating point",
                "CD sensitivity records: 126 per method",
                "Supplementary 20x CD check: 126 observations per method on a non-PSF-calibration frame",
                "Final figures: 31 PNG files; all at least 900 x 700 pixels",
                "Final tables: 19 CSV files",
                "Obsolete two-way artifacts: absent",
                "Independent metric recomputation: passed",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    print(report)


if __name__ == "__main__":
    main()
