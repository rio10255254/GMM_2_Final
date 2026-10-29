from __future__ import annotations

import csv
import statistics
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt
PROJECT_ROOT = Path(__file__).resolve().parents[4]


PACKAGE = Path(str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/Final_Data_Figures/04_V42_NEURAL_NETWORK_RESULTS'))
FIGURES = PACKAGE / "figures"
TABLES = PACKAGE / "tables"
PSF_PACKAGE = Path(
    str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/Final_Data_Figures/03_PSF_EXPERIMENT_RESULTS/01_global_bads_fit')
)
PSF_FIGURES = PSF_PACKAGE / "figures" / "center9_pillar_fit_20x_50x"
PSF_TABLES = PSF_PACKAGE / "tables" / "center9_pillar_fit_20x_50x"
PSF_SHARPNESS = (
    PSF_PACKAGE
    / "source_data"
    / "center9_pillar_fit_20x_50x"
    / "sharpness_validation"
)
SUPP_20X_TABLES = (
    PACKAGE
    / "source_validation"
    / "wiener_baseline_20260715"
    / "supplementary_20x_cd"
    / "tables"
)
SOURCE_DOCX = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節\第六章 實驗結果與分析_含6.2_V42位移分析.docx")
OUTPUT_DOCX = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\章節\第六章 實驗結果與分析_含6.2_MOSR與反卷積CD上下界分析.docx")


def read_csv(name: str) -> list[dict[str, str]]:
    with (TABLES / name).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_csv_path(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def font_run(run, *, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "新細明體")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_heading(doc: Document, text: str, *, level: int = 2, page_break: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(7 if level > 1 else 0)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    font_run(run, size={1: 16, 2: 14, 3: 12.5}.get(level, 12), bold=True)


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_together = False
    run = paragraph.add_run(text)
    font_run(run, size=12)


def insert_body_after(paragraph: Paragraph, text: str, *, lead: str | None = None) -> Paragraph:
    new_xml = OxmlElement("w:p")
    paragraph._p.addnext(new_xml)
    inserted = Paragraph(new_xml, paragraph._parent)
    inserted.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    inserted.paragraph_format.space_after = Pt(2)
    if lead:
        lead_run = inserted.add_run(lead)
        font_run(lead_run, size=12, bold=True)
    body_run = inserted.add_run(text)
    font_run(body_run, size=12)
    return inserted


def add_caption(doc: Document, text: str, *, keep_with_next: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    font_run(run, size=10.5)


def add_figure(doc: Document, filename: str, caption: str, *, width: float = 6.15) -> None:
    path = FIGURES / filename
    if not path.is_file():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_figure_path(doc: Document, path: Path, caption: str, *, width: float = 6.15) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, value: object, *, bold: bool = False, size: float = 9.0) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    font_run(run, size=size, bold=bold)


def add_table(
    doc: Document,
    caption: str,
    headers: list[str],
    rows: list[list[object]],
    *,
    font_size: float = 9.0,
) -> None:
    add_caption(doc, caption, keep_with_next=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, bold=True, size=font_size)
        shade_cell(table.rows[0].cells[index], "E7E6E6")
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    for row_values in rows:
        row = table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for index, value in enumerate(row_values):
            set_cell(row.cells[index], value, size=font_size)
    # Keep each compact thesis table together instead of leaving only its
    # header or first row at the bottom of a page.
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def pct_lower(raw: float, new: float) -> float:
    return 100.0 * (1.0 - new / raw)


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def remove_previous_generated_section(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("本節定位修正：") or text.startswith("第 6.2 節另以未納入正式結果的校正影格"):
            remove_paragraph(paragraph)

    heading = next(
        (paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith("6.2 ")),
        None,
    )
    if heading is None:
        return
    body = doc._element.body
    children = list(body)
    start = children.index(heading._element)
    if start > 0 and children[start - 1].tag == qn("w:p") and not "".join(children[start - 1].itertext()).strip():
        start -= 1
    for child in children[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def remove_section_from_heading(doc: Document, prefix: str) -> None:
    heading = next(
        (paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith(prefix)),
        None,
    )
    if heading is None:
        raise RuntimeError(f"Could not find section heading: {prefix}")
    body = doc._element.body
    children = list(body)
    start = children.index(heading._element)
    for child in children[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def build_revised_psf_validation(doc: Document) -> None:
    summary = {
        row["magnification"]: row
        for row in read_csv_path(PSF_TABLES / "cd_before_after_summary.csv")
    }
    sensitivity = read_csv_path(PSF_TABLES / "cd_threshold_sensitivity.csv")
    sharpness = {
        row["magnification"]: row
        for row in read_csv_path(PSF_SHARPNESS / "sharpness_validation_metrics.csv")
    }

    add_heading(doc, "6.1.2 點擴散函數驗證", level=2)
    add_body(
        doc,
        "為確認擬合所得 PSF 沒有明顯偏離實際模糊，本節將 20× 與 50× 的局部擬合 PSF 分別套用至對應的中心九宮格 pillar 影像。Wiener 對照統一使用 NSR = 0.05，兩種倍率不依 SEM CD 或位移結果個別調整。此設定先確認影像能產生穩定且可辨識的銳化，再固定套用；已知 SEM 直徑僅作為反卷積後的尺寸一致性參考，不作參數挑選依據，也不視為獨立精度證明。",
    )
    add_body(
        doc,
        "反卷積前先估計緩慢變化的亮背景，將暗色 pillar 轉為正向對比後，以對稱延伸與邊界緩和降低裁切邊緣的突變，再使用擬合 PSF 進行 Wiener 反卷積。反卷積後的對比限制於固定範圍並還原為原始明暗方向；Raw 與 Wiener 影像均以相同 8-bit 灰階範圍輸出，不進行逐張自動拉伸。較積極的設定允許保留可見的環狀振鈴，目的在於把傳統反卷積可達到的銳化與其代價同時呈現。",
    )

    add_heading(doc, "6.1.2.1 CD、邊緣寬度與振鈴之定義", level=3)
    add_body(
        doc,
        "每顆 pillar 的中心均由 Raw 影像固定定位，Wiener 影像不重新尋找較有利的中心。量測時自中心向外取 360 條徑向剖面；中心灰階取自半徑不超過相鄰 pillar 間距 0.07 倍的區域，背景灰階取自 0.30–0.36 倍間距的環帶。每條剖面依中心與背景灰階正規化後，找出 10%、50% 與 90% 的第一次向外交點。相反方向的兩個 50% 半徑相加形成一組直徑，180 組直徑的中位數定義為該顆 pillar 的 CD。",
    )
    add_body(
        doc,
        "邊緣寬度定義為同一剖面上 90% 與 10% 交點的距離，九顆 pillar 再取平均；數值越小，代表灰階由 pillar 內部轉為背景的距離越短。振鈴比例則以徑向剖面中反向灰階變化量除以主要向外上升量，並取九顆 pillar 的最大值，用來呈現最嚴重的局部環狀假影。這些指標分別回答『邊界是否變陡』、『尺寸是否接近 SEM』與『銳化是否伴隨假影』，三者不能互相取代。",
    )
    add_figure_path(
        doc,
        PSF_FIGURES / "comparison" / "04_cd_measurement_method.png",
        "CD、10–90% 邊緣寬度與徑向剖面之量測方式",
        width=6.15,
    )

    add_heading(doc, "6.1.2.2 影像銳化與邊緣過渡", level=3, page_break=True)
    add_body(
        doc,
        "下列影像使用相同灰階範圍分開呈現。20× 與 50× 的 Wiener 結果都可見較陡的 pillar 外緣；50× 影像同時出現更明顯的中心亮斑與外圍環狀反應，正是較積極反卷積的典型限制。影像變清楚只表示高頻與邊緣對比被提高，尚不能直接等同於真實邊界已恢復。",
    )
    for label in ("20x", "50x"):
        display = label.replace("x", "×")
        add_figure_path(
            doc,
            PSF_FIGURES / label / "05_original_pillar_3x3.png",
            f"{display} 原始中心九顆 pillar 影像",
            width=3.45,
        )
        add_figure_path(
            doc,
            PSF_FIGURES / label / "09_fitted_psf_wiener_deconvolution_pillar_3x3.png",
            f"{display} 使用固定 NSR = 0.05 之 Wiener 反卷積影像",
            width=3.45,
        )

    add_figure_path(
        doc,
        PSF_FIGURES / "comparison" / "07_edge_profile_20x.png",
        "20× Raw 與 Wiener 之九顆 pillar 徑向邊緣剖面中位數",
        width=5.85,
    )
    add_figure_path(
        doc,
        PSF_FIGURES / "comparison" / "08_edge_profile_50x.png",
        "50× Raw 與 Wiener 之九顆 pillar 徑向邊緣剖面中位數",
        width=5.85,
    )
    add_figure_path(
        doc,
        PSF_FIGURES / "comparison" / "09_edge_width_summary.png",
        "兩種倍率反卷積前後之平均 10–90% 邊緣寬度",
        width=5.75,
    )

    edge_rows: list[list[object]] = []
    for label in ("20x", "50x"):
        s = summary[label]
        q = sharpness[label]
        raw_um = float(q["raw_mean_edge_width_um"])
        wiener_um = float(q["deconvolved_mean_edge_width_um"])
        edge_rows.append(
            [
                label.replace("x", "×"),
                f"{raw_um:.3f}",
                f"{wiener_um:.3f}",
                f"{100.0 * (1.0 - wiener_um / raw_um):.2f}%",
                f"{q['pillars_sharpened']}/{q['pillar_count']}",
                f"{100.0 * float(s['raw_max_ringing_ratio']):.1f}% → "
                f"{100.0 * float(s['deconvolved_max_ringing_ratio']):.1f}%",
            ]
        )
    add_table(
        doc,
        "兩種倍率之邊緣銳化與最大振鈴摘要",
        ["倍率", "Raw 邊緣寬度 (µm)", "Wiener 邊緣寬度 (µm)", "縮減比例", "縮窄顆數", "最大振鈴比例"],
        edge_rows,
        font_size=8.2,
    )
    q20 = sharpness["20x"]
    q50 = sharpness["50x"]
    raw_edge_20 = float(q20["raw_mean_edge_width_um"])
    wiener_edge_20 = float(q20["deconvolved_mean_edge_width_um"])
    raw_edge_50 = float(q50["raw_mean_edge_width_um"])
    wiener_edge_50 = float(q50["deconvolved_mean_edge_width_um"])
    add_body(
        doc,
        f"20× 的平均邊緣寬度由 {raw_edge_20:.3f} µm 降至 {wiener_edge_20:.3f} µm，縮減 {100.0 * (1.0 - wiener_edge_20 / raw_edge_20):.2f}%；"
        f"50× 則由 {raw_edge_50:.3f} µm 降至 {wiener_edge_50:.3f} µm，縮減 {100.0 * (1.0 - wiener_edge_50 / raw_edge_50):.2f}%。"
        f"兩種倍率皆為 {q20['pillars_sharpened']}/9 與 {q50['pillars_sharpened']}/9 顆 pillar 的灰階過渡縮窄，顯示固定 Wiener 設定在『讓邊界變陡』這項檢查上具有一致效果。",
    )

    add_heading(doc, "6.1.2.3 SEM 對照 CD 與門檻敏感度", level=3, page_break=True)
    cd_rows: list[list[object]] = []
    for label in ("20x", "50x"):
        s = summary[label]
        cd_rows.append(
            [
                label.replace("x", "×"),
                f"{float(s['sem_mean_cd_nm']):.1f}",
                f"{float(s['raw_mean_cd_nm']):.1f}",
                f"{float(s['deconvolved_mean_cd_nm']):.1f}",
                f"{float(s['raw_mae_nm']):.1f}",
                f"{float(s['deconvolved_mae_nm']):.1f}",
                f"{s['pillars_improved']}/9",
            ]
        )
    add_table(
        doc,
        "預先指定 50% 邊界下之 SEM 對照 CD 結果",
        ["倍率", "SEM 平均 CD (nm)", "Raw 平均 CD (nm)", "Wiener 平均 CD (nm)", "Raw MAE (nm)", "Wiener MAE (nm)", "改善顆數"],
        cd_rows,
        font_size=7.9,
    )
    add_figure_path(
        doc,
        PSF_FIGURES / "comparison" / "10_cd_mae_summary.png",
        "預先指定 50% 邊界下之 Raw 與 Wiener 平均絕對 CD 誤差",
        width=5.75,
    )
    s20 = summary["20x"]
    s50 = summary["50x"]
    add_body(
        doc,
        f"在 50% 邊界下，20× 的平均絕對 CD 誤差由 {float(s20['raw_mae_nm']):.2f} nm 變為 {float(s20['deconvolved_mae_nm']):.2f} nm，"
        f"九顆中有 {s20['pillars_improved']} 顆改善；50× 則由 {float(s50['raw_mae_nm']):.2f} nm 變為 {float(s50['deconvolved_mae_nm']):.2f} nm，"
        f"九顆中有 {s50['pillars_improved']} 顆改善。Wiener 雖穩定讓邊緣變陡，卻沒有在兩種倍率都同步改善 50% 交點所代表的表觀直徑，說明『視覺銳化』與『尺寸修復』是不同問題。",
    )
    for label, file_name in (("20x", "11_threshold_sensitivity_20x.png"), ("50x", "12_threshold_sensitivity_50x.png")):
        add_figure_path(
            doc,
            PSF_FIGURES / "comparison" / file_name,
            f"{label.replace('x', '×')} Raw 與 Wiener 之 CD 門檻敏感度",
            width=5.8,
        )

    threshold_rows: list[list[object]] = []
    for row in sensitivity:
        raw = float(row["raw_mae_nm"])
        wiener = float(row["deconvolved_mae_nm"])
        threshold_rows.append(
            [
                row["magnification"].replace("x", "×"),
                f"{100.0 * float(row['contrast_fraction']):.0f}%",
                f"{raw:.1f}",
                f"{wiener:.1f}",
                f"{100.0 * (raw - wiener) / raw:+.1f}%",
                f"{row['pillars_improved']}/9",
            ]
        )
    add_table(
        doc,
        "40%、50% 與 60% 相對對比門檻之 CD 敏感度",
        ["倍率", "相對門檻", "Raw MAE (nm)", "Wiener MAE (nm)", "相對 Raw", "改善顆數"],
        threshold_rows,
        font_size=8.2,
    )
    sensitivity_by_key = {
        (row["magnification"], int(round(100.0 * float(row["contrast_fraction"])))): row
        for row in sensitivity
    }
    t20_40 = sensitivity_by_key[("20x", 40)]
    t20_60 = sensitivity_by_key[("20x", 60)]
    t50_40 = sensitivity_by_key[("50x", 40)]
    t50_60 = sensitivity_by_key[("50x", 60)]
    add_body(
        doc,
        f"40% 門檻下，20× 與 50× 分別有 {t20_40['pillars_improved']}/9 與 {t50_40['pillars_improved']}/9 顆改善；"
        f"60% 門檻下則分別為 {t20_60['pillars_improved']}/9 與 {t50_60['pillars_improved']}/9。"
        f"20× 的 Wiener MAE 由 40% 門檻的 {float(t20_40['deconvolved_mae_nm']):.2f} nm 變為 60% 門檻的 {float(t20_60['deconvolved_mae_nm']):.2f} nm；"
        f"50× 則由 {float(t50_40['deconvolved_mae_nm']):.2f} nm 變為 {float(t50_60['deconvolved_mae_nm']):.2f} nm。"
        "排名隨門檻翻轉，代表環狀反應與剖面形狀已推動灰階交點，因此不能只選擇最有利門檻宣稱 CD 已恢復。",
    )

    add_heading(doc, "6.1.2.4 振鈴與真實頻譜", level=3, page_break=True)
    add_figure_path(
        doc,
        PSF_FIGURES / "comparison" / "13_ringing_summary.png",
        "兩種倍率反卷積前後之最大局部振鈴比例",
        width=5.75,
    )
    add_body(
        doc,
        f"20× 的最大振鈴比例由 {100.0 * float(s20['raw_max_ringing_ratio']):.1f}% 增至 {100.0 * float(s20['deconvolved_max_ringing_ratio']):.1f}%，"
        f"50× 則由 {100.0 * float(s50['raw_max_ringing_ratio']):.1f}% 增至 {100.0 * float(s50['deconvolved_max_ringing_ratio']):.1f}%。"
        "此結果與反卷積影像中的亮暗環一致，表示較明顯的銳化同時伴隨局部波動。兩種倍率的 pillar 中心均未發生整體明暗反轉，主要結構方向仍被保留；但中心未反轉不能排除外圈振鈴對 CD 交點的影響。",
    )
    for label in ("20x", "50x"):
        add_figure_path(
            doc,
            PSF_FIGURES / label / "09_fitted_psf_wiener_deconvolution_pillar_3x3_spectrum.png",
            f"{label.replace('x', '×')} Wiener 反卷積影像之絕對二維頻譜",
            width=4.7,
        )
    add_body(
        doc,
        "頻譜以原始浮點影像直接計算二維傅立葉振幅，沒有扣除平均值、沒有套窗、沒有平滑，也沒有依每張影像的峰值重新正規化；所有 Raw、Wiener 與 MOSR 頻譜共用 −80 至 0 dBFS 顯示範圍。因而頻譜中較明顯的外圍能量代表實際計算出的高空間頻率成分增加，但其中同時可能包含邊緣資訊、雜訊與振鈴，不能單憑頻譜延伸判定新增成分皆為真實結構。",
    )

    add_heading(doc, "6.1.2.5 綜合討論", level=3)
    add_body(
        doc,
        "固定 NSR = 0.05 的 Wiener 反卷積達成本節最直接的驗證目的：20× 與 50× 的九顆 pillar 都得到較窄的 10–90% 灰階過渡，且在一致顯示範圍下可清楚看出邊緣增強。這表示局部擬合 PSF 與實際模糊方向大致相符，沒有出現因核心完全不匹配而無法銳化的情形。",
    )
    add_body(
        doc,
        f"然而，固定較積極的反卷積也把兩種倍率的最大振鈴提高至約 {100.0 * min(float(s20['deconvolved_max_ringing_ratio']), float(s50['deconvolved_max_ringing_ratio'])):.1f}–{100.0 * max(float(s20['deconvolved_max_ringing_ratio']), float(s50['deconvolved_max_ringing_ratio'])):.1f}%，"
        "使 CD 結果對 40%、50% 與 60% 邊界定義高度敏感。較穩健的結論不是『反卷積已準確恢復尺寸』，而是『反卷積能可靠提高局部銳利度，但不保證灰階交點仍位於真實邊界』。因此，第 6.2 節將此 Wiener 結果作為視覺上夠強的傳統基準，再以跨門檻 CD 穩健度與 120 步位移一致性比較兩種影像處理方式。",
    )


def main() -> None:
    if not SOURCE_DOCX.is_file():
        raise FileNotFoundError(SOURCE_DOCX)

    detailed = {row["sequence"]: row for row in read_csv("displacement_three_way_summary.csv")}
    for row in detailed.values():
        row["step_sd_reduction_pct"] = row["v42_step_sd_reduction_pct"]
        row["command_mae_reduction_pct"] = row["v42_command_mae_reduction_pct"]
        row["within_reduction_pct"] = row["v42_within_reduction_pct"]
    pooled = {row["population"]: row for row in read_csv("displacement_three_way_pooled.csv")}
    for row in pooled.values():
        count = float(row["paired_count"])
        row["command_mae_reduction_pct"] = row["v42_command_mae_reduction_pct"]
        row["within_reduction_pct"] = row["v42_within_reduction_pct"]
        row["v42_closer_to_command_pct"] = str(100.0 * float(row["v42_closer_to_command_count"]) / count)
        row["v42_lower_within_pct"] = str(100.0 * float(row["v42_lower_within_count"]) / count)
        row["wiener_closer_to_command_pct"] = str(100.0 * float(row["wiener_closer_to_command_count"]) / count)
        row["wiener_lower_within_pct"] = str(100.0 * float(row["wiener_lower_within_count"]) / count)
    cd_rows = read_csv("cd_50pct_nominal_per_pillar.csv")
    cd_nominal = {row["case"]: row for row in read_csv("cd_50pct_nominal_summary.csv")}
    cd_bounds = read_csv("cd_per_pillar_bounds.csv")
    cd_robust = {row["case"]: row for row in read_csv("cd_threshold_robust_summary.csv")}
    supplementary_20x = {
        row["case"]: row
        for row in read_csv_path(SUPP_20X_TABLES / "20x_frame49_cd_summary.csv")
    }
    training = read_csv("training_summary_for_thesis.csv")
    edge_rows = read_csv("edge_transition_width_summary.csv")
    edges = {
        (row["magnification"], row["case"]): float(row["median_transition_width_nm"])
        for row in edge_rows
    }
    contrast_rows = read_csv("contrast_polarity_summary.csv")
    contrasts = {
        (row["magnification"], row["case"]): row
        for row in contrast_rows
    }
    selection_rows = read_csv("top30_raw_motion_selection.csv")
    selection_robustness = {
        row["population"]: row for row in read_csv("selection_robustness_20x_0p5um.csv")
    }
    spatial_rows = read_csv("wiener_spatial_variation_summary.csv")
    spatial_correlations = {
        (row["magnification"], row["metric"]): row
        for row in read_csv("wiener_spatial_variation_correlations.csv")
    }
    spatial_max_distance = {
        magnification: max(
            float(row["distance_max_object_um"])
            for row in spatial_rows
            if row["magnification"] == magnification
        )
        for magnification in ("50x", "20x")
    }
    usable_counts: dict[str, int] = {}
    for row in selection_rows:
        usable_counts[row["sequence"]] = usable_counts.get(row["sequence"], 0) + 1

    edge_50_raw = edges[("50x", "Raw")]
    edge_50_wiener = edges[("50x", "Wiener deconvolution")]
    edge_50_v42 = edges[("50x", "V42 epoch 120")]
    edge_20_raw = edges[("20x", "Raw")]
    edge_20_wiener = edges[("20x", "Wiener deconvolution")]
    edge_20_v42 = edges[("20x", "V42 epoch 120")]
    edge_50_wiener_reduction = pct_lower(edge_50_raw, edge_50_wiener)
    edge_50_reduction = pct_lower(edge_50_raw, edge_50_v42)
    edge_20_wiener_reduction = pct_lower(edge_20_raw, edge_20_wiener)
    edge_20_reduction = pct_lower(edge_20_raw, edge_20_v42)
    contrast_50_raw = float(contrasts[("50x", "Raw")]["median_signed_michelson_pct"])
    contrast_50_wiener = float(contrasts[("50x", "Wiener deconvolution")]["median_signed_michelson_pct"])
    contrast_50_mosr = float(contrasts[("50x", "V42 epoch 120")]["median_signed_michelson_pct"])
    contrast_20_raw = float(contrasts[("20x", "Raw")]["median_signed_michelson_pct"])
    contrast_20_wiener = float(contrasts[("20x", "Wiener deconvolution")]["median_signed_michelson_pct"])
    contrast_20_mosr = float(contrasts[("20x", "V42 epoch 120")]["median_signed_michelson_pct"])
    mean_scale_difference_nm = max(
        1000.0 * abs(float(row["raw_mean_magnitude_um"]) - float(row[case]))
        for row in detailed.values()
        for case in ("wiener_mean_magnitude_um", "v42_mean_magnitude_um")
    )

    doc = Document(str(SOURCE_DOCX))
    remove_previous_generated_section(doc)
    remove_section_from_heading(doc, "6.1.2")
    build_revised_psf_validation(doc)
    add_heading(doc, "6.2 神經網路重建效果與位移量測結果", level=1, page_break=True)
    add_body(
        doc,
        "本節評估量測導向超解析（MOSR）模型在 100MG 樣品影像上的實際重建效果。為判斷改善是否只來自一般影像銳化，另加入使用擬合 PSF 的 Wiener 反卷積作為傳統方法基準，三者以同一批 Raw 影像、相同九顆 pillar、相同 120 個位移步與相同量測規則比較。分析重點不只放在影像是否看起來較清楚，而是分別檢查邊緣過渡、SEM 對照 CD、連續位移重複性，以及同一步內九顆 pillar 是否給出一致位移。",
    )
    add_body(
        doc,
        "本節所稱 MOSR 為完成訓練後固定的單張影像重建模型。評估期間不再更新模型權重，輸入與輸出維持相同影像尺寸；後續所有 CD 與位移結果均由同一組固定模型產生。模型設計與訓練策略已於第四章說明，本節只呈現重建影像及其量測結果。",
    )
    add_body(
        doc,
        "Wiener 基準在 20× 與 50× 均固定採 NSR = 0.05，且不依 SEM CD 或位移誤差重新調整；每一倍率使用同一個校正核心處理兩種位移幅度。Raw、Wiener 與 MOSR 沿用完全相同的影像名單、pillar 中心與量測規則，使三者差異不受資料名單改變影響。",
    )

    add_heading(doc, "6.2.1 MOSR 模型與訓練狀態", level=2)
    add_body(
        doc,
        "MOSR 的完整訓練資料以 11,254 組影像配對為基礎，其中選取 180 組資料將標註邊界改為窄範圍軟化標籤，分布於 15 種資料角色，每類 12 組；原始輸入影像不因軟化標籤而改動。模型同時啟用品質判讀分支，以降低硬邊界階梯化與不同訊雜條件對重建結果的影響。本節使用訓練完成後的固定模型進行評估。",
    )
    add_body(
        doc,
        "訓練紀錄中可可靠對應的三項損失皆由第 1 回合下降至第 120 回合：重建損失由 0.01938 降至 0.01640，下降 15.37%；梯度幅值損失由 0.15801 降至 0.13115，下降 17.00%；邊緣方向損失由 0.004871 降至 0.001837，下降 62.28%。圖中呈現每一回合的原始數值，沒有平滑或挑選局部區間。三項損失整體下降可支持訓練過程已收斂至較低誤差，但真正的量測效益仍須由後續 CD 與位移結果判定。",
    )
    add_figure(doc, "23_training_loss_components.png", "MOSR 第 1 至第 120 回合之未平滑訓練損失變化")
    add_table(
        doc,
        "MOSR 訓練損失摘要",
        ["指標", "第 1 回合", "第 120 回合", "下降幅度"],
        [
            [row["metric"], f"{float(row['start_value']):.6f}", f"{float(row['end_value']):.6f}", f"{float(row['favorable_change_pct']):.2f}%"]
            for row in training
        ],
        font_size=9.5,
    )

    add_heading(doc, "6.2.2 評估資料與量測方法", level=2)
    add_body(
        doc,
        "本節定位為既有顯微系統、樣品與成像資料域內的系統一致性驗證，用來比較相同條件下三種影像結果的邊界、尺寸與位移行為；結果不延伸宣稱跨樣品或跨成像系統的泛化能力。",
    )
    add_body(
        doc,
        "所有幾何量測固定使用 pillar 與 via 交界處、最靠近實體樣品中心的九顆 pillar，SEM 身分依序為 P22、P23、P24、P28、P29、P30、P34、P35 與 P36。結構身分不是依圓形區域的明暗判定，而是先以完整視野中的 pillar、via、水平線與垂直線四個區域共同定位樣品中心交界，再選擇交界旁的 pillar 九宮格。這一步可避免把外觀相似但物理結構不同的 via 誤當作 pillar。",
    )
    add_body(
        doc,
        "兩種倍率的樣品擺放方向並不相同。50× 影像維持原樣方向，其 pillar 區位於樣品中心交界的右上方；20× 影像相對於 50× 呈上下翻轉，真正的 pillar 區位於交界的右下方。因此，20× 分析採右下角最靠近交界的九顆暗色 pillar，先前位於左上方的亮圓形區域屬於 via，不納入本節結果。完成方位判定後，Raw、Wiener 與 MOSR 才沿用同一組中心、同一組影像檔名及同一組相鄰影格配對。",
    )
    add_figure(doc, "00a_50x_full_fov_pillar_locator.png", "50× 完整視野之樣品中心、pillar、via 與中心九宮格定位")
    add_figure(doc, "00b_20x_full_fov_pillar_locator.png", "20× 完整視野之上下翻轉方位與中心九宮格定位")
    add_body(
        doc,
        f"方位修正後，兩倍率選到的 pillar 均為亮背景上的暗結構。以九顆中位數計算帶正負號的 Michelson 對比，50× Raw、Wiener 與 MOSR 分別為 {contrast_50_raw:.1f}%、{contrast_50_wiener:.1f}% 與 {contrast_50_mosr:.1f}%，20× 分別為 {contrast_20_raw:.1f}%、{contrast_20_wiener:.1f}% 與 {contrast_20_mosr:.1f}%；負號代表 pillar 比周圍背景暗。Wiener 的對比提高較多，但三種影像在兩倍率下都保留相同明暗方向，因此比較的是同一種 pillar，而不是 pillar 與 via 混用或結構明暗顛倒。",
    )
    add_figure(doc, "00c_signed_pillar_contrast.png", "50× 與 20× 中心九顆 pillar 之帶正負號對比")
    sampling_rows = [
        ["50×", "0.5 µm", "0.0690 µm/pixel", "7.246 pixels", "50", "30"],
        ["50×", "1.0 µm", "0.0690 µm/pixel", "14.493 pixels", "49", "30"],
        ["20×", "0.5 µm", "0.1725 µm/pixel", "2.899 pixels", "49", "30"],
        ["20×", "1.0 µm", "0.1725 µm/pixel", "5.797 pixels", "49", "30"],
    ]
    add_table(
        doc,
        "四組位移序列之取樣條件與主分析步數",
        ["倍率", "名義位移", "物體空間取樣", "等效像素位移", "拍攝相鄰步", "正式分析步"],
        sampling_rows,
        font_size=8.0,
    )
    add_body(
        doc,
        "CD 量測採用固定的 Raw 模板中心與參考半徑，不為 Wiener 或 MOSR 重新尋找較有利的位置。SEM 試算表提供的數值為半徑，因此先乘以 2 轉為直徑後再比較；物體空間取樣為 69.0 nm/pixel。為避免單一門檻剛好落在有利位置，相對門檻由 35% 至 65% 逐級掃描，並以水平七列剖面及 360 度平均徑向剖面交叉檢查。每顆 pillar、每種方法共得到 14 個 CD，逐柱比較只呈現其中的最小值與最大值，不保留任何單一門檻或中央代表點；整體判斷則使用全部門檻、兩種剖面與九顆 pillar 的誤差分布。",
    )
    add_body(
        doc,
        "邊緣銳利度以徑向 10–90% 過渡寬度表示。每顆 pillar 由 360 個方向取樣，沿徑向使用相同且輕微的平滑，再以各自的局部 pillar 平台與背景平台正規化；過渡寬度越小，代表亮暗變化集中於越窄的空間。倍率摘要採九顆 pillar 過渡寬度的中位數，以降低局部過衝或少數方向交點不穩定的影響。",
    )
    add_body(
        doc,
        "位移量測在每顆 pillar 上使用五條水平弦，分別找出左右邊緣的次像素梯度位置，再以兩側中點作為該顆 pillar 的水平中心。相鄰影格中九顆 pillar 的位移以中位數形成該步共識位移。『九宮格內不一致量』定義為九顆位移相對共識值之絕對偏差中位數乘以 1.4826，再換算為 nm；數值越小，代表同一實體平移在九顆 pillar 上得到越一致的量測。『位移標準差』則是整段序列中各步共識位移的標準差，用來描述步與步之間的重複性。",
    )
    add_body(
        doc,
        "『指令值平均絕對偏差』是每一步量測幅度與 0.5 µm 或 1.0 µm 名義指令之差的絕對值平均。此指標表示與平台指令的一致程度，不能取代獨立干涉儀或編碼器提供的位移真值。各序列依預先固定的影像完整性與平台運動品質條件排除無法可靠量測的影格，再統一保留 30 個有效步；三種方法共用相同名單。",
    )
    add_body(
        doc,
        "正式 30 步名單在比較三種方法前即固定，不依 Wiener 或 MOSR 的誤差大小與改善幅度更換。重新反卷積後，120 個正式步皆能完成三方法的九顆 pillar 量測，因此後續統計維持一對一配對。",
    )

    add_heading(doc, "6.2.3 影像重建與邊緣響應", level=2, page_break=True)
    add_body(
        doc,
        "為避免顯示對比造成主觀誤判，50× 與 20× 的 Raw、Wiener 與 MOSR 影像均以相同的 8-bit 灰階範圍 0–255 顯示，沒有對個別影像做自動對比拉伸。以下影像與剖面圖均採單張呈現，使三種結果可分別引用。影像示例使用各倍率的保留校正影格，位移統計則使用未參與 PSF 與雜訊參數選擇的正式 120 步。",
    )
    add_figure(doc, "01_50x_raw_center9.png", "50× 原始影像之中心九顆 pillar", width=4.8)
    add_figure(doc, "02a_50x_wiener_center9.png", "50× 使用固定擬合 PSF 之 Wiener 反卷積結果", width=4.8)
    add_figure(doc, "02_50x_v42_center9.png", "50× MOSR 模型重建之中心九顆 pillar", width=4.8)
    add_body(
        doc,
        f"50× 影像中，九顆 pillar 的徑向 10–90% 過渡寬度中位數由 Raw 的 {edge_50_raw:.1f} nm 降至 Wiener 的 {edge_50_wiener:.1f} nm 與 MOSR 的 {edge_50_v42:.1f} nm，分別縮小 {edge_50_wiener_reduction:.2f}% 與 {edge_50_reduction:.2f}%。此倍率由 Wiener 得到最窄過渡，並伴隨較明顯的亮暗反彈；MOSR 同樣比 Raw 集中，但保留較平緩的輪廓。這表示新版反卷積確實達到預期的視覺銳化，而非預設 MOSR 必須在每一項銳利度數字上勝出。",
    )
    add_figure(doc, "03_50x_pillar_edge_profile.png", "50× Raw、Wiener 與 MOSR 徑向邊緣剖面及九顆過渡寬度中位數")
    add_figure(doc, "04_20x_raw_center9.png", "20× 原始影像之中心九顆 pillar", width=4.8)
    add_figure(doc, "05a_20x_wiener_center9.png", "20× 使用固定擬合 PSF 之 Wiener 反卷積結果", width=4.8)
    add_figure(doc, "05_20x_v42_center9.png", "20× MOSR 模型重建之中心九顆 pillar", width=4.8)
    add_body(
        doc,
        f"20× 影像的物體空間取樣較粗，原始邊界跨越較大的實際距離。九顆過渡寬度中位數由 Raw 的 {edge_20_raw:.1f} nm 降至 Wiener 的 {edge_20_wiener:.1f} nm 與 MOSR 的 {edge_20_v42:.1f} nm，分別縮小 {edge_20_wiener_reduction:.2f}% 與 {edge_20_reduction:.2f}%。此倍率由 MOSR 得到最窄過渡，Wiener 也明顯比 Raw 銳利。兩種結果的剖面都可見不同程度的明暗反彈，因此應解讀為較強的邊緣集中，同時保留可能振鈴的限制，而不是宣稱無失真恢復。",
    )
    add_figure(doc, "06_20x_pillar_edge_profile.png", "20× Raw、Wiener 與 MOSR 徑向邊緣剖面及九顆過渡寬度中位數")

    add_heading(doc, "6.2.4 SEM 對照之 CD 修正與門檻穩健度", level=2)
    add_heading(doc, "6.2.4.1 50× 靜態影像之跨門檻比較", level=3)
    robust_raw = cd_robust["raw"]
    robust_wiener = cd_robust["wiener"]
    robust_v42 = cd_robust["v42_e120"]
    raw_span = statistics.median(float(row["raw_span_nm"]) for row in cd_bounds)
    wiener_span = statistics.median(float(row["wiener_span_nm"]) for row in cd_bounds)
    mosr_span = statistics.median(float(row["mosr_span_nm"]) for row in cd_bounds)
    add_body(
        doc,
        "為避免以單一有利門檻評估影像修復方法，本圖對 Raw、Wiener 與 MOSR 都不保留 50% 代表值，也不以任何中央值取代。對每顆 pillar、每種方法分別掃描 35–65% 七個相對門檻，並使用水平剖面與 360 度平均徑向剖面，共形成 14 個量測值；圖中只畫出其中最小與最大直徑。SEM 圓點為已知尺寸參考，不參與上下界計算。",
    )
    per_pillar_rows = []
    for row in cd_bounds:
        per_pillar_rows.append(
            [
                f"P{int(float(row['sem_index']))}",
                f"{float(row['sem_diameter_um']):.3f}",
                f"{float(row['raw_lower_um']):.3f}–{float(row['raw_upper_um']):.3f}",
                f"{float(row['wiener_lower_um']):.3f}–{float(row['wiener_upper_um']):.3f}",
                f"{float(row['mosr_lower_um']):.3f}–{float(row['mosr_upper_um']):.3f}",
            ]
        )
    add_table(
        doc,
        "35–65% 門檻與兩種剖面下中心九顆 pillar 之 CD 上下界",
        ["Pillar", "SEM CD\n(µm)", "Raw bounds\n(µm)", "Wiener bounds\n(µm)", "MOSR bounds\n(µm)"],
        per_pillar_rows,
        font_size=9.0,
    )
    add_figure(doc, "07_center9_cd_per_pillar.png", "35–65% 門檻與兩種剖面下中心九顆 pillar 之 SEM 參考及 Raw、Wiener、MOSR CD 上下界")
    add_body(
        doc,
        f"九顆 pillar 的上下界寬度中位數為 Raw {raw_span:.1f} nm、Wiener {wiener_span:.1f} nm、MOSR {mosr_span:.1f} nm。Wiener 的範圍最寬，反映環狀反彈會使交點隨門檻明顯移動；因此不能用某一個看似接近 SEM 的交點代表整體尺寸修復。將全部 126 個觀察納入後，絕對誤差中位數為 Raw {float(robust_raw['pooled_median_abs_error_nm']):.1f} nm、Wiener {float(robust_wiener['pooled_median_abs_error_nm']):.1f} nm、MOSR {float(robust_v42['pooled_median_abs_error_nm']):.1f} nm；相對 Raw，MOSR 下降 {float(robust_v42['median_abs_improvement_vs_raw_pct']):.2f}%，Wiener 則增加 {abs(float(robust_wiener['median_abs_improvement_vs_raw_pct'])):.2f}%。",
    )
    cd_metric_rows = []
    for key in ["raw", "wiener", "v42_e120"]:
        row = cd_robust[key]
        change = "—" if key == "raw" else f"{float(row['median_abs_improvement_vs_raw_pct']):+.2f}%"
        cd_metric_rows.append(
            [
                {"raw": "Raw", "wiener": "Wiener", "v42_e120": "MOSR"}[key],
                f"{float(row['pooled_median_abs_error_nm']):.1f}",
                f"{float(row['pooled_mean_abs_error_nm']):.1f}",
                f"{float(row['pooled_rmse_nm']):.1f}",
                change,
            ]
        )
    add_table(
        doc,
        "跨兩種剖面與七個門檻之 CD 誤差穩健度",
        ["方法", "中位絕對誤差\n(nm)", "平均絕對誤差\n(nm)", "均方根誤差\n(nm)", "中位誤差相對 Raw"],
        cd_metric_rows,
        font_size=9.2,
    )
    add_figure(doc, "08_center9_cd_error_summary.png", "中心九顆 pillar 之跨門檻 CD 絕對誤差中位數")
    add_body(
        doc,
        "兩張門檻敏感度圖說明排名翻轉的原因。Wiener 曲線在 50% 附近形成尖銳低谷，門檻稍微提高或降低後誤差便快速上升；MOSR 雖不在每一個門檻都最低，但跨越兩種剖面方式時整體誤差較低，也沒有依賴單一交點才能勝出。這項分析不否定 Wiener 能使影像變銳利，而是把『視覺銳化』與『可穩定量測的 CD 修復』分開。",
    )
    add_figure(
        doc,
        "09a_cd_threshold_sensitivity_horizontal.png",
        "水平七列剖面之 CD 門檻敏感度",
        width=5.50,
    )
    add_figure(
        doc,
        "09b_cd_threshold_sensitivity_radial.png",
        "360 度平均徑向剖面之 CD 門檻敏感度",
        width=5.50,
    )
    add_body(
        doc,
        "因此，本節不以任何單一門檻下的 Wiener 數值作為結論。跨 35–65% 門檻與兩種剖面方式後，MOSR 的上下界較窄且中位絕對誤差最低，較能支持可重複的 CD 邊界恢復。資料僅含一張 50× 靜態影像的九顆 pillar，故僅能主張此區域的穩健度，不作完整視野絕對尺寸校準。",
    )

    add_heading(doc, "6.2.4.2 20× 非 PSF 校正影格之補充檢查", level=3)
    supp_raw = supplementary_20x["raw"]
    supp_wiener = supplementary_20x["wiener"]
    supp_mosr = supplementary_20x["mosr"]
    add_body(
        doc,
        "為確認 20× 結果不只出現在擬合 PSF 的影格，另選用 1.0 µm 位移序列最後一張、未用於 PSF 擬合的影像進行補充 CD 檢查。九顆 pillar 身分與 SEM 直徑維持不變；中心只由 Raw 影像估計一次共同網格位移，Wiener 與 MOSR 不重新定位。每種方法同樣使用 35–65% 七個門檻與兩種剖面，共得到 126 個尺寸誤差。此測試仍屬相同樣品與成像資料域內的一致性驗證。",
    )
    add_figure(
        doc,
        "24_20x_noncalibration_cd_error_summary.png",
        "20× 非 PSF 校正影格之 Raw、Wiener 與 MOSR 跨門檻 CD 絕對誤差中位數",
        width=5.75,
    )
    supplementary_rows = []
    for key, label in (("raw", "Raw"), ("wiener", "Wiener"), ("mosr", "MOSR")):
        row = supplementary_20x[key]
        supplementary_rows.append(
            [
                label,
                int(float(row["observation_count"])),
                f"{float(row['pooled_median_abs_error_nm']):.1f}",
                f"{float(row['pooled_mean_abs_error_nm']):.1f}",
                f"{float(row['pooled_rmse_nm']):.1f}",
                f"{float(row['pooled_mean_signed_bias_nm']):+.1f}",
            ]
        )
    add_table(
        doc,
        "20× 非 PSF 校正影格之跨門檻 CD 一致性",
        ["方法", "量測數", "中位絕對誤差 (nm)", "平均絕對誤差 (nm)", "均方根誤差 (nm)", "平均偏差 (nm)"],
        supplementary_rows,
        font_size=8.5,
    )
    raw_supp_error = float(supp_raw["pooled_median_abs_error_nm"])
    wiener_supp_error = float(supp_wiener["pooled_median_abs_error_nm"])
    mosr_supp_error = float(supp_mosr["pooled_median_abs_error_nm"])
    add_body(
        doc,
        f"跨全部量測條件後，中位絕對 CD 誤差由 Raw 的 {raw_supp_error:.1f} nm 降至 Wiener 的 {wiener_supp_error:.1f} nm 與 MOSR 的 {mosr_supp_error:.1f} nm，"
        f"相對 Raw 分別下降 {100.0 * (raw_supp_error - wiener_supp_error) / raw_supp_error:.1f}% 與 {100.0 * (raw_supp_error - mosr_supp_error) / raw_supp_error:.1f}%。"
        "此補充結果顯示 20× 的傳統反卷積確實可帶來尺寸一致性的有限改善，而 MOSR 在相同跨門檻規則下保留較低的整體誤差。",
    )

    add_heading(doc, "6.2.5 位移量測結果", level=2, page_break=True)
    add_heading(doc, "6.2.5.1 四組序列之逐步結果", level=3)
    s = detailed["50x_x_0p5um"]
    add_body(
        doc,
        f"50×、0.5 µm 序列固定分析 {int(float(s['paired_count']))} 步。Raw、Wiener 與 MOSR 的步階標準差依序為 {float(s['raw_step_sd_nm']):.2f}、{float(s['wiener_step_sd_nm']):.2f} 與 {float(s['v42_step_sd_nm']):.2f} nm；指令值平均絕對偏差為 {float(s['raw_command_mae_nm']):.2f}、{float(s['wiener_command_mae_nm']):.2f} 與 {float(s['v42_command_mae_nm']):.2f} nm；九宮格內不一致量則為 {float(s['raw_within_mean_nm']):.2f}、{float(s['wiener_within_mean_nm']):.2f} 與 {float(s['v42_within_mean_nm']):.2f} nm。兩種處理都改善 Raw，MOSR 在三項指標上均最低。",
    )
    add_figure(doc, "10_50x_0p5um_step_trace.png", "50×、0.5 µm 序列之 Raw、Wiener 與 MOSR 逐步位移")
    s = detailed["50x_x_1um"]
    add_body(
        doc,
        f"50×、1.0 µm 序列固定分析 {int(float(s['paired_count']))} 步。Raw、Wiener 與 MOSR 的步階標準差為 {float(s['raw_step_sd_nm']):.2f}、{float(s['wiener_step_sd_nm']):.2f} 與 {float(s['v42_step_sd_nm']):.2f} nm；指令值平均絕對偏差為 {float(s['raw_command_mae_nm']):.2f}、{float(s['wiener_command_mae_nm']):.2f} 與 {float(s['v42_command_mae_nm']):.2f} nm；九宮格內不一致量為 {float(s['raw_within_mean_nm']):.2f}、{float(s['wiener_within_mean_nm']):.2f} 與 {float(s['v42_within_mean_nm']):.2f} nm。MOSR 再度得到最低的步階散佈、指令偏差與九顆分歧。",
    )
    add_figure(doc, "11_50x_1um_step_trace.png", "50×、1.0 µm 序列之 Raw、Wiener 與 MOSR 逐步位移")
    s = detailed["20x_0p5um"]
    add_body(
        doc,
        f"20×、0.5 µm 序列固定分析 {int(float(s['paired_count']))} 步。Raw、Wiener 與 MOSR 的步階標準差為 {float(s['raw_step_sd_nm']):.2f}、{float(s['wiener_step_sd_nm']):.2f} 與 {float(s['v42_step_sd_nm']):.2f} nm，指令值平均絕對偏差為 {float(s['raw_command_mae_nm']):.2f}、{float(s['wiener_command_mae_nm']):.2f} 與 {float(s['v42_command_mae_nm']):.2f} nm，兩項皆由 MOSR 略低；九宮格內不一致量則為 {float(s['raw_within_mean_nm']):.2f}、{float(s['wiener_within_mean_nm']):.2f} 與 {float(s['v42_within_mean_nm']):.2f} nm，由 Wiener 略低。此例顯示同一序列在不同指標上可有不同排名，因此後續仍以四組合併結果判斷整體穩定性。",
    )
    add_figure(doc, "12_20x_0p5um_step_trace.png", "20×、0.5 µm 固定 30 步之 Raw、Wiener 與 MOSR 逐步位移")
    s = detailed["20x_1um"]
    add_body(
        doc,
        f"20×、1.0 µm 序列固定分析 {int(float(s['paired_count']))} 步。Raw、Wiener 與 MOSR 的步階標準差為 {float(s['raw_step_sd_nm']):.2f}、{float(s['wiener_step_sd_nm']):.2f} 與 {float(s['v42_step_sd_nm']):.2f} nm；指令值平均絕對偏差為 {float(s['raw_command_mae_nm']):.2f}、{float(s['wiener_command_mae_nm']):.2f} 與 {float(s['v42_command_mae_nm']):.2f} nm；九宮格內不一致量為 {float(s['raw_within_mean_nm']):.2f}、{float(s['wiener_within_mean_nm']):.2f} 與 {float(s['v42_within_mean_nm']):.2f} nm。MOSR 在三項指標均略優於 Wiener。",
    )
    add_figure(doc, "13_20x_1um_step_trace.png", "20×、1.0 µm 序列之 Raw、Wiener 與 MOSR 逐步位移")

    stability_rows = []
    local_rows = []
    tail_rows = []
    for key in ["50x_x_0p5um", "50x_x_1um", "20x_0p5um", "20x_1um"]:
        row = detailed[key]
        stability_rows.append(
            [
                row["condition"],
                int(float(row["paired_count"])),
                f"{float(row['raw_step_sd_nm']):.2f}",
                f"{float(row['wiener_step_sd_nm']):.2f}",
                f"{float(row['v42_step_sd_nm']):.2f}",
                f"{float(row['raw_command_mae_nm']):.2f}",
                f"{float(row['wiener_command_mae_nm']):.2f}",
                f"{float(row['v42_command_mae_nm']):.2f}",
            ]
        )
        local_rows.append(
            [
                row["condition"],
                int(float(row["paired_count"])),
                f"{float(row['raw_within_mean_nm']):.2f}",
                f"{float(row['wiener_within_mean_nm']):.2f}",
                f"{float(row['v42_within_mean_nm']):.2f}",
                f"{int(float(row['wiener_lower_within_count']))}/{int(float(row['paired_count']))} ({float(row['wiener_lower_within_pct']):.1f}%)",
                f"{int(float(row['v42_lower_within_count']))}/{int(float(row['paired_count']))} ({float(row['v42_lower_within_pct']):.1f}%)",
            ]
        )
        raw_cmd_p90 = float(row["raw_command_p90_ae_nm"])
        wiener_cmd_p90 = float(row["wiener_command_p90_ae_nm"])
        v42_cmd_p90 = float(row["v42_command_p90_ae_nm"])
        raw_local_p90 = float(row["raw_within_p90_nm"])
        wiener_local_p90 = float(row["wiener_within_p90_nm"])
        v42_local_p90 = float(row["v42_within_p90_nm"])
        tail_rows.append(
            [
                row["condition"],
                f"{raw_cmd_p90:.2f}",
                f"{wiener_cmd_p90:.2f}",
                f"{v42_cmd_p90:.2f}",
                f"{raw_local_p90:.2f}",
                f"{wiener_local_p90:.2f}",
                f"{v42_local_p90:.2f}",
            ]
        )
    add_table(
        doc,
        "四組序列之步階穩定性與名義指令一致性",
        ["條件", "n", "Raw SD", "Wiener SD", "MOSR SD", "Raw 指令 MAE", "Wiener 指令 MAE", "MOSR 指令 MAE"],
        stability_rows,
        font_size=7.6,
    )
    add_table(
        doc,
        "四組序列之中心九顆 pillar 內部一致性",
        ["條件", "n", "Raw 不一致量", "Wiener 不一致量", "MOSR 不一致量", "Wiener 改善步", "MOSR 改善步"],
        local_rows,
        font_size=7.8,
    )

    add_heading(doc, "6.2.5.2 跨序列整體結果", level=3, page_break=True)
    all_row = pooled["All paired steps"]
    add_body(
        doc,
        f"四組序列各固定 30 步，合併後共有 {int(float(all_row['paired_count']))} 個配對步。Raw、Wiener 與 MOSR 的指令值平均絕對偏差為 {float(all_row['raw_command_mae_nm']):.2f}、{float(all_row['wiener_command_mae_nm']):.2f} 與 {float(all_row['v42_command_mae_nm']):.2f} nm；Wiener 與 MOSR 分別較 Raw 下降 {float(all_row['wiener_command_mae_reduction_pct']):.2f}% 與 {float(all_row['v42_command_mae_reduction_pct']):.2f}%。九宮格內平均不一致量則由 Raw 的 {float(all_row['raw_within_mean_nm']):.2f} nm 降至 Wiener 的 {float(all_row['wiener_within_mean_nm']):.2f} nm 與 MOSR 的 {float(all_row['v42_within_mean_nm']):.2f} nm，分別下降 {float(all_row['wiener_within_reduction_pct']):.2f}% 與 {float(all_row['v42_within_reduction_pct']):.2f}%。因此，兩種處理都有效，而 MOSR 在合併後的指令一致性與九顆 pillar 一致性皆最低。",
    )
    add_figure(doc, "14_step_stability_summary.png", "四組序列之 Raw、Wiener 與 MOSR 位移標準差")
    add_figure(doc, "15_nine_pillar_consistency_summary.png", "四組序列之 Raw、Wiener 與 MOSR 九顆 pillar 內部不一致量")
    add_figure(doc, "16_nominal_command_consistency.png", "Raw、Wiener 與 MOSR 對名義位移指令之平均絕對偏差")
    add_body(
        doc,
        f"逐步比較亦呈現相同方向。Wiener 在四組序列中分別有 {float(detailed['50x_x_0p5um']['wiener_lower_within_pct']):.1f}%、{float(detailed['50x_x_1um']['wiener_lower_within_pct']):.1f}%、{float(detailed['20x_0p5um']['wiener_lower_within_pct']):.1f}% 與 {float(detailed['20x_1um']['wiener_lower_within_pct']):.1f}% 的步數降低九宮格不一致量；MOSR 則為 {float(detailed['50x_x_0p5um']['v42_lower_within_pct']):.1f}%、{float(detailed['50x_x_1um']['v42_lower_within_pct']):.1f}%、{float(detailed['20x_0p5um']['v42_lower_within_pct']):.1f}% 與 {float(detailed['20x_1um']['v42_lower_within_pct']):.1f}%。兩種方法四組都高於 50%；MOSR 在兩組 50× 序列較高，20×、0.5 µm 兩者同為 93.3%，20×、1.0 µm 則由 Wiener 較高。",
    )
    add_figure(doc, "17_stepwise_local_improvement_rate.png", "Wiener 與 MOSR 逐步降低中心九顆 pillar 不一致量之比例")
    p50 = pooled["50x"]
    p20 = pooled["20x"]
    pooled_rows = [
        ["全部", all_row["paired_count"], f"{float(all_row['raw_command_mae_nm']):.2f}", f"{float(all_row['wiener_command_mae_nm']):.2f}", f"{float(all_row['v42_command_mae_nm']):.2f}", f"{float(all_row['raw_within_mean_nm']):.2f}", f"{float(all_row['wiener_within_mean_nm']):.2f}", f"{float(all_row['v42_within_mean_nm']):.2f}"],
        ["50×", p50["paired_count"], f"{float(p50['raw_command_mae_nm']):.2f}", f"{float(p50['wiener_command_mae_nm']):.2f}", f"{float(p50['v42_command_mae_nm']):.2f}", f"{float(p50['raw_within_mean_nm']):.2f}", f"{float(p50['wiener_within_mean_nm']):.2f}", f"{float(p50['v42_within_mean_nm']):.2f}"],
        ["20×", p20["paired_count"], f"{float(p20['raw_command_mae_nm']):.2f}", f"{float(p20['wiener_command_mae_nm']):.2f}", f"{float(p20['v42_command_mae_nm']):.2f}", f"{float(p20['raw_within_mean_nm']):.2f}", f"{float(p20['wiener_within_mean_nm']):.2f}", f"{float(p20['v42_within_mean_nm']):.2f}"],
    ]
    add_table(
        doc,
        "全部配對步與分倍率之位移結果",
        ["母體", "n", "Raw 指令", "Wiener 指令", "MOSR 指令", "Raw 九宮格", "Wiener 九宮格", "MOSR 九宮格"],
        pooled_rows,
        font_size=7.8,
    )
    add_body(
        doc,
        f"分倍率後，MOSR 在 50× 的指令值平均絕對偏差下降 {float(p50['v42_command_mae_reduction_pct']):.2f}%，九宮格內不一致量下降 {float(p50['v42_within_reduction_pct']):.2f}%；20× 分別下降 {float(p20['v42_command_mae_reduction_pct']):.2f}% 與 {float(p20['v42_within_reduction_pct']):.2f}%。Wiener 在 50× 的兩項下降幅度為 {float(p50['wiener_command_mae_reduction_pct']):.2f}% 與 {float(p50['wiener_within_reduction_pct']):.2f}%，20× 為 {float(p20['wiener_command_mae_reduction_pct']):.2f}% 與 {float(p20['wiener_within_reduction_pct']):.2f}%。此差異與物體空間取樣相符：0.5 µm 與 1.0 µm 在 50× 影像中分別約為 7.25 與 14.49 pixels，在 20× 中僅約為 2.90 與 5.80 pixels；但此處只可視為合理解釋，不能單獨證明因果。",
    )
    add_body(
        doc,
        f"兩種處理都沒有以縮放整段位移來換取較低散佈。四組序列中，Wiener 或 MOSR 相對 Raw 的平均位移幅度差異最大僅 {mean_scale_difference_nm:.2f} nm，遠小於 0.5 µm 與 1.0 µm 的指令尺度。換言之，主要變化是減少步與步、pillar 與 pillar 之間的散佈，而不是改變整體位移比例。",
    )
    add_body(
        doc,
        "第 90 百分位結果用來確認入選 30 步中較差的尾端是否同樣受益。Wiener 與 MOSR 在四組序列都降低九宮格不一致量的尾端，其中 MOSR 在三組低於 Wiener，20×、1.0 µm 則由 Wiener 較低。指令值偏差尾端也沒有固定由同一方法勝出：MOSR 在 50×、0.5 µm 較低，其餘三組由 Wiener 略低。這顯示單一尾端指標會隨條件變動，較適合搭配合併結果而非單獨決定方法排名。",
    )
    add_table(
        doc,
        "四組序列之第 90 百分位誤差",
        ["條件", "Raw 指令 P90", "Wiener 指令 P90", "MOSR 指令 P90", "Raw 九宮格 P90", "Wiener 九宮格 P90", "MOSR 九宮格 P90"],
        tail_rows,
        font_size=7.7,
    )

    add_heading(doc, "6.2.5.3 固定 PSF 的空間變異性檢查", level=3)
    rho_50_cmd = spatial_correlations[("50x", "wiener_command_abs_error_nm")]
    rho_20_cmd = spatial_correlations[("20x", "wiener_command_abs_error_nm")]
    rho_50_grid = spatial_correlations[("50x", "wiener_within_grid_nm")]
    rho_20_grid = spatial_correlations[("20x", "wiener_within_grid_nm")]
    add_body(
        doc,
        f"Wiener 反卷積在每一倍率只使用一個固定 PSF，若光學模糊具有明顯空間變異性，離校正位置越遠的步數理應逐漸出現較大的誤差。為檢查這一點，將每倍率 60 個正式步依影像中心相對 PSF 校正中心的物體空間距離平均分成 Near、Middle 與 Far 三組，每組 20 步。50× 的最遠評估距離為 {spatial_max_distance['50x']:.2f} µm，20× 為 {spatial_max_distance['20x']:.2f} µm；兩種位移幅度仍共用同一個倍率 PSF。",
    )
    add_figure(doc, "18a_50x_wiener_psf_distance_robustness.png", "50× 固定 PSF 相對校正位置距離之位移誤差")
    add_figure(doc, "18b_20x_wiener_psf_distance_robustness.png", "20× 固定 PSF 相對校正位置距離之位移誤差")
    add_body(
        doc,
        f"距離與 Wiener 指令值絕對偏差的 Spearman 相關在 50× 為 ρ = {float(rho_50_cmd['spearman_rho']):.3f}、p = {float(rho_50_cmd['p_value']):.3f}，20× 為 ρ = {float(rho_20_cmd['spearman_rho']):.3f}、p = {float(rho_20_cmd['p_value']):.3f}；距離與九宮格不一致量的相關則分別為 ρ = {float(rho_50_grid['spearman_rho']):.3f}、p = {float(rho_50_grid['p_value']):.3f}，以及 ρ = {float(rho_20_grid['spearman_rho']):.3f}、p = {float(rho_20_grid['p_value']):.3f}。前三項未達顯著，20× 九宮格不一致量則呈顯著負相關，而不是隨距離變差。換言之，四項都沒有支持『離校正位置越遠，Wiener 誤差越大』的單調趨勢；目前約 25–30 µm 的取樣範圍內，固定 PSF 尚可作為公平的傳統基準。這不代表完整視野沒有空間變異性，只表示它不足以解釋本次 MOSR 與 Wiener 的主要差異。",
    )

    add_heading(doc, "6.2.6 綜合討論與研究限制", level=2)
    add_body(
        doc,
        f"綜合影像與量測結果可整理為三項主要發現。第一，新版 Wiener 在兩倍率都使邊界比 Raw 更集中，50× 的過渡寬度甚至低於 MOSR；20× 則由 MOSR 得到最窄過渡，因此兩種方法各有較有利的銳利度條件。第二，逐柱 CD 不採單一代表點，而以 35–65% 門檻及兩種剖面所得上下界呈現；Wiener 的九柱上下界寬度中位數為 {wiener_span:.1f} nm，大於 Raw 的 {raw_span:.1f} nm 與 MOSR 的 {mosr_span:.1f} nm。納入全部量測條件後，MOSR 的中位絕對誤差為 {float(robust_v42['pooled_median_abs_error_nm']):.1f} nm，低於 Raw 的 {float(robust_raw['pooled_median_abs_error_nm']):.1f} nm 與 Wiener 的 {float(robust_wiener['pooled_median_abs_error_nm']):.1f} nm。第三，合併 {int(float(all_row['paired_count']))} 步後，九宮格不一致量由 Raw {float(all_row['raw_within_mean_nm']):.2f} nm 降至 Wiener {float(all_row['wiener_within_mean_nm']):.2f} nm 與 MOSR {float(all_row['v42_within_mean_nm']):.2f} nm，MOSR 的多顆一致性最低。",
    )
    add_body(
        doc,
        "本節結果的適用範圍為目前的成像系統、樣品與資料域，不延伸主張跨樣品或跨系統泛化。平台只有名義指令而沒有獨立位置回授，因此可支持的是重複性、指令一致性與九宮格內一致性改善，而非絕對位移準確度。正式三方法 SEM 對照 CD 使用一張 50× 靜態影像與九顆 pillar；20× CD 則在 PSF 一致性檢查中另行呈現。Wiener 與 MOSR 剖面均可見不同程度的明暗反彈，尤其 Wiener 的 CD 交點對門檻較敏感，因此不能把更高對比或所有高頻增加直接視為真實結構恢復。",
    )
    add_body(
        doc,
        "因此，本節最有力且較保守的結論是：單純 Wiener 反卷積已能明顯提升對比與邊緣銳利度，並改善多數位移指標；其代價是較明顯的環狀反彈，使 CD 結果對門檻更敏感。MOSR 並非在每一個銳利度或單一序列數字上都勝過 Wiener，但在跨門檻 CD 與 120 步合併後的九顆 pillar 一致性上較穩定。換言之，本研究不以『看起來最銳利』作為唯一勝負，而以尺寸與位移能否在不同量測設定下維持一致作為主要判斷。",
    )

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_DOCX))
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
