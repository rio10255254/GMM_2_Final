"""Build a single self-contained Word report from frozen R140v2 evidence."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[2]
HERE=Path(__file__).resolve().parent
DATA=json.loads((HERE/"doc_data.json").read_text(encoding="utf-8"))
OUT=ROOT/"reports/R140v2_Full_Model_Training_Validation_Report_20260923.docx"
FIG=HERE/"figures"

doc=Document()
sec=doc.sections[0]
sec.page_height,sec.page_width=Cm(29.7),Cm(21.0)
sec.top_margin,sec.bottom_margin=Cm(1.78),Cm(1.65)
sec.left_margin,sec.right_margin=Cm(1.82),Cm(1.82)
sec.header_distance,sec.footer_distance=Cm(.7),Cm(.8)

def set_font(style,size,bold=False,color=(0,0,0)):
    s=doc.styles[style]
    s.font.name="Microsoft JhengHei"
    s.font.size=Pt(size)
    s.font.bold=bold
    s.font.color.rgb=RGBColor(*color)
    rpr=s._element.get_or_add_rPr()
    fonts=rpr.rFonts
    if fonts is None:
        fonts=OxmlElement("w:rFonts");rpr.insert(0,fonts)
    fonts.set(qn("w:eastAsia"),"Microsoft JhengHei")
    fonts.set(qn("w:ascii"),"Microsoft JhengHei")
    fonts.set(qn("w:hAnsi"),"Microsoft JhengHei")

set_font("Normal",10.3)
set_font("Title",21,True)
set_font("Subtitle",10.5,False,(65,76,84))
set_font("Heading 1",15,True)
set_font("Heading 2",11.8,True)
set_font("Caption",8.8,False,(72,82,89))
title_ppr=doc.styles["Title"]._element.get_or_add_pPr()
title_border=title_ppr.find(qn("w:pBdr"))
if title_border is not None:title_ppr.remove(title_border)
normal=doc.styles["Normal"]
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.18
doc.styles["Heading 1"].paragraph_format.space_before=Pt(16)
doc.styles["Heading 1"].paragraph_format.space_after=Pt(7)
doc.styles["Heading 1"].paragraph_format.keep_with_next=True
doc.styles["Heading 2"].paragraph_format.space_before=Pt(11)
doc.styles["Heading 2"].paragraph_format.space_after=Pt(5)
doc.styles["Heading 2"].paragraph_format.keep_with_next=True
doc.styles["Caption"].paragraph_format.space_before=Pt(2)
doc.styles["Caption"].paragraph_format.space_after=Pt(9)

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
rr=footer.add_run("R140v2  模型訓練與真實影像驗證  ·  2026年9月23日")
rr.font.name="Microsoft JhengHei";rr.font.size=Pt(8);rr.font.color.rgb=RGBColor(95,104,110)

def para(text="",bold_prefix=None):
    p=doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True;p.add_run(text[len(bold_prefix):])
    else:p.add_run(text)
    return p

def heading(text,level=1):
    p=doc.add_paragraph(text,style=f"Heading {level}")
    return p

def small_note(text):
    p=doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after=Pt(8)
    r=p.add_run(text);r.italic=True;r.font.size=Pt(8.8);r.font.color.rgb=RGBColor(78,86,94)
    return p

def fmt(x,n=3):return f"{float(x):.{n}f}"

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();x=tcPr.find(qn("w:shd"))
    if x is None:x=OxmlElement("w:shd");tcPr.append(x)
    x.set(qn("w:fill"),fill)

def cell_border(cell):
    tcPr=cell._tc.get_or_add_tcPr();b=tcPr.find(qn("w:tcBorders"))
    if b is None:b=OxmlElement("w:tcBorders");tcPr.append(b)
    for side in ("top","left","bottom","right"):
        e=b.find(qn("w:"+side))
        if e is None:e=OxmlElement("w:"+side);b.append(e)
        e.set(qn("w:val"),"single");e.set(qn("w:sz"),"3");e.set(qn("w:color"),"D9D9D9")

def table(headers,rows,widths=None,font_size=8.35):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    header=t.rows[0]
    for i,label in enumerate(headers):
        c=header.cells[i];c.text=str(label);shade(c,"26394A");cell_border(c)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:c.width=Cm(widths[i])
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
            for r in p.runs:r.bold=True;r.font.size=Pt(font_size);r.font.color.rgb=RGBColor(255,255,255)
    trPr=header._tr.get_or_add_trPr();rep=OxmlElement("w:tblHeader");rep.set(qn("w:val"),"true");trPr.append(rep)
    for ri,values in enumerate(rows):
        row=t.add_row()
        for i,value in enumerate(values):
            c=row.cells[i];c.text=str(value);cell_border(c)
            if widths:c.width=Cm(widths[i])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2:shade(c,"F4F7F9")
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(.6);p.paragraph_format.space_before=Pt(.6)
                p.paragraph_format.line_spacing=1.03
                for r in p.runs:r.font.name="Microsoft JhengHei";r.font.size=Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def figure(path,caption,width_cm=16.7):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.keep_with_next=True
    inline=p.add_run().add_picture(str(path),width=Cm(width_cm))
    inline._inline.docPr.set("descr",caption)
    c=doc.add_paragraph(caption,style="Caption")
    c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p,c

def by(rows,model,kind=None):
    return next(x for x in rows if x["model"]==model and (kind is None or x.get("kind")==kind))

title=doc.add_paragraph(style="Title")
title.add_run("R140v2 模型設計 訓練與真實影像驗證總報告")
title.paragraph_format.space_after=Pt(4)
sub=doc.add_paragraph(style="Subtitle")
sub.add_run("完成於 2026 年 9 月 23 日  ·  30 趟實拍序列  ·  930 張影像  ·  角落九顆柱體")
sub.paragraph_format.space_after=Pt(15)
para("本文件供論文模型評估與後續寫作使用，集中記錄 R140v2 為何重做、網路如何運作、訓練資料與條件、固定真實影像驗證方法，以及所有主要結果。R140v2 已完成從零訓練與全數實拍量測；它改善相鄰短步位移，但沒有同時成為尺寸、輪廓與 SEM 對照的全面最佳模型。")
heading("先看結論",1)
lead=[
    ("短步位移","R140v2 相鄰 PI 平均絕對差 3.640 nm、重拍標準差 3.530 nm；R137 為 3.917 與 3.906 nm。"),
    ("尺寸與輪廓","R140v2 直徑重拍 18.581 nm、輪廓重拍 18.430 nm；R137 分別為 14.262 與 17.710 nm。"),
    ("PSF 實際作用","開關 PSF 對輸出有極小差異；930 張影像的八顆候選權重一律為 12.5%，沒有證據支持逐影像自動挑選。"),
    ("適用判斷","可視為正常條件下短步位移的研究候選；長距離、絕對尺寸與獨立 SEM 準確度尚不支持全面替換。"),
]
table(["面向","核對後的結果"],lead,[3.0,13.2],9.2)
small_note("以下『誤差』為相對 PI 的操作性比較；『重拍』衡量同條件重拍的一致性，並非絕對真值準確度。")

heading("一 研究範圍與修正原因",1)
para("上一版 R139 雖完成 8192 次更新與 930 張影像驗證，審查發現 U-Net 在第 128 次前的修正輸出已完全飽和；之後主網路權重不再有效更新，最後幾乎只把原圖整體調暗 2.287 個八位元灰階，輪廓與位移接近未處理影像。V42 的卷積區塊有正規化，R139 的較深 U-Net 卻省略該穩定措施而沿用 R130 的最高學習率 0.001。這是一個可信的設計失配線索，但本次同時修改多個條件，不能把原因或改善唯一歸給某一模組。")
para("R140 初版短測通過後，在正式更新前因資料讀取及損失程式與原件的位元組雜湊不符而停止。R140v2 另建獨立包，逐位元複製原件、重做短測、重新鎖定來源，沒有改寫前一版的失敗校驗或歷史結果。")
table(["核對對象","實際數量與狀態"],[
    ("合成母場景","訓練 512、開發 128、生成測試 128；共 768 個不同母場景與幾何。"),
    ("正式訓練","16 輪、8192 次獨立更新，最後第 8192 次模型固定使用。"),
    ("真實顯微影像","30 趟重拍 × 每趟 31 張 = 930 張，原尺寸 1544 × 2064。"),
    ("輪廓與位移","新模型 8370 個圓輪廓；七方法合計 12600 筆整體位移，逐筆重算通過。"),
    ("完整性","930 張輸出雜湊、來源程式與既有證據均核對；正式基準未變。"),
],[4.0,12.2],9.0)

heading("二 模型架構與影像處理",1)
figure(FIG/"architecture.png","圖 1  R140v2 的共享四次修正流程。光學模組只在第四次的半解析度解碼位置介入。")
para("每次修正都使用同一組網路權重。輸入由目前估計影像、原始影像、兩者差，以及原圖的水平梯度、垂直梯度、拉普拉斯訊號和局部平均組成七個通道。U-Net 的通道數依序為 32、64、96、128；三次降採樣與三次上採樣均有抗混疊處理，並以跳接保留對應解析度的局部資訊。")
para("網路在每個像素的通道內做正規化，避免以全圖統計改變局部亮度基準。最後的修正頭先乘 0.1 再經過 tanh；每次步幅起始約 0.05、上限 0.25，累積修正限制在 ±0.35，最終輸出仍以原圖為錨並限制在 0 至 1。這些界線用來防止 R139 曾出現的巨大輸出和訓練失活。")
table(["構件","規格","設計作用"],[
    ("編碼與解碼","四層通道 32／64／96／128，三次縮放，對應跳接","組合局部邊界與較大範圍上下文"),
    ("抗混疊","降採樣前固定二項式模糊；雙線性上採樣後再模糊及卷積","減少縮放造成的格點假訊號"),
    ("通道正規化","每個像素跨通道計算，不使用全圖空間統計","限制中間特徵放大"),
    ("反覆修正","同一主網路共用四次；最後一次插入 PSF","保留可檢查的中間狀態"),
    ("原圖錨定","有界殘差疊回原圖；輸出範圍 0 至 1","限制無根據的整幅偏移"),
    ("模型規模","1,446,366 個參數，FP32 推論","比較時應連同實際時間評估"),
],[3.0,5.2,8.0],8.4)
heading("光學模組的接入",2)
para("光學資料包含 25 顆 101 × 101 的量測 PSF。固定的八顆候選編號是 0、3、7、10、14、17、21、24；在半尺寸解碼特徵上做保守的頻域處理，再以可學權重混合。光學分支只在第四次修正使用，最終原始強度 0.021359、比例 0.28，兩者乘積約 0.005981。這是 V42 類型模組的接入，但 PSF 座標原點的處理與原 V42 程式並非逐行相同。")
small_note("PSF 來源只移轉校準緩衝資料，沒有沿用 V42 已訓練網路權重；其校準來源可能與舊 SEM 資料相關，因此這裡不把 SEM 對照稱作完全獨立驗證。")

heading("三 訓練資料與完整訓練方式",1)
para("訓練採用與 R130 同一份合成影像陣列，並用完全相同位元組的資料讀取與損失程式。每個母場景包含兩張對應影像，原陣列為 384 × 384；訓練時裁成 320 × 320，只以中央 128 × 128 計算監督。資料標記顯示沒有直接放入真實顯微影像、SEM 答案或外部來源影像；這是資料包與讀取器可驗證的範圍，不能推論光學校準完全與 SEM 無關。")
catnames={"corners_lines_junctions":"轉角線段與交會","textured_edges":"紋理邊緣","flat_negative":"平坦負例",
          "irregular_curved":"不規則曲線","dense_close_structures":"密集近距結構","clean_identity":"乾淨原樣",
          "noisy_identity":"帶噪原樣"}
cats={(x["split"],x["category"]):x["count"] for x in DATA["categories"]}
table(["合成場景類型","訓練","開發","生成測試"],
      [(name,cats.get(("train",key),0),cats.get(("dev",key),0),cats.get(("test",key),0))
       for key,name in catnames.items()],[7.2,3.0,3.0,3.0],8.8)
para("16 輪中，每輪把 512 個訓練母場景各使用一次，順序由固定種子重排；訓練噪聲依輪次重新產生。開發與生成測試資料固定，不以測試成績挑選保存點。最後只使用第 8192 次模型驗證實拍，沒有把真實影像、PI 位移或 SEM 尺寸作為訓練目標。")
table(["訓練項目","固定條件"],[
    ("最佳化","AdamW；最高學習率 0.0003，最低 0.000003，前 128 次暖機後餘弦下降。"),
    ("正則與穩定性","權重衰減 0.00001；梯度範數上限 1；FP32，不啟用 TF32。"),
    ("損失函數","影像項 1.0 + 梯度項 0.2 + 雙張配對殘差項 0.2。"),
    ("PSF 訓練時序","第 1、2 輪關閉；第 3 輪比例 0.12，至第 9 輪增至 0.28；光學參數學習率為主網路 5%。"),
    ("可重現設定","模型種子 1309001、順序種子 1309002、噪聲種子 1309003；每 128 次保存。"),
    ("模型選取","固定最後第 8192 次，不按實拍結果挑權重。"),
],[4.0,12.2],8.8)
para("影像項抑制重建灰階差；梯度項約束邊緣變化；配對殘差項利用已知合成位移，要求兩張圖的修正量在平移後相容。這個配對限制不是拿真實 PI 位移當答案。")
figure(ROOT/"reports/figures/training.png","圖 2  8192 次從零訓練的平均損失與固定開發資料走勢。單次訓練損失因不同場景起伏，判讀以固定開發資料為主。")
para(f"固定開發資料總誤差從 5.724764 降至第 8192 次的 {DATA['dev'][-1]['total']:.6f}；最終既有生成測試為 {DATA['test']['total']:.6f}。整個訓練含生成評估約 {DATA['training']['elapsed_seconds']/60:.1f} 分鐘。早期及 PSF 啟用後的六個預定門檻均通過。")
heading("預定穩定性門檻",2)
table(["更新次數","固定樣本誤差","修正輸出 99 百分位","飽和比例","結果"],
      [(x["step"],fmt(x["dev0_loss"],3),fmt(x["logit_abs_p99"],3),fmt(x["tanh_saturation_fraction"],3),"通過" if x["passed"] else "失敗")
       for x in DATA["gates"]],[2.2,3.3,4.8,2.5,3.4],8.6)
small_note("飽和定義為修正輸出絕對值大於 8。R139 早期固定樣本曾有 100% 飽和；R140v2 六個預定檢查均為 0%。")

heading("四 真實影像與 SEM 的驗證方法",1)
para("實拍資料是同一批柱體的 30 次完整重拍；每趟 31 張、九顆固定位置，共 930 張。對 Raw 及每個已保存模型使用相同的追蹤種子、ACM 輪廓擬合與 LSC 位移計算。每張影像九個圓輪廓以 64 個輪廓點讀出，位移先取九顆圓心差的分量中位數，再沿固定軸投影。相鄰位移比較位置 j 與 j−1；起點位移直接比較位置 j 與 0，不累加前面 30 段誤差。")
para("方向與尺度僅以最前六趟 Raw 與 PI 對應資料設定，標稱尺度 0.1725 μm/px；其餘 24 趟可作校準後檢查。這是 PI 參考的操作性尺度，尚非獨立可追溯的 SI 絕對尺寸校準。結果主表保留所有 30 趟，同時另列後 24 趟，以免混淆校準趟次與後續趟次。")
para("直徑重拍指每個位置與柱體在 30 趟之間的直徑標準差，再對位置與柱體彙總。輪廓重拍先相對各自擬合圓心轉成 720 個角向半徑，再計算同位置與柱體跨趟標準差。去尺寸形狀重拍還會扣除該輪廓的平均半徑。九顆柱體和 720 個插值方向不是彼此獨立的重拍；配對區間以完整一趟為重抽樣單位。")
para("SEM 對照沿用既有九顆區域與兩個方向候選，沒有為 R140v2 重新挑方向或尺度。它可指出本批資料的直徑差與輪廓差，但目前的個別柱體方向對應、光學 PSF 校準沿革及獨立絕對真值資格仍有限制。")
figure(FIG/"qualitative_crop.png","圖 3  同一張實拍影像的同一區域，Raw、R137 與 R140v2 採相同灰階範圍；青色線為各方法的 ACM 輪廓。這只是單張示例，統計結論以 930 張量測為準。")

heading("五 實拍驗證結果",1)
heading("相鄰與起點位移",2)
figure(FIG/"adjacent_comparison.png","圖 4  六個有效處理條件的相鄰 PI 平均絕對差與重拍標準差。R139 退化結果在下表完整保留。")
names=["Raw","R129","R130-Clean","R137-WeakPSF","R138-V42PSF8","R139-UNetRecurrent","R140v2-StableUNet"]
short={"Raw":"Raw","R129":"R129","R130-Clean":"R130","R137-WeakPSF":"R137","R138-V42PSF8":"R138","R139-UNetRecurrent":"R139","R140v2-StableUNet":"R140v2"}
rows=[]
for m in names:
    a=by(DATA["displacement"],m,"adjacent");s=by(DATA["displacement"],m,"start")
    rows.append((short[m],fmt(a["mae_nm"]),fmt(a["repeatability_nm"]),fmt(a["maximum_error_nm"]),
                 fmt(s["mae_nm"]),fmt(s["repeatability_nm"])))
table(["模型","相鄰差 nm","相鄰重拍 nm","相鄰最大 nm","起點差 nm","起點重拍 nm"],rows,
      [2.3,2.75,2.75,2.75,2.75,2.9],8.25)
para("R140v2 相鄰平均絕對差比 R137 低 0.276669 nm，30 趟中有 23 趟改善；以整趟配對重抽樣得到的 95% 區間為 −0.432 至 −0.119 nm。起點位移平均差僅低 0.004770 nm，配對區間 −0.198 至 +0.186 nm，沒有同等明確的優勢。長距離終點與跨樣本泛化不應由短步結果推論。")
ci_rows=[]
for baseline in ("Raw","R129","R130-Clean","R137-WeakPSF","R138-V42PSF8"):
    for kind,display in (("adjacent_mae_nm","相鄰"),("start_mae_nm","起點")):
        x=next(v for v in DATA["paired"] if v["reference"]==baseline and v["metric"]==kind)
        ci_rows.append((short[baseline],display,fmt(x["difference_nm"]),
                        f"{fmt(x['ci95_low'])} 至 {fmt(x['ci95_high'])}",int(x["improved_scans"])))
table(["相對模型","位移類型","平均差 nm","95% 整趟配對區間 nm","改善趟數"],ci_rows,
      [2.4,2.2,2.7,5.9,3.0],8.35)
small_note("差值定義為 R140v2 減相對模型；負值代表 R140v2 的平均絕對差較低。區間以 30 個完整趟次重抽樣，不能當成 9 顆或 720 個方向的獨立樣本區間。")
figure(FIG/"position_repeatability.png","圖 5  各位置的相鄰位移跨 30 趟重拍標準差。數值以固定九顆圓心的整體位移計算。")
heading("只看校準後的二十四趟",2)
late_names=["Raw","R129","R137-WeakPSF","R140v2-StableUNet"]
late_rows=[]
for m in late_names:
    a=by(DATA["late24"],m,"adjacent");s=by(DATA["late24"],m,"start")
    late_rows.append((short[m],fmt(a["mae_nm"]),fmt(a["repeatability_nm"]),fmt(s["mae_nm"]),fmt(s["repeatability_nm"])))
table(["模型","相鄰差 nm","相鄰重拍 nm","起點差 nm","起點重拍 nm"],late_rows,
      [2.7,3.3,3.4,3.3,3.5],8.7)
para("後 24 趟的相鄰平均差，R140v2 為 3.762 nm，R137 為 4.103 nm；R140v2 在其中 19 趟較好。這 24 趟沒有參與 Raw 的方向與尺度設定，但這一批真實影像在歷代模型研究中已反覆被檢視，不能稱為未見的模型確認集。")

heading("尺寸與輪廓重拍",2)
figure(FIG/"shape_repeatability.png","圖 6  直徑、輪廓與去尺寸形狀的跨趟重拍標準差。三個指標應分開判讀。")
geom_rows=[]
for m in names:
    x=by(DATA["geometry"],m)
    geom_rows.append((short[m],fmt(x["diameter_mean_nm"]),fmt(x["diameter_repeat_sd_nm"]),
                      fmt(x["contour_repeat_sd_nm"]),fmt(x["shape_repeat_sd_nm"])))
table(["模型","平均直徑 nm","直徑重拍 nm","輪廓重拍 nm","去尺寸形狀 nm"],geom_rows,
      [2.6,3.3,3.4,3.3,3.6],8.4)
para("R140v2 的直徑重拍與輪廓重拍遠優於 Raw，卻分別比 R137 高 4.319 與 0.720 nm。對 R137 的整趟配對區間，直徑差為 +2.338 至 +5.534 nm，輪廓差為 +0.022 至 +1.228 nm；因此不能用短步位移的勝出代替尺寸或輪廓勝出。R140v2 與 R129 的輪廓重拍幾乎相同，兩者差 −0.006 nm 的區間跨越零。")

heading("SEM 方向候選與尺寸偏差",2)
figure(FIG/"sem_comparison.png","圖 7  SEM 方向候選二的平均直徑有號差與去尺寸形狀差；圖中的兩種量不能合成單一排名。")
sem_rows=[]
for m in names:
    one=next(x for x in DATA["sem"] if x["model"]==m and x["mapping_candidate"]==1)
    two=next(x for x in DATA["sem"] if x["model"]==m and x["mapping_candidate"]==2)
    sem_rows.append((short[m],fmt(one["mean_diameter_difference_nm"]),fmt(one["full_radial_difference_nm"]),
                     fmt(two["full_radial_difference_nm"]),fmt(one["de_size_shape_difference_nm"]),
                     fmt(two["de_size_shape_difference_nm"])))
table(["模型","直徑差 nm","全半徑 1","全半徑 2","去尺寸 1","去尺寸 2"],sem_rows,
      [2.3,2.65,2.65,2.65,2.95,2.95],8.0)
para("SEM 九顆的平均直徑為 2614.439 nm。R140v2 的平均直徑為 2842.878 nm，有號差 +228.440 nm，較 R137 的 +222.383 nm 略大。兩個方向候選的全半徑差亦沒有改善；去尺寸形狀差則隨方向候選變化。這些數值只能在固定舊 SEM 比法下比較，不能據此宣稱絕對邊界或尺寸已恢復真值。")

heading("PSF 是否真的參與",2)
figure(FIG/"psf_activity.png","圖 8  八顆候選在全部 930 張影像均維持平均權重；四張固定實拍開關 PSF 的輸出差雖可測，幅度極小。")
para("最後模型確有一個微小的 PSF 輸出效應：四張固定實拍的開關輸出均方根差為 0.0044 至 0.0054 個八位元灰階，最大單點差約 0.106 至 0.134 灰階。但全部 930 張的八顆候選權重逐張都是 0.125，沒有可見的逐影像選擇。32 組生成測試上，PSF 開與關的平均損失分別是 0.900999 與 0.900952，開啟沒有顯示可驗證的損失收益。這不是重新訓練的消融實驗，僅能說模組在前向輸出上有微小作用。")
table(["光學檢查","結果"],[
    ("八候選編號","0、3、7、10、14、17、21、24；930 張各候選權重均為 0.125。"),
    ("四張固定實拍","PSF 開關輸出 RMS 差 0.0044、0.0044、0.0054、0.0051 個八位元灰階。"),
    ("32 組生成資料","開啟平均損失 0.900999；關閉 0.900952；差 +0.000047。"),
    ("推論界線","證實微小輸出參與，不支持『自動挑較好 PSF』或『PSF 提升最終量測』。"),
],[4.0,12.2],8.8)

heading("運行時間與完整性",2)
para(f"930 張實拍的完整尺寸 GPU 重建，單張純推論中位 {DATA['runtime']['total_reconstruction_seconds']['median']:.3f} 秒，含讀寫中位 {DATA['runtime']['end_to_end_load_infer_save_seconds']['median']:.3f} 秒；後續輪廓讀取中位 {DATA['measurement_runtime']['median_seconds']:.3f} 秒。這些是本輪實測時間，沒有同時控制其他 GPU 程序的跨模型速度實驗，故不作固有效率排名。")
para("最終模型的 8192 次資料母場景、順序與噪聲記錄逐筆與 R139 核對一致；930 張輸出逐檔核對雜湊，8370 個新輪廓另做數值重擬合，七方法 12600 筆位移獨立重新計算，與保存數值最大差約 3.56 × 10⁻¹² nm。核對支持這批結果確實重現，不代表實驗設計的所有科學限制已解決。")

heading("六 結果判讀與論文使用界線",1)
para("若主要任務是固定拍攝條件下的相鄰短步位移，R140v2 在這一批重拍資料中比 R129、R137、R138 的平均誤差及重拍波動更低，是可繼續評估的研究候選。它相對 R137 的直徑與輪廓重拍略差，SEM 平均尺寸偏差也較大；起點位移幾乎無差。不能把它稱作全指標最強，也不能因模組名稱而宣稱學會了物理 PSF 選擇。")
para("目前缺口包括：新樣本與不同拍攝場景、與光學校準真正獨立的同批 SEM 真值、個別柱體方向對應、完整量測不確定度、以及對穩定措施與 PSF 的乾淨重訓消融。現有 30 趟可支持本批重拍一致性與 PI 參考比較，不支持絕對 SI 尺寸或廣泛泛化結論。正式基準和論文模型並未因本次流程完成而自動替換。")

heading("附錄一 三十趟逐趟位移",1)
para("下表每格是該趟 30 個位置的九顆整體位移平均絕對差，單位 nm。前六趟參與 Raw 方向與尺度設定；第七至三十趟為校準後資料。")
scan_rows=[(x["scan"],x["run"],fmt(x["new_adj"]),fmt(x["r137_adj"]),fmt(x["new_start"]),fmt(x["r137_start"])) for x in DATA["scan_rows"]]
table(["趟","資料代號","R140v2 相鄰","R137 相鄰","R140v2 起點","R137 起點"],scan_rows,
      [1.1,4.1,2.75,2.6,2.9,2.75],7.8)

heading("附錄二 三十位置重拍與偏差",1)
para("位置 j 的相鄰重拍欄為 30 趟九顆整體位移誤差的標準差。平均有號誤差可為正或負，不能當平均絕對差。")
position_rows=[(x["position"],fmt(x["pi_um"],2),fmt(x["new_adj_bias"]),fmt(x["new_adj_repeat"]),
                fmt(x["r137_adj_repeat"]),fmt(x["new_start_bias"])) for x in DATA["position_rows"]]
table(["位置","PI 相鄰 μm","R140v2 相鄰偏差","R140v2 相鄰重拍","R137 相鄰重拍","R140v2 起點偏差"],
      position_rows,[1.15,2.5,3.2,3.25,3.0,3.1],7.8)

heading("附錄三 數值來源與重現核對",1)
para("本文件所有表格與圖表均由 R140v2 獨立實驗包的固定輸出重新讀取，未改動訓練模型或舊實驗。若要重查細目，以該包為根目錄，訓練紀錄在 runs/formal_8192，實拍重建與輪廓在 results，七方法統計在 results/comparison；原始自動結果在 reports/R140v2_RESULTS.md。")
table(["證據","可核對內容"],[
    ("固定訓練模型","第 8192 次保存點；SHA256："+DATA["training"]["final_checkpoint"]["sha256"]),
    ("訓練與來源","8192 次更新；512／128／128 分割；來源鎖定與歷史來源雜湊通過。"),
    ("完整實拍","930 張重建檔雜湊通過；30 趟 × 31 張；8370 個新模型圓重擬合。"),
    ("統計核對","七方法 12600 筆位移重算；SEM 兩個既定方向候選保留。"),
    ("模型地位","完整研究結果，不自動升級正式基準或替換論文模型。"),
],[4.0,12.2],8.5)

doc.core_properties.title="R140v2 模型設計 訓練與真實影像驗證總報告"
doc.core_properties.subject="顯微影像重建與柱體位移量測"
doc.core_properties.keywords="R140v2 MOSR U Net LSC ACM PSF"
doc.save(OUT)
print("DOCX_READY",OUT)
