from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from build_thesis_v11 import document_math_fingerprint
from build_thesis_v12 import formula_table_audit, sha256
PROJECT_ROOT = Path(__file__).resolve().parents[3]


DOCX = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v12.docx")
BASE_AUDIT = Path(str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/_paper_editing/qa/v12_build_audit.json'))
OUT = Path(str(PROJECT_ROOT / 'Thesis_MOSR_Evidence/_paper_editing/qa/v12_after_word_audit.json'))


def main() -> None:
    expected = json.loads(BASE_AUDIT.read_text(encoding="utf-8"))
    doc = Document(DOCX)
    math = document_math_fingerprint(doc)
    formula_tables = formula_table_audit(doc)
    text = "\n".join(p.text for p in doc.paragraphs)
    required = [
        "5.4.7 現代重建模型與跨形貌評估配置",
        "6.2.7 跨形貌、現代模型與計算效能補充驗證",
        "6.2.7.1 20× via之跨形貌與現代模型比較",
        "6.2.7.2 線條形貌之邊界集中與週期保持",
        "6.2.7.3 位移結果之影格選取敏感度與統計檢查",
        "6.2.7.4 整張影像運算效能",
        "6.2.7.5 補充驗證小結",
        "6.2.8 綜合討論",
    ]
    result = {
        "path": str(DOCX),
        "sha256": sha256(DOCX),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "math": math,
        "expected_math": expected["math_after"],
        "math_exact_match": math == expected["math_after"],
        "formula_tables": formula_tables,
        "required_heading_counts": {s: text.count(s) for s in required},
        "internal_code_hits": sum(text.count(s) for s in ("V42", "v42", "R61", "20260719")),
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))

    assert math["count"] == expected["math_after"]["count"]
    assert math["text_sha256"] == expected["math_after"]["text_sha256"]
    assert not formula_tables["visible_border_table_indices"]
    assert all(v == 1 for v in result["required_heading_counts"].values())


if __name__ == "__main__":
    main()
