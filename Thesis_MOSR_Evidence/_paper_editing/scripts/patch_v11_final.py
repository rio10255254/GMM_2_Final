from copy import deepcopy
from pathlib import Path
import hashlib

from docx import Document
from docx.oxml.ns import qn


DOCX = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v11.docx")
TEXT = (
    "後續工作可先擴充不同批次樣品、pillar尺寸、via、線條與方形結構，建立跨樣品測試；"
    "再於不同視野位置、焦平面、曝光與照明條件下增加PSF校正密度，以檢查固定光學先驗在更大範圍內的適用性"
    "[50,54,55,56]。若更換相機、物鏡或照明模組，應重新建立物體空間尺度與PSF資料，再判斷模型需完整重訓或少量校正"
    "[75,77,90]。"
)


def formula_text_hash(doc: Document) -> tuple[int, str]:
    values = []
    for node in doc._element.xpath(".//m:oMath"):
        values.append("".join((el.text or "") for el in node.iter(qn("m:t"))))
    digest = hashlib.sha256("\u241e".join(values).encode("utf-8")).hexdigest()
    return len(values), digest


def replace_text(paragraph, text: str) -> None:
    rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            rpr = deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def main() -> None:
    doc = Document(DOCX)
    if len(doc.paragraphs) != 1296 or "後續工作可先擴充" not in doc.paragraphs[1195].text:
        raise RuntimeError("Unexpected v11 document structure")
    before = formula_text_hash(doc)
    replace_text(doc.paragraphs[1195], TEXT)
    doc.save(DOCX)
    check = Document(DOCX)
    after = formula_text_hash(check)
    if before != after:
        raise RuntimeError(f"Formula text changed: before={before}, after={after}")
    if check.paragraphs[1195].text != TEXT:
        raise RuntimeError("Final citation distribution was not saved")
    print({"formula_fingerprint": after, "paragraph_1195": check.paragraphs[1195].text})


if __name__ == "__main__":
    main()
