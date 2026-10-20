from __future__ import annotations

import json
import os
import shutil
from copy import deepcopy
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_thesis_v11 import clear_and_set, document_math_fingerprint, enable_field_updates


DOCUMENT = Path(
    r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v11.docx"
)
RESULTS = Path(
    r"D:\GMM_2_Final\Thesis_MOSR_Evidence\LabSystem_MOSR\Experiments"
    r"\BaseUNet_Baseline_20260718\results_full_data_fixed_psf_v2"
)
SUMMARY = RESULTS / "analysis" / "learning_baseline_method_summary.csv"
TRAINING = RESULTS / "training_summary.json"
FIGURE = RESULTS / "figures" / "base_unet_50x_fixed_scale_comparison.png"
QA = Path(r"D:\GMM_2_Final\Thesis_MOSR_Evidence\_paper_editing\qa")
BACKUP = QA / "v11_before_base_unet_integration.docx"
AUDIT = QA / "v11_base_unet_integration_audit.json"


def find_exact(doc: Document, text: str):
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {text!r}, found {len(matches)}")
    return matches[0]


def find_startswith(doc: Document, prefix: str):
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 55, start: int = 55, bottom: int = 55, end: int = 55) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float = 8.5, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def populate_cell(cell, text: str, *, bold: bool = False, size: float = 8.5) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    parts = str(text).split("\n")
    for index, part in enumerate(parts):
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=bold)
        if index < len(parts) - 1:
            run.add_break()
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_borders(table) -> None:
    """Use restrained horizontal rules and no vertical rules."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    definitions = {
        "top": ("single", "8", "000000"),
        "bottom": ("single", "8", "000000"),
        "insideH": ("single", "4", "BFBFBF"),
        "left": ("nil", None, None),
        "right": ("nil", None, None),
        "insideV": ("nil", None, None),
    }
    for name, (value, size, color) in definitions.items():
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), value)
        if size is not None:
            edge.set(qn("w:sz"), size)
            edge.set(qn("w:space"), "0")
            edge.set(qn("w:color"), color)


def insert_picture_before(anchor, path: Path, width_inches: float = 6.25):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run().add_picture(str(path), width=Inches(width_inches))
    return paragraph


def add_shared_ruler_table(doc: Document, anchor, rows: list[dict]) -> None:
    caption = anchor.insert_paragraph_before(
        "表 6-15 Base U-Net comparison under the shared metrology ruler",
        style="Table Caption",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    headers = [
        "Method",
        "50× CD median\nabsolute error (nm)",
        "20× CD median\nabsolute error (nm)",
        "PI-feedback displacement\nMAE (nm)",
        "Nine-pillar\ninconsistency (nm)",
    ]
    keys = [
        "method",
        "cd_50x_median_abs_error_nm",
        "cd_20x_median_abs_error_nm",
        "displacement_reference_mae_nm",
        "nine_pillar_inconsistency_mean_nm",
    ]
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    widths = [0.95, 1.28, 1.28, 1.48, 1.28]
    for column, width in zip(table.columns, widths):
        for cell in column.cells:
            cell.width = Inches(width)
    for column_index, header in enumerate(headers):
        populate_cell(table.cell(0, column_index), header, bold=True, size=8.0)
        set_cell_shading(table.cell(0, column_index), "D9D9D9")

    minima = {
        key: min(float(row[key]) for row in rows)
        for key in keys[1:]
    }
    for row_index, row in enumerate(rows, start=1):
        populate_cell(table.cell(row_index, 0), row["method"], bold=False, size=8.5)
        for column_index, key in enumerate(keys[1:], start=1):
            value = float(row[key])
            populate_cell(
                table.cell(row_index, column_index),
                f"{value:.1f}",
                bold=abs(value - minima[key]) < 1.0e-9,
                size=8.5,
            )
    set_table_borders(table)
    anchor._p.addprevious(table._tbl)


def update_existing_text(doc: Document, base: dict, mosr: dict, training: dict) -> None:
    english_metrics = find_startswith(doc, "Experiments were conducted on an HS-100MG sample")
    clear_and_set(
        english_metrics,
        english_metrics.text.replace("feedback reference truth", "PI feedback truth"),
    )
    method_flow = find_startswith(doc, "研究流程依序包含光學校正、監督資料建立、模型訓練")
    clear_and_set(
        method_flow,
        method_flow.text.replace("PI閉迴路位移參考", "PI閉迴路位移真值"),
    )
    chapter4_summary = find_startswith(doc, "Raw、Wiener、MOSR與各消融條件使用相同座標")
    clear_and_set(
        chapter4_summary,
        chapter4_summary.text.replace("PI閉迴路位移參考比較", "PI閉迴路位移真值比較"),
    )
    ch6_scope = find_startswith(doc, "本節在同一顯微系統與HS-100MG樣品下")
    clear_and_set(
        ch6_scope,
        ch6_scope.text.replace("作為參考真值", "作為真值"),
    )
    metric_definition = find_startswith(doc, "『PI回授真值平均絕對偏差』是每一步影像量測幅度")
    clear_and_set(
        metric_definition,
        metric_definition.text.replace("本實驗位移參考真值", "本實驗位移真值"),
    )
    chapter6_limit = find_startswith(doc, "本節結果建立於目前顯微系統、HS-100MG樣品")
    clear_and_set(
        chapter6_limit,
        chapter6_limit.text.replace("位移參考真值", "位移真值"),
    )
    chapter7_scope = find_startswith(doc, "本研究的實證範圍為同一顯微系統")
    clear_and_set(
        chapter7_scope,
        chapter7_scope.text.replace("位移參考真值", "位移真值"),
    )

    section = find_exact(doc, "6.2.6 模組與損失函數消融實驗")
    clear_and_set(section, "6.2.6 基礎模型對照、模組與損失函數消融實驗")

    intro = find_startswith(doc, "前述Raw、Wiener與MOSR的比較可判斷完整重建方法是否有助於量測；")
    clear_and_set(
        intro,
        "前述Raw、Wiener與MOSR的比較可判斷完整重建方法是否有助於量測；本節另以參數量相近的基礎U-Net檢查一般編碼器—解碼器是否已足以維持尺寸與位移，再固定PSF、資料、訓練預算與評估規則，分別檢查抗混疊上採樣、輸入參考式殘差輸出、量測導向訓練目標及其他輸出模組。",
    )

    old_summary_heading = find_exact(doc, "6.2.6.6 消融結果與主要方法貢獻")
    clear_and_set(old_summary_heading, "6.2.6.7 消融與基礎模型對照之綜合判讀")
    summary_text = find_startswith(doc, "消融結果把兩項主要創新與量測現象直接連結。")
    clear_and_set(
        summary_text,
        "消融結果把兩項主要創新與量測現象直接連結。抗混疊上採樣降低不同pillar之間的位移分歧，主要改善跨空間位置的幾何一致性；輸入參考式殘差輸出維持Raw的低頻位置與亮暗拓樸，主要抑制輪廓漂移與振鈴，並在20倍CD上呈現明顯效果。基礎U-Net雖可降低位移波動，卻未同時維持兩倍率的CD，說明一般重建能力不能取代量測導向限制。",
    )

    discussion = find_startswith(doc, "綜合三方法比較與消融實驗，可回答本研究的三項問題。")
    clear_and_set(
        discussion,
        "綜合三方法、基礎U-Net與消融實驗，可回答本研究的三項問題。第一，單張同尺寸重建能使邊界集中，但最窄邊緣不必然帶來最穩定CD；MOSR是比較方法中唯一在50倍與20倍都相對Raw降低CD誤差者。基礎U-Net在20倍得到最低CD誤差，50倍卻明顯惡化，顯示一般重建的改善具有倍率依賴性。第二，MOSR在120個配對步中同時降低PI回授真值平均絕對偏差與九顆pillar不一致量；基礎U-Net的兩項位移彙整值更低，說明一般U-Net亦能穩定連續影格，但不能據此推論跨倍率尺寸同樣可靠。第三，抗混疊上採樣使九顆pillar不一致量較一般上採樣降低14.31%；輸入參考式殘差輸出則抑制兩倍率振鈴，並使20倍CD誤差維持在較低範圍。",
    )

    abstract_zh = find_startswith(doc, "結果顯示，Wiener能產生較窄的部分邊緣")
    clear_and_set(
        abstract_zh,
        abstract_zh.text
        + (
            f" 參數量相近且使用相同固定PSF先驗的基礎U-Net在20倍CD與位移彙整值較低，"
            f"但50倍CD中位絕對誤差增至{base['cd_50x_median_abs_error_nm']:.1f} nm；"
            "MOSR則在兩倍率都相對Raw降低CD誤差，顯示一般視覺重建的改善具有條件依賴性。"
        ),
    )
    abstract_en = find_startswith(doc, "Wiener produced narrower edges in some conditions")
    clear_and_set(
        abstract_en,
        abstract_en.text
        + (
            " A parameter-matched base U-Net using the same fixed PSF prior achieved lower 20x CD error and displacement aggregates, "
            f"but its 50x median absolute CD error increased to {base['cd_50x_median_abs_error_nm']:.1f} nm. "
            "MOSR reduced CD error relative to Raw at both magnifications, showing that the benefits of generic reconstruction were condition-dependent."
        ),
    )

    contribution = find_startswith(doc, "貢獻三：建立由量測失效形式導出的共同訓練目標與配對驗證。")
    clear_and_set(
        contribution,
        contribution.text
        + " 另以參數量相近、保留相同固定PSF先驗的基礎U-Net作為學習式對照，用來區分一般影像重建能力與量測導向限制的差異。",
    )

    ch7_conclusion = find_startswith(doc, "第二與第三個研究問題關注平移一致性及邊界保護。")
    clear_and_set(
        ch7_conclusion,
        ch7_conclusion.text
        + " 基礎U-Net的結果則顯示，20倍CD與位移可以同時改善，50倍CD卻反向惡化；因此模型必須同時接受跨倍率尺寸與位移檢查。",
    )
    ch7_contribution = find_startswith(doc, "影像內容、邊界幾何、運動一致性與輸出保守性四組目標")
    clear_and_set(
        ch7_contribution,
        ch7_contribution.text
        + " 參數量相近的基礎U-Net對照進一步顯示，一般U-Net可在部分條件取得更低誤差，但其效果不會自然延伸至另一倍率；量測導向設計的重點是降低這種條件依賴性。",
    )
    ch7_quant = find_startswith(doc, "在四組位移序列共120個配對步中，Raw、Wiener與MOSR")
    clear_and_set(
        ch7_quant,
        ch7_quant.text
        + (
            f" 基礎U-Net的PI回授真值平均絕對偏差與九顆pillar不一致量為"
            f"{base['displacement_reference_mae_nm']:.2f}與"
            f"{base['nine_pillar_inconsistency_mean_nm']:.2f} nm，但50倍與20倍CD中位絕對誤差為"
            f"{base['cd_50x_median_abs_error_nm']:.1f}與{base['cd_20x_median_abs_error_nm']:.1f} nm。"
        ),
    )
    future = find_startswith(doc, "量測追溯方面，可使用雷射干涉儀與PI閉迴路電容式感測回授同步量測")
    clear_and_set(
        future,
        "量測追溯方面，可使用雷射干涉儀與PI閉迴路電容式感測回授同步量測，進一步交叉驗證位移真值並建立包含重複性、像素尺度、SEM參考、平台回授及模型變動的完整不確定度預算[84,85]。模型驗證亦可增加訓練重複、跨樣品測試及跨條件分層統計，使量測導向設計的適用範圍更加明確。",
    )
    ch7_summary = find_startswith(doc, "本研究證明，顯微影像重建的有效性不能只由清晰度判斷。")
    clear_and_set(
        ch7_summary,
        ch7_summary.text.replace(
            "後續若完成",
            "基礎U-Net對照亦顯示，一般學習式重建不能保證跨倍率CD都維持相同改善方向。後續若完成",
        ),
    )


def insert_chapter5_baseline(doc: Document, training: dict) -> None:
    anchor = find_exact(doc, "5.5 小結")
    heading = anchor.insert_paragraph_before("5.4.6 基礎 U-Net 對照配置", style="Heading 3")
    heading.paragraph_format.keep_with_next = True
    anchor.insert_paragraph_before(
        "為判斷量測改善是否只是一般U-Net重建能力所造成，本研究另建立參數量相近的三層編碼器—解碼器對照。此模型使用一般雙線性上採樣與跳接，不含抗混疊上採樣、輸入參考式殘差輸出、邊界殘差、品質門控及量測導向損失；固定PSF光學先驗、訓練資料分組與內部驗證名單則與完整模型保持一致，因此比較不以移除PSF作為變因。",
        style="Normal",
    )
    anchor.insert_paragraph_before(
        (
            f"基礎U-Net共有{int(training['parameter_count']):,}個參數，與完整模型的4,021,214個參數相近。"
            f"訓練採L1與SSIM的普通影像重建損失，共訓練{int(training['config']['epochs'])}回合；"
            f"每回合遍歷{int(training['train_count']):,}組訓練資料，並由{int(training['validation_count']):,}組固定內部驗證資料選出第{int(training['best_epoch'])}回合模型。"
            "此對照的目的不是取代模組消融，而是檢查一般學習式銳化是否能在相同量測規則下同時維持CD與微位移。"
        ),
        style="Normal",
    )


def insert_chapter6_baseline(doc: Document, rows: list[dict], base: dict, training: dict) -> None:
    mosr = next(row for row in rows if row["method"] == "MOSR")
    anchor = find_exact(doc, "6.2.6.7 消融與基礎模型對照之綜合判讀")
    heading = anchor.insert_paragraph_before(
        "6.2.6.6 基礎 U-Net 對照：視覺銳化與量測可用性的差異",
        style="Heading 4",
    )
    heading.paragraph_format.keep_with_next = True
    anchor.insert_paragraph_before(
        (
            "為排除改善僅來自一般編碼器—解碼器的可能，本研究使用相同固定PSF先驗與資料分組，另訓練參數量相近的基礎U-Net。"
            "該模型採一般雙線性上採樣與普通影像重建損失，不含抗混疊上採樣、輸入參考式殘差輸出及量測導向限制。"
            "所有結果仍使用與Raw、Wiener及MOSR相同的九顆pillar、物體空間尺度、35%至65%門檻、兩種CD剖面及固定30步位移母體。"
        ),
        style="Normal",
    )
    anchor.insert_paragraph_before(
        (
            f"基礎U-Net完成{int(training['config']['epochs'])}回合訓練，並由固定內部驗證集選出第{int(training['best_epoch'])}回合。"
            "圖6-53使用相同九宮格裁切範圍與0至255灰階顯示，不對各方法逐張拉伸。"
            "基礎U-Net的整體灰階較低，pillar外圍形成狹窄環狀邊界，內部對比也與Raw不同。"
            "由於三者未分別拉伸，這些差異反映模型實際輸出，而不是顯示強化；50倍CD量測會受到外圍環狀邊界影響。"
        ),
        style="Normal",
    )
    insert_picture_before(anchor, FIGURE)
    caption = anchor.insert_paragraph_before(
        "圖 6-53 Fixed-scale visual comparison of Raw, Base U-Net, and MOSR",
        style="Figure Caption",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_shared_ruler_table(doc, anchor, rows)
    anchor.insert_paragraph_before(
        (
            f"基礎U-Net的PI回授真值平均絕對偏差為{base['displacement_reference_mae_nm']:.2f} nm，"
            f"九顆pillar不一致量為{base['nine_pillar_inconsistency_mean_nm']:.2f} nm，"
            f"20倍CD中位絕對誤差為{base['cd_20x_median_abs_error_nm']:.1f} nm；三項均低於MOSR，"
            "顯示一般學習式重建在部分條件下可以得到很好的數值。"
            f"然而，基礎U-Net的50倍CD中位絕對誤差為{base['cd_50x_median_abs_error_nm']:.1f} nm，"
            f"明顯高於Raw的95.3 nm與MOSR的{mosr['cd_50x_median_abs_error_nm']:.1f} nm。"
            "因此，本研究不把基礎U-Net描述為全面較差，而將結果解讀為其改善具有倍率依賴性；MOSR的優勢是兩倍率CD都相對Raw改善，並同時維持低位移誤差。"
        ),
        style="Normal",
    )


def formula_table_audit(doc: Document) -> dict:
    count = 0
    visible = []
    for index, table in enumerate(doc.tables):
        if not table._tbl.xpath(".//m:oMath"):
            continue
        count += 1
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is not None:
            values = [edge.get(qn("w:val"), "") for edge in list(borders)]
            if any(value not in {"nil", "none"} for value in values):
                visible.append(index)
    return {"count": count, "visible_border_table_indices": visible}


def main() -> None:
    for path in (DOCUMENT, SUMMARY, TRAINING, FIGURE):
        if not path.exists():
            raise FileNotFoundError(path)
    QA.mkdir(parents=True, exist_ok=True)
    frame = pd.read_csv(SUMMARY)
    expected_methods = ["Raw", "Wiener", "Base U-Net", "MOSR"]
    if frame["method"].tolist() != expected_methods:
        raise RuntimeError(f"Unexpected method order: {frame['method'].tolist()}")
    rows = frame.to_dict(orient="records")
    base = next(row for row in rows if row["method"] == "Base U-Net")
    training = json.loads(TRAINING.read_text(encoding="utf-8"))

    doc = Document(DOCUMENT)
    if any("6.2.6.6 基礎 U-Net 對照" in paragraph.text for paragraph in doc.paragraphs):
        raise RuntimeError("Base U-Net section is already present; refusing to duplicate it")
    before_math = document_math_fingerprint(doc)
    before_formula_tables = formula_table_audit(doc)
    if before_formula_tables["visible_border_table_indices"]:
        raise RuntimeError("Formula borders were visible before integration")

    shutil.copy2(DOCUMENT, BACKUP)
    mosr = next(row for row in rows if row["method"] == "MOSR")
    update_existing_text(doc, base, mosr, training)
    insert_chapter5_baseline(doc, training)
    insert_chapter6_baseline(doc, rows, base, training)
    enable_field_updates(doc)

    temporary = DOCUMENT.with_name(DOCUMENT.stem + "_integrating.docx")
    doc.save(temporary)
    verification = Document(temporary)
    after_math = document_math_fingerprint(verification)
    after_formula_tables = formula_table_audit(verification)
    if before_math != after_math:
        raise RuntimeError(f"Equation fingerprint changed: {before_math} -> {after_math}")
    if after_formula_tables["visible_border_table_indices"]:
        raise RuntimeError("At least one formula table regained a visible border")
    if not any(
        paragraph.text.strip().startswith("6.2.6.6 基礎 U-Net 對照")
        for paragraph in verification.paragraphs
    ):
        raise RuntimeError("Inserted Base U-Net section not found after save")
    os.replace(temporary, DOCUMENT)

    audit = {
        "document": str(DOCUMENT),
        "backup": str(BACKUP),
        "result_summary": str(SUMMARY),
        "figure": str(FIGURE),
        "training": {
            "parameter_count": int(training["parameter_count"]),
            "best_epoch": int(training["best_epoch"]),
            "best_validation": training["best_validation"],
        },
        "base_unet_metrics": base,
        "math_before": before_math,
        "math_after": after_math,
        "formula_tables_before": before_formula_tables,
        "formula_tables_after": after_formula_tables,
        "paragraph_count_after": len(verification.paragraphs),
        "table_count_after": len(verification.tables),
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
