from __future__ import annotations

import hashlib
import json
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from lxml import etree


SOURCE = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v10.docx")
TARGET = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v11.docx")
AUDIT = Path(r"D:\GMM_2_Final\Thesis_MOSR_Evidence\_paper_editing\qa\v11_build_audit.json")

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS = {"m": M_NS, "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def paragraph_signature(doc: Document) -> None:
    expected = {
        7: "點擴散函數校正與物理約束式神經網路之",
        18: "摘要",
        127: "緒論",
        517: "研究方法",
        938: "實驗結果與分析",
        1180: "未來展望",
        1200: "參考文獻",
    }
    if len(doc.paragraphs) != 1296:
        raise RuntimeError(f"Unexpected paragraph count: {len(doc.paragraphs)}")
    for idx, anchor in expected.items():
        if anchor not in doc.paragraphs[idx].text:
            raise RuntimeError(f"Paragraph {idx} no longer matches expected anchor: {doc.paragraphs[idx].text!r}")


def clear_and_set(paragraph, text: str) -> None:
    """Replace paragraph text while preserving paragraph style and first-run formatting."""
    run_props = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            run_props = deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run()
    if run_props is not None:
        run._r.insert(0, run_props)
    parts = text.split("\n")
    for i, part in enumerate(parts):
        if i:
            run.add_break()
        run.add_text(part)


def iter_unique_table_paragraphs(doc: Document):
    seen_cells: set[int] = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                key = id(cell._tc)
                if key in seen_cells:
                    continue
                seen_cells.add(key)
                yield from cell.paragraphs


def replace_phrase_everywhere(doc: Document, old: str, new: str) -> int:
    changed = 0
    paras = list(doc.paragraphs) + list(iter_unique_table_paragraphs(doc))
    for p in paras:
        if old in p.text:
            clear_and_set(p, p.text.replace(old, new))
            changed += 1
    return changed


def set_table_no_borders(table) -> None:
    try:
        table.style = "Normal Table"
    except KeyError:
        style = table._tbl.tblPr.find(qn("w:tblStyle"))
        if style is not None:
            table._tbl.tblPr.remove(style)
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")
        for attr in ("sz", "space", "color"):
            key = qn(f"w:{attr}")
            if key in edge.attrib:
                del edge.attrib[key]

    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side_name in ("top", "bottom"):
        side = margins.find(qn(f"w:{side_name}"))
        if side is None:
            side = OxmlElement(f"w:{side_name}")
            margins.append(side)
        side.set(qn("w:w"), "0")
        side.set(qn("w:type"), "dxa")


def compact_formula_tables(doc: Document) -> int:
    count = 0
    for table in doc.tables:
        if not table._tbl.xpath(".//m:oMath"):
            continue
        count += 1
        set_table_no_borders(table)
        for p_el in table._tbl.iter(qn("w:p")):
            p = Paragraph(p_el, table)
            fmt = p.paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
            fmt.line_spacing = 1.0
    return count


def document_math_fingerprint(doc: Document) -> dict:
    root = doc._element
    nodes = root.xpath(".//m:oMath")
    text_values = []
    structures = []
    for node in nodes:
        text_values.append("".join((el.text or "") for el in node.iter(qn("m:t"))))
        skeleton = "|".join(etree.QName(el).localname for el in node.iter())
        structures.append(skeleton)
    payload = "\u241e".join(text_values).encode("utf-8")
    structure_payload = "\u241e".join(structures).encode("utf-8")
    return {
        "count": len(nodes),
        "text_sha256": hashlib.sha256(payload).hexdigest(),
        "structure_sha256": hashlib.sha256(structure_payload).hexdigest(),
    }


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, TARGET)
    doc = Document(TARGET)
    paragraph_signature(doc)
    before_math = document_math_fingerprint(doc)
    doc.core_properties.title = "結合點擴散函數先驗之量測導向顯微影像重建：關鍵尺寸與微位移量測"
    doc.core_properties.subject = "量測導向顯微影像重建、關鍵尺寸與微位移量測"

    replacements = {
        7: "結合點擴散函數先驗之量測導向顯微影像重建：\n關鍵尺寸與微位移量測",
        8: "Measurement-Oriented Microscopy Reconstruction with PSF Priors\nfor Critical-Dimension and Micro-Displacement Metrology",
        19: "隨著半導體製程與精密製造的特徵尺寸持續縮小，顯微成像除了必須提供足夠的空間解析能力，也必須維持關鍵尺寸與微位移量測的穩定性。影像變得更銳利並不必然代表幾何量測更正確；過度銳化、反卷積振鈴與重建造成的輪廓漂移，都可能改變表觀邊界位置，使視覺品質與量測品質產生落差。",
        20: "本研究提出量測導向超解析（Measurement-Oriented Super-Resolution, MOSR），將單張顯微影像重建視為量測前處理，並依可觀察的量測失效設計模型。針對解碼上採樣對取樣相位的敏感性，模型加入抗混疊上採樣，以降低相同結構位於不同像素相位或視野位置時的幾何差異；針對高頻增強可能改變低頻輪廓與中心位置的問題，輸出端採用輸入參考式殘差輸出，僅將幅度受限的高頻修正加入原始影像。影像內容、邊界幾何、運動一致性與輸出保守性四組目標共同限制可接受的重建結果。局部與多位置點擴散函數則分別提供Wiener反卷積基準與固定光學先驗。模型使用時只需輸入單張影像，且輸入與輸出尺寸相同。",
        21: "實驗以20倍與50倍物鏡拍攝HS-100MG樣品，並在相同九顆pillar、相同物體空間尺度及相同後端量測規則下比較Raw、Wiener與MOSR。50倍影像的徑向10–90%表觀邊緣過渡寬度中位數由620.1 nm降至Wiener的467.7 nm與MOSR的508.5 nm；20倍則由802.2 nm降至493.1 nm與246.1 nm。跨越35%至65%門檻及兩種剖面定義後，三種方法的CD中位絕對誤差分別為95.3、123.7與76.1 nm。在由PI閉迴路電容式感測回授確認的0.5 µm與1.0 µm步階條件下，合併四組序列共120個配對步後，Raw、Wiener與MOSR相對回授參考真值的平均絕對偏差分別為9.54、7.46與6.74 nm，九顆pillar內部不一致量則為17.63、7.92與6.47 nm。",
        22: "結果顯示，Wiener能產生較窄的部分邊緣，卻具有最大的跨門檻CD誤差；MOSR並未以所有條件下的最大銳利度為目標，而是在CD門檻穩健性與多顆pillar位移一致性之間取得較佳工作點。相較Raw，MOSR使CD中位絕對誤差降低20.16%，並使九顆pillar不一致量降低63.29%；相較Wiener，兩項指標分別降低38.48%與18.31%。消融實驗顯示，以一般上採樣取代抗混疊上採樣後，四組序列合併的九顆pillar不一致量增加14.31%；移除輸入參考式殘差輸出後，20倍CD中位絕對誤差增加73.97%，兩種倍率的振鈴亦明顯增加。上述結果說明，最小邊緣過渡寬度不能單獨代表量測品質；有效的顯微影像重建必須同時維持邊界、尺寸與平移關係。",
        23: "關鍵字：量測導向超解析、點擴散函數先驗、抗混疊上採樣、輸入參考式殘差輸出、關鍵尺寸、微位移量測。",
        25: "As feature sizes continue to shrink in semiconductor manufacturing and precision engineering, microscopy-based metrology must provide both adequate spatial resolution and stable critical-dimension and micro-displacement measurements. A sharper image is not necessarily a more accurate image: excessive sharpening, deconvolution ringing, and reconstruction-induced contour drift can shift apparent boundaries and create a discrepancy between visual quality and measurement quality.",
        26: "This thesis proposes Measurement-Oriented Super-Resolution (MOSR), which treats single-image microscopy reconstruction as a metrology preprocessing step and derives the network design from observable measurement failures. Anti-aliased upsampling is introduced to reduce geometry changes caused by sampling phase and spatial position during decoding. An input-referenced residual output then adds only amplitude-limited high-frequency corrections to the raw image, thereby preserving its low-frequency contour and center position. Four groups of objectives—image content, boundary geometry, motion consistency, and conservative output behavior—jointly constrain acceptable reconstructions. Local and multi-position point spread functions provide the Wiener baseline and a fixed optical prior, respectively. At inference, the model accepts one image and returns an image of the same size.",
        27: "Experiments were conducted on an HS-100MG sample using 20× and 50× objectives. Raw images, Wiener deconvolution, and MOSR were compared using the same nine pillars, object-space scale, and downstream measurement rules. For the 50× images, the median radial 10–90% apparent edge-transition width decreased from 620.1 nm in Raw to 467.7 nm with Wiener and 508.5 nm with MOSR; for the 20× images, it decreased from 802.2 nm to 493.1 nm and 246.1 nm, respectively. Across thresholds from 35% to 65% and two profile definitions, the median absolute CD errors were 95.3, 123.7, and 76.1 nm. Under 0.5 µm and 1.0 µm step conditions confirmed by PI closed-loop capacitive-sensor feedback, the mean absolute deviations from the feedback reference truth across 120 paired steps were 9.54, 7.46, and 6.74 nm, whereas the within-nine-pillar inconsistencies were 17.63, 7.92, and 6.47 nm.",
        28: "Wiener produced narrower edges in some conditions but also the largest cross-threshold CD error. MOSR did not target maximum sharpness under every condition; instead, it provided a better operating point between CD threshold robustness and multi-pillar displacement consistency. Relative to Raw, MOSR reduced the median absolute CD error by 20.16% and the nine-pillar inconsistency by 63.29%; relative to Wiener, the reductions were 38.48% and 18.31%. In ablation experiments, replacing anti-aliased upsampling with conventional upsampling increased the pooled nine-pillar disagreement by 14.31%. Removing the input-referenced residual output increased the 20× median absolute CD error by 73.97% and substantially increased ringing at both magnifications. These results show that minimum edge-transition width alone does not represent metrology quality; effective microscopy reconstruction must preserve boundaries, dimensions, and translation relationships together.",
        29: "Keywords: Measurement-Oriented Super-Resolution; Point Spread Function Prior; Anti-aliased Upsampling; Input-referenced Residual Output; Critical Dimension; Micro-displacement Measurement.",

        133: "顯微影像重建的核心困難，不只是恢復高頻資訊，而是增強過程可能改變量測所依賴的幾何關係。PSF描述光學系統如何擴散空間訊號[13,50]；反卷積與學習式重建則試圖由受模糊與雜訊影響的影像恢復結構[46,52]。若重建只追求像素誤差、感知品質或主觀清晰度，即使輸出外觀改善，仍可能產生邊界偏移、振鈴或影格依賴性[63,64]。",
        134: "本研究因此將影像重建定義為量測前處理，並把『邊界是否更容易定位、尺寸是否接近獨立參考、位移是否更穩定』列為主要判斷依據。此一量測導向超解析（Measurement-Oriented Super-Resolution, MOSR）不增加輸出像素數，也不以任意生成細節為目的，而是在與輸入相同的取樣網格上，恢復原始影像中已存在但不易定位的邊界資訊。",
        135: "方法設計由兩種可觀察的量測失效導出。第一，特徵圖的下採樣與上採樣會使網路對像素相位與空間位置敏感[39,58,59,61,91]，使同一實體平移在不同pillar上得到不同結果；因此解碼器採用抗混疊上採樣。第二，高頻增強可能同時改變低頻輪廓、主要亮暗關係與結構中心；因此輸出端採用輸入參考式殘差輸出，只允許幅度受限的高頻修正。影像、邊界、運動與輸出保守性四組目標再共同限制模型的可行解[62,67,68,69,70,92]。PSF在所有模型條件中作為固定光學先驗，不作為單獨變動的研究因素。",
        136: "依上述問題定義，本研究提出三個相互銜接的研究問題：",
        137: "第一，在不改變輸出取樣尺寸的條件下，單張影像重建能否使pillar邊界更集中，並降低CD對門檻與剖面定義的敏感性？",
        138: "第二，抗混疊上採樣能否降低解碼器對像素相位與視野位置的敏感性，使同一影格中多顆pillar對同一實體位移得到更一致的估計？",
        139: "第三，輸入參考式殘差輸出與量測導向訓練目標，能否在增加可用邊界資訊時抑制輪廓漂移與振鈴，並使CD與微位移量測優於Raw及Wiener基準？",
        141: "本研究的主要貢獻集中於『量測失效如何轉化為模型設計』，以及這些設計能否由配對量測證據直接檢查。PSF校正、多來源資料及共同後端量測提供研究基礎；學術主張則放在抗混疊上採樣、輸入參考式殘差輸出與量測導向訓練目標如何共同維持邊界、尺寸與平移關係。",
        142: "貢獻一：提出面向微位移量測的抗混疊解碼方式。抗混疊與平移穩定網路已有理論基礎[39,58,61]，上採樣的頻譜穩定性與影像重建平移等變性亦已有相關研究[60,91,92]。本研究進一步將這些原理落實於顯微重建解碼器，並以同一影格九顆pillar對同一實體位移的一致性直接驗證。配對消融顯示，抗混疊上採樣使四組真實位移序列合併後的九顆pillar不一致量較一般上採樣降低14.31%。",
        143: "貢獻二：提出輸入參考式殘差輸出。殘差學習與輸入跳接已有成熟研究基礎[42,43,44]；本研究將其重新界定為量測邊界的保護機制，以Raw提供低頻位置與大尺度亮暗基準，再加入幅度受限的高頻邊界修正。移除此機制後，20倍CD中位絕對誤差增加73.97%，50倍與20倍的外圍振鈴分別增加246.16%與392.37%，顯示輸入位置基準能有效限制輪廓與亮暗拓樸偏離原始量測。",
        144: "貢獻三：建立由量測失效形式導出的共同訓練目標與配對驗證。影像內容、邊界幾何、運動一致性與輸出保守性四組目標，分別限制外觀失真、輪廓偏移、影格依賴與不必要改動[62,67,68,69,70]。Raw、Wiener、MOSR及各消融條件共用影像名單、pillar身分、物體空間尺度、中心、門檻與統計規則，讓邊緣、CD與位移結果可在同一座標下比較[80,81,82,83,84,85]。",
        145: "本研究的創新不建立在單一既有元件的重新命名，而在於以可量測的失效機制連結模型設計、訓練限制與後端證據。抗混疊上採樣主要對應多位置的相對幾何一致性；輸入參考式殘差輸出主要對應輪廓與亮暗關係的保留；四組訓練目標則共同界定可接受的重建範圍。第六章以完整模型與配對消融分別檢查整體效果及兩項主要模組的作用。",
        146: "現有驗證涵蓋同一顯微系統、HS-100MG樣品、20倍與50倍影像、九顆SEM參考pillar及四組PI閉迴路位移序列。此範圍足以回答本研究的三項問題；跨樣品、跨系統及更完整的量測追溯，則列入第七章的未來工作。",
        149: "第一章為緒論，說明研究背景、量測導向問題定義、三項研究問題，以及抗混疊上採樣、輸入參考式殘差輸出與共同量測驗證等主要貢獻。",
        150: "第二章為文獻回顧，回顧PSF模型、校正方法、反卷積演算法、資料驅動超解析，以及量測導向影像重建與平移等變網路的相關研究。",
        151: "第三章為理論背景，涵蓋純量繞射理論、非同調成像PSF、像差模型、卷積神經網路與U-Net基礎，以及平移等變性與量測不確定度。",
        152: "第四章說明量測導向重建方法，包括成像與物體空間座標、固定PSF光學先驗、抗混疊上採樣、輸入參考式殘差輸出、量測導向訓練目標，以及Raw、Wiener與MOSR共用的尺寸與位移量測規則。",
        153: "第五章介紹顯微光學系統、PI閉迴路位移平台、樣品與SEM參考量測、軟硬體環境、模型架構、訓練資料分組、訓練設定及消融實驗配置。",
        154: "第六章呈現PSF校正與反卷積結果，比較Raw、Wiener與MOSR在邊緣、CD及位移量測上的差異，並以配對消融檢查抗混疊上採樣、輸入參考式殘差輸出與共同訓練目標的作用。",
        155: "第七章彙整研究結論與主要量化結果，說明研究適用範圍，並提出跨樣品泛化、量測追溯與系統部署等未來工作。",

        573: "本研究不以單一影像分數判定重建成功，而是要求結果通過三個相互關聯的量測檢查。首先，邊緣過渡應更集中，但不得以明顯振鈴形成新的交點；其次，CD應在不同門檻與剖面定義下維持穩定，並接近SEM參考；最後，同一實體平移在九顆pillar上應得到一致的位移。三者共同構成量測導向的判斷標準。",
        574: "方法設計由上述失效形式直接導出。像素相位與空間位置造成的多pillar差異，由抗混疊上採樣處理；高頻增強造成的低頻輪廓與中心漂移，由輸入參考式殘差輸出限制；影像、邊界、運動與輸出保守性四組訓練目標，則共同限制模型不得以假影或非必要改動換取表面上的清晰度。",
        576: "研究流程依序包含光學校正、監督資料建立、模型訓練、單張同尺寸重建與後端量測。各階段共用像素座標、物體空間尺度及樣品身分，使光學資訊、模型輸出、SEM尺寸與PI閉迴路位移參考可以相互核對。",
        577: "第一項主要設計為抗混疊上採樣。解碼器放大特徵圖時，插值與取樣相位可能使相同結構在不同位置產生不一致的邊界偏移。本研究先以雙線性插值恢復尺寸，再以低通濾波抑制混疊，最後利用卷積整合特徵，以提高不同pillar之間的空間一致性。",
        578: "第二項主要設計為輸入參考式殘差輸出。最終輸出以Raw維持低頻輪廓、主要亮暗與中心位置，只加入比例受限的高頻邊界修正，避免網路完整預測直接取代輸入影像並改變量測基準。",
        579: "第三項設計為量測導向訓練目標。影像內容、邊界幾何、運動一致性與輸出保守性四組目標，分別限制外觀失真、輪廓偏移、影格間位置不一致與不必要的影像改動。這四組目標共同界定模型可以接受的重建範圍。",
        580: "局部PSF用於特定九宮格的Wiener基準，多位置PSF則以固定形式提供網路光學先驗。PSF路徑、邊界殘差分支與品質門控在完整模型及各消融條件中維持不變，使主要比較集中於抗混疊上採樣與輸入參考式殘差輸出的作用。",
        581: "整體而言，本研究以固定光學先驗描述成像模糊，以兩項主要模組維持空間與邊界關係，再以四組量測導向目標限制訓練結果。第六章使用相同後端量測規則，依序檢查完整方法的效果及兩項主要模組的作用。",
        624: "局部PSF適合描述特定九宮格，但不足以代表整個視野。為提供模型一致的光學背景，本研究將多個視野位置的PSF對齊至共同中心並正規化為單位能量，再依倍率建立固定光學先驗。此先驗在完整模型、訓練重複與所有消融條件中保持相同，不作為單獨變動的研究因素。[50,54,55,56]",
        634: "此模組把局部校正、共同光學先驗與依輸入調整的融合權重分開處理，使PSF資訊以幅度受限的殘差方式加入網路，而不是以強反卷積直接取代Raw。就本研究的主張而言，它是固定的光學條件；主要待驗證的設計仍為抗混疊上採樣與輸入參考式殘差輸出。",
        636: "本研究的重建模型以量測可用性為目標。在Raw邊界受到模糊與雜訊影響時，模型需要提高灰階轉換的集中程度；但任何增強若同時使pillar輪廓向內、向外或同方向偏移，便會改變CD與中心位置。因此，模型設計必須同時處理邊界可辨識度與幾何保留。",
        637: "抗混疊上採樣與輸入參考式殘差輸出分別作用於兩個容易改變位置的階段。前者在解碼特徵放大時降低像素相位與空間位置造成的差異；後者在形成最終影像時保留Raw的低頻輪廓、亮暗拓樸與中心基準。固定PSF先驗、邊界殘差上限及品質門控則提供光學方向與修正幅度限制。",
        638: "模型採單張輸入、同尺寸輸出，所有處理沿用相同像素座標與物體空間尺度。輸出不是一幅與輸入無關的新影像，而是以Raw為基準的受限重建；這使後端可以對Raw、Wiener與MOSR使用相同的pillar中心、剖面、門檻與位移追蹤規則。",
        639: "抗混疊、殘差連接與平移等變已有相關研究[39,42,43,58,61,91,92]。本研究的重點是把這些原理依量測失效重新配置：抗混疊上採樣對應多位置的相對位移分歧，輸入參考式殘差輸出對應低頻輪廓與中心漂移，並以配對CD、振鈴與九顆pillar位移一致性檢查其作用。",
        716: "本研究的訓練目標不是單獨降低輸出與參考影像的像素差異，而是把可用重建限制在同時符合影像、邊界與運動關係的範圍。像素與感知品質並非完全一致[30,31]；頻域損失可補充不同空間頻率的學習[67,68]，梯度與任務特定損失可強化結構及下游需求[69,70]，等變約束則用於維持轉換前後的幾何關係[62,92]。",
        717: "單獨強調像素誤差容易產生過度平滑；單獨強調梯度或高頻則可能產生過度銳化、振鈴與人工紋理。兩種結果都可能使表觀邊界偏離原始結構。因此，本研究不把任一單項分數作為唯一目標，而是用互補限制共同縮小可接受解的範圍。",
        718: "總損失分為影像內容、邊界幾何、運動一致性與輸出保守性四組。四組目標分別限制影像外觀、輪廓與中心位置、相鄰影格的運動關係，以及非量測區域的不必要修改。",
        730: "運動一致性損失的目的，不只是使平均位移接近PI閉迴路回授參考真值，更重要的是降低模型自身造成的影格依賴性。若同一結構只因落在不同像素相位或不同影格中，就被重建成不同尺寸或中心位置，便不適合作為微位移量測依據。",
        733: "四組目標形成互補限制。影像內容維持真實顯微外觀；邊界幾何約束尺寸、輪廓與中心；運動一致性維持影格間的位置關係；輸出保守性則防止非目標區域被不必要地改動。模型必須同時滿足這些條件，而不是用單一指標的改善交換另一項量測能力。",
        734: "量測導向訓練目標將一般影像相似度延伸到CD與微位移所需的幾何關係。完整模型使用四組目標共同訓練；第六章的單項移除結果則用來觀察補償關係，而不把整體效果勉強歸因於單一損失。",
        800: "本章建立一套由量測失效導出的單張顯微影像重建方法。固定PSF先驗描述成像模糊；抗混疊上採樣處理解碼時的像素相位與空間位置敏感性；輸入參考式殘差輸出限制低頻輪廓、亮暗拓樸與中心位置漂移；四組量測導向目標共同限制外觀、邊界、運動與非必要改動。",
        801: "光學部分以中心九宮格的逐顆SEM參考幾何擬合局部PSF，並以多位置PSF建立固定光學先驗。局部PSF提供Wiener對照，多位置PSF在所有模型及消融條件中保持相同，因此不會成為模組比較的額外變因。",
        802: "模型採單張、同尺寸輸入與輸出。抗混疊上採樣維持解碼過程的空間對應，輸入參考式殘差輸出使Raw保留為低頻位置基準；邊界殘差、品質門控與PSF修正則限制或引導局部修正。",
        803: "Raw、Wiener、MOSR與各消融條件使用相同座標、pillar身分、中心、物體空間尺度、門檻與統計規則。第六章將先以邊緣、跨門檻CD及PI閉迴路位移參考比較完整方法，再用其他條件一致的消融實驗檢查兩項主要模組。",

        1029: "本節在同一顯微系統與HS-100MG樣品下，使用相同物體空間尺度與後端規則比較三種影像的邊界、尺寸與位移。CD以九顆SEM參考尺寸為真值；位移則以PI閉迴路電容式感測回授提供的步階位置作為參考真值。",
        1044: "『回授參考平均絕對偏差』是每一步影像量測幅度與PI閉迴路電容式感測回授所確認之0.5 µm或1.0 µm步階的絕對差平均。此指標直接表示影像量測與本實驗位移參考真值的一致程度。各序列依預先固定的影像完整性與平台運動品質條件排除無法可靠量測的影格，再統一保留30個有效步；Raw、Wiener與MOSR共用相同名單。",
        1073: "九顆pillar的上下界寬度中位數為Raw 354.5 nm、Wiener 429.6 nm、MOSR 143.3 nm。Wiener的範圍最寬，反映環狀反彈會使交點隨門檻明顯移動，因此不能用某一個看似接近SEM的交點代表整體尺寸修復。將全部126個觀察納入後，絕對誤差中位數為Raw 95.3 nm、Wiener 123.7 nm、MOSR 76.1 nm。MOSR相較Raw降低20.16%，相較Wiener降低38.48%；這表示可量測的邊界需要同時具備集中度與門檻穩健性。",
        1110: "四組序列各固定30步，合併後共有120個配對步。Raw、Wiener與MOSR的回授參考平均絕對偏差為9.54、7.46與6.74 nm；Wiener與MOSR分別較Raw下降21.79%與29.36%。九宮格內平均不一致量則由Raw的17.63 nm降至Wiener的7.92 nm與MOSR的6.47 nm，分別下降55.08%與63.29%。因此，兩種處理均改善位移量測，而MOSR在合併後的回授參考一致性與九顆pillar一致性皆為最佳。",
        1133: "距離與Wiener回授參考絕對偏差的Spearman相關在50倍為ρ = 0.074、p = 0.576，20倍為ρ = -0.014、p = 0.915；距離與九宮格不一致量的相關則分別為ρ = 0.138、p = 0.292，以及ρ = -0.319、p = 0.013。前三項未達顯著，20倍九宮格不一致量則呈顯著負相關，而不是隨距離變差。四項結果都不支持『離校正位置越遠，Wiener誤差越大』的單調趨勢；在本次約25–30 µm取樣範圍內，固定PSF可作為條件一致的傳統基準。",
        1136: "前述Raw、Wiener與MOSR的比較可判斷完整重建方法是否有助於量測；本節再固定PSF、資料、訓練預算與評估規則，分別檢查抗混疊上採樣、輸入參考式殘差輸出、量測導向訓練目標及其他輸出模組。",
        1150: "九顆pillar的空間一致性呈現最明顯的改善。四組序列的不一致量均降低；合併120步後，九宮格不一致量由Raw的17.63 nm降至6.96 nm，下降60.52%，回授參考平均絕對偏差亦由9.54 nm降至6.70 nm，下降29.84%。兩組隨機初始化均重現主要實驗的改善方向。",
        1153: "6.2.6.3 輸入參考式殘差輸出對邊界的影響",
        1154: "輸入參考式殘差輸出以Raw維持低頻輪廓、主要亮暗與中心位置，只加入幅度受限的模型高頻成分。消融時移除此機制並以相同訓練預算重新訓練，用來檢查保留輸入位置基準是否必要。",
        1156: "圖 6-50 輸入參考式殘差輸出對兩倍率SEM對照CD誤差之影響",
        1157: "20倍移除此機制後，CD中位絕對誤差增加88.94 nm，即73.97%；兩組初始化分別增加78.44與99.44 nm，95%配對區間為49.29至125.29 nm。50倍的單一CD誤差方向不同，移除後降低15.09 nm；因此，本研究以跨倍率一致的振鈴結果與20倍CD結果共同判斷其作用，而不把單一指標延伸為所有倍率的通則。",
        1158: "振鈴在兩個倍率呈現一致方向：移除輸入參考式殘差輸出後，50倍與20倍的外圍振鈴分別增加246.16%與392.37%。若只在推論時直接關閉此機制，兩組靜態CD輸出亦因徑向對比不符合量測條件而失效，較大合成位移也明顯惡化。這些結果顯示，該機制的主要作用是維持輸入影像的邊界位置與亮暗拓樸，並抑制高頻增強造成的環狀假影。",
        1168: "單項移除後，其餘目標與模型模組可在重新訓練時部分補償，因此個別差異較小並不等於該目標無關。現有結果顯示，四組量測導向目標適合作為共同的可行域限制；可明確獨立歸因的效果則來自抗混疊上採樣與輸入參考式殘差輸出兩項主要模組。",
        1173: "消融結果把兩項主要創新與量測現象直接連結。抗混疊上採樣降低不同pillar之間的位移分歧，主要改善跨空間位置的幾何一致性；輸入參考式殘差輸出維持Raw的低頻位置與亮暗拓樸，主要抑制輪廓漂移與振鈴，並在20倍CD上呈現明顯效果。四組量測導向目標則共同限制模型不得以單一外觀改善交換其他量測能力。",
        1174: "PSF融合在完整模型及所有消融條件中維持固定，用來提供一致的光學先驗。多來源監督、邊界殘差與品質門控共同支援訓練與輸出限制；其角色不會改變本節對兩項主要模組的歸因。",
        1176: "綜合三方法比較與消融實驗，可回答本研究的三項問題。第一，單張同尺寸重建能使邊界集中，但最窄邊緣不必然帶來最穩定CD；MOSR在跨門檻CD上優於Raw與Wiener。第二，MOSR在120個配對步中具有最低的回授參考平均絕對偏差與九顆pillar不一致量。第三，抗混疊上採樣使九顆pillar不一致量較一般上採樣降低14.31%；輸入參考式殘差輸出則抑制兩倍率振鈴，並使20倍CD誤差維持在較低範圍。",
        1177: "本節結果建立於目前顯微系統、HS-100MG樣品、兩種倍率與四組位移條件。PI平台的閉迴路電容式感測回授提供本實驗的位移參考真值，使影像位移可與實際回授步階比較。現階段的主要適用範圍是同系統、同類結構的量測；跨樣品、跨系統與更高層級的外部追溯仍需後續驗證。消融使用兩組固定初始化，足以檢查方向一致性，更多重複可再縮小估計區間。",
        1178: "綜合現有證據，MOSR的主要價值不是單純產生最銳利影像，而是把邊界增強限制在可量測的範圍內。抗混疊上採樣維持多位置的平移關係，輸入參考式殘差輸出保護輪廓與亮暗拓樸，四組量測導向目標共同限制外觀、尺寸與運動。這三層設計使模型在CD與微位移兩類任務上形成一致的量測主軸。",

        1180: "結論與未來工作",
        1181: "本研究以量測導向超解析為主軸，完成光學先驗建立、單張同尺寸重建、SEM對照CD量測、PI閉迴路微位移量測及配對消融實驗。本章先彙整研究問題、主要方法貢獻與量化結果，再說明研究適用範圍及後續工作。",
        1182: "7.1 研究結論",
        1183: "第一個研究問題關注單張同尺寸重建能否改善邊界與CD。實驗顯示，Raw、Wiener與MOSR的邊緣過渡、門檻敏感度及SEM對照CD並非完全同向。Wiener可得到較窄的部分邊緣，卻具有最大的跨門檻CD誤差與最寬CD上下界；MOSR則同時降低跨門檻CD誤差與上下界寬度。這證明銳利度只能描述灰階變化速度，不能單獨代表量測正確性。",
        1184: "第二與第三個研究問題關注平移一致性及邊界保護。MOSR在四組PI閉迴路序列合併後，具有最低的回授參考平均絕對偏差與九顆pillar不一致量。消融結果進一步顯示，抗混疊上採樣負責降低不同pillar的位移分歧；輸入參考式殘差輸出則負責維持Raw的低頻位置與亮暗拓樸，抑制高頻增強造成的輪廓漂移與振鈴。",
        1185: "7.2 主要方法貢獻",
        1186: "本研究的第一項主要貢獻，是將解碼上採樣造成的像素相位與空間位置敏感性，轉化為可由九顆pillar共同位移直接檢查的模型問題，並以抗混疊上採樣改善跨位置幾何一致性。第二項主要貢獻，是將殘差輸出重新界定為量測邊界的保護機制，以Raw提供低頻位置基準，只允許受限的高頻修正。兩項設計分別處理解碼階段與輸出階段的幾何風險。",
        1187: "影像內容、邊界幾何、運動一致性與輸出保守性四組目標，進一步把外觀、輪廓、影格關係與非必要改動納入同一訓練限制。局部PSF提供Wiener對照，多位置PSF在模型及消融中作為固定光學先驗。此安排使研究主張集中於量測失效對應的模組與訓練設計，而不是光學元件與網路的單純串接。",
        1188: "7.3 主要量化結果",
        1189: "在CD量測中，Raw、Wiener與MOSR跨七個門檻與兩種剖面的中位絕對誤差分別為95.3、123.7與76.1 nm；MOSR相較Raw降低20.16%，相較Wiener降低38.48%。九顆pillar的CD上下界寬度中位數則由Raw的354.5 nm與Wiener的429.6 nm降至MOSR的143.3 nm，顯示MOSR的表觀邊界對量測規則較不敏感。",
        1190: "在四組位移序列共120個配對步中，Raw、Wiener與MOSR的回授參考平均絕對偏差為9.54、7.46與6.74 nm，九顆pillar不一致量為17.63、7.92與6.47 nm。抗混疊上採樣相較一般上採樣使合併不一致量降低14.31%；移除輸入參考式殘差輸出後，20倍CD中位絕對誤差增加73.97%，50倍與20倍振鈴分別增加246.16%與392.37%。",
        1191: "7.4 研究範圍與量測可信度",
        1192: "本研究的實證範圍為同一顯微系統、HS-100MG樣品、20倍與50倍物鏡、中心九顆pillar及四組PI閉迴路位移序列。SEM提供逐顆pillar尺寸參考真值；PI平台的閉迴路電容式感測回授提供位移參考真值。所有方法共用影像、中心、物體空間尺度、門檻與統計規則，因此方法間差異可由同一量測鏈直接比較。",
        1193: "現有結果尚不直接代表不同樣品、不同光學系統或完整視野的普遍表現。兩組固定初始化可確認主要消融方向，但增加訓練重複可進一步縮小估計區間；50倍MOSR仍存在較Raw明顯的外圍振鈴，也說明邊界集中、CD穩健度與假影必須共同檢查。這些範圍不影響本研究在既定系統與量測條件下的比較結論。",
        1194: "7.5 未來工作",
        1195: "後續工作可先擴充不同批次樣品、pillar尺寸、via、線條與方形結構，建立跨樣品測試；再於不同視野位置、焦平面、曝光與照明條件下增加PSF校正密度，以檢查固定光學先驗在更大範圍內的適用性[50,54,55,56]。若更換相機、物鏡或照明模組，應重新建立物體空間尺度與PSF資料，再判斷模型需完整重訓或少量校正[75,77,90]。",
        1196: "量測追溯方面，可使用雷射干涉儀與PI閉迴路電容式感測回授同步量測，進一步交叉驗證位移參考並建立包含重複性、像素尺度、SEM參考、平台回授及模型變動的完整不確定度預算[84,85]。模型驗證亦可增加訓練重複、參數量相近的基礎U-Net對照及跨條件分層統計，使量測導向設計相對一般影像重建的差異更加明確。",
        1197: "7.6 小結",
        1198: "本研究證明，顯微影像重建的有效性不能只由清晰度判斷。當模型以抗混疊上採樣維持跨位置的平移關係、以輸入參考式殘差輸出保護輪廓與亮暗拓樸，並以量測導向目標共同限制外觀、尺寸與運動時，重建影像可在既定系統下同時改善CD門檻穩健度與微位移一致性。後續若完成跨樣品、跨系統與更高層級量測追溯，MOSR可進一步形成可部署的精密顯微量測方法。",
    }

    for idx, text in replacements.items():
        clear_and_set(doc.paragraphs[idx], text)

    phrase_replacements = {
        "原始影像位置保留機制": "輸入參考式殘差輸出",
        "原始影像位置保留": "輸入參考式殘差輸出",
        "指令值平均絕對偏差": "回授參考平均絕對偏差",
        "Wiener 指令值絕對偏差": "Wiener 回授參考絕對偏差",
        "對名義位移指令": "相對PI閉迴路回授參考",
        "名義指令一致性": "回授參考一致性",
        "名義指令的平均絕對偏差": "回授參考平均絕對偏差",
        "名義指令": "PI閉迴路回授參考",
        "平台指令": "PI閉迴路回授參考",
        "指令偏差": "回授參考偏差",
        "指令一致性": "回授參考一致性",
    }
    phrase_counts = {}
    for old, new in phrase_replacements.items():
        phrase_counts[old] = replace_phrase_everywhere(doc, old, new)

    normalization_replacements = {
        "平台PI閉迴路回授參考": "PI閉迴路回授真值",
        "PI閉迴路回授參考真值": "PI回授真值",
        "相對回授參考真值": "相對PI回授真值",
        "相對PI閉迴路回授參考": "相對PI回授真值",
        "回授參考平均絕對偏差": "PI回授真值平均絕對偏差",
        "Wiener回授參考絕對偏差": "Wiener PI回授真值絕對偏差",
        "回授參考絕對偏差": "PI回授真值絕對偏差",
        "回授參考偏差": "PI回授真值偏差",
        "回授參考一致性": "PI回授真值一致性",
        "指令值偏差": "PI回授真值偏差",
        "指令尺度": "PI回授步階尺度",
    }
    normalization_counts = {}
    for old, new in normalization_replacements.items():
        normalization_counts[old] = replace_phrase_everywhere(doc, old, new)

    formula_table_count = compact_formula_tables(doc)
    enable_field_updates(doc)
    doc.save(TARGET)

    check = Document(TARGET)
    after_math = document_math_fingerprint(check)
    if after_math != before_math:
        raise RuntimeError(f"Formula content changed: before={before_math}, after={after_math}")
    if formula_table_count != 81:
        raise RuntimeError(f"Expected 81 formula tables, found {formula_table_count}")
    if len(check.paragraphs) != 1296:
        raise RuntimeError(f"Paragraph count changed after save: {len(check.paragraphs)}")

    remaining_formula_grid = 0
    remaining_formula_borders = 0
    noncompact_formula_paragraphs = 0
    for table in check.tables:
        if not table._tbl.xpath(".//m:oMath"):
            continue
        if table.style is not None and table.style.name == "Table Grid":
            remaining_formula_grid += 1
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is None or any(edge.get(qn("w:val")) != "nil" for edge in borders):
            remaining_formula_borders += 1
        for p_el in table._tbl.iter(qn("w:p")):
            p = Paragraph(p_el, table)
            fmt = p.paragraph_format
            if fmt.space_before != Pt(0) or fmt.space_after != Pt(0):
                noncompact_formula_paragraphs += 1

    if remaining_formula_grid or remaining_formula_borders or noncompact_formula_paragraphs:
        raise RuntimeError(
            "Formula-format validation failed: "
            f"grid={remaining_formula_grid}, borders={remaining_formula_borders}, "
            f"spacing={noncompact_formula_paragraphs}"
        )

    audit = {
        "source": str(SOURCE),
        "target": str(TARGET),
        "source_size": SOURCE.stat().st_size,
        "target_size": TARGET.stat().st_size,
        "paragraph_count": len(check.paragraphs),
        "table_count": len(check.tables),
        "formula_tables_cleaned": formula_table_count,
        "formula_fingerprint_before": before_math,
        "formula_fingerprint_after": after_math,
        "remaining_formula_table_grid_styles": remaining_formula_grid,
        "remaining_formula_visible_borders": remaining_formula_borders,
        "formula_paragraph_spacing_issues": noncompact_formula_paragraphs,
        "rewritten_paragraphs": sorted(replacements.keys()),
        "global_phrase_replacements": phrase_counts,
        "terminology_normalization": normalization_counts,
        "title_zh": check.paragraphs[7].text,
        "title_en": check.paragraphs[8].text,
        "chapter_7_title": check.paragraphs[1180].text,
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
