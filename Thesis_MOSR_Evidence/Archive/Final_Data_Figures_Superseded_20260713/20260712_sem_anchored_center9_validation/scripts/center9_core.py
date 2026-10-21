from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import openpyxl
from openpyxl.utils import get_column_letter


QuadrantLayout = tuple[tuple[str, str], tuple[str, str]]


@dataclass(frozen=True)
class Center9Identity:
    sem_excel_cell: str
    sem_index: int
    sem_layout_row: int
    sem_layout_col: int
    optical_grid_row: int
    optical_grid_col: int
    sem_radius_nm: float
    sem_diameter_um: float
    capture_transform: str


def _as_square(layout: Sequence[Sequence[str]]) -> tuple[tuple[str, ...], ...]:
    result = tuple(tuple(str(value) for value in row) for row in layout)
    if not result or any(len(row) != len(result) for row in result):
        raise ValueError("layout must be a non-empty square matrix")
    return result


def _apply_layout_transform(
    layout: tuple[tuple[str, ...], ...], transform: str
) -> tuple[tuple[str, ...], ...]:
    n = len(layout)
    if transform == "identity":
        return layout
    if transform == "rot90_cw":
        return tuple(tuple(layout[n - 1 - col][row] for col in range(n)) for row in range(n))
    if transform == "rot180":
        return tuple(tuple(layout[n - 1 - row][n - 1 - col] for col in range(n)) for row in range(n))
    if transform == "rot270_cw":
        return tuple(tuple(layout[col][n - 1 - row] for col in range(n)) for row in range(n))
    if transform == "flip_x":
        return tuple(tuple(reversed(row)) for row in layout)
    if transform == "flip_y":
        return tuple(reversed(layout))
    if transform == "transpose":
        return tuple(tuple(layout[col][row] for col in range(n)) for row in range(n))
    if transform == "anti_transpose":
        return tuple(
            tuple(layout[n - 1 - col][n - 1 - row] for col in range(n))
            for row in range(n)
        )
    raise ValueError(f"unknown transform: {transform}")


TRANSFORMS = (
    "identity",
    "rot90_cw",
    "rot180",
    "rot270_cw",
    "flip_x",
    "flip_y",
    "transpose",
    "anti_transpose",
)


def infer_quadrant_transform(
    physical_layout: Sequence[Sequence[str]],
    observed_layout: Sequence[Sequence[str]],
) -> str:
    physical = _as_square(physical_layout)
    observed = _as_square(observed_layout)
    if len(physical) != len(observed):
        raise ValueError("physical and observed layouts must have equal size")
    matches = [
        name
        for name in TRANSFORMS
        if _apply_layout_transform(physical, name) == observed
    ]
    if len(matches) != 1:
        raise ValueError(
            f"quadrant transform must be unique; matched {len(matches)}: {matches}"
        )
    return matches[0]


def transform_grid_coordinate(
    row: int, col: int, *, grid_size: int, transform: str
) -> tuple[int, int]:
    values = {"row": row, "col": col, "grid_size": grid_size}
    for label, value in values.items():
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ValueError(f"{label} must be an integer")
        if not float(value).is_integer():
            raise ValueError(f"{label} must be an integer")
    n = int(grid_size)
    r = int(row)
    c = int(col)
    if not (1 <= r <= n and 1 <= c <= n):
        raise ValueError(f"grid coordinate outside 1..{n}: {(r, c)}")
    if transform == "identity":
        return r, c
    if transform == "rot90_cw":
        return c, n + 1 - r
    if transform == "rot180":
        return n + 1 - r, n + 1 - c
    if transform == "rot270_cw":
        return n + 1 - c, r
    if transform == "flip_x":
        return r, n + 1 - c
    if transform == "flip_y":
        return n + 1 - r, c
    if transform == "transpose":
        return c, r
    if transform == "anti_transpose":
        return n + 1 - c, n + 1 - r
    raise ValueError(f"unknown transform: {transform}")


def center9_identity_map(
    workbook_path: str | Path, *, capture_transform: str
) -> tuple[Center9Identity, ...]:
    path = Path(workbook_path)
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["SEM data"]
        identities: list[Center9Identity] = []

        for layout_row, excel_row in zip(range(4, 7), range(8, 11)):
            for layout_col, excel_col in zip(range(1, 4), range(10, 13)):
                cell = f"{get_column_letter(excel_col)}{excel_row}"
                radius_nm = float(sheet[cell].value)
                sem_index = (layout_row - 1) * 6 + (7 - layout_col)
                listed_index = int(sheet.cell(sem_index + 2, 17).value)
                listed_radius_nm = float(sheet.cell(sem_index + 2, 19).value)
                if listed_index != sem_index or abs(listed_radius_nm - radius_nm) > 1e-9:
                    raise ValueError(
                        f"SEM layout/index mismatch at {cell}: "
                        f"layout={radius_nm}, index={sem_index}, listed={listed_radius_nm}"
                    )
                optical_row, optical_col = transform_grid_coordinate(
                    layout_row,
                    layout_col,
                    grid_size=6,
                    transform=capture_transform,
                )
                identities.append(
                    Center9Identity(
                        sem_excel_cell=cell,
                        sem_index=sem_index,
                        sem_layout_row=layout_row,
                        sem_layout_col=layout_col,
                        optical_grid_row=optical_row,
                        optical_grid_col=optical_col,
                        sem_radius_nm=radius_nm,
                        sem_diameter_um=radius_nm * 2.0 / 1000.0,
                        capture_transform=capture_transform,
                    )
                )
        return tuple(identities)
    finally:
        workbook.close()
