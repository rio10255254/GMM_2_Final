from pathlib import Path
import re

from docx import Document

import build_chapter4_reframe as base


OUT = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\第四章_研究方法_量測導向完整擴寫稿_樸素圖表版.docx")
base.FIG_DIR = Path(r"D:\GMM_2_Final\Thesis_MOSR_Evidence\Final_Data_Figures\05_CHAPTER4_METHOD_PLAIN")


def body(doc, text, **kwargs):
    return base.add_body(doc, text, **kwargs)


def eq(doc, text):
    return base.add_equation(doc, text)


def fig(doc, filename, caption, width=14.4):
    return base.add_figure(doc, filename, caption, width)


def table(doc, caption, headers, rows, widths, font_size=8.7):
    return base.add_table(doc, caption, headers, rows, widths, font_size=font_size)


def build():
    doc = Document()
    base.configure_document(doc)

    doc.add_heading("第4章　研究方法", level=1)

    body(doc, "本研究的目標並不是單純把顯微影像變得更銳利，而是讓影像在變清楚之後，仍能可靠地用來量測結構尺寸與微小位移。對量測而言，影像看起來清楚不一定代表結果正確：過度銳化可能把邊界推向錯誤位置，反卷積振鈴可能產生不存在的亮暗輪廓，而逐張自動拉伸對比也可能讓不同方法看似有明顯差異，實際上只是顯示方式不同。因此，本章把光學校正、影像重建與後端量測視為同一條方法鏈，所有設計均以「邊界是否更容易定位、尺寸是否更接近獨立參考、位移是否更穩定」作為判斷依據。")
    body(doc, "本研究將此概念稱為量測導向超解析（Measurement-Oriented Super-Resolution, MOSR）。此處的超解析並非宣稱任意創造超出光學系統可傳遞範圍的細節，而是利用光學點擴散函數、真實顯微影像、已知幾何與已知位移資料，恢復原始影像中已存在但因模糊與雜訊而不易定位的邊界資訊。模型於實際使用時只輸入單張影像，輸出尺寸與輸入相同；連續影像與位移資訊主要用於訓練與驗證，並不要求部署時堆疊多張影格。")

    doc.add_heading("4.1 研究問題與整體架構", level=2)
    doc.add_heading("4.1.1 顯微成像與量測問題", level=3)
    body(doc, "在近似線性且空間局部不變的範圍內，實際拍攝影像可表示為理想物體與點擴散函數的二維卷積，再加上量測雜訊：")
    eq(doc, "g(x,y) = [o*h](x,y) + n(x,y).                                                   (4-1)")
    body(doc, "式中 o(x,y) 為理想物體在物體空間中的強度分布，h(x,y) 為該拍攝位置的點擴散函數（point spread function, PSF），星號表示二維卷積，n(x,y) 則包含相機讀出雜訊、光子雜訊、照明不均與其他未建模擾動。PSF 可理解為光學系統把一個理想點擴散成多大、多偏斜以及往哪個方向拖曳；當 PSF 越寬或越不對稱，原本陡峭的 pillar 邊界就越容易變成緩慢的灰階過渡。")
    body(doc, "式（4-1）的空間不變假設只在有限區域內成立。實際顯微系統的離軸像差、樣品傾斜、局部對焦差與照明分布會使 PSF 隨視野位置改變，因此本研究不把單一 PSF 視為整個視野的唯一答案。局部擬合 PSF 用來說明特定九宮格附近的模糊並建立傳統反卷積對照；多位置 PSF 則用於模型中的固定光學先驗，使重建不過度依賴單一位置。")

    doc.add_heading("4.1.2 物體空間尺度與座標定義", level=3)
    body(doc, "相機紀錄的是感測器平面的離散像素，但本研究最終關心的是樣品表面的尺寸與位移，因此所有像素量測都必須換算到物體空間。若相機像素尺寸為 pcam，物鏡倍率為 M，則忽略額外中繼倍率時，物體空間每像素尺寸 pobj 為：")
    eq(doc, "p_obj = p_cam / M.                                                               (4-2)")
    body(doc, "本研究使用 Basler ace acA2040-120um 工業相機，其感測器像素尺寸為 3.45 µm。由式（4-2）可得，50 倍物鏡影像的物體空間取樣為 0.069 µm/pixel，即 69 nm/pixel；20 倍物鏡影像為 0.1725 µm/pixel，即 172.5 nm/pixel。PSF 矩陣的索引仍位於相機影像的取樣網格上，當剖面寬度、半高全寬或位移需要用物理單位表示時，再乘上對應倍率的物體空間像素尺寸。")
    table(doc,
          "Table 4-1 Imaging scale and coordinate convention",
          ["Item", "Setting", "Meaning in this study"],
          [
              ["Camera", "Basler ace acA2040-120um", "Monochrome industrial camera used for optical acquisition"],
              ["Sensor pixel size", "3.45 µm/pixel", "Native sampling pitch on the sensor plane"],
              ["50× object-space scale", "69 nm/pixel", "3.45 µm divided by 50"],
              ["20× object-space scale", "172.5 nm/pixel", "3.45 µm divided by 20"],
              ["PSF matrix", "101 × 101 samples", "Discrete PSF support on the captured-image sampling grid"],
              ["Reported CD and motion", "Object space", "Pixel measurements multiplied by the magnification-specific scale"],
          ], [3.7, 4.3, 6.6], font_size=8.5)
    body(doc, "上述換算只改變單位，不會自動提升光學解析度。20 倍影像每個像素代表的樣品距離較大，因此同一顆約 2.63 µm 的 pillar 在影像上只跨越約 15.2 像素；50 倍影像則約跨越 38.1 像素。這也是兩種倍率不能直接以像素數比較 CD 誤差的原因：所有結果必須先轉為奈米，再與同一顆 pillar 的 SEM 尺寸比較。")

    doc.add_heading("4.1.3 量測導向的問題定義", level=3)
    body(doc, "令輸入影像為 x，重建模型為 fθ，輸出影像為 ŷ，後端量測方法為 M。模型輸出與量測值的關係可寫成：")
    eq(doc, "ŷ = f_θ(x),                    m̂ = M(ŷ).                                      (4-3)")
    body(doc, "其中 θ 表示模型參數，m̂ 可代表單顆 pillar 的 CD、相鄰影像間的位移，或同一步進下九顆 pillar 的局部位移集合。式（4-3）強調影像不是研究終點，而是量測流程的中間證據。即使 ŷ 的對比更高，若 M(ŷ) 對門檻非常敏感、逐顆尺寸偏離 SEM，或連續步進忽大忽小，該重建仍不能視為量測改善。")
    body(doc, "依此定義，本研究設定四項完成條件。第一，重建後的 10–90% 邊緣寬度應縮小，但振鈴不可明顯增加。第二，九顆 pillar 的 CD 必須逐顆對應 SEM，而不是只讓整體平均值接近。第三，相鄰影像的位移估計應在固定 30 步中呈現較小的步間波動與九宮格內不一致。第四，所有改善必須在相同裁切、相同方向、相同像素尺度、相同中心與相同量測規則下成立。")

    doc.add_heading("4.1.4 整體研究流程", level=3)
    fig(doc, "fig_4_1_overall_method_workflow.png", "Figure 4-1 Measurement-oriented reconstruction workflow")
    body(doc, "如圖 4-1 所示，研究流程由五個階段組成。第一階段先辨識 pillar 與 via 的實際區域，選出靠近樣品中心的 pillar 九宮格，校正不同倍率的拍攝方向，並以 SEM 尺寸擬合局部 PSF。第二階段建立包含已知幾何、真實顯微紋理、SEM 參考與已知位移的訓練資料。第三階段利用影像內容、邊界幾何、位移一致性與輸出保守性共同訓練模型。第四階段以單張影像完成同尺寸重建。第五階段才進行邊緣、CD、位移與條件穩健性驗證。")
    table(doc,
          "Table 4-2 Research objectives, method design, and acceptance evidence",
          ["Research objective", "Method design", "Acceptance evidence"],
          [
              ["Describe optical blur", "Local parametric PSF fitting with SEM geometry", "PSF profile, fitted-image agreement, and matched Wiener baseline"],
              ["Reduce position-dependent bias", "Fixed fusion of normalized multi-position PSFs", "Stable single-image physical prior across the field"],
              ["Preserve geometry", "Bounded boundary residual and Raw low-frequency anchor", "Narrower edges without larger ringing or uncontrolled CD shift"],
              ["Preserve motion", "Known-shift supervision and translation consistency", "Lower step variation and lower nine-pillar disagreement"],
              ["Prevent favorable reprocessing", "Raw-locked centers, steps, scale, and thresholds", "Identical evaluation protocol for Raw, Wiener, and proposed model"],
          ], [4.1, 5.4, 5.1], font_size=8.2)

    doc.add_heading("4.2 局部 PSF 校正與多位置光學先驗", level=2)
    doc.add_heading("4.2.1 Pillar 區域、中心九宮格與方向校正", level=3)
    fig(doc, "fig_4_2_center_nine_and_sem_ground_truth.png", "Figure 4-2 Center-nine pillar selection, orientation correction, and SEM ground truth")
    body(doc, "PSF 擬合不能只看到九個近似圓形的影像就直接使用，因為樣品上同時存在 pillar 與 via 兩大結構區域，而且不同倍率拍攝時樣品擺放方向可能不同。本研究先依整片樣品的版圖位置判定 pillar/via 交界，再由交界處往樣品中心方向選取最近的 pillar 3×3 九宮格。圖 4-2(a) 中的星號表示樣品中心方向，虛線框則表示實際採用的九顆 pillar。")
    body(doc, "Pillar 與 via 的判定不使用影像明暗。明場影像的明暗極性會隨倍率、曝光、照明與顯示方式改變，同一種結構在不同資料中可能呈現亮底暗圓或暗底亮圓；若把極性當成結構身分，便可能在 20 倍與 50 倍資料中選到不同區域。因此，結構身分以樣品區域與交界位置決定，極性只在後續邊緣梯度搜尋時用來決定由暗到亮或由亮到暗的方向。")
    body(doc, "兩種倍率的原始拍攝方向也不相同。50 倍影像維持拍攝方向；20 倍影像先做垂直翻轉，使九顆 pillar 的排列與 SEM 真實方向一致。方向校正完成後才配置逐顆 SEM 尺寸。此順序很重要，因為若先套用尺寸再翻轉，九個數值雖然平均相同，逐顆誤差卻會被錯誤配對。")

    doc.add_heading("4.2.2 SEM 半徑轉換與逐顆幾何真值", level=3)
    body(doc, "SEM 提供的九個數值為半徑，不是直徑。故第 j 顆 pillar 的 SEM 臨界尺寸定義為 CDSEM,j=2rSEM,j。九個數值按照真實樣品方向排列，矩陣左下角（R3C1）最靠近樣品中心。轉換後的平均 CD 為 2629.737 nm，樣本標準差為 24.859 nm，最小值與最大值分別為 2598.752 nm 與 2676.797 nm。")
    sem_values = [
        ("R1C1", 1310.107187), ("R1C2", 1310.008186), ("R1C3", 1321.157790),
        ("R2C1", 1310.214279), ("R2C2", 1304.359769), ("R2C3", 1329.306897),
        ("R3C1 (nearest sample center)", 1310.889999),
        ("R3C2", 1299.375827), ("R3C3", 1338.398285),
    ]
    table(doc,
          "Table 4-3 Per-pillar SEM radius and ground-truth critical dimension",
          ["Pillar position", "SEM radius (nm)", "Ground-truth CD (nm)"],
          [[p, f"{r:.6f}", f"{2*r:.6f}"] for p, r in sem_values],
          [5.2, 4.5, 4.7], font_size=8.6)
    body(doc, "本研究保留表 4-3 的逐顆差異，不以單一平均圓取代九顆 pillar。原因是 SEM 已顯示不同位置存在約 78 nm 的真實尺寸範圍；若全部使用平均直徑建立 Ground Truth，模型或反卷積造成的誤差會與樣品本身的製程差異混在一起。逐顆比較可讓第 j 顆光學 CD 只對應第 j 顆 SEM CD，避免用一顆偏大的 pillar 補償另一顆偏小的 pillar。")

    doc.add_heading("4.2.3 參數化 PSF 與 Zernike 像差", level=3)
    body(doc, "PSF 擬合以單位圓光瞳上的 Zernike 多項式描述波前像差。令 ρ 為正規化光瞳半徑、φ 為方位角、aj 為第 j 項像差係數，則等效光瞳函數與強度 PSF 可概念性表示為：")
    eq(doc, "P(ρ,φ) = A(ρ) exp[i2πΣ_j a_j Z_j(ρ,φ)],      h = |F{P}|² / Σ|F{P}|².          (4-4)")
    body(doc, "A(ρ) 表示孔徑振幅，Z_j 為 Zernike 基底，F{·} 表示由光瞳傳播到像面的傅立葉轉換。分母使 PSF 的總能量正規化為 1，因此不同候選 PSF 的亮度尺度不會混入形狀比較。Zernike 多項式的好處是每個係數對應常見像差形態，可用有限參數描述失焦、散光、彗差與高階像差，而不必直接自由調整 101×101 個像素。")
    table(doc,
          "Table 4-4 Zernike terms used in the local PSF fitting",
          ["Terms", "Optical meaning", "Treatment in fitting"],
          [
              ["Z2–Z3", "Horizontal and vertical tilt", "Fixed at zero after subpixel pre-alignment"],
              ["Z4 and Z6", "Primary astigmatism", "Optimized"],
              ["Z5", "Defocus", "Optimized"],
              ["Z7 and Z10", "Trefoil", "Optimized"],
              ["Z8 and Z9", "Coma", "Optimized"],
              ["Z11 and Z15", "Quadrafoil", "Optimized"],
              ["Z12 and Z14", "Secondary astigmatism", "Optimized"],
              ["Z13", "Spherical aberration", "Optimized"],
              ["Amplitude", "Overall fitted intensity scale", "Optimized separately from unit-energy PSF"],
          ], [3.1, 5.2, 6.0], font_size=8.2)
    body(doc, "Z2 與 Z3 代表光瞳傾斜，在影像中主要表現為位置平移。若未先對位便讓這兩項自由變動，最佳化可能把裁切偏移誤當成光學像差。因此，本研究先在 ±2 像素範圍內，以 0.25 像素間隔搜尋亞像素平移，完成預對位後將 Z2、Z3 固定為 0，只搜尋 Z4–Z15 與振幅。這樣可把「影像位置」與「PSF 形狀」分開處理。")

    doc.add_heading("4.2.4 BADS 目標函數與擬合程序", level=3)
    fig(doc, "fig_4_3_psf_fitting_process.png", "Figure 4-3 Parametric PSF fitting procedure")
    body(doc, "如圖 4-3 所示，擬合流程先將逐顆 SEM 幾何形成理想 Ground Truth，再與候選 PSF 卷積得到預測影像 Ipred。預測影像與實拍影像 Iobs 只在中心九宮格的有效遮罩內比較。遮罩可避免裁切邊界、非目標結構與無關背景主導目標函數。")
    eq(doc, "I_pred(θ) = I_GT * h(θ),                    Σ_u,v h(u,v;θ) = 1.             (4-5)")
    body(doc, "本研究使用 Bayesian Adaptive Direct Search（BADS）搜尋像差係數。BADS 適合處理無解析梯度、每次評估需要重新產生 PSF 與卷積影像的問題。目標函數同時包含遮罩內灰階均方誤差與正規化梯度均方誤差：")
    eq(doc, "J(θ) = MSE_mask(I_obs,I_pred) + 8×10⁴ MSE_mask[N(∇I_obs),N(∇I_pred)].      (4-6)")
    body(doc, "第一項要求預測影像維持整體亮暗與模糊分布；第二項比較邊界方向與陡峭程度。N(·) 表示在比較前把梯度尺度正規化，避免較高對比影像只因梯度振幅較大就取得較高權重。梯度項係數 8×10⁴ 是為了讓邊界資訊在數值尺度上足以與強度項共同影響搜尋，而不是把此數值解讀為梯度比亮度重要八萬倍。")
    body(doc, "每輪 BADS 會根據目前候選係數產生單位能量 PSF、卷積 Ground Truth、計算遮罩損失，再提出下一組係數。當目標函數不再改善或達到停止條件時，保留最低損失的 PSF。最終 PSF 的峰值、剖面、半高全寬與頻率響應必須由未再強化的 PSF 矩陣直接計算，以避免顯示用對比調整影響物理分析。")

    doc.add_heading("4.2.5 固定 25-PSF 融合與保守物理先驗", level=3)
    body(doc, "局部 20 倍與 50 倍 PSF 用於特定九宮格的光學分析與 Wiener 對照；模型內部則使用不同位置的 25 張 PSF 描述視野變化。每張 PSF 皆為 101×101 且先獨立正規化，再以等權平均形成固定融合 PSF：")
    eq(doc, "h̄ = (1/25) Σ_k=1^25 h_k,                    Σ_u,v h̄(u,v) = 1.              (4-7)")
    body(doc, "等權平均不是假設所有位置完全相同，而是建立一個不偏向單一位置的保守模糊先驗。若直接用某一角落或某一中心位置的 PSF 強反卷積，空間不匹配可能在其他區域產生振鈴；固定融合可先提供共同的去模糊方向，再由網路利用影像內容調整局部細節。")
    eq(doc, "x_phys = x + α{F⁻¹[H̄*X/(|H̄|²+γ)] − x},     α=0.10, γ=10.              (4-8)")
    body(doc, "式中 X 與 H̄ 分別為輸入影像與融合 PSF 的頻譜，H̄* 為複共軛，γ 為抑制噪聲放大的正則化量，α 則控制物理殘差只以多少比例加入原圖。本研究使用 α=0.10、γ=10，目的不是讓此步驟單獨完成強反卷積，而是向後續網路提供溫和且有物理方向的初始修正。")

    doc.add_heading("4.2.6 可切換的自適應多 PSF 策略", level=3)
    body(doc, "固定平均具有穩定性，但若未來需要進一步處理局部模糊差異，可使用四候選 PSF 的自適應加權策略。四個代表性 PSF 分別產生補償特徵 Dk，再由輸入特徵預測非負權重 wk，使權重總和為 1：")
    eq(doc, "D_mix = Σ_k=1^4 w_kD_k,      Σ_kw_k=1;      F_out = F + β(D_mix−F).         (4-9)")
    body(doc, "β 用來限制自適應分支的最大影響，避免權重判斷錯誤時直接改寫整張影像。此策略保留於完整方法架構中，但本章後續以固定 25-PSF 融合描述現行重建路徑；自適應權重若啟用，必須使用相符的訓練狀態，並另以相同量測規則驗證，不把其潛在效果混入現有結果。")
    table(doc,
          "Table 4-5 Roles of the three PSF representations",
          ["PSF representation", "Primary role", "What it must not be confused with"],
          [
              ["Local 20× or 50× fitted PSF", "Optical analysis and matched Wiener baseline near the selected center-nine region", "It is not assumed to describe the entire field"],
              ["Fixed mean of 25 PSFs", "Active conservative physical prior used by the reconstruction model", "It is not a strong standalone deconvolution result"],
              ["Adaptive four-candidate PSFs", "Switchable extension for location-dependent compensation", "Current evaluation is not attributed to adaptive weighting"],
          ], [4.2, 5.5, 5.5], font_size=8.3)

    doc.add_heading("4.3 量測導向單張影像重建模型", level=2)
    doc.add_heading("4.3.1 同尺寸編碼器—解碼器", level=3)
    fig(doc, "fig_4_4_measurement_oriented_network.png", "Figure 4-4 Measurement-oriented single-image reconstruction network")
    body(doc, "圖 4-4 顯示模型的主要資料流。輸入為一張 H×W 的單通道顯微影像，先通過固定 25-PSF 物理先驗，再進入編碼器—解碼器。編碼端逐層擴大感受範圍，用來辨識背景、pillar 週期與較大尺度的模糊；解碼端逐步恢復空間解析，跳接則把淺層的精確位置資訊直接帶到相同尺度的解碼層，降低下採樣造成的位置損失。")
    body(doc, "輸出仍為 H×W，而不是把影像放大成更多像素。這種同尺寸設計可避免把插值後較密的像素誤認為真正解析度提升，也使輸出能直接沿用原始相機的物體空間尺度。部署時不堆疊多張影像，因此每一張重建結果可獨立對應原始拍攝時刻，適合後續逐步位移追蹤。")

    doc.add_heading("4.3.2 方向感知且受限的邊界殘差", level=3)
    body(doc, "主幹網路後加入局部邊界修正。修正器同時讀取目前預測、水平與垂直梯度、梯度幅值及解碼特徵，判斷每個位置需要增加或減少多少局部細節。由於 pillar 可能是暗圓或亮圓，修正方向依梯度符號與局部極性決定，而不是預先假定固定亮暗。")
    eq(doc, "y_b = y_0 + g(x)·clip[r(x),−0.025,0.025].                                  (4-10)")
    body(doc, "式中 y0 為主幹輸出，r(x) 為預測殘差，g(x) 為 0 至 1 的局部門控。clip 限制每個像素的絕對殘差不超過 0.025（灰階已正規化至 0–1）。此上限的意義是限制模型只能微調邊界附近的灰階，不能自由改寫大面積幾何；門控則讓平坦背景與高不確定區域使用更小修正。")

    doc.add_heading("4.3.3 原圖低頻位置錨定", level=3)
    body(doc, "僅限制單像素殘差仍不足以保證整體位置不動，因為許多小殘差累積後仍可能形成輪廓平移。因此，本研究把原始影像當作低頻位置錨點。先以 σ=2.0 像素的高斯濾波取得預測影像低頻，再把預測與低頻的差視為細節，只取其中 0.35 倍加回原始影像：")
    eq(doc, "y_anchor = x + 0.35[y_b − G_σ(y_b)],             σ=2.0 pixels.              (4-11)")
    body(doc, "式（4-11）保留 x 的背景起伏、平均亮度與大尺度位置，模型主要提供 yb 中超過 σ=2.0 像素尺度的局部細節。0.35 並非對影像做全域對比放大，而是限制高頻殘差的加入比例。完成錨定後再校正輸出平均值，使重建不因網路偏移造成整張影像變亮或變暗。")

    doc.add_heading("4.3.4 輸入品質門控", level=3)
    body(doc, "20 倍與 50 倍影像的取樣密度、邊界訊號與噪聲不同，即使相同倍率，不同曝光或焦點也可能使影像品質改變。為避免固定修正量在低品質影像中放大雜訊，本研究以邊界訊號 S 與平坦區噪聲 N 建立 0 至 1 的品質分數 q：")
    eq(doc, "q = S/(S+3N+ε),              s = 0.78 + 0.30q.                              (4-12)")
    body(doc, "S 由梯度幅值較高的像素估計可用邊界訊號，N 由局部平坦區相對 5×5 平滑影像的偏差估計。當 q 接近 0，修正尺度 s 接近 0.78，輸出較保守；當 q 接近 1，s 最多為 1.08，較清楚的輸入可保留更多細節。門控只縮放重建與原圖之差，不改變像素座標。")
    fig(doc, "fig_4_5_output_safeguards.png", "Figure 4-5 Output safeguards for geometric preservation")
    body(doc, "圖 4-5 將三道保護機制分開說明：受限殘差控制單點修正幅度，原圖低頻錨定控制大尺度位置與背景，輸入品質門控則依影像可靠程度調整整體修正。三者共同作用，使模型的主要自由度集中在可量測邊界，而不是任意重新生成整張顯微影像。")
    table(doc,
          "Table 4-6 Main configuration of the reconstruction model",
          ["Component", "Fixed setting", "Measurement purpose"],
          [
              ["Input and output", "One H×W image to one H×W image", "Preserve the original sampling grid and acquisition time"],
              ["Physical prior", "Mean of 25 normalized PSFs", "Provide a field-robust deblurring direction"],
              ["Physical residual", "α=0.10; γ=10", "Avoid strong inverse-filter amplification"],
              ["Boundary residual", "Absolute cap=0.025", "Limit local geometric rewriting"],
              ["Raw anchor", "Gaussian σ=2.0 pixels; detail gain=0.35", "Tie low-frequency position and mean intensity to Raw"],
              ["Quality gate", "Correction scale 0.78–1.08", "Use conservative correction for noisy inputs"],
          ], [3.6, 4.6, 6.5], font_size=8.3)

    doc.add_heading("4.4 訓練資料與監督來源", level=2)
    doc.add_heading("4.4.1 為何需要多種真值", level=3)
    fig(doc, "fig_4_6_training_supervision_design.png", "Figure 4-6 Training supervision design")
    body(doc, "單一資料來源無法同時提供真實外觀、正確尺寸與精確位移。合成幾何具有明確邊界，但缺少真實相機雜訊；真實影像保留照明與紋理，卻不一定具有獨立尺寸真值；SEM 可提供幾何尺度，但影像外觀與光學相機不同；已知位移配對可約束平移，卻不能單獨決定絕對 CD。因此，本研究依功能整合多種監督，並明確區分每一類資料能證明什麼。")
    body(doc, "圖 4-6 上方五類來源分別提供已知形狀、SEM 幾何、真實顯微外觀、已知平移與跨條件穩定性；下方四組訓練目標則把這些資訊轉成影像內容、邊界幾何、運動一致性與保守性限制。SEM 參考定義幾何，真實影像參考定義可接受的外觀，兩者不被視為同一種真值。")

    doc.add_heading("4.4.2 訓練配對組成", level=3)
    body(doc, "完整訓練資料包含 11,254 組影像配對。為避免以內部資料夾或實驗代號描述，本研究按監督功能重新整理為九類，如表 4-7。最大一類為真實紋理上的精確位移配對，占 30.82%；其餘資料補足合成運動、真實邊緣、實際序列、穩定性、幾何、步階保留、一般形狀與 SEM 錨點。")
    table(doc,
          "Table 4-7 Training-pair composition by supervision function",
          ["Supervision function", "Pairs", "Share", "Primary contribution"],
          [
              ["Exact-shift real textures", "3,468", "30.82%", "Translation consistency with real camera texture"],
              ["Synthetic motion", "1,520", "13.51%", "Controlled subpixel displacement"],
              ["Edge-refined real images", "1,470", "13.06%", "Boundary recovery under real blur and noise"],
              ["Real-sequence motion", "1,176", "10.45%", "Captured temporal and mechanical variation"],
              ["Stability real textures", "800", "7.11%", "Conservative response to texture changes"],
              ["Synthetic geometry", "760", "6.75%", "Exact shape and boundary location"],
              ["Step preservation", "760", "6.75%", "Protection against displacement-scale distortion"],
              ["General shape", "660", "5.86%", "Coverage beyond a single circular pattern"],
              ["SEM-referenced geometry", "640", "5.69%", "Independent physical size anchor"],
          ], [4.8, 1.5, 1.7, 6.5], font_size=7.8)
    body(doc, "資料比例不等於證據強度。例如 SEM 參考只有 5.69%，但它是少數具有獨立物理尺度的監督；真實紋理配對數量較多，主要作用是讓模型適應相機噪聲與實際外觀，不能因此被解讀為有 3,468 組獨立 SEM 真值。這項區分可避免把訓練資料的影像相似性誤寫成量測正確性。")

    doc.add_heading("4.4.3 邊界軟標籤", level=3)
    body(doc, "二值圓形在像素網格上會形成階梯狀邊界，若把硬式 0/1 邊界直接當成唯一真值，模型可能學到鋸齒或對單一像素門檻過度敏感。為提供連續的亞像素監督，本研究在 180 組代表性覆蓋資料上，只於距離邊界 2 像素的窄帶內施加 σ=0.75 像素的高斯軟化。物體中心與遠離邊界的背景保持不變。")
    body(doc, "軟標籤的功能不是讓 Ground Truth 變模糊，而是把真實邊界可能落在兩個像素中心之間的情況表示成連續過渡。訓練時再搭配邊界遮罩、中心線與距離圖，使模型分別學習「邊界在哪裡」、「物體中心是否穩定」與「每個像素離輪廓多遠」。這些資訊對 CD 與位移定位比單純像素均方誤差更直接。")

    doc.add_heading("4.4.4 分組切分與適用範圍", level=3)
    body(doc, "相鄰影格通常只差微小位移，若隨機逐張分割訓練與驗證資料，同一序列的近似影像可能同時出現在兩邊，造成驗證結果過度樂觀。因此，本研究以原始序列、倍率與形狀群組為單位切分，確保同一段連續影像不跨越訓練與驗證集合。所有資料在切分前即固定群組，後續不因單一結果好壞重新分配。")
    body(doc, "目前資料涵蓋本研究使用的顯微系統、100MG 樣品、兩種倍率與多種幾何／位移條件，因此可支持預定量測情境下的功能驗證。若要主張跨設備、跨照明光譜或跨材料泛化，仍需使用未參與資料建立的新樣品與新系統，在不重新訓練的情況下測試。本章不把尚未完成的外部泛化當成既有結論。")

    doc.add_heading("4.5 訓練目標與分階段微調", level=2)
    doc.add_heading("4.5.1 四組功能性損失", level=3)
    body(doc, "本研究不以單一像素損失決定模型，而將損失依量測目的分成影像內容、邊界幾何、運動一致性與輸出保守性四組。總損失可寫成：")
    eq(doc, "L_total = L_content + L_boundary + L_motion + L_guard.                      (4-13)")
    body(doc, "Lcontent 維持灰階、結構相似度、梯度與頻帶能量，使輸出保留真實影像外觀；Lboundary 利用遮罩、中心線、距離圖與邊緣方向約束輪廓；Lmotion 使相同物體在相鄰影格或已知位移下保持尺寸與位移關係；Lguard 則懲罰低頻偏移、非邊界區域改動與過大局部殘差。四組損失同時使用，可避免模型只在一張圖上看起來清楚，卻破壞連續量測。")
    table(doc,
          "Table 4-8 Functional loss groups and their interpretation",
          ["Loss group", "Information constrained", "Failure mode reduced"],
          [
              ["Image content", "Intensity, structure, gradient, and usable frequency content", "Unrealistic appearance or global contrast drift"],
              ["Boundary geometry", "Mask, centerline, distance, and edge direction", "Boundary displacement and shape distortion"],
              ["Motion consistency", "Known shifts, adjacent-frame centers, and step scale", "Frame-dependent position or size changes"],
              ["Safeguards", "Low-frequency anchor, locality, and bounded residual", "Hallucinated detail and large-area rewriting"],
          ], [3.3, 5.5, 6.2], font_size=8.3)

    doc.add_heading("4.5.2 平移一致性", level=3)
    body(doc, "對已知位移 Δ，理想模型應滿足「先位移再重建」與「先重建再位移」得到相同結果。令 SΔ 表示亞像素平移算子，則平移一致性損失為：")
    eq(doc, "L_eq = ||f_θ(S_Δx) − S_Δf_θ(x)||_1.                                        (4-14)")
    body(doc, "若模型在不同影格中任意移動邊界，式（4-14）的兩條路徑會不一致。此損失可抑制與輸入位置相關的人工變形，但它只保證相對一致，不能單獨決定絕對尺寸，因此仍需與 SEM、合成幾何與原圖低頻錨定共同使用。")

    doc.add_heading("4.5.3 分階段微調與資料數量", level=3)
    fig(doc, "fig_4_7_translation_consistency_and_finetuning.png", "Figure 4-7 Translation consistency and staged fine-tuning")
    body(doc, "如圖 4-7 所示，訓練以已有基本重建能力的模型為起點，保留主要特徵提取層，只更新輸出側解碼、邊界修正與品質控制模組。這種分階段微調可把有限的更新集中在量測相關輸出，降低重新訓練整個模型造成既有影像內容能力大幅漂移的風險。")
    body(doc, "微調共進行 120 輪，學習率為 2×10⁻⁷；運動監督包含 2,400 組序列配對與最多 3,600 組位移配對，另以 1,000 組序列配對驗證。模型選擇使用固定的最終訓練狀態，後續 20 倍、50 倍、CD 與位移分析不混用不同訓練輪次。訓練損失只用來確認收斂，方法是否有效仍由第 4.6 節定義的物理量測判斷。")

    doc.add_heading("4.6 公平比較與量測程序", level=2)
    doc.add_heading("4.6.1 Raw、Wiener 與所提模型的共同規則", level=3)
    body(doc, "驗證比較三種影像：未重建的原始影像（Raw）、使用該區域局部擬合 PSF 的 Wiener 反卷積，以及所提模型輸出。Raw 提供實際拍攝基準；Wiener 代表單純使用已知 PSF 的傳統方法；所提模型則代表固定多 PSF 先驗與量測保護機制的整體效果。")
    body(doc, "三種方法使用完全相同的九宮格裁切、方向、物體空間像素尺寸、pillar 編號、Raw 鎖定中心、參考半徑、量測門檻與位移步驟。中心與步驟只從 Raw 決定，不為 Wiener 或模型重新置中，也不為某一方法排除不利影像。這項規則可避免方法輸出只是因為重新找了更合適的中心或挑了不同步驟而看起來更好。")
    body(doc, "影像顯示與數值計算分開處理。論文圖可使用同一組強度範圍方便閱讀，但不得逐張自動拉伸；所有 CD、邊緣與頻譜數值均由未額外銳化或對比強化的資料計算。頻譜使用 20log10 振幅並固定顯示 −80 至 0 dB，因此不同方法的亮度差異反映實際相對頻率能量，不是各自強化後的視覺效果。")

    doc.add_heading("4.6.2 邊緣銳利度與振鈴", level=3)
    body(doc, "邊緣銳利度以正規化灰階剖面的 10–90% edge width 表示。對每一條由背景穿越 pillar 內部的剖面，先依亮暗極性轉成由 0 到 1 的共同方向，再以內插求得 10% 與 90% 交點。兩交點在物體空間的距離為：")
    eq(doc, "w_10–90 = |x_90 − x_10|·p_obj.                                               (4-15)")
    body(doc, "w10–90 越小代表邊界過渡越陡。由於強反卷積可用過衝換取較小寬度，本研究同時檢查邊界外側是否超出局部背景—pillar 的正常灰階範圍，並將最大過衝／欠衝記為振鈴幅度。只有在 edge width 下降且振鈴沒有明顯惡化時，才解讀為有效銳利化。")

    doc.add_heading("4.6.3 CD 門檻與邊界取得", level=3)
    fig(doc, "fig_4_8_edge_and_cd_measurement.png", "Figure 4-8 Edge sharpness and critical-dimension measurement")
    body(doc, "CD 量測首先固定每顆 pillar 的 Raw 中心與參考半徑。對第 j 顆 pillar，在其局部區域估計背景代表值 Ibg,j 與結構內部代表值 Ipillar,j。比例門檻 τ 不是固定灰階數，而是由背景往 pillar 內部前進的相對比例：")
    eq(doc, "T_j(τ) = I_bg,j + τ(I_pillar,j − I_bg,j),      τ∈{0.35,0.40,…,0.65}.         (4-16)")
    body(doc, "若 pillar 為暗結構，Ipillar<Ibg，式（4-16）仍會自動形成由亮背景向暗內部的門檻；若極性相反，同一公式亦適用。因此門檻的 50% 指局部背景與 pillar 內部的中點，不是整張影像固定灰階 0.5。主要結果採 τ=0.50，並掃描 0.35–0.65 以檢查結論是否依賴單一門檻。")
    body(doc, "本研究使用兩種互補 CD 模式。水平模式將中心上下各 3 列，共 7 列剖面取平均，再找左右門檻交點；徑向模式則由固定中心向 360 個方向搜尋門檻交點，取對向半徑形成直徑並彙整。水平模式易於解釋左右邊界，徑向模式可平均多方向雜訊。兩種模式均沿用相同中心，不對重建影像個別重新置中。")
    eq(doc, "CD_optical,j = d_j(pixel)·p_obj,        e_j = CD_optical,j − CD_SEM,j.          (4-17)")
    body(doc, "第 j 顆 pillar 的像素直徑 dj 乘上倍率對應的 pobj 後得到奈米 CD，再與表 4-3 的同一顆 SEM CD 相減。此逐顆對應保留 SEM 真實製程差異；不使用九顆平均值先做整體縮放，也不以重建後影像的最佳門檻回頭調整模型。")

    doc.add_heading("4.6.4 CD 誤差指標與解讀", level=3)
    body(doc, "為避免單一統計量掩蓋問題，本研究同時使用平均絕對誤差（MAE）、絕對誤差中位數（MdAE）、均方根誤差（RMSE）、平均帶號偏差（Bias）與最大絕對誤差（MaxAE）。令 N=9，則：")
    eq(doc, "MAE=(1/N)Σ|e_j|,    RMSE=√[(1/N)Σe_j²],    Bias=(1/N)Σe_j.                    (4-18)")
    eq(doc, "MdAE=median(|e_j|),                       MaxAE=max(|e_j|).                    (4-19)")
    body(doc, "MAE 表示九顆 pillar 的平均誤差；MdAE 對單一離群值較不敏感，代表典型 pillar；RMSE 會加重較大誤差；Bias 可判斷方法是否整體高估或低估 CD；MaxAE 則顯示最差一顆是否失控。理想結果不是只讓 Bias 接近 0，因為正負誤差可能互相抵銷；應同時看到 MAE 或 MdAE 下降，且 MaxAE 不惡化。")
    body(doc, "門檻掃描的用途也不是從七個門檻中挑一個最好看的數字，而是將每個 τ、兩種 CD 模式與九顆 pillar 的結果完整保留。若某方法只在 τ=0.50 改善、其餘門檻大幅變差，表示邊界仍對門檻敏感；若多數門檻與兩種模式都呈現一致改善，才可說重建提高 CD 量測穩健性。")

    doc.add_heading("4.6.5 逐顆邊界位移追蹤", level=3)
    body(doc, "位移主軸不使用單一整張影像平移量取代九顆 pillar，而是先逐顆追蹤邊界。對每顆 pillar，於垂直半徑比例 −0.55、−0.275、0、0.275 與 0.55 取得五條水平弦。每條剖面先平滑，再依暗／亮極性決定預期梯度方向，於左右搜尋窗中找出梯度峰值。峰值以相鄰三點拋物線內插，得到亞像素左、右邊界。")
    eq(doc, "c_j,k = (x_left,j,k + x_right,j,k)/2,       c_j = median_k(c_j,k).             (4-20)")
    body(doc, "式中 k=1…5 為五條弦，cj,k 為該弦左右邊界的中點，cj 為五條弦中位數。以左右邊界中點表示位置，可降低單側邊緣模糊或振鈴造成的偏差；以五條弦中位數彙整，則可降低局部污染或單一剖面失敗的影響。相鄰影格的第 j 顆位移為 cj(t+1)−cj(t)。")
    body(doc, "除邊界追蹤外，局部二維模板匹配用作 Raw 候選步驟的資格交叉檢查，確認九宮格確實朝預期方向移動且沒有大幅錯配；正式逐顆位移仍以五弦邊界方法計算。這可把「步驟是否可用」與「每顆 pillar 位移多少」分成兩個層次。")

    doc.add_heading("4.6.6 九顆共識與空間一致性", level=3)
    body(doc, "同一平台步進理論上應使九顆 pillar 移動相近距離，因此每一步的代表位移取九顆中位數，而非平均值；中位數較不受單一追蹤失敗影響。九宮格內不一致量使用經常態尺度校正的中位數絕對偏差：")
    eq(doc, "d_t = median_j[d_t,j],       u_t = 1.4826·median_j|d_t,j−d_t|.                (4-21)")
    body(doc, "ut 越小，表示同一步在九個空間位置得到的估計越接近。係數 1.4826 使資料近似常態時的 MAD 可與標準差尺度比較，但此指標仍保留中位數對離群值的穩健性。它不代表平台絕對誤差，而是影像在同一視野內的空間一致性。")

    doc.add_heading("4.6.7 Raw 鎖定的最佳 30 步", level=3)
    fig(doc, "fig_4_9_displacement_tracking_and_selection.png", "Figure 4-9 Nine-pillar displacement tracking and Raw-locked 30-step selection")
    body(doc, "每組序列接近 50 個相鄰步，是因實驗條件有限，壓電平台可能出現錯步、跨工作階段中斷或目標離開視野，因此額外影格作為備用。正式比較固定取品質最佳的 30 步，但挑選只看 Raw：先排除跨工作階段、九顆不完整、水平位移偏離名義指令超過 25%，或垂直串擾超過名義值 25% 的候選，再依 Raw 的局部二維運動品質排序。")
    body(doc, "一旦 30 步確定，Raw、Wiener 與所提模型都使用完全相同的相鄰影格編號。若先看重建結果再各自挑 30 步，模型可能只因排除了較差輸出而得到不公平優勢；Raw 鎖定可確保挑選反映拍攝與平台品質，而不是某一重建方法的結果。四組正式序列包含 20 倍與 50 倍，以及名義 0.5 µm 與 1.0 µm 步進。")

    doc.add_heading("4.6.8 位移指標與真值界線", level=3)
    body(doc, "位移分析至少報告四項結果：30 步代表位移相對名義指令的 MAE、帶號位移標準差、九宮格內不一致量，以及垂直串擾。帶號標準差保留正負方向，可反映步間波動；若先取絕對值再算標準差，正負錯誤可能被隱藏。逐步改善比例則計算模型在同一步是否比 Raw 更接近名義值，用來確認平均改善不是由少數步驟主導。")
    body(doc, "壓電平台指令是名義參考，不是獨立位移真值。平台可能存在遲滯、非線性與實際錯步，因此即使影像估計與指令有差異，也不能全部歸因於影像方法。沒有外部編碼器或干涉儀時，本研究主要主張重建是否降低步間波動、九顆分歧與跨軸誤差；對絕對位移精度的結論保持保守。")

    doc.add_heading("4.6.9 空間變異與反卷積解讀", level=3)
    body(doc, "Wiener 反卷積假設所用 PSF 與目標區域相符。若九宮格遠離 PSF 擬合位置，或不同倍率的 PSF 來自方向錯配區域，PSF mismatch 可能把原本模糊轉成振鈴與低頻擴散，造成反卷積後看似更糊或 CD 偏差。故本研究的 Wiener 對照使用最接近目標區域的局部擬合 PSF，並保留正則化以避免過度放大高頻雜訊。")
    body(doc, "即使使用局部 PSF，傳統反卷積只解決卷積模型中的一部分問題，仍可能受到空間變異、雜訊、背景不均與非線性對比影響。因此，Wiener 若在個別條件下取得較低 CD 誤差，不代表它在整體量測上優於模型；需同時檢查邊緣振鈴、門檻穩健性、四組位移序列與九宮格內一致性。所提模型也不能只以視覺銳利度獲勝，必須在同一套量測規則下證明較穩定。")

    table(doc,
          "Table 4-9 Measurement metrics and interpretation rules",
          ["Metric", "Definition", "Desired direction", "Interpretation safeguard"],
          [
              ["10–90% edge width", "Object-space distance between normalized 10% and 90% crossings", "Lower", "Must be read together with ringing"],
              ["Ringing amplitude", "Overshoot or undershoot beyond the local intensity range", "Not increased", "Prevents artificial edge sharpening"],
              ["CD MAE / MdAE", "Per-pillar optical CD versus its own SEM CD", "Lower", "No cancellation of positive and negative errors"],
              ["CD Bias", "Mean signed per-pillar CD error", "Near zero", "Cannot be used alone because errors may cancel"],
              ["CD MaxAE", "Largest absolute per-pillar CD error", "Not increased", "Detects a locally failed pillar"],
              ["Threshold robustness", "CD behavior across 35–65% and two measurement modes", "Stable", "Prevents selecting one favorable threshold"],
              ["Command MAE", "Difference from nominal platform step", "Lower", "Command is a reference, not independent truth"],
              ["Step SD", "Standard deviation of signed 30-step motion", "Lower", "Measures temporal repeatability"],
              ["Within-grid disagreement", "1.4826×MAD of nine pillar motions", "Lower", "Measures spatial consistency"],
              ["Cross-axis motion", "Estimated orthogonal displacement", "Near zero", "Detects axis coupling or tracking failure"],
          ], [3.0, 5.3, 2.3, 4.3], font_size=7.6)

    doc.add_heading("4.7 本章方法的可驗證主張", level=2)
    body(doc, "綜合上述設計，本研究的第一項主張是：以正確 pillar 區域、真實方向與逐顆 SEM 幾何建立局部 PSF，可得到可用於該區域光學分析與傳統反卷積的物理模型。此主張需要由 PSF 形貌、剖面、半高全寬、擬合影像與實拍影像的一致性支持，而不能只展示一張經過色階強化的 PSF 圖。")
    body(doc, "第二項主張是：固定多位置 PSF、受限邊界殘差、Raw 低頻錨定與品質門控，可在單張同尺寸重建中增加可用邊界資訊，同時限制大尺度位置與幾何被任意改寫。此主張需由 10–90% edge width、振鈴、逐顆 CD 與門檻掃描共同支持。")
    body(doc, "第三項也是本研究的主要主張：加入已知位移監督與平移一致性後，重建影像應使連續位移更容易追蹤。此主張不以單一平均位移證明，而以 Raw 鎖定的相同 30 步、逐顆五弦邊界、九顆中位數、九宮格 MAD、步間標準差與跨軸運動完整呈現。如此才能區分真正的位移穩定化與單純的視覺銳化。")
    body(doc, "本章也界定不應過度延伸的結論。現有資料可驗證預定顯微系統與樣品條件下的量測功能，但尚不能取代外部設備的絕對位移真值，也不能自動證明跨設備與跨材料泛化。多 PSF 自適應策略保留為可切換延伸，現有結果則歸因於固定 25-PSF 先驗與已驗證的輸出保護路徑。")

    doc.add_heading("4.8 小結", level=2)
    body(doc, "本章由顯微成像模型出發，先定義感測器像素與物體空間的換算，再以 pillar/via 交界、中心九宮格、方向校正與 SEM 半徑轉直徑建立逐顆幾何真值。局部 PSF 以 Zernike 像差與 BADS 擬合，25 張正規化 PSF 則形成模型使用的固定光學先驗，並保留可切換的自適應候選策略。")
    body(doc, "重建模型採單張同尺寸輸出，利用編碼器—解碼器恢復邊界，並以最大 0.025 的局部殘差、σ=2.0 像素與 0.35 細節增益的 Raw 低頻錨定，以及 0.78–1.08 的品質門控限制幾何改動。訓練資料同時提供已知幾何、真實紋理、SEM 尺寸與已知位移，使模型學習的不只是影像外觀，也包含尺寸與平移一致性。")
    body(doc, "後續實驗將依本章固定的共同規則比較 Raw、Wiener 與所提模型：CD 使用同中心、同門檻與逐顆 SEM 對應；位移使用 Raw 鎖定的最佳 30 步、五弦亞像素邊界與九顆共識。藉由邊緣、振鈴、CD 誤差、門檻穩健性、步間波動與空間一致性，可判斷影像變清楚後是否真正提高量測可靠度。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)

    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(cell.text for tbl in doc.tables for row in tbl.rows for cell in row.cells)
    audit_text = (text + "\n" + table_text).replace("acA2040-120um", "")
    forbidden = re.findall(r"\b(?:V\d+|R\d+|e\d{2,4})\b|20\d{2}[-_/]\d{1,2}", audit_text, flags=re.I)
    print(f"Saved: {OUT}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Inline figures: {len(doc.inline_shapes)}")
    print(f"Forbidden internal labels found: {forbidden}")


if __name__ == "__main__":
    build()
