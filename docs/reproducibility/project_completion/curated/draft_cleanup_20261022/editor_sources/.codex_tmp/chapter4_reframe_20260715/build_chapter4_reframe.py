from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\第四章_研究方法_量測導向重整稿.docx")
FIG_DIR = Path(r"D:\GMM_2_Final\Thesis_MOSR_Evidence\Final_Data_Figures\05_CHAPTER4_METHOD_REFRAME")

FONT_EA = "標楷體"
FONT_LATIN = "Times New Roman"
BLUE = "2E6FAE"
TEAL = "268C83"
PURPLE = "7355A8"
ORANGE = "E9822B"
LIGHT_BLUE = "E9F1F8"
LIGHT_GRAY = "F2F3F5"
GRID = "000000"


def set_run_font(run, size=12, bold=None, italic=None, color=None, latin=FONT_LATIN, east_asia=FONT_EA):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, 10)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EA)
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.85)
    normal.paragraph_format.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT_LATIN
    h1._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EA)
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string("1A1F24")
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(16)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT_LATIN
    h2._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EA)
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string("1A1F24")
    h2.paragraph_format.first_line_indent = Cm(0)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.line_spacing = 1.5
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    h3.font.name = FONT_LATIN
    h3._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EA)
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string("1A1F24")
    h3.paragraph_format.first_line_indent = Cm(0)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(2)
    h3.paragraph_format.keep_with_next = True

    if "Equation" not in [s.name for s in doc.styles]:
        eq = doc.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        eq = doc.styles["Equation"]
    eq.font.name = FONT_LATIN
    eq._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EA)
    eq.font.size = Pt(11)
    eq.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq.paragraph_format.first_line_indent = Cm(0)
    eq.paragraph_format.space_before = Pt(4)
    eq.paragraph_format.space_after = Pt(4)
    eq.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    caption.font.name = FONT_LATIN
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EA)
    caption.font.size = Pt(10)
    caption.font.color.rgb = RGBColor.from_string("1A1F24")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = False

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("第4章　研究方法")
    set_run_font(hr, 10)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_body(doc, text, first_indent=True, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if not first_indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, 12, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, 12)
    else:
        r = p.add_run(text)
        set_run_font(r, 12)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph(style="Equation")
    parts = re.split(r"(_[A-Za-z0-9]+)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("_"):
            r = p.add_run(part[1:])
            set_run_font(r, 9.5)
            r.font.subscript = True
        else:
            r = p.add_run(part)
            set_run_font(r, 11)
    return p


def add_figure(doc, filename, caption, width_cm=14.4):
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cp = doc.add_paragraph(style="Caption")
    r = cp.add_run(caption)
    set_run_font(r, 10)
    return cp


def add_table(doc, caption, headers, rows, widths=None, header_fill="F2F2F2", font_size=9.4):
    cp = doc.add_paragraph(style="Caption")
    cp.paragraph_format.keep_with_next = True
    r = cp.add_run(caption)
    set_run_font(r, 10)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(str(text))
        set_run_font(rr, font_size, bold=True, color="000000", east_asia=FONT_EA)
        if widths:
            cell.width = Cm(widths[i])

    for ridx, row in enumerate(rows):
        tr = table.add_row()
        prevent_row_split(tr)
        for i, value in enumerate(row):
            cell = tr.cells[i]
            set_cell_shading(cell, "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            rr = p.add_run(str(value))
            set_run_font(rr, font_size, east_asia=FONT_EA)
            if widths:
                cell.width = Cm(widths[i])
    for row in table.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.first_line_indent = Cm(0)
    return table


def add_section_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build():
    doc = Document()
    configure_document(doc)

    doc.add_heading("第4章　研究方法", level=1)

    add_body(doc, "明場顯微影像的邊界品質會同時受到物鏡數值孔徑、視野像差、感測器取樣與影像雜訊影響。若只以視覺清晰度判斷影像復原效果，過度銳化、振鈴或對比重新分配都可能使物體邊界看似更清楚，卻讓臨界尺寸與位移量測產生偏差。因此，本研究以量測需求為出發點，結合點擴散函數校正、物理先驗與深度學習，建立一套兼顧影像銳利度與幾何位置保真的重建方法。")
    add_body(doc, "本章所稱量測導向超解析（Measurement-Oriented Super-Resolution, MOSR），是指重建影像能提供較穩定的臨界尺寸與亞像素位移估測，而不是僅依同尺寸輸出便宣稱突破物鏡的繞射極限。模型在使用時只需輸入單張影像，並輸出相同尺寸的重建影像；多張影像所提供的位移資訊則用於訓練與驗證，使模型學會在影像變清楚的同時保留原有幾何位置。")

    doc.add_heading("4.1 研究方法概述", level=2)
    doc.add_heading("4.1.1 量測導向的問題定義", level=3)
    add_body(doc, "令原始顯微影像為 x，影像重建模型為 fθ，後端量測方法為 M。一般影像復原多著重於重建影像與參考影像之間的像素差異；本研究則進一步檢查重建後的量測值，是否能在不同影像品質、視野位置與連續位移條件下維持一致。其關係可表示為：")
    add_equation(doc, "ŷ = fθ(x),                    m̂ = M(ŷ).                                      (4-1)")
    add_body(doc, "其中 m̂ 可代表柱狀結構（pillar）的臨界尺寸、相鄰影像之間的位移，或同一步進下九顆 pillar 所估得的局部位移。由此可知，影像只是量測流程的中間結果；模型是否有效，應同時由邊界是否變窄、尺寸是否接近掃描式電子顯微鏡（SEM）結果，以及位移估測是否穩定來判斷。")
    add_body(doc, "為避免神經網路任意改變物體位置，本研究採用三項設計原則。第一，以多個視野位置的點擴散函數建立光學先驗，降低單一位置校正結果的限制。第二，以原始影像保留低頻位置，只允許模型對邊界附近的高頻細節做小幅修正。第三，將已知位移與連續影像配對加入訓練，使相同物體經過位移前後仍維持一致的形狀與尺度。")

    doc.add_heading("4.1.2 整體研究流程", level=3)
    add_body(doc, "本研究流程如圖 4-1 所示，可分為光學校正、訓練資料建立、模型訓練、單張影像重建與量測評估五個階段。光學校正先確認樣品結構、拍攝方向與中心九顆 pillar，再以 SEM 尺寸建立幾何真值並擬合局部點擴散函數。訓練階段將真實影像、合成幾何、SEM 尺寸與已知位移配對共同使用；重建階段則以單張影像完成同尺寸輸出。最後以完全相同的裁切範圍與量測規則，比較原始影像、維納反卷積與所提模型的邊緣、臨界尺寸及位移結果。")
    add_figure(doc, "fig_4_1_overall_method_workflow.png", "Figure 4-1 Measurement-oriented reconstruction workflow", 14.4)
    add_table(doc,
        "Table 4-1 Research objectives, method design, and evaluation",
        ["Research objective", "Method design", "Evaluation"],
        [
            ["Account for field-dependent blur", "Fusion of normalized PSFs from multiple field positions", "PSF profile, MTF, and matched Wiener restoration"],
            ["Preserve geometric position", "Raw-image low-frequency anchor and bounded boundary correction", "Edge width, ringing, and SEM-referenced CD"],
            ["Stabilize displacement measurement", "Exact-shift pairs and motion-consistency supervision", "Command consistency, step spread, and nine-pillar agreement"],
            ["Maintain robustness across conditions", "Input-quality control and identical measurement rules", "20×/50× comparison and threshold sweep"],
        ], [4.1, 5.8, 5.0], font_size=8.8)
    add_body(doc, "本研究中的點擴散函數擬合、維納反卷積與神經重建各有不同用途。局部擬合結果用來描述特定拍攝位置的光學模糊，並作為傳統反卷積的物理對照；神經模型則利用多位置點擴散函數形成較保守的光學先驗，再由資料學習邊界修正。三種影像只有在區域、方向、像素換算與量測方式一致時，才進行數值比較。")

    doc.add_heading("4.2 點擴散函數校正與多 PSF 融合", level=2)
    doc.add_heading("4.2.1 中心九顆 pillar 與 SEM 幾何真值", level=3)
    add_body(doc, "局部 PSF 擬合的區域不是影像中任意九個圓形結構，而是 pillar 與 via 兩大區域交界處、最靠近樣品中心的 pillar 3×3 九宮格。選區時先依整片樣品版圖判定 pillar/via，再依交界與中心位置確認九顆目標；其後才進行影像裁切與方向校正。如此可避免把 via 當成 pillar，或因倍率改變與樣品旋轉而錯配 SEM 尺寸。")
    add_body(doc, "兩種倍率的拍攝方向並不相同。50 倍影像沿用原始方向；20 倍影像先做垂直翻轉，使九顆 pillar 的排列與 SEM 真實方向一致。此步驟在配置逐顆尺寸前完成，且 pillar 與 via 的判斷依據為樣品位置與結構區域，而不是影像的明暗極性。")
    add_body(doc, "SEM 提供的九個數值為半徑而非直徑，因此建立幾何真值與比較光學臨界尺寸時，均先將半徑乘以 2。九顆 pillar 的平均 SEM 臨界尺寸為 2629.74 nm，樣本標準差為 24.86 nm，範圍為 2598.75 至 2676.80 nm。表 4-2 保留每一顆 pillar 的實際差異，而不是以單一平均圓取代，以免將真實製程差異誤判為光學模糊。")
    sem_rows = []
    radii = [
        ("R1C1", 1310.107187), ("R1C2", 1310.008186), ("R1C3", 1321.157790),
        ("R2C1", 1310.214279), ("R2C2", 1304.359769), ("R2C3", 1329.306897),
        ("R3C1", 1310.889999), ("R3C2", 1299.375827), ("R3C3", 1338.398285),
    ]
    for pos, radius in radii:
        sem_rows.append([pos, f"{radius:.6f}", f"{2*radius:.6f}"])
    add_table(doc, "Table 4-2 SEM-derived ground-truth dimensions",
              ["Pillar position", "SEM radius (nm)", "Ground-truth CD (nm)"],
              sem_rows, [4.0, 5.0, 5.0], font_size=9.2)

    doc.add_heading("4.2.2 參數化 PSF 與最佳化擬合", level=3)
    add_body(doc, "擬合時先以澤尼克多項式描述光學像差，再將幾何真值與候選點擴散函數進行二維卷積，得到預測影像。為避免傾斜像差吸收原本應由影像對位處理的位移，流程先以亞像素搜尋完成預對位，並將 Z2 與 Z3 傾斜項固定為 0；實際搜尋 Z4 至 Z15 的像差係數與一個整體振幅係數。每次建立的點擴散函數均以總能量為 1 進行正規化。")
    add_equation(doc, "I_pred(θ) = I_GT * h(θ),                    Σu,v h(u,v;θ) = 1.                 (4-2)")
    add_body(doc, "對中心九顆 pillar 的有效區域，貝葉斯自適應直接搜尋法（Bayesian Adaptive Direct Search, BADS）的目標函數同時衡量灰階強度與邊緣梯度。梯度在比較前分別正規化，再提高其權重，使最佳化過程更重視輪廓寬度與邊界方向：")
    add_equation(doc, "J(θ) = MSE_mask(I_obs, I_pred) + 8×10⁴ MSE_mask[N(∇I_obs), N(∇I_pred)].       (4-3)")
    add_body(doc, "式中第一項維持整體灰階與對比關係，第二項使 pillar 邊界的寬度與方向接近實拍影像。BADS 不需計算目標函數的解析梯度，適合用來同時搜尋多個像差係數；預對位、固定傾斜項，以及以前一位置結果作為下一次搜尋起點，則可降低位置偏差與相鄰視野間的參數跳動。")

    doc.add_heading("4.2.3 固定多 PSF 融合與物理先驗", level=3)
    add_body(doc, "由於點擴散函數會隨視野位置改變，若只使用單一位置的校正結果，可能使模型過度適應局部模糊。本研究因此使用 5×5 視野網格中的 25 張點擴散函數，每張大小為 101×101。各張點擴散函數先分別正規化，使其能量總和為 1，再以等權平均形成固定融合結果：")
    add_equation(doc, "h̄ = (1/25) Σₖ₌₁²⁵ hₖ,                    Σu,v h̄(u,v) = 1.                  (4-4)")
    add_body(doc, "固定融合 PSF 進入保守式 Wiener 型前處理。若 X 為輸入影像頻譜、H̄ 為 h̄ 的頻率響應，則先計算逆濾波估計，再以小比例 α 與原圖混合：")
    add_equation(doc, "x_phys = x + α{F⁻¹[H̄*X/(|H̄|²+γ)] − x}.                                  (4-5)")
    add_body(doc, "本研究設定 α=0.10、γ=10，使物理前處理只提供溫和的去模糊方向，而不將強反卷積結果直接當成最終輸出。局部擬合的 20 倍與 50 倍點擴散函數主要用於特定位置的光學分析與維納反卷積對照；上述 25 張點擴散函數則用於模型內部的多位置物理先驗，兩者用途不同。")

    doc.add_heading("4.2.4 可切換的自適應多 PSF 策略", level=3)
    add_body(doc, "固定平均能提高整體穩定性，但無法依單張影像調整不同模糊型態。為保留更細緻的空間適應能力，本研究在固定融合之外，另設計可切換的四候選 PSF 加權策略。此策略從多位置點擴散函數中選出四個代表性候選，分別對中間特徵進行保守補償，再由輸入特徵預測四個權重。令第 k 個候選的補償特徵為 Dk、權重為 wk，則融合結果可表示為：")
    add_equation(doc, "D_mix = Σₖ₌₁⁴ wₖDₖ,     Σₖwₖ = 1;       F_out = F + β(D_mix − F).          (4-6)")
    add_body(doc, "β 用來限制自適應補償的最大影響，因此可在固定融合的基礎上逐步加入位置適應能力，而不讓候選 PSF 直接改寫整張影像。此機制屬於完整架構中的可切換策略；實際啟用時需使用相符的訓練權重，並以相同量測流程單獨比較其影響。")

    doc.add_heading("4.3 量測導向影像重建網路", level=2)
    add_body(doc, "所提網路的資料流如圖 4-2 所示。輸入影像先經過固定多 PSF 物理前處理，再由 U 形編碼器—解碼器擷取多尺度特徵，最後依序進行局部邊界修正、原圖低頻位置錨定與影像品質調節。各模組的共同目的，是讓模型恢復可量測的邊界細節，同時限制其對物體大尺度位置與平均亮度的改動。")
    add_figure(doc, "fig_4_2_measurement_oriented_reconstruction_path.png", "Figure 4-2 Measurement-oriented reconstruction network and output safeguards", 14.4)
    doc.add_heading("4.3.1 同尺寸編碼器—解碼器", level=3)
    add_body(doc, "網路輸入與輸出皆為單通道影像，且維持相同的 H×W 尺寸。編碼端逐層取得不同尺度的邊界與紋理特徵，解碼端利用跳接結合淺層的位置資訊與深層特徵。上採樣採用雙線性插值、平滑與卷積的組合，以降低小位移經過上採樣後產生棋盤格或相位跳動的機會。此架構不是增加輸出像素數，而是在原取樣網格上恢復與量測相關的邊界資訊。")

    doc.add_heading("4.3.2 方向感知且受限的邊界修正", level=3)
    add_body(doc, "主幹輸出後加入局部邊界修正器。修正器同時讀取解碼特徵、目前輸出，以及由 Sobel X/Y 梯度與梯度幅值構成的方向資訊，再預測一個局部殘差與門控值。本研究將殘差最大幅度限制為 0.025（以 0–1 灰階表示），使模型只能做小幅、局部且與邊界方向相關的修正。這個上限可降低模型為追求視覺銳利度而移動整個輪廓的風險。")

    doc.add_heading("4.3.3 原圖低頻位置錨定", level=3)
    add_body(doc, "邊界修正後的影像不直接作為最終輸出，而先通過原圖位置錨定。模型以 σ=2.0 像素的高斯濾波將預測影像分成低頻與高頻，只取 0.35 倍的預測高頻加回原始影像，最後再校正全圖平均值：")
    add_equation(doc, "y_anchor = x + 0.35[ypred − Gσ(ypred)],      σ = 2.0 pixels.                 (4-7)")
    add_body(doc, "由於輸出低頻基底直接來自 x，物體的大尺度位置、背景起伏與平均亮度較不容易被神經網路重新解釋；模型主要貢獻被限制在邊界附近的細節。此設計是本研究「量測安全」概念的核心，亦使重建後位移改善可合理解讀為邊界追蹤變穩定，而不是整段座標被任意縮放。")

    doc.add_heading("4.3.4 影像品質調節", level=3)
    add_body(doc, "不同倍率與曝光條件會造成影像訊號與雜訊程度不同，因此模型會依輸入品質調節最終修正量。首先以梯度幅值最高的 5% 像素估計有效邊界訊號 S，再以平坦區域相對 5×5 局部平均的差值估計雜訊 N，形成介於 0 與 1 之間的品質分數 q：")
    add_equation(doc, "q = S/(S + 3N + ε),       s = 0.78 + (1.08 − 0.78)q.                         (4-8)")
    add_body(doc, "最終輸出為 y=x+s(yanchor−x)。當影像品質較差時，s 接近 0.78，模型修正較保守；品質較好時，s 最多增加至 1.08，使模型保留可用高頻。門控只調整原圖與重建差值的強度，不直接改變空間座標，因此適合用於不同倍率與訊雜條件。")

    add_table(doc,
        "Table 4-3 Main configuration of the proposed reconstruction model",
        ["Component", "Setting", "Measurement purpose"],
        [
            ["PSF prior", "25 × 101×101 PSFs; normalized mean", "Reduce dependence on a single field position"],
            ["Physical residual", "α = 0.10; γ = 10", "Provide conservative deblurring direction"],
            ["Boundary correction", "Orientation inputs; maximum residual = 0.025", "Sharpen locally without moving the whole object"],
            ["Raw anchor", "Detail mode; σ = 2.0; gain = 0.35", "Preserve low-frequency position and mean intensity"],
            ["Quality gate", "8 hidden channels; scale = 0.78–1.08", "Adapt correction strength to input quality"],
            ["Inference size", "One H×W image → one H×W image", "Single-frame deployment without frame stacking"],
        ], [3.2, 5.2, 6.0], font_size=8.8)

    doc.add_heading("4.4 訓練資料與監督設計", level=2)
    add_figure(doc, "fig_4_3_supervision_and_evidence_logic.png", "Figure 4-3 Training supervision and evaluation design", 14.4)
    doc.add_heading("4.4.1 資料組成與功能分群", level=3)
    add_body(doc, "本研究共使用 11,254 組影像配對。為清楚說明各類資料的用途，將其分成五個功能群：具已知幾何的合成資料、SEM 幾何錨點、保留真實顯微紋理的重建參考、具有已知位移的自位移或實際序列配對，以及跨倍率與形狀的穩定性資料。這五類資料分別提供形狀、尺寸、真實影像外觀、移動一致性與保守修正的監督。")
    add_table(doc,
        "Table 4-4 Training-pair composition by supervision source",
        ["Label source", "Pairs", "Share", "Primary function"],
        [
            ["Exact-shift pairs on real textures", "3,468", "30.82%", "Translation consistency on real image texture"],
            ["Synthetic motion pairs", "1,520", "13.51%", "Controlled displacement supervision"],
            ["Edge-refined real-image pairs", "1,470", "13.06%", "Boundary restoration on real images"],
            ["Real-sequence motion pairs", "1,176", "10.45%", "Motion geometry from captured sequences"],
            ["Stability pairs on real textures", "800", "7.11%", "Shape stability under texture variation"],
            ["Synthetic geometry pairs", "760", "6.75%", "Absolute shape and boundary anchor"],
            ["Step-preservation pairs", "760", "6.75%", "Prevent displacement-scale distortion"],
            ["General shape pairs", "660", "5.86%", "Broaden shape coverage"],
            ["SEM-referenced pairs", "640", "5.69%", "Independent geometric anchor"],
        ], [5.2, 1.7, 1.7, 5.8], font_size=8.0)

    doc.add_heading("4.4.2 軟標籤與邊界監督", level=3)
    add_body(doc, "若直接以硬式二值邊界作為唯一標籤，模型容易學到階梯狀或鋸齒狀輪廓。為提供連續的亞像素監督，本研究在各類資料中選取 12 組、合計 180 組樣本，於距離邊界 2 像素的窄帶內施加 σ=0.75 像素的高斯軟化。此操作只調整邊界附近，不改變物體中心與大面積背景，因此目的不是模糊影像，而是讓邊界位置能在像素之間平順變化。")
    add_body(doc, "每組資料另包含邊界遮罩、中心線與距離圖，使模型分別學習輪廓位置、幾何骨架與離邊界的距離。SEM 參考資料雖占總資料的 5.69%，但它提供獨立的幾何尺度；其餘真實影像參考則主要讓模型接觸實際的相機雜訊、照明與樣品紋理。兩者用途不同，因此不能將全部影像配對都視為具有獨立量測真值。")

    doc.add_heading("4.4.3 序列分組與資料使用範圍", level=3)
    add_body(doc, "訓練與驗證資料依原始影像序列、倍率與形狀種類分組，避免相鄰影格或同一物體的高度相似裁切同時出現在兩個集合中。由於資料集包含與後續評估相同拍攝系統及樣品條件的影像，因此目前結果主要代表預定量測條件下的功能驗證。若要進一步證明跨樣品、跨照明或跨設備的泛化能力，仍需使用未參與資料建立的樣品，在固定模型參數下另行測試。")

    doc.add_heading("4.5 訓練目標與訓練策略", level=2)
    doc.add_heading("4.5.1 四組功能性損失", level=3)
    add_body(doc, "本研究依用途將訓練目標整理為四組：影像內容與低頻保真、邊界幾何、位移一致性，以及保守性限制。總損失可表示為：")
    add_equation(doc, "L_total = L_content + L_boundary + L_motion + L_guard.                           (4-9)")
    add_body(doc, "內容損失維持重建影像的灰階、結構相似度、梯度與頻帶能量；邊界損失利用遮罩、中心線與距離圖約束輪廓位置；位移損失要求相鄰影像中的同一物體維持中心、半徑與厚度；保守性損失則限制低頻內容與不需修正區域的改變。如此可避免模型只追求單張影像變清楚，卻在連續影像中改變物體尺寸或位置。")

    doc.add_heading("4.5.2 平移一致性與位移配對", level=3)
    add_body(doc, "對具有精確位移 Δ 的影像，模型需滿足先位移再重建與先重建再位移的結果相近：")
    add_equation(doc, "L_eq = ‖fθ(SΔx) − SΔfθ(x)‖₁.                                                     (4-10)")
    add_body(doc, "式 4-10 只保證兩條處理路徑的一致性，因此仍需與幾何真值、邊界及低頻損失共同使用。實際訓練另加入 2,400 組序列配對與最多 3,600 組位移配對，並以 1,000 組序列配對進行驗證。自位移資料是在真實顯微紋理上施加已知的亞像素位移，因此位移量具有明確標籤，同時保留相機與樣品的真實影像特性。")

    doc.add_heading("4.5.3 分階段微調", level=3)
    add_body(doc, "模型不是由隨機參數重新開始訓練，而是以已具備基本重建能力的預訓練模型為起點。微調時保留大部分特徵提取層，只更新影像品質調節、邊界修正、最後一段解碼與輸出端，使模型的改變集中在量測相關的邊界與輸出控制。微調共進行 120 輪，學習率設定為 2×10⁻⁷，正式分析一律使用同一組最終模型權重，避免不同圖表混用不同訓練階段。")
    add_body(doc, "訓練與驗證採固定的群組切分，並依資料功能、倍率與形狀種類維持比例平衡。模型選擇同時觀察重建、梯度與邊緣方向損失，但這些訓練數值只用來確認收斂；方法是否真正有效，仍以臨界尺寸與位移量測結果判斷。")

    doc.add_heading("4.6 量測與驗證方法", level=2)
    doc.add_heading("4.6.1 三種影像的同條件比較", level=3)
    add_body(doc, "驗證時比較原始影像、使用局部擬合 PSF 的維納反卷積影像，以及所提模型的重建影像。三者使用相同裁切範圍、pillar 編號、方向校正、物體空間像素換算與量測方法。顯示影像可使用共同的對比範圍協助觀察，但所有數值均由未額外強化的影像計算。頻譜圖統一採用 20log10 振幅，並固定顯示範圍為 −80 至 0 dB，避免每張圖各自拉伸後造成高頻改善的錯覺。")

    doc.add_heading("4.6.2 邊緣銳利度與振鈴", level=3)
    add_body(doc, "每顆 pillar 先依局部亮暗極性取得由背景到結構內部的正規化邊緣剖面，再計算由 10% 上升至 90% 所需的物體空間距離，記為 10–90% edge width。此數值越小代表轉換越陡。由於單純反卷積可能以過衝換取窄邊緣，另計算輪廓兩側超出局部亮暗範圍的振鈴幅度；只有邊緣變窄且振鈴未明顯上升時，才可解讀為有效銳利化。")

    doc.add_heading("4.6.3 臨界尺寸量測與門檻定義", level=3)
    add_body(doc, "CD 量測先以相鄰 pillar 的週期位置建立九個局部區域，再依每顆的亮暗極性定義局部門檻。門檻比例 τ 表示由背景代表值向 pillar 內部代表值前進的比例：")
    add_equation(doc, "T(τ) = I_bg + τ(I_pillar − I_bg),              τ ∈ {0.35,0.40,…,0.65}.       (4-11)")
    add_body(doc, "主要結果使用 τ=0.50，並同時以水平剖面與多方向徑向平均求得邊界直徑。水平剖面可直接對應一維臨界尺寸，徑向平均則可降低單一方向雜訊。量測結果轉為奈米後，逐顆與表 4-2 的 SEM 臨界尺寸比較，並計算平均絕對誤差、絕對誤差中位數、均方根誤差、平均帶號偏差與最大絕對誤差。再彙整 τ=0.35 至 0.65 的結果，用來檢查方法是否只在單一門檻下偶然較好。")

    doc.add_heading("4.6.4 位移序列與固定 30 步", level=3)
    add_body(doc, "位移驗證包含 50 倍與 20 倍物鏡，以及名義 0.5 µm 與 1.0 µm 步進，共四組序列。每組拍攝接近 50 張，是為壓電平台可能出現的異常步進保留備用影像；正式比較固定使用品質最佳的 30 個相鄰步。步數挑選只依原始影像進行：先排除跨工作階段、九顆邊界不完整、水平位移偏離名義值超過 25%，或垂直串擾超過名義值 25% 的候選，再依原始影像的運動品質排序。維納反卷積與神經重建均使用完全相同的 30 步，避免事後挑選對某一方法較有利的影像。")
    add_body(doc, "對每一步 i，九顆 pillar 分別得到水平位移 d(i,j)。該步共識位移取九顆中位數；九宮格內不一致量則以中位數絕對偏差的常態尺度估計：")
    add_equation(doc, "d̄ᵢ = medianⱼ[d(i,j)],       uᵢ = 1.4826·medianⱼ|d(i,j)−d̄ᵢ|.                  (4-12)")
    add_body(doc, "uᵢ 越小，代表同一平台步進在九顆 pillar 上得到的估計越一致。另計算各步與名義指令的平均絕對偏差、30 步的帶號步階標準差、垂直串擾與逐步改善比例。壓電平台指令只提供名義參考；在沒有獨立編碼器或干涉儀的情況下，不能將指令值視為完全無誤的位移真值。因此，本研究主要判斷重複性與內部一致性是否改善，而不宣稱完成平台的絕對精度校正。")

    doc.add_heading("4.6.5 空間變異與條件穩健性", level=3)
    add_body(doc, "為檢查點擴散函數的空間變異是否影響反卷積，記錄每個九宮格中心與 PSF 校正位置之間的感測器距離及物體空間距離，並分成近、中、遠三組。維納反卷積的名義指令偏差與九宮格不一致量，會進一步檢查是否隨距離增加而惡化。所提模型則同時比較兩種倍率、兩種步進、不同門檻與逐步結果，避免只用單一平均值判斷穩定性。")
    add_table(doc,
        "Table 4-5 Measurement metrics used in the evaluation",
        ["Metric", "Definition", "Desired direction", "Primary risk controlled"],
        [
            ["10–90% edge width", "Distance between normalized 10% and 90% crossings", "Lower", "Blurred transition"],
            ["Ringing amplitude", "Overshoot/undershoot outside local intensity range", "Not increased", "Artificial sharpening"],
            ["CD MAE / median AE", "Per-pillar optical CD versus SEM CD", "Lower", "Systematic or typical size error"],
            ["Threshold robustness", "Pooled error across τ=0.35–0.65 and two modes", "Lower and stable", "Cherry-picking a threshold"],
            ["Command MAE", "Mean absolute difference from nominal PI step", "Lower", "Command inconsistency"],
            ["Step SD", "Standard deviation of signed 30-step displacement", "Lower", "Temporal instability"],
            ["Within-grid disagreement", "1.4826×MAD of nine pillar displacements per step", "Lower", "Spatial inconsistency"],
            ["Cross-axis motion", "Estimated vertical motion in a horizontal sequence", "Near zero", "Axis coupling or tracking error"],
        ], [3.2, 5.2, 2.2, 3.8], font_size=8.2)

    doc.add_heading("4.7 小結", level=2)
    add_body(doc, "本章建立了一套以量測保真為核心的影像重建方法。光學校正以中心九顆 pillar、SEM 半徑轉換與參數化 PSF 擬合建立可追溯的幾何基準；多位置 PSF 先形成固定物理先驗，並保留可切換的自適應加權策略。神經網路再利用受限邊界修正、原圖低頻位置錨定與影像品質調節，控制可改動的影像內容。")
    add_body(doc, "訓練資料同時包含 SEM 幾何、合成真值、真實顯微紋理與已知位移配對，使模型不只學習影像外觀，也學習尺寸與平移一致性。後續實驗將以相同區域與相同量測規則，比較原始影像、維納反卷積與所提模型在邊緣、臨界尺寸、門檻穩健性及連續位移上的表現，藉此判斷影像變清楚後是否真正提升量測可靠度。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
