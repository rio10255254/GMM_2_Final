from __future__ import annotations

import io
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageChops
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v16_修訂版.docx")
OUTPUT = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v17_修訂版.docx")
WORK = Path(r"D:\GMM_2_Final\PSF_Fitting\.codex_docx_work")
EARLY = WORK / "v17_chinese_figures" / "early"
WIENER = WORK / "v17_chinese_figures" / "wiener" / "figures"
R63 = WORK / "v17_chinese_figures" / "r63"
CROSS = WORK / "v17_chinese_figures" / "cross_morphology"
PREPARED = WORK / "v17_media_prepared"


def set_paragraph_text(paragraph, text: str) -> None:
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = OxmlElement("w:r")
    if first_rpr is not None:
        run.append(first_rpr)
    for no, part in enumerate(text.split("\n")):
        if no:
            run.append(OxmlElement("w:br"))
        node = OxmlElement("w:t")
        if part.startswith(" ") or part.endswith(" "):
            node.set(qn("xml:space"), "preserve")
        node.text = part
        run.append(node)
    paragraph._p.append(run)


def replace_text_nodes(paragraph, old: str, new: str) -> None:
    nodes = paragraph._p.xpath(".//w:t")
    full = "".join(node.text or "" for node in nodes)
    start = full.find(old)
    if start < 0:
        raise AssertionError(f"Missing caption text: {old!r} in {full!r}")
    end = start + len(old)
    positions = []
    cursor = 0
    for node in nodes:
        length = len(node.text or "")
        positions.append((cursor, cursor + length, node))
        cursor += length
    first_index = next(i for i, (a, b, _) in enumerate(positions) if a <= start < b)
    last_index = next(i for i, (a, b, _) in enumerate(positions) if a < end <= b)
    first_a, _, first = positions[first_index]
    last_a, _, last = positions[last_index]
    before = (first.text or "")[: start - first_a]
    after = (last.text or "")[end - last_a :]
    if first_index == last_index:
        first.text = before + new + after
    else:
        first.text = before + new
        for i in range(first_index + 1, last_index):
            positions[i][2].text = ""
        last.text = after


def set_cell_text(cell, text: str, bold: bool | None = None) -> None:
    template_ppr = deepcopy(cell.paragraphs[0]._p.pPr) if cell.paragraphs and cell.paragraphs[0]._p.pPr is not None else None
    template_rpr = None
    for p in cell.paragraphs:
        for run in p.runs:
            if run._r.rPr is not None:
                template_rpr = deepcopy(run._r.rPr)
                break
        if template_rpr is not None:
            break
    tc_pr = cell._tc.tcPr
    for child in list(cell._tc):
        if child is not tc_pr:
            cell._tc.remove(child)
    p = OxmlElement("w:p")
    if template_ppr is not None:
        p.append(template_ppr)
    p_pr = p.get_or_add_pPr()
    jc = p_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        p_pr.append(jc)
    jc.set(qn("w:val"), "center")
    run = OxmlElement("w:r")
    if template_rpr is not None:
        run.append(template_rpr)
    if bold is not None:
        r_pr = run.get_or_add_rPr()
        b_node = r_pr.find(qn("w:b"))
        if b_node is None:
            b_node = OxmlElement("w:b")
            r_pr.append(b_node)
        b_node.set(qn("w:val"), "1" if bold else "0")
    for no, part in enumerate(text.split("\n")):
        if no:
            run.append(OxmlElement("w:br"))
        node = OxmlElement("w:t")
        node.text = part
        run.append(node)
    p.append(run)
    cell._tc.append(p)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_table_widths(table, widths: list[int]) -> None:
    grid = table._tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    if len(cols) != len(widths):
        raise AssertionError("Table grid column count does not match width list")
    for col, width in zip(cols, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        if len(row.cells) != len(widths):
            continue
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def trim_and_fit(new_path: Path, target_size: tuple[int, int]) -> bytes:
    image = Image.open(new_path).convert("RGB")
    background = Image.new("RGB", image.size, "white")
    diff = ImageChops.difference(image, background).convert("L")
    mask = diff.point(lambda p: 255 if p > 7 else 0)
    bbox = mask.getbbox()
    if bbox:
        margin = max(12, round(min(image.size) * 0.008))
        left = max(0, bbox[0] - margin)
        top = max(0, bbox[1] - margin)
        right = min(image.width, bbox[2] + margin)
        bottom = min(image.height, bbox[3] + margin)
        image = image.crop((left, top, right, bottom))
    canvas = Image.new("RGB", target_size, "white")
    scale = min(target_size[0] / image.width, target_size[1] / image.height)
    resized = image.resize(
        (max(1, round(image.width * scale)), max(1, round(image.height * scale))),
        Image.Resampling.LANCZOS,
    )
    canvas.paste(resized, ((target_size[0] - resized.width) // 2, (target_size[1] - resized.height) // 2))
    stream = io.BytesIO()
    canvas.save(stream, format="PNG", optimize=True)
    return stream.getvalue()


def patch_media(docx_path: Path, replacements: dict[str, Path]) -> None:
    PREPARED.mkdir(parents=True, exist_ok=True)
    temp = docx_path.with_suffix(".media-patched.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(temp, "w") as zout:
        names = set(zin.namelist())
        missing = sorted(set(replacements) - names)
        if missing:
            raise AssertionError(f"Missing media parts after docx save: {missing}")
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in replacements:
                with Image.open(io.BytesIO(data)) as original:
                    target_size = original.size
                if item.filename == "word/media/image98.png":
                    # Figure 6-53 originally used a cropped source rectangle.
                    # The crop is removed below, so match the media canvas to
                    # the existing Word picture frame to avoid distortion.
                    target_size = (1530, 977)
                data = trim_and_fit(replacements[item.filename], target_size)
                (PREPARED / Path(item.filename).name).write_bytes(data)
            zout.writestr(item, data)
    temp.replace(docx_path)


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    replacement_paragraphs = {
        21: "實驗以20倍與50倍物鏡拍攝HS-100MG樣品，並在相同九顆pillar、相同物體空間尺度及相同後端量測規則下比較Raw、Wiener與MOSR。50倍影像的徑向10–90%表觀邊緣過渡寬度中位數由620.1 nm降至Wiener的467.7 nm與MOSR的508.5 nm；20倍則由802.2 nm降至493.1 nm與246.1 nm。跨越35%至65%閾值及兩種剖面定義後，三種方法的CD絕對誤差中位數分別為95.3、123.7與76.1 nm。PI閉迴路平台分別執行名目0.5 µm與1.0 µm逐步位移，並以平台實際回授位置作為位移參考；四組序列共120個配對步中，Raw、Wiener與MOSR相對平台實際回授位移的平均絕對偏差分別為9.54、7.46與6.74 nm，九顆pillar內部位移離散度則為17.63、7.92與6.47 nm。",
        27: "Experiments were conducted on an HS-100MG sample using 20× and 50× objectives. Raw images, Wiener deconvolution, and MOSR were compared using the same nine pillars, object-space scale, and downstream measurement rules. For the 50× images, the median radial 10–90% apparent edge-transition width decreased from 620.1 nm in Raw to 467.7 nm with Wiener and 508.5 nm with MOSR; for the 20× images, it decreased from 802.2 nm to 493.1 nm and 246.1 nm, respectively. Across thresholds from 35% to 65% and two profile definitions, the median absolute CD errors were 95.3, 123.7, and 76.1 nm. The PI closed-loop stage executed nominal 0.5 µm and 1.0 µm steps, and the measured stage feedback positions were used as the displacement reference. Across 120 paired steps, the mean absolute deviations from this feedback reference were 9.54, 7.46, and 6.74 nm, whereas the within-nine-pillar inconsistencies were 17.63, 7.92, and 6.47 nm.",
        536: "整章的方法可先以一段文字概述。先以SEM已知的九顆pillar尺寸反推該區域的局部PSF，用它描述系統的模糊；此 PSF 有兩項用途，一是作為傳統Wiener反卷積的對照組，二是做為重建網路的光學先驗。接著以單張輸入、輸入輸出同尺寸的編碼器—解碼器重建影像，並用兩項設計防止網路損壞邊界：抗混疊上取樣負責化解特徵放大時的像素相位差異，以輸入影像為基準的殘差輸出則讓原圖保住低頻輪廓與中心位置，模型只補上比例受限的高頻。訓練階段以四組量測導向損失函數，搭配分工明確的多來源資料進行。最後，原始影像（Raw）、Wiener反卷積與本文模型（MOSR）三者跑完全相同的後端量測程序，比較邊緣過渡寬度與振鈴、逐顆CD相對SEM參考值的誤差，以及九顆pillar位移相對PI平台實際回授位移的一致性。",
        564: "由這四項原則往回推，研究流程也就定了下來：光學校正、監督資料建立、模型訓練、單張等尺寸重建、後端量測。各階段共用同一組像素座標、物體空間尺度與樣品身分，光學資訊、模型輸出、SEM尺寸與PI平台實際回授位置因而可以互相核對。",
        737: "運動一致性損失的目的不只是讓平均位移貼近PI平台實際回授位移，更重要的是降低模型自身造成的影格相依性。同一結構若僅因落在不同像素相位或不同影格中，就被重建成不同的尺寸或中心位置，便不宜作為微位移量測的依據。",
        802: "位移結果則報告四項：代表位移與PI平台實際回授位移的差距、帶號位移序列的步間波動、陣列內位移離散度，以及跨軸位移；亞像素位移的估計基礎見文獻[80,81]。所有方法使用相同的影格對與相同的彙整規則，並依共同的量測與不確定度原則解讀[84,85]，以免因資料篩選不同而製造出表面上的改善。",
        803: "本研究以平台實際回授位置作為位移參考。在相同影格、相同中心與相同彙整規則下，比較各方法相對於此位移參考的偏差、步間波動、跨軸變化與陣列內一致性。",
        811: "Raw、Wiener、MOSR與各消融條件共用相同的座標、pillar身分、中心、物體空間尺度、閾值與統計規則。第六章將先以邊緣、跨閾值CD及PI平台實際回授位移比較完整方法，再以其他條件一致的消融實驗檢驗兩項主要模組。",
        1034: "本節在同一顯微系統與HS-100MG樣品下，使用相同物體空間尺度與後端規則比較三種影像的邊界、尺寸與位移。CD以九顆SEM尺寸作為逐顆參考；PI閉迴路平台分別執行名目0.5 µm與1.0 µm逐步位移，並以平台實際回授位置作為共同位移參考。",
        1049: "『相對PI平台實際回授位移之平均絕對偏差』是每一步影像量測位移幅度與該步平台實際回授位移之絕對差平均。此指標表示三種影像方法與同一平台位移參考的一致程度。各序列依預先固定的影像完整性與平台運動品質條件排除無法可靠量測的影格，再統一保留30個有效步；Raw、Wiener與MOSR共用相同名單。",
        1099: "50×、0.5 µm序列固定分析30步。Raw、Wiener與MOSR的逐次位移量標準差依序為7.04、4.94與3.61 nm；相對PI平台實際回授位移之平均絕對偏差為5.25、4.07與3.07 nm；3×3 pillar陣列內位移離散度則為13.03、6.86與4.48 nm。兩種處理都改善Raw，MOSR在三項指標上均最低。該序列的逐次位移量如圖6-37所示。",
        1102: "50×、1.0 µm序列固定分析30步。Raw、Wiener與MOSR的逐次位移量標準差為7.69、4.36與3.72 nm；相對PI平台實際回授位移之平均絕對偏差為6.98、3.89與3.36 nm；3×3 pillar陣列內位移離散度為12.06、7.07與3.73 nm。MOSR再度得到最低的逐次位移量離散程度、平台實際回授位移偏差與3×3 pillar陣列內位移離散度。該序列的逐次位移量如圖6-38所示。",
        1105: "20×、0.5 µm序列固定分析30步。Raw、Wiener與MOSR的逐次位移量標準差為16.25、15.19與14.73 nm，相對PI平台實際回授位移之平均絕對偏差為13.13、12.78與11.74 nm，兩項皆由MOSR略低；3×3 pillar陣列內位移離散度則為25.14、8.70與9.13 nm，由Wiener略低。此例顯示同一序列在不同指標上可有不同排名，因此後續仍以四組合併結果判斷整體穩定性。該序列的逐次位移量如圖6-39所示。",
        1108: "20×、1.0 µm序列固定分析30步。Raw、Wiener與MOSR的逐次位移量標準差為13.52、7.46與6.94 nm；相對PI平台實際回授位移之平均絕對偏差為12.82、9.13與8.81 nm；3×3 pillar陣列內位移離散度為20.28、9.03與8.53 nm。MOSR在三項指標均略優於Wiener。該序列的逐次位移量如圖6-40所示。",
        1116: "四組序列各固定30步，合併後共有120個配對步。Raw、Wiener與MOSR相對PI平台實際回授位移之平均絕對偏差為9.54、7.46與6.74 nm；Wiener與MOSR分別較Raw下降21.79%與29.36%。3×3 pillar陣列內平均位移離散度則由Raw的17.63 nm降至Wiener的7.92 nm與MOSR的6.47 nm，分別下降55.08%與63.29%。因此，兩種處理均改善位移量測，而MOSR在合併後的平台實際回授位移一致性與九顆pillar一致性皆為最佳。四組序列的逐次位移量標準差、九顆pillar內部位移離散度及相對PI平台實際回授位移的平均絕對偏差，分別整理於圖6-41至圖6-43。",
        1127: "分倍率後，MOSR在50×相對PI平台實際回授位移的平均絕對偏差下降47.51%，3×3 pillar陣列內位移離散度下降67.26%；20×分別下降20.80%與61.10%。Wiener在50×的兩項下降幅度為34.98%與44.47%，20×為15.57%與60.94%。0.5 µm與1.0 µm步階在50×影像中分別約為7.25與14.49 pixels，在20×中約為2.90與5.80 pixels；較密的物體空間取樣與50×較大的改善幅度相符。",
        1128: "兩種處理都沒有以縮放整段位移來換取較低散佈。四組序列中，Wiener或MOSR相對Raw的平均位移幅度差異最大僅1.44 nm，遠小於0.5 µm與1.0 µm的名目位移步長。主要變化因此是減少步與步、pillar與pillar之間的散佈，而不是改變整體位移比例。",
        1138: "距離與Wiener相對PI平台實際回授位移之絕對偏差的Spearman相關在50倍為ρ = 0.074、p = 0.576，20倍為ρ = -0.014、p = 0.915；距離與3×3 pillar陣列位移離散度的相關則分別為ρ = 0.138、p = 0.292，以及ρ = -0.319、p = 0.013。前三項未達顯著，20倍3×3 pillar陣列位移離散度則呈顯著負相關，而不是隨距離變差。四項結果都不支持『離校正位置越遠，Wiener誤差越大』的單調趨勢；在本次約25–30 µm取樣範圍內，固定PSF可作為條件一致的傳統基準。",
        1155: "九顆pillar的空間一致性呈現最明顯的改善。四組序列的位移離散度均降低；合併120步後，3×3 pillar陣列位移離散度由Raw的17.63 nm降至6.96 nm，下降60.52%，相對PI平台實際回授位移之平均絕對偏差亦由9.54 nm降至6.70 nm，下降29.84%。兩組隨機初始化均重現主要實驗的改善方向。",
        1194: "為確認位移改善不只出現在正式的Raw定義30步名單，另以時間規律性、取樣區中央及全部同次拍攝影格重新計算；每種規則均由時間資訊或Raw資料決定。表6-16列出Raw與MOSR相對PI平台實際回授位移的平均絕對偏差，並於規則名稱中標示pillar與via的共同有效樣本數。",
        1198: "運算效能以相同2048×1536影像、相同顯示卡、相同自動混合精度及整張影像直接推論量測。表6-17列出參數量相近的NAFNet、EDSR-x1與MOSR，時間為兩次暖機後六次推論的中位數與第95百分位。",
        1200: "NAFNet全解析度影像推論的中位時間為152.7 ms，第95百分位為155.3 ms，吞吐量為20.60 MP/s，為三種模型中最快。MOSR的中位時間為400.2 ms，第95百分位為406.8 ms，吞吐量為7.86 MP/s，相當於約2.50 frames/s；其中位時間仍較EDSR-x1縮短56.3%。結合前述via CD與量測穩健性結果，MOSR雖非速度最快，仍在次秒級推論下取得較佳量測結果，顯示運算速度與量測品質需一併評估。",
        1202: "跨形貌與跨模型結果將主要pillar證據延伸至亮via與線條。20倍via的SEM對照CD、閾值區間與邊緣寬度均由MOSR最低；水平與垂直線條在邊緣過渡集中後仍維持原有週期；四種客觀影格規則下，pillar與via的位移平均絕對偏差均相對Raw下降。全解析度直接推論結果顯示NAFNet速度最快，MOSR則在保持較佳via量測結果的同時維持次秒級推論。",
        1205: "本節在同一受控顯微平台，以pillar、via、水平與垂直線條、兩種倍率及兩種位移步距，對照Raw、Wiener、EDSR、NAFNet與MOSR。SEM提供圓形結構尺寸，PI平台實際回授位置提供共同位移參考；各方法共用影像、結構身分、物體空間尺度、閾值與統計規則，使方法差異能由同一量測鏈直接比較。",
        1211: "第二與第三個研究問題關注平移一致性及邊界保護。在主要Raw、Wiener與MOSR比較中，MOSR於四組序列合併後取得最低的PI平台實際回授位移平均絕對偏差與九顆pillar位移離散度。配對式消融比較顯示，抗混疊上取樣降低不同pillar的位移離散度；相較於移除以輸入影像為基準的殘差輸出的模型，完整模型則降低跨倍率振鈴，並在20倍CD上得到明顯改善。兩項模組分別對應解碼階段的跨位置一致性與輸出階段的輪廓穩定性。",
        1218: "在四組位移序列共120個配對步中，Raw、Wiener與MOSR相對PI平台實際回授位移的平均絕對偏差為9.54、7.46與6.74 nm，九顆pillar位移離散度為17.63、7.92與6.47 nm。抗混疊上取樣相較一般上取樣使合併位移離散度降低14.31%；移除以輸入影像為基準的殘差輸出後，20倍CD絕對誤差中位數增加73.97%，50倍與20倍振鈴分別增加246.15%與392.61%。",
        1221: "本研究在同一受控顯微平台中，使用HS-100MG樣品、20倍與50倍物鏡、名目0.5 µm與1.0 µm位移步長，以及pillar、via與水平、垂直線條完成驗證。SEM提供九顆pillar與九顆via的逐顆尺寸，PI平台實際回授位置提供共同位移參考；所有方法共用影像、中心、物體空間尺度、閾值與統計規則，因此方法間差異可由同一量測鏈直接比較。",
        1222: "量測可信度由逐顆SEM尺寸、PI平台實際回授位移、七個CD閾值、兩種剖面、固定影格名單、影格選取敏感度、成對自助法信賴區間及兩組固定初始化共同建立。主文保留合併結果與第90百分位摘要，完整條件結果則提供不同閾值、影格選取與訓練變動下的交叉檢查。",
        1225: "後續可使用雷射干涉儀與PI平台實際回授位置同步量測，進一步建立外部量測追溯關係，並將重複性、像素尺度、SEM參考、平台回授與模型變動整合為完整不確定度預算[84,85]。模型驗證亦可增加不同批次樣品、跨系統測試及跨條件分層統計，進一步擴充MOSR的操作範圍。",
    }
    for index, new_text in replacement_paragraphs.items():
        set_paragraph_text(doc.paragraphs[index], new_text)

    caption_replacements = {
        1111: ("四組序列之逐次位移量標準差與PI閉迴路位移參考值一致性", "四組序列之逐次位移量標準差與PI平台實際回授位移一致性"),
        1122: ("Raw、Wiener 與 MOSR 相對PI閉迴路位移參考值之平均絕對偏差", "Raw、Wiener與MOSR相對PI平台實際回授位移之平均絕對偏差"),
        1186: ("20× via measurement results under the shared metrology ruler", "20倍via在共同量測規則下的結果"),
        1195: ("Displacement MAE sensitivity to frame-selection rules", "不同影格選取規則下的位移平均絕對偏差（單位：nm；改善率除外）"),
        1199: ("Full-frame direct-inference performance on the same GPU", "相同顯示卡下的全解析度直接推論效能"),
    }
    for index, (old, new) in caption_replacements.items():
        replace_text_nodes(doc.paragraphs[index], old, new)

    table_15 = doc.tables[111]
    headers_15 = ["方法", "CD中點\n平均絕對誤差\n(nm)", "CD閾值區間寬度\n(nm)", "10–90%邊緣寬度\n(nm)"]
    for cell, text in zip(table_15.rows[0].cells, headers_15):
        set_cell_text(cell, text, bold=True)

    table_16 = doc.tables[112]
    headers_16 = ["選取規則\n(pillar/via步數)", "Pillar\nRaw", "Pillar\nMOSR", "Pillar\n改善率", "Via\nRaw", "Via\nMOSR", "Via\n改善率"]
    for cell, text in zip(table_16.rows[0].cells, headers_16):
        set_cell_text(cell, text, bold=True)
    rules = [
        "Raw固定選取30步\n(120/120)",
        "等時間間隔30步\n(113/120)",
        "中央拍攝區段30步\n(114/120)",
        "同批次全部影格\n(178/192)",
    ]
    for row, rule in zip(table_16.rows[1:], rules):
        set_cell_text(row.cells[0], rule, bold=False)
    set_table_widths(table_16, [2401, 1083, 1083, 1083, 1083, 1083, 1083])

    table_17 = doc.tables[113]
    headers_17 = ["方法", "參數量\n(M)", "中位時間\n(ms)", "第95百分位時間\n(ms)", "吞吐量\n(MP/s)"]
    for cell, text in zip(table_17.rows[0].cells, headers_17):
        set_cell_text(cell, text, bold=True)
    set_table_widths(table_17, [1600, 1500, 1850, 2400, 1823])
    edsr_tr = table_17.rows[1]._tr
    naf_tr = deepcopy(edsr_tr)
    table_17._tbl.insert(table_17._tbl.index(edsr_tr) + 1, naf_tr)
    naf_row = table_17.rows[2]
    for cell, text in zip(naf_row.cells, ["NAFNet", "4.182", "152.7", "155.3", "20.60"]):
        set_cell_text(cell, text, bold=False)

    for table in doc.tables:
        if table.rows:
            set_repeat_header(table.rows[0])
    for table in (table_15, table_16, table_17):
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Remove the old crop retained by the Figure 6-53 picture frame.  The
    # replacement image already contains the complete Chinese legend and axes.
    found_figure_53 = False
    for paragraph in doc.paragraphs:
        for blip in paragraph._p.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            part = paragraph.part.related_parts.get(rid)
            if part is not None and str(part.partname).endswith("/image98.png"):
                blip_fill = blip.getparent()
                for src_rect in blip_fill.xpath("./a:srcRect"):
                    for attr in ("l", "t", "r", "b"):
                        src_rect.set(attr, "0")
                found_figure_53 = True
    if not found_figure_53:
        raise AssertionError("Could not locate the Figure 6-53 picture frame")

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(OUTPUT)

    replacements = {
        "word/media/image45.png": EARLY / "fig6-01_zh.png",
        "word/media/image47.png": EARLY / "fig6-02_zh.png",
        "word/media/image49.png": EARLY / "fig6-04_zh.png",
        "word/media/image50.png": EARLY / "fig6-05_zh.png",
        "word/media/image51.png": EARLY / "fig6-06_zh.png",
        "word/media/image56.png": EARLY / "fig6-11_zh.png",
        "word/media/image57.png": EARLY / "fig6-12_zh.png",
        "word/media/image58.png": EARLY / "fig6-13_zh.png",
        "word/media/image59.png": EARLY / "fig6-14_zh.png",
        "word/media/image60.png": EARLY / "fig6-15_zh.png",
        "word/media/image61.png": EARLY / "fig6-16_zh.png",
        "word/media/image62.png": EARLY / "fig6-17_zh.png",
        "word/media/image63.png": EARLY / "fig6-18_zh.png",
        "word/media/image64.png": EARLY / "fig6-19_zh.png",
        "word/media/image65.png": EARLY / "fig6-20_zh.png",
        "word/media/image66.png": EARLY / "fig6-21_zh.png",
        "word/media/image67.png": EARLY / "fig6-22_zh.png",
        "word/media/image68.png": WIENER / "00c_signed_pillar_contrast.png",
        "word/media/image69.png": WIENER / "01_50x_raw_center9.png",
        "word/media/image70.png": WIENER / "02a_50x_wiener_center9.png",
        "word/media/image71.png": WIENER / "02_50x_v42_center9.png",
        "word/media/image72.png": WIENER / "03_50x_pillar_edge_profile.png",
        "word/media/image73.png": WIENER / "04_20x_raw_center9.png",
        "word/media/image74.png": WIENER / "05a_20x_wiener_center9.png",
        "word/media/image75.png": WIENER / "05_20x_v42_center9.png",
        "word/media/image76.png": WIENER / "06_20x_pillar_edge_profile.png",
        "word/media/image77.png": WIENER / "07_center9_cd_per_pillar.png",
        "word/media/image78.png": WIENER / "08_center9_cd_error_summary.png",
        "word/media/image79.png": WIENER / "09a_cd_threshold_sensitivity_horizontal.png",
        "word/media/image80.png": WIENER / "09b_cd_threshold_sensitivity_radial.png",
        "word/media/image81.png": WIENER / "19_20x_noncalibration_cd_error_summary.png",
        "word/media/image82.png": WIENER / "10_50x_0p5um_step_trace.png",
        "word/media/image83.png": WIENER / "11_50x_1um_step_trace.png",
        "word/media/image84.png": WIENER / "12_20x_0p5um_step_trace.png",
        "word/media/image85.png": WIENER / "13_20x_1um_step_trace.png",
        "word/media/image86.png": WIENER / "14_step_stability_summary.png",
        "word/media/image87.png": WIENER / "15_nine_pillar_consistency_summary.png",
        "word/media/image88.png": WIENER / "16_nominal_command_consistency.png",
        "word/media/image89.png": WIENER / "17_stepwise_local_improvement_rate.png",
        "word/media/image90.png": WIENER / "18a_50x_wiener_psf_distance_robustness.png",
        "word/media/image91.png": WIENER / "18b_20x_wiener_psf_distance_robustness.png",
        "word/media/image92.png": R63 / "15_full_vs_raw_cd_error.png",
        "word/media/image93.png": R63 / "16_full_vs_raw_step_stability.png",
        "word/media/image94.png": R63 / "17_full_vs_raw_grid_agreement.png",
        "word/media/image95.png": R63 / "18_raw_position_anchor_cd_error.png",
        "word/media/image96.png": R63 / "19_upsampling_grid_agreement.png",
        "word/media/image97.png": R63 / "20_full_vs_raw_ringing_tradeoff.png",
        "word/media/image98.png": CROSS / "fig6-53_zh.png",
        "word/media/image99.png": CROSS / "fig6-54_zh.png",
    }
    for path in replacements.values():
        if not path.is_file():
            raise FileNotFoundError(path)
    patch_media(OUTPUT, replacements)
    print(OUTPUT)
    print(f"replaced_media={len(replacements)}")


if __name__ == "__main__":
    main()
