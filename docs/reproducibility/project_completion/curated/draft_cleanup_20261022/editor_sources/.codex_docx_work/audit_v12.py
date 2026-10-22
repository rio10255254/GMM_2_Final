from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


TERMS = (
    "Base U-Net",
    "BaseU-Net",
    "BaseUNet",
    "base_unet",
    "base unet",
    "base u-net",
    "基礎 U-Net",
    "基礎U-Net",
    "基礎模型對照",
)


def iter_block_items(doc: Document):
    parent = doc.element.body
    for child in parent.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def has_term(text: str) -> bool:
    low = text.casefold()
    return any(term.casefold() in low for term in TERMS)


def paragraph_drawings(paragraph: Paragraph) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    for drawing in paragraph._p.xpath(".//w:drawing"):
        for doc_pr in drawing.xpath(".//wp:docPr"):
            item = {
                    "id": doc_pr.get("id", ""),
                    "name": doc_pr.get("name", ""),
                    "descr": doc_pr.get("descr", ""),
                    "title": doc_pr.get("title", ""),
                }
            blips = drawing.xpath(".//a:blip")
            if blips:
                rid = blips[0].get(qn("r:embed"), "")
                item["rid"] = rid
                if rid and rid in paragraph.part.rels:
                    item["target"] = str(paragraph.part.rels[rid].target_ref)
            out.append(item)
    for pict in paragraph._p.xpath(".//w:pict"):
        out.append({"id": "", "name": "VML", "descr": "", "title": ""})
    return out


def main() -> None:
    path = Path(os.environ["DOCX_PATH"])
    doc = Document(path)
    blocks = list(iter_block_items(doc))

    print(f"DOCX={path}")
    print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} BLOCKS={len(blocks)}")

    print("\n=== BODY HITS WITH CONTEXT ===")
    hit_indices: list[int] = []
    for idx, block in enumerate(blocks):
        if isinstance(block, Paragraph):
            text = block.text.strip()
        else:
            text = " | ".join(cell.text.strip().replace("\n", " / ") for row in block.rows for cell in row.cells)
        if has_term(text):
            hit_indices.append(idx)
    shown: set[int] = set()
    for hit in hit_indices:
        lo, hi = max(0, hit - 3), min(len(blocks), hit + 4)
        for idx in range(lo, hi):
            if idx in shown:
                continue
            shown.add(idx)
            block = blocks[idx]
            if isinstance(block, Paragraph):
                style = block.style.name if block.style else ""
                drawings = paragraph_drawings(block)
                print(f"B{idx:04d} P style={style!r} drawings={drawings} :: {block.text.strip()}")
            else:
                preview = " || ".join(
                    " | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells)
                    for row in block.rows
                )
                print(f"B{idx:04d} T rows={len(block.rows)} cols={len(block.columns)} :: {preview}")
        print("---")

    print("\n=== FIGURE/CAPTION REGION 5.4-7.3 ===")
    active = False
    for idx, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            if active:
                print(f"B{idx:04d} TABLE rows={len(block.rows)} cols={len(block.columns)}")
            continue
        text = block.text.strip()
        if text.startswith("5.4"):
            active = True
        if active:
            style = block.style.name if block.style else ""
            drawings = paragraph_drawings(block)
            if text or drawings or style.lower().startswith("heading") or "Caption" in style:
                print(f"B{idx:04d} P style={style!r} drawings={drawings} :: {text}")
        if text.startswith("7.4"):
            break

    print("\n=== PACKAGE PART HITS ===")
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if not name.lower().endswith((".xml", ".rels", ".txt")):
                continue
            raw = zf.read(name)
            text = raw.decode("utf-8", errors="ignore")
            matches = [term for term in TERMS if term.casefold() in text.casefold()]
            if matches:
                print(f"{name}: {matches}")

        print("\n=== MEDIA INVENTORY ===")
        for name in zf.namelist():
            if name.startswith("word/media/"):
                info = zf.getinfo(name)
                print(f"{name}\t{info.file_size}")

    print("\n=== FIELD INVENTORY ===")
    for idx, p in enumerate(doc.paragraphs):
        instr = " ".join(node.text or "" for node in p._p.xpath(".//w:instrText")).strip()
        if instr:
            print(f"P{idx:04d}\t{instr}\t{p.text.strip()}")


if __name__ == "__main__":
    main()
