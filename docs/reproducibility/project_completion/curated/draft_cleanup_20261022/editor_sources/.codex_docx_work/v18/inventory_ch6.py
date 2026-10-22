from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(r"C:\Users\Rio\OneDrive\學業\碩班\論文\可用\論文完整版_20260716_v17_修訂版.docx")
OUT = Path(r"D:\GMM_2_Final\PSF_Fitting\.codex_docx_work\v18\chapter6_inventory.json")


def paragraph_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def image_parts(paragraph) -> list[str]:
    parts: list[str] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        part = paragraph.part.related_parts.get(rid)
        if part is not None:
            parts.append(str(part.partname))
    return parts


def main() -> None:
    doc = Document(SOURCE)
    figures = []
    tables = []
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph_text(paragraph)
        if paragraph.style.name == "Figure Caption" and text.startswith("圖 6-"):
            adjacent = []
            for j in range(max(0, i - 2), min(len(doc.paragraphs), i + 3)):
                parts = image_parts(doc.paragraphs[j])
                if parts:
                    adjacent.extend(parts)
            figures.append({"paragraph_index": i, "caption": text, "media": sorted(set(adjacent))})
        if paragraph.style.name == "Table Caption" and text.startswith("表 6-"):
            tables.append({"paragraph_index": i, "caption": text})
    OUT.write_text(json.dumps({"figures": figures, "tables": tables}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"figures={len(figures)} tables={len(tables)}")
    print(OUT)


if __name__ == "__main__":
    main()
