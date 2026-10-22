from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def first_run_props(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return copy.deepcopy(run._r.rPr)
    return None


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    props = first_run_props(paragraph)
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if props is not None:
        run._r.insert(0, copy.deepcopy(props))


def find_one(doc: Document, prefix: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph beginning {prefix!r}; found {len(matches)}")
    return matches[0]


def replace_prefix(doc: Document, prefix: str, text: str) -> None:
    set_paragraph_text(find_one(doc, prefix), text)


def copy_paragraph_format(target: Paragraph, source: Paragraph) -> None:
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, copy.deepcopy(source._p.pPr))


def insert_paragraph_after(anchor: Paragraph, source_format: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    copy_paragraph_format(paragraph, source_format)
    set_paragraph_text(paragraph, text)
    return paragraph


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table: Table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])


def copy_cell_appearance(target, source, text: str, bold: bool | None = None) -> None:
    source_tc_pr = source._tc.tcPr
    if source_tc_pr is not None:
        target_tc_pr = target._tc.get_or_add_tcPr()
        for child in list(target_tc_pr):
            if child.tag != qn("w:tcW"):
                target_tc_pr.remove(child)
        for child in source_tc_pr:
            if child.tag != qn("w:tcW"):
                target_tc_pr.append(copy.deepcopy(child))

    source_p = source.paragraphs[0]
    target_p = target.paragraphs[0]
    copy_paragraph_format(target_p, source_p)
    props = first_run_props(source_p)
    clear_paragraph(target_p)
    run = target_p.add_run(text)
    if props is not None:
        run._r.insert(0, copy.deepcopy(props))
    if bold is not None:
        run.bold = bold
    target_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def insert_zernike_table(doc: Document) -> None:
    if any(p.text.strip().startswith("表 3-1 OSA/ANSI") for p in doc.paragraphs):
        return

    anchor = find_one(doc, "基於波前分析與像差校正的需求")
    caption_template = find_one(doc, "表 5-1 實驗系統元件與編號")
    caption = insert_paragraph_after(
        anchor,
        caption_template,
        "表 3-1 OSA/ANSI 標準 Zernike 模態順序（j = 0 至 14）",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True

    ref_table = doc.tables[81]
    table = doc.add_table(rows=9, cols=6)
    table.style = ref_table.style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    caption._p.addnext(table._tbl)

    widths = [700, 900, 2548, 700, 900, 2548]
    set_table_widths(table, widths)

    header = ["j", "(n, m)", "像差名稱", "j", "(n, m)", "像差名稱"]
    modes = [
        ("0", "(0, 0)", "活塞項（piston）"),
        ("1", "(1, -1)", "傾斜項（sine tilt）"),
        ("2", "(1, 1)", "傾斜項（cosine tilt）"),
        ("3", "(2, -2)", "散光（sine astigmatism）"),
        ("4", "(2, 0)", "失焦（defocus）"),
        ("5", "(2, 2)", "散光（cosine astigmatism）"),
        ("6", "(3, -3)", "三葉像差（sine trefoil）"),
        ("7", "(3, -1)", "彗差（sine coma）"),
        ("8", "(3, 1)", "彗差（cosine coma）"),
        ("9", "(3, 3)", "三葉像差（cosine trefoil）"),
        ("10", "(4, -4)", "四葉像差（sine quadrafoil）"),
        ("11", "(4, -2)", "次級散光（sine）"),
        ("12", "(4, 0)", "主球差（primary spherical）"),
        ("13", "(4, 2)", "次級散光（cosine）"),
        ("14", "(4, 4)", "四葉像差（cosine quadrafoil）"),
    ]

    for col, value in enumerate(header):
        copy_cell_appearance(table.rows[0].cells[col], ref_table.rows[0].cells[col % 4], value, bold=True)

    left = modes[:8]
    right = modes[8:]
    for row_idx in range(1, 9):
        values = list(left[row_idx - 1])
        values += list(right[row_idx - 1]) if row_idx - 1 < len(right) else ["", "", ""]
        for col, value in enumerate(values):
            copy_cell_appearance(
                table.rows[row_idx].cells[col],
                ref_table.rows[min(row_idx, len(ref_table.rows) - 1)].cells[col % 4],
                value,
                bold=False,
            )

    for row in table.rows:
        row._tr.get_or_add_trPr()


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: revise_thesis_v7.py input.docx output.docx")
    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    doc = Document(input_path)

    replace_prefix(
        doc,
        "基於波前分析與像差校正的需求",
        "基於波前分析與像差校正的需求，澤尼克多項式（Zernike polynomials）因其在單位圓域內具備正交性，可將不同型態的波前變形以有限組係數表示，因此被廣泛用於光瞳平面的像差分解與估測。若將光瞳座標歸一化至單位圓（如圖 3-9），即可用澤尼克基底描述光瞳上的波前誤差 [15], [16]。本研究採 OSA/ANSI 標準順序，使用 j = 0 至 14 的前 15 項模態建立有限階參數化；各模態之徑向階數、方位階數與常用像差名稱列於表 3-1，其光程差分布示意如圖 3-10 所示。",
    )
    insert_zernike_table(doc)

    replace_prefix(
        doc,
        "程式內部採與OSA/ANSI模態順序一致",
        "本研究以 OSA/ANSI 標準模態順序表示等效波前。j = 0 為活塞項，j = 1 與 j = 2 為兩個正交傾斜項，其後依序涵蓋散光、失焦、彗差、三葉像差、球差及其他高階模態。活塞項不改變正規化後的強度 PSF，傾斜項則主要表現為整體影像位移，因此本研究先完成亞像素幾何預對位，再估計會改變 PSF 形貌的高階模態係數與整體振幅。",
    )
    replace_prefix(
        doc,
        "為避免此問題，本研究先獨立處理幾何對位",
        "為避免將結構位置偏差吸收到像差係數中，本研究先在水平與垂直方向搜尋相對平移，使 SEM 參考幾何形成的理想影像與實拍九宮格完成亞像素預對位。完成預對位後，再估計高階像差係數與整體影像振幅，將「九宮格位於影像中的位置」與「光學模糊具有何種形貌」分成兩個問題處理。",
    )

    replace_prefix(
        doc,
        "上述來源依功能分工",
        "上述來源依功能分工：真實影像提供外觀，合成幾何與 SEM 提供輪廓及尺度，已知位移與連續影格提供相對位置，跨條件資料提供穩定性。模型不依賴單一全能的參考影像，而是整合多種各自具有明確用途的監督訊息。",
    )
    replace_prefix(
        doc,
        "本研究的量測導向訓練不是從完全隨機",
        "本研究採同一網路架構的連續分階段訓練。基礎階段先建立一般顯微影像的內容重建能力；量測導向階段則延續既有參數，針對邊界、尺寸與位移一致性進行加強訓練。",
    )
    replace_prefix(
        doc,
        "採用既有模型作為起點的原因",
        "兩個階段使用相同的主體架構與顯微影像資料體系，不引入外部模型或不同架構的預訓練權重。連續訓練使前段特徵提取層保留對亮暗分布、背景變化、局部紋理及結構排列的基本表徵，再由後段量測導向訓練強化邊界與位置關係。",
    )
    replace_prefix(
        doc,
        "因此，微調時保留主要的前段特徵提取能力",
        "量測導向階段保留主要的前段特徵提取能力，將參數更新集中於品質門控、邊界修正與輸出相關層。此設計使模型在維持原有影像內容能力的同時，重新學習哪些細節可被增強，以及哪些幾何關係應保持穩定。",
    )
    replace_prefix(
        doc,
        "微調資料仍依原始序列與物理來源分組",
        "訓練與內部驗證資料依原始序列及物理來源分組，使同一來源的相鄰影格與衍生配對不跨越兩個集合。正式比較則固定使用同一個完成訓練的模型權重，不因倍率、CD 或位移指標而切換模型。",
    )

    replace_prefix(
        doc,
        "模型由既有穩定權重開始微調",
        "模型採同一架構的連續分階段訓練：先由基礎階段建立顯微影像重建能力，再延續其參數進行量測導向加強訓練。第二階段保留大部分主幹參數，主要更新品質門控、邊界修正與輸出相關層。最佳化採 Adam，初始學習率為 2×10⁻⁷，影像批次為 8，驗證批次與序列配對批次皆為 2，共訓練 120 回合。訓練過程同步記錄重建、梯度幅值、邊緣方向與平移一致性相關項目，並以固定監測影像檢查灰階範圍、邊緣寬度及高頻比例。",
    )

    replace_prefix(
        doc,
        "本實驗使用BudgetSensors HS-100MG樣品",
        "本實驗使用 BudgetSensors HS-100MG 樣品進行點擴散函數（Point Spread Function, PSF）之擬合與校正。為建立具有明確尺寸基準的擬合目標，本研究選取樣品中已透過掃描式電子顯微鏡（Scanning Electron Microscope, SEM）量測關鍵尺寸（Critical Dimension, CD），且位於樣品中心附近的 pillar 九宮格結構作為校正對象。依據逐顆 SEM 尺寸建立對應的參考幾何影像，作為光學成像模型與 PSF 參數擬合的比較基準。",
    )
    replace_prefix(
        doc,
        "本研究首先依據SEM量測所得的pillar實際尺寸建立Ground Truth影像",
        "本研究首先依據 SEM 量測所得的 pillar 實際尺寸建立參考幾何影像，並將實驗拍攝影像與模擬成像結果進行比較。接著使用 BADS 演算法搜尋最佳化參數，透過調整 Zernike 多項式所描述的等效像差係數，使參考幾何經候選 PSF 成像後盡可能接近對應視野位置的實驗影像。每一個視野區域皆具有獨立的最佳化結果，因此可分別描述光學系統在不同視野位置下的等效模糊特性。",
    )
    replace_prefix(
        doc,
        "本研究分別使用 Olympus MPLN-50X",
        "本研究分別使用 Olympus MPLN-50X 與 Olympus MPLN-20X 物鏡所拍攝的 pillar 九宮格影像進行點擴散函數擬合。兩組影像皆選取 pillar 與 via 交界附近的九顆 pillar 作為擬合目標。SEM 提供的九個數值為半徑，範圍為 1299.376-1338.398 nm；換算為 CD 直徑後，範圍為 2598.752-2676.797 nm，平均為 2629.737 nm。SEM 參考幾何依各 pillar 的實際半徑與物體空間取樣尺寸建立，並根據 50 倍與 20 倍影像的拍攝方向調整排列，使影像位置與同一顆實體 pillar 正確對應。",
    )
    replace_prefix(
        doc,
        "本研究使用Noll編號Z4至Z15",
        "本研究依 OSA/ANSI 標準順序，以有限階 Zernike 基底表示等效波前像差，並透過 BADS 演算法搜尋高階模態係數與振幅參數。最佳化過程先由候選 PSF 與 SEM 參考幾何完成成像運算，再比較模擬影像與實際拍攝影像的亮度及邊緣梯度差異。由於 pillar 邊界位置直接影響後續關鍵尺寸量測，目標函數除考量影像亮度差異外，也提高邊緣梯度誤差的權重，使擬合結果同時描述模糊程度與邊界變化。",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    check = Document(output_path)
    text = "\n".join(p.text for p in check.paragraphs)
    forbidden = ["使用Noll編號Z4至Z15", "程式內部採與OSA/ANSI", "既有穩定權重開始微調", "Ground Truth"]
    remaining = [term for term in forbidden if term in text]
    if remaining:
        raise RuntimeError(f"Forbidden stale wording remains: {remaining}")
    if not any(p.text.startswith("表 3-1 OSA/ANSI") for p in check.paragraphs):
        raise RuntimeError("Table 3-1 caption was not inserted")


if __name__ == "__main__":
    main()
