from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def replace_text(paragraph, text: str) -> None:
    run_props = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            run_props = copy.deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_props is not None:
        run._r.insert(0, run_props)


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(path)
    replacements = {
        "2.1點擴散函數（Point Spread Function, PSF）模型": "2.1 點擴散函數（Point Spread Function, PSF）模型",
        "5.1.5 XY軸壓電位移平台": "5.1.5 XY 軸壓電位移平台",
        "5.1.6 Z軸壓電位移平台": "5.1.6 Z 軸壓電位移平台",
        "5.3 MATLAB環境": "5.3 MATLAB 環境",
        "第六章固定使用第120回合權重進行評估，不依CD或位移結果事後挑選較有利的回合。訓練曲線僅用來確認最佳化過程是否收斂；模型是否具有量測效益，仍由第六章的邊緣、CD與位移結果判定。":
            "第六章的所有正式比較均固定使用第120回合權重，並在不同CD與位移指標間維持同一模型，不另外更換回合。訓練曲線僅用來確認最佳化過程是否收斂；模型是否具有量測效益，仍由第六章的邊緣、CD與位移結果判定。",
        "模型部分則補充實際資料數量、分組切分、邊界軟標籤及微調設定。由於最終評估模型在查看第六章量測結果前即固定，後續Raw、Wiener與MOSR的比較可使用同一影像名單與量測規則進行。":
            "模型部分則補充實際資料數量、分組切分、邊界軟標籤及微調設定。最終比較固定採用同一模型權重，因此Raw、Wiener與MOSR可使用相同影像名單與量測規則進行對照。",
        "圖 3-13 Unet網路架構示意圖": "圖 3-13 U-Net網路架構示意圖",
        "圖 3-17 BADS模型更新與尋找候選點策略示意圖": "圖 3-17 BADS最佳化之模型更新與候選點搜尋策略",
    }
    found = set()
    for paragraph in doc.paragraphs:
        if paragraph.text in replacements:
            found.add(paragraph.text)
            replace_text(paragraph, replacements[paragraph.text])
    missing = set(replacements) - found
    if missing:
        raise RuntimeError(f"Missing expected paragraphs: {sorted(missing)}")
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
