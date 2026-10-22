from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def first_run_props(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return copy.deepcopy(run._r.rPr)
    return None


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_text(
    paragraph: Paragraph,
    text: str,
    *,
    style: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    template: Paragraph | None = None,
) -> Paragraph:
    props = first_run_props(paragraph) or (first_run_props(template) if template else None)
    clear_paragraph(paragraph)
    if style:
        paragraph.style = style
    run = paragraph.add_run(text)
    if props is not None:
        run._r.insert(0, copy.deepcopy(props))
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return paragraph


def set_lines(paragraph: Paragraph, lines: list[str], *, template: Paragraph | None = None) -> None:
    props = first_run_props(paragraph) or (first_run_props(template) if template else None)
    clear_paragraph(paragraph)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        if props is not None:
            run._r.insert(0, copy.deepcopy(props))
        if index < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def append_text(paragraph: Paragraph, text: str) -> None:
    props = first_run_props(paragraph)
    run = paragraph.add_run(text)
    if props is not None:
        run._r.insert(0, copy.deepcopy(props))


def copy_paragraph_properties(target: Paragraph, source: Paragraph) -> None:
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, copy.deepcopy(source._p.pPr))


def format_chapter_heading(target: Paragraph, source: Paragraph, title: str) -> None:
    copy_paragraph_properties(target, source)
    source_run_props = first_run_props(source)
    clear_paragraph(target)
    run = target.add_run(title)
    if source_run_props is not None:
        run._r.insert(0, copy.deepcopy(source_run_props))


def insert_paragraph_before(
    anchor: Paragraph,
    text: str = "",
    *,
    style: str | None = None,
    template: Paragraph | None = None,
) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    if template is not None and template._p.pPr is not None:
        paragraph._p.insert(0, copy.deepcopy(template._p.pPr))
    if style:
        paragraph.style = style
    if text:
        set_text(paragraph, text, style=style, template=template)
    return paragraph


def insert_paragraph_after(
    anchor: Paragraph,
    text: str = "",
    *,
    style: str | None = None,
    template: Paragraph | None = None,
) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    if template is not None and template._p.pPr is not None:
        paragraph._p.insert(0, copy.deepcopy(template._p.pPr))
    if style:
        paragraph.style = style
    if text:
        set_text(paragraph, text, style=style, template=template)
    return paragraph


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_caption(paragraph: Paragraph, text: str) -> None:
    set_text(paragraph, text, style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = True


def set_table_caption(paragraph: Paragraph, text: str) -> None:
    set_text(paragraph, text, style="Table Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        delete_paragraph(extra)


def replace_in_runs(paragraph: Paragraph, replacements: list[tuple[str, str]]) -> None:
    for run in paragraph.runs:
        original = run.text
        if not original:
            continue
        value = original
        for old, new in replacements:
            value = value.replace(old, new)
        if value != original:
            run.text = value


def add_page_break_before(anchor: Paragraph, template: Paragraph) -> Paragraph:
    paragraph = insert_paragraph_before(anchor, template=template)
    clear_paragraph(paragraph)
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    return paragraph


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(path)
    p = list(doc.paragraphs)
    tables = list(doc.tables)

    if "Table Caption" not in [style.name for style in doc.styles]:
        table_caption_style = doc.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
        table_caption_style.base_style = doc.styles["Caption"]
        table_caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    body_template = p[443]
    heading1_template = p[57]
    heading2_template = p[58]
    heading3_template = p[105]
    heading4_template = p[1040]
    reference_template = p[1128]

    # Cover page: balanced title wrapping and standard English date.
    set_lines(p[7], ["點擴散函數校正與物理約束式神經網路之", "平移定位研究"])
    set_text(p[17], "July 2026")

    # Chinese abstract aligned with the evidence presented in Chapters 4–6.
    set_text(
        p[19],
        "隨著半導體製程與精密製造持續朝微小尺度發展，顯微成像式量測系統不僅需要提供足夠的空間解析能力，也必須確保後續尺寸與位移量測具有穩定性與可重複性。對量測而言，影像看起來更清楚並不必然代表結果更正確；過度銳化、反卷積振鈴或逐張對比調整，都可能改變邊界位置，使視覺改善與幾何量測產生落差。",
    )
    set_text(
        p[20],
        "本研究提出量測導向超解析（Measurement-Oriented Super-Resolution, MOSR）作為精密顯微量測的單張影像重建框架。該方法將光學點擴散函數、已知結構幾何、真實顯微影像與連續位移資訊整合於同一訓練與驗證流程，並以邊界可定位性、關鍵尺寸及位移一致性作為主要判斷依據。模型在實際使用時仍以單張影像作為輸入，輸出尺寸與輸入相同；多影格與位移資訊僅用於訓練及驗證階段的幾何一致性監督。局部與多位置點擴散函數則分別提供傳統反卷積基準及模型的光學先驗策略。",
    )
    set_text(
        p[21],
        "實驗以20倍與50倍物鏡拍攝HS-100MG樣品，並在相同九顆pillar與相同量測規則下比較原始影像、Wiener反卷積及所提MOSR模型。50倍影像的徑向10–90%邊緣過渡寬度中位數由620.1 nm降至Wiener的589.5 nm與MOSR的508.5 nm；20倍則由802.2 nm降至675.8 nm與246.1 nm。跨越35%至65%門檻及兩種剖面方式後，CD中位絕對誤差分別為95.3、151.0與76.1 nm。合併四組位移序列共120個配對步後，三者對名義指令的平均絕對偏差分別為9.54、6.80與6.74 nm，而九顆pillar內部不一致量則由17.63 nm降至10.03與6.47 nm。",
    )
    set_text(
        p[22],
        "結果顯示，Wiener反卷積能提升對比並改善部分量測指標，但其CD結果較容易受到門檻與振鈴影響；MOSR的主要優勢則表現在較集中的邊緣、較穩健的CD邊界，以及更低的多pillar位移分歧。由於位移實驗使用平台名義指令而非獨立位移真值，且CD驗證仍集中於同一資料域與有限區域，本研究不宣稱完成絕對位移校準或跨系統泛化，而是證明以量測可靠性作為重建目標，能比單純追求視覺銳利度更直接地支援後端尺寸與位移分析。",
    )
    set_text(
        p[23],
        "關鍵字：量測導向超解析、點擴散函數校正、物理約束式神經網路、關鍵尺寸、平移定位、Wiener反卷積、量測穩定性。",
    )

    # English abstract mirrors the revised Chinese abstract.
    set_text(
        p[26],
        "As semiconductor manufacturing and precision engineering continue to advance toward smaller feature sizes, microscopy-based metrology systems must provide not only sufficient spatial resolution but also stable and repeatable dimensional and displacement measurements. For metrology, a visually sharper image is not necessarily a more accurate image. Excessive sharpening, deconvolution ringing, and independently stretched display contrast can all shift apparent boundaries and create a discrepancy between visual improvement and geometric measurement.",
    )
    set_text(
        p[27],
        "This thesis proposes Measurement-Oriented Super-Resolution (MOSR) as a single-image reconstruction framework for precision microscopy. The framework integrates optical point spread functions, known structural geometry, real microscope images, and continuous-displacement information within a unified training and evaluation pipeline. Boundary localizability, critical dimension, and displacement consistency are treated as the primary evaluation targets. At inference time, the model accepts one image and produces an output of the same size; multi-frame and displacement information is used only during training and validation as geometric-consistency supervision. Local and multi-position PSFs respectively support the conventional deconvolution baseline and the optical-prior strategy of the reconstruction model.",
    )
    set_text(
        p[28],
        "Experiments were conducted on an HS-100MG sample using 20× and 50× objectives. Raw images, Wiener deconvolution, and the proposed MOSR model were compared using the same nine pillars and the same measurement rules. For the 50× images, the median radial 10–90% edge-transition width decreased from 620.1 nm in the raw images to 589.5 nm with Wiener deconvolution and 508.5 nm with MOSR; for the 20× images, it decreased from 802.2 nm to 675.8 nm and 246.1 nm, respectively. Across thresholds from 35% to 65% and two profile definitions, the median absolute CD errors were 95.3, 151.0, and 76.1 nm. Across 120 paired displacement steps, the mean absolute deviations from the nominal command were 9.54, 6.80, and 6.74 nm, whereas the within-nine-pillar inconsistency decreased from 17.63 nm to 10.03 and 6.47 nm.",
    )
    set_text(
        p[29],
        "The results show that Wiener deconvolution can improve contrast and selected measurement indicators, but its CD estimates are more sensitive to threshold choice and ringing. The primary advantage of MOSR is the combination of concentrated edges, a more robust CD boundary, and lower displacement disagreement among multiple pillars. Because the displacement experiment uses nominal stage commands rather than an independent displacement reference, and the CD validation remains limited to the same data domain and a local region, the thesis does not claim absolute displacement calibration or cross-system generalization. Instead, it demonstrates that optimizing reconstruction for measurement reliability provides more direct support for dimensional and displacement analysis than optimizing visual sharpness alone.",
    )
    set_text(
        p[30],
        "Keywords: Measurement-Oriented Super-Resolution; Point Spread Function Calibration; Physics-Constrained Neural Network; Critical Dimension; Translational Localization; Wiener Deconvolution; Measurement Stability.",
    )

    # Chapter 1: objectives, contributions, and chapter map aligned with the newer method/results chapters.
    set_text(
        p[67],
        "建立以量測可靠性為核心的影像重建與評估框架，使原始影像、傳統反卷積與神經網路重建能在相同邊界、尺寸及位移規則下比較。",
    )
    set_text(
        p[68],
        "結合局部與多位置點擴散函數先驗，建立單張輸入、同尺寸輸出的顯微影像重建模型，降低任意銳化造成幾何位置改寫的風險。",
    )
    set_text(
        p[69],
        "以20倍與50倍影像的邊緣過渡、SEM對照關鍵尺寸、連續位移重複性及九顆pillar空間一致性，驗證重建結果的實際量測效益與適用界線。",
    )
    set_text(
        p[71],
        "本研究的貢獻可歸納為三項彼此連續的成果：首先重新定義顯微影像重建的評估目標；其次將光學先驗、真實影像與幾何監督整合為單張影像重建方法；最後以尺寸與位移量測建立同一條驗證鏈。三項貢獻共同服務於『重建後影像是否仍可可靠量測』這一核心問題。",
    )
    set_text(
        p[72],
        "貢獻一：提出量測導向超解析（MOSR）問題框架。相較於以PSNR、SSIM或主觀清晰度為主要目標的影像重建，MOSR將邊界可定位性、關鍵尺寸誤差、位移重複性與多位置一致性列為主要評估依據，使影像品質的判斷直接連結到後端量測結果。",
    )
    set_text(
        p[73],
        "為使不同方法能被公平比較，本研究建立共同量測流程：所有方法使用相同影像名單、相同pillar身分與中心、相同物體空間換算、相同邊界定義及相同統計方式。另以門檻敏感度、振鈴與多顆結構分歧檢查單一數值改善是否具有穩健性，避免將顯示對比或特定操作點誤當成普遍修復效果。",
    )
    set_text(
        p[74],
        "貢獻二：建立結合光學先驗與幾何穩定性的單張影像重建方法。方法以局部PSF支援區域校正與傳統反卷積比較，並以多位置PSF描述實際系統可能出現的模糊形態；模型透過同尺寸編碼器—解碼器、位移穩定上採樣、受限邊界殘差、原圖基準與品質門控，在提升邊界可辨識性的同時限制非物理高頻與位置漂移。",
    )
    set_text(
        p[75],
        "貢獻三：以尺寸與位移形成端到端驗證。研究不只比較影像外觀，而是先以SEM尺寸檢查CD邊界，再以四組連續位移序列檢查步間散佈、名義指令一致性及九顆pillar的空間分歧。此設計顯示傳統反卷積與神經重建可能在不同單一指標上各有優勢，而所提模型的主要價值在於跨門檻與跨位置時較低的量測敏感性。",
    )
    set_text(
        p[76],
        "上述貢獻的適用範圍以本研究實際資料為界。現有結果可支持同資料域內的邊界集中、局部CD穩健度與位移一致性改善，但不等同於絕對光學解析度提升、獨立位移真值校準或跨樣品泛化。這些限制將在第六章明確說明，並於第七章作為後續研究方向。",
    )
    set_text(
        p[82],
        "第四章說明量測導向重建的研究方法，依序建立成像與物體空間座標、局部與多位置PSF策略、單張影像重建模型、功能分離的訓練資料與損失，以及Raw、Wiener與所提模型共用的尺寸與位移量測規則。",
    )
    set_text(
        p[83],
        "第五章描述顯微光學系統、樣品與SEM參考量測、軟硬體環境、模型架構、訓練資料分組與實際微調設定。",
    )
    set_text(
        p[84],
        "第六章呈現PSF校正與反卷積驗證，並比較原始影像、Wiener反卷積及MOSR模型在邊緣、CD、位移重複性與九宮格空間一致性上的結果。",
    )
    set_text(
        p[85],
        "第七章提出未來展望，說明空間變異光學模型、跨樣品泛化、獨立位移真值、量測不確定度及系統部署等後續方向。",
    )

    # Chapter 2: replace citation placeholders with verified literature and add missing model citations.
    append_text(p[139], " [40]")
    append_text(p[140], " [41]")
    append_text(p[142], " [42]")
    append_text(p[143], " [43]")
    append_text(p[145], " [21]")
    append_text(p[147], " [44]")
    append_text(p[149], " [45]")
    set_text(
        p[153],
        "雖然卷積神經網路已能在視覺品質上顯著改善超解析影像，多數評估仍以均方誤差、峰值訊噪比與結構相似度為主。這些指標可描述像素或整體結構差異，卻不一定能預測重建影像對後續量測、辨識或定位任務的實際效益。感知品質與失真之間可能存在取捨，而顯微重建方法的公平比較也需要與定位精度及軟體行為共同檢查，因此近年研究逐漸將評估延伸至下游任務與定位表現 [30–33]。",
    )
    set_text(
        p[154],
        "對結構化幾何目標與精密位移量測而言，關鍵問題不是影像內容的分類，而是邊界是否能在亞像素尺度被穩定定位。PSNR或SSIM無法直接區分『銳利但位置不穩定』與『視覺改善較小但量測更一致』的重建結果。因而，本研究將影像重建視為量測鏈中的中間步驟，並把尺寸誤差、位移重複性及多位置一致性作為核心判斷依據 [31–33]。",
    )
    set_text(
        p[157],
        "為使深度學習模型在有限資料下保持合理的物理行為，既有研究已採用偏微分方程殘差、成像正向模型、資料一致性或已知光學特性限制模型輸出。傳統物理資訊神經網路以方程殘差約束解空間；在顯微影像重建中，則更常將成像算子、雜訊模型或結構先驗直接納入訓練與驗證 [34–37]。本研究沿用後者思路，將重建影像應滿足的幾何與頻域條件轉寫為可計算的限制。",
    )
    set_text(
        p[158],
        "其中一項重要條件為平移等變性。理想情況下，輸入影像平移後，模型輸出也應產生相同幅度的平移；若下採樣、上採樣或邊界處理破壞此關係，小幅位移便可能被轉換成非預期的形狀或強度變化。既有研究可從群組等變卷積或抗混疊下採樣改善此問題 [38, 39]；本研究則進一步利用已知位移影像對，於訓練與驗證階段直接檢查跨影格幾何關係。",
    )
    set_text(
        p[159],
        "除平移關係外，本研究亦限制高頻能量與邊界微分行為。前者避免模型在光學系統缺乏支持的頻帶中任意增加能量，後者則用來抑制過衝、振鈴與不合理的邊界方向變化。這些限制並不保證輸出必然等同於真實物體，而是縮小模型可能產生的非物理解，使後續尺寸與位移量測較容易維持一致 [37]。",
    )
    set_text(
        p[162],
        "綜合前述各節，PSF建模與校正能提供明確的光學解釋，但傳統反卷積對PSF準確度與雜訊設定敏感；資料驅動模型具備推論效率與非線性表達能力，若只追求視覺指標，卻可能改變量測邊界。量測導向重建因此需要同時結合光學先驗、幾何一致性與後端量測。第四章將據此建立單張影像重建方法，並以多影格位移關係作為訓練與驗證階段的幾何監督，以緩解銳利度與定位穩定性之間的衝突。",
    )

    # Chapter 3: correct the scope description and add a concise bridge to Chapter 4.
    set_text(
        p[164],
        "本章由純量繞射與傅立葉光學出發，說明點擴散函數、像差與系統解析能力的理論基礎，並介紹本研究使用的U-Net架構與BADS最佳化工具。後半部聚焦於平移等變性，說明理想卷積何以保留位移關係，以及下採樣、上採樣、插值與邊界處理如何引入定位誤差，作為第四章重建模型與位移一致性設計的理論基礎。",
    )
    h39 = insert_paragraph_before(p[441], "3.9 小結", style="Heading 2", template=heading2_template)
    b39a = insert_paragraph_before(
        p[441],
        "本章建立由光學成像到數位重建的理論連結。PSF與Zernike像差模型用來描述影像模糊，BADS提供低維參數的實驗校正方式；U-Net則提供多尺度影像特徵整合的基礎。",
        style="Normal",
        template=body_template,
    )
    b39b = insert_paragraph_before(
        p[441],
        "對本研究最重要的是，重建模型即使在單張影像上運作，也必須維持輸入位移與輸出位移之間可預測的關係。第四章將把上述光學模型、平移一致性與後端量測規則整合為量測導向的研究方法。",
        style="Normal",
        template=body_template,
    )
    _ = (h39, b39a, b39b)

    # Normalize chapter headings while preserving the document's established numbering format.
    format_chapter_heading(p[442], heading1_template, "研究方法")
    format_chapter_heading(p[739], heading1_template, "實驗系統與模型實作")
    format_chapter_heading(p[845], heading1_template, "實驗結果與分析")

    # Chapter 5: correct visible numbering, wording, and reproducibility details.
    set_text(p[776], "5.1.4 照明模組與光源", style="Heading 3")
    set_text(
        p[752],
        "然而，數位取樣間距僅代表感測器在物體空間中的取樣尺度，並不等同於系統可分辨的最小結構尺寸。實際解析能力仍受到物鏡數值孔徑、照明光譜、繞射效應與像差影響。為便於比較，本研究以450 nm作為解析度估算的參考波長；該數值不是指實驗使用單一450 nm窄頻光源。依瑞利準則估算，5倍、20倍與50倍物鏡的理論橫向解析度分別約為1830、686.3與366 nm。三種倍率的數位取樣間距均小於相應估算值，因此感測器取樣不是本系統解析能力的唯一限制。",
    )
    set_text(
        p[758],
        "圖5-5為本研究所使用的Olympus U-TLU 1×筒鏡。系統採用同系列筒鏡與物鏡，以維持無限遠校正光路的相容性；其焦距、視場數、入瞳直徑與可見光適用範圍依製造商規格整理於表5-4。",
    )
    set_text(p[777], p[777].text.replace("圖5-9", "圖5-10"))
    set_text(p[779], p[779].text.replace("圖5-4", "圖5-11"))
    set_text(p[780], p[780].text.replace("圖5-5", "圖5-12"))
    set_text(p[785], p[785].text.replace("表5-7", "表5-8"))
    set_text(
        p[790],
        p[790].text.replace("圖5-10", "圖5-14").replace("[70]", ""),
    )
    set_text(p[791], p[791].text.replace("表5-6", "表5-9"))
    set_text(
        p[796],
        p[796].text.replace("[72]", "").replace("[73]", "").replace("圖5-12", "圖5-15"),
    )
    set_text(
        p[800],
        "本節說明研究使用的標準樣品與獨立參考儀器。HS-100MG提供規則排列的微結構，場發射槍掃描式電子顯微鏡則提供較高解析度的邊界與尺寸參考，兩者共同支援後續PSF擬合及CD驗證。",
    )
    set_text(p[802], p[802].text.replace("圖5-15", "圖5-16"))
    set_text(p[803], p[803].text.replace("圖5-16", "圖5-17"))
    set_text(p[808], "5.2.2 場發射槍掃描式電子顯微鏡", style="Heading 3")
    set_text(p[809], p[809].text.replace("圖5-19", "圖5-18"))
    set_text(p[814], "5.3 MATLAB環境", style="Heading 2")
    set_text(p[815], p[815].text.replace("表5-10", "表5-11"))
    set_text(p[819], p[819].text.replace("表5-11", "表5-12"))
    set_text(p[825], p[825].text.replace("圖5-22", "圖5-19"))
    set_text(p[826], p[826].text.replace("表5-14", "表5-13"))
    set_text(p[829], p[829].text.replace("圖5-23", "圖5-20"))

    # Remove duplicated model-description paragraphs.
    delete_paragraph(p[834])
    delete_paragraph(p[835])

    # Fill the missing camera entry in the system-component table.
    for row in tables[82].rows:
        if row.cells and row.cells[0].text.strip() == "相機":
            set_cell_text(row.cells[1], "Basler ace acA2040-120um")

    # Chapter 5 figure captions, keeping one figure per visual and a restrained academic style.
    set_caption(p[743], "圖 5-1 實驗系統整體架構")
    set_caption(p[746], "圖 5-2 實驗系統各元件配置")
    ch5_inserted_captions = {
        754: "圖 5-3 Basler ace acA2040-120um工業相機",
        756: "圖 5-4 Sony IMX252LLR-C感測器光譜響應",
        760: "圖 5-5 Olympus U-TLU 1×筒鏡",
        763: "圖 5-6 Olympus U-TLU筒鏡可見光穿透率",
        769: "圖 5-7 Olympus MPLFLN5X2物鏡",
        771: "圖 5-8 Olympus MPLN20X物鏡",
        774: "圖 5-9 Olympus MPLN50X物鏡",
        778: "圖 5-10 Olympus BX3M-RLA-S反射光照明模組",
        781: "圖 5-11 Olympus BX3M-LEDR白光LED光源",
        782: "圖 5-12 Olympus BX3M-LEDR光源光譜分布",
        786: "圖 5-13 XY軸壓電位移平台",
        792: "圖 5-14 PI PIFOC P-725.4CD Z軸壓電位移平台",
        797: "圖 5-15 樣品粗調與傾斜校正平台",
        804: "圖 5-16 BudgetSensors HS-100MG標準樣品",
        806: "圖 5-17 HS-100MG微結構分布",
        811: "圖 5-18 JEOL JSM-7800F Prime場發射槍掃描式電子顯微鏡",
        823: "圖 5-19 所提MOSR模型整體架構",
        840: "圖 5-20 殘差式PSF修正模組",
    }
    for index, caption in ch5_inserted_captions.items():
        new_caption = insert_paragraph_after(p[index], caption, style="Caption", template=p[97])
        new_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Chapter 5 table captions.
    set_table_caption(p[747], "表 5-1 實驗系統元件與編號")
    set_table_caption(p[748], "表 5-2 實驗系統元件製造商與型號")
    table_caption_map = {
        755: "表 5-3 Basler ace acA2040-120um相機規格",
        762: "表 5-4 Olympus U-TLU筒鏡規格",
        770: "表 5-5 Olympus MPLFLN5X2物鏡規格",
        773: "表 5-6 Olympus MPLN20X物鏡規格",
        775: "表 5-7 Olympus MPLN50X物鏡規格",
        788: "表 5-8 XY軸壓電位移平台規格",
        794: "表 5-9 Z軸壓電位移平台規格",
        813: "表 5-10 JEOL JSM-7800F Prime規格",
        816: "表 5-11 MATLAB與工具箱版本",
        820: "表 5-12 Python與深度學習套件版本",
        830: "表 5-13 編碼器層級參數",
        831: "表 5-14 瓶頸層參數",
        832: "表 5-15 解碼器、PSF修正與輸出層參數",
    }
    for index, caption in table_caption_map.items():
        set_table_caption(p[index], caption)

    # Repurpose existing blank paragraphs and add the missing concrete training description.
    set_text(p[842], "5.4.3 訓練資料組成與分組切分", style="Heading 3", template=heading3_template)
    set_text(
        p[843],
        "模型使用11,254組影像配對進行微調，其中9,857組作為訓練資料，1,397組作為驗證資料。資料依功能涵蓋外觀重建、SEM幾何錨點、連續位移、合成運動、真實多倍率影像與保守性樣本；不同來源分別提供外觀、邊界、運動與不破壞原圖等監督。為降低相鄰影格或同源合成樣本同時進入訓練與驗證集合造成的資料洩漏，切分時依資料角色、樣品來源、倍率與結構類型共同分組，並固定兩集合的索引名單。",
        style="Normal",
        template=body_template,
    )
    set_text(
        p[844],
        "完整資料中另選取180組樣本進行邊界軟標籤調整，涵蓋15類資料、每類12組。軟化只作用於標註邊界，原始輸入影像保持不變；處理採2-pixel邊界帶與0.75-pixel標準差的局部平滑，以降低理想硬邊界在較粗取樣影像中形成階梯狀監督的風險。這180組不是額外測試資料，而是既有訓練標籤的局部修正。",
        style="Normal",
        template=body_template,
    )
    insert_paragraph_before(
        p[845], "5.4.4 微調設定與模型選擇", style="Heading 3", template=heading3_template
    )
    insert_paragraph_before(
        p[845],
        "模型由既有穩定權重開始微調，保留大部分主幹參數，主要更新品質門控、邊界修正與輸出相關層，以降低有限真實資料使整體網路過度漂移的風險。最佳化採Adam，初始學習率為2×10⁻⁷，影像批次為8，驗證批次與序列配對批次皆為2，共訓練120回合。訓練過程同步記錄重建、梯度幅值、邊緣方向與平移一致性相關項目，並以固定監測影像檢查灰階範圍、邊緣寬度及高頻比例。",
        style="Normal",
        template=body_template,
    )
    insert_paragraph_before(
        p[845],
        "第六章固定使用第120回合權重進行評估，不依CD或位移結果事後挑選較有利的回合。訓練曲線僅用來確認最佳化過程是否收斂；模型是否具有量測效益，仍由第六章的邊緣、CD與位移結果判定。",
        style="Normal",
        template=body_template,
    )
    insert_paragraph_before(p[845], "5.5 小結", style="Heading 2", template=heading2_template)
    insert_paragraph_before(
        p[845],
        "本章說明實驗所使用的顯微光學系統、標準樣品、SEM參考量測與軟硬體環境，並將相機取樣尺度、物鏡條件與白光照明對後續PSF及尺寸分析的影響加以區分。",
        style="Normal",
        template=body_template,
    )
    insert_paragraph_before(
        p[845],
        "模型部分則補充實際資料數量、分組切分、邊界軟標籤及微調設定。由於最終評估模型在查看第六章量測結果前即固定，後續Raw、Wiener與MOSR的比較可使用同一影像名單與量測規則進行。",
        style="Normal",
        template=body_template,
    )

    # Chapter 6: formalize headings and remove editorial/internal naming.
    set_text(p[846], "6.1 點擴散函數校正", style="Heading 2")
    set_text(p[858], "6.1.1 點擴散函數擬合結果", style="Heading 3")
    set_text(p[865], p[865].text.replace("圖 6-X", "圖6-4與圖6-5"))
    set_text(
        p[873],
        "本節先以SEM尺寸錨點參與參數校準的設定，檢查局部擬合PSF在原擬合區域內可達到的反卷積修復效果。由於SEM尺寸已參與雜訊參數選擇，此結果屬於校正內部的可行性驗證，不用來與所提MOSR模型進行獨立排名，也不延伸宣稱反卷積普遍優於神經重建。",
    )
    set_text(
        p[874],
        "第6.2節另以未納入正式位移結果的校正影格決定Wiener參數，20倍與50倍的固定NSR分別為0.101895與0.079656，選擇過程不使用SEM CD或位移結果。兩節分別回答『局部PSF是否具有修復潛力』及『在保留資料上三種方法如何比較』，因此使用不同目的的參數設定。",
    )
    set_text(p[875], "6.1.2.1 驗證設定與判讀原則", style="Heading 4", template=heading4_template)
    for index, label in {
        881: "CD與邊緣寬度的量測方式",
        888: "對比門檻的意義",
        892: "評估指標與判讀方式",
        915: "20倍影像",
        925: "50倍影像",
        936: "兩種倍率的合併結果",
    }.items():
        set_text(p[index], label, style="Normal", bold=True)
    for index, label in {
        902: "6.1.2.2 銳利度驗證：反卷積前後的邊緣過渡寬度",
        912: "6.1.2.3 反卷積前後的CD量測結果",
        944: "6.1.2.4 不同對比門檻下的結果穩定性",
        954: "6.1.2.5 反卷積振鈴現象",
        964: "6.1.2.6 綜合討論與研究限制",
    }.items():
        set_text(p[index], label, style="Heading 4", template=heading4_template)

    set_text(
        p[974],
        "本節評估所提MOSR模型在HS-100MG樣品影像上的實際重建效果，並固定使用完成訓練後的第120回合權重。為判斷改善是否只來自一般影像銳化，另加入使用擬合PSF的Wiener反卷積作為傳統基準。三種方法使用同一批原始影像、相同九顆pillar、相同120個位移步與相同量測規則，分別檢查邊緣過渡、SEM對照CD、連續位移重複性及同一步內九顆pillar的一致性。",
    )
    set_text(
        p[975],
        "本節模型採用第五章所述的固定權重與推論設定。第6.1節的局部PSF擬合用於光學校正與Wiener對照；模型內部使用的光學先驗則依第四、五章的方法建立。兩者提供互補的光學校正與影像重建證據，但本文不將第6.1節的局部擬合PSF與模型輸出建立未經驗證的直接因果關係。",
    )
    set_text(
        p[976],
        "Wiener基準的PSF與雜訊參數只由未納入正式30步結果的校正影格決定，選擇時不使用SEM CD或位移誤差。20倍與50倍各使用一個固定PSF套用至兩種位移幅度；MOSR與Wiener均不重新挑選正式分析步數，使後續差異不受資料名單改變影響。",
    )
    set_text(p[977], "6.2.1 模型與訓練狀態", style="Heading 3")
    set_text(
        p[978],
        "模型訓練資料與切分方式已於第五章說明。本節固定使用第120回合模型，並不將其描述為每一項指標上的最佳回合。模型的品質判讀與邊界修正分支用於降低硬邊界階梯化及不同訊雜條件造成的輸出偏移；其效果仍須由後續影像、CD與位移結果共同判定。",
    )
    set_text(p[1039], "6.2.5 位移量測結果", style="Heading 3", template=heading3_template)

    # Replace remaining internal model code in body text and tables.
    visible_replacements = [
        ("V42 SoftLabelSNR Full Standalone", "所提MOSR"),
        ("V42 epoch 120", "MOSR（最終模型）"),
        ("V42 第 120 回合模型", "MOSR模型"),
        ("V42 第 1 至第 120 回合", "MOSR模型第1至第120回合"),
        ("V42", "MOSR"),
    ]
    for paragraph in doc.paragraphs:
        replace_in_runs(paragraph, visible_replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_runs(paragraph, visible_replacements)

    # Chapter 6 figure captions.
    ch6_figure_caption_indices = [
        850, 852, 862, 868, 870, 887, 910, 921, 923, 932, 934, 943, 961,
        981, 988, 990, 993, 1005, 1007, 1009, 1012, 1014, 1016, 1018,
        1021, 1027, 1032, 1035, 1037, 1043, 1046, 1049, 1052, 1061,
        1063, 1065, 1068, 1079, 1081, 1087, 1090, 1092,
    ]
    for number, index in enumerate(ch6_figure_caption_indices, start=1):
        title = re.sub(r"^圖\s*6[-－]?\w*\s*", "", p[index].text).strip()
        set_caption(p[index], f"圖 6-{number} {title}")

    # Chapter 6 table captions.
    ch6_table_captions = {
        920: "表 6-1 20倍影像反卷積前後之CD誤差摘要",
        930: "表 6-2 50倍影像反卷積前後之CD誤差摘要",
        952: "表 6-3 不同對比門檻下之CD平均絕對誤差變化與改善數量",
        966: "表 6-4 兩種倍率之反卷積驗證結果摘要",
        982: "表 6-5 MOSR模型訓練損失摘要",
        994: "表 6-6 四組位移序列之取樣條件與主分析步數",
        1024: "表 6-7 50%名義門檻下中心九顆pillar之CD比較",
        1029: "表 6-8 跨兩種剖面與七個門檻之CD誤差穩健度",
        1053: "表 6-9 四組序列之步階穩定性與名義指令一致性",
        1055: "表 6-10 四組序列之中心九顆pillar內部一致性",
        1069: "表 6-11 全部配對步與分倍率之位移結果",
        1074: "表 6-12 四組序列之第90百分位誤差",
    }
    for index, caption in ch6_table_captions.items():
        set_table_caption(p[index], caption)

    # Add a dedicated Chapter 6 summary before the future-work chapter.
    set_text(p[1097], "6.3 本章小結", style="Heading 2", template=heading2_template)
    insert_paragraph_before(
        p[1098],
        "局部PSF反卷積結果證明擬合核具有使邊緣集中與降低部分CD誤差的潛力，但其效果會受到振鈴、門檻選擇與空間位置影響。因此，PSF驗證的合理結論是確認局部修復能力與限制，而不是宣稱完成無失真的絕對尺寸還原。",
        style="Normal",
        template=body_template,
    )
    insert_paragraph_before(
        p[1098],
        "在相同120個位移步的比較中，Wiener與MOSR對名義指令的平均偏差接近，但MOSR將九顆pillar內部不一致量進一步降低至6.47 nm，形成比單張銳利度更直接的量測證據。現有結果支持同資料域內的邊界、CD與位移穩定性改善；跨樣品泛化與絕對位移準確度則留待後續研究。",
        style="Normal",
        template=body_template,
    )

    # Chapter 7: future outlook, intentionally distinct from the Chapter 6 conclusion.
    chapter7_heading = insert_paragraph_before(p[1099], template=heading1_template)
    format_chapter_heading(chapter7_heading, heading1_template, "未來展望")
    insert_paragraph_before(
        p[1099],
        "本研究已建立由光學校正、影像重建到尺寸與位移量測的完整方法鏈，並在HS-100MG樣品上驗證其局部可行性。後續工作的重點不是單純追求更強的視覺銳化，而是擴大光學條件、樣品種類與獨立真值，使量測結果能由同資料域驗證逐步走向可追溯的實際應用。",
        style="Normal",
        template=body_template,
    )
    future_sections = [
        (
            "7.1 空間變異光學模型的完整量化",
            [
                "目前局部PSF與多位置PSF提供了描述成像模糊的兩種尺度。未來可在不同視野位置、焦平面、曝光與照明條件下建立更密集的校正資料，並以重複拍攝估計每一位置PSF參數的變異範圍。如此可區分真正的空間像差、短期對焦變動與擬合不確定度，避免將所有差異都歸因於單一模糊核。",
                "後續亦可將多PSF融合權重與影像位置、局部清晰度及量測殘差共同記錄，檢查模型是否在合理條件下選擇相符的光學先驗。評估仍應以邊界、CD與位移穩定性為主，而不是只比較PSF外觀或頻譜能量。",
            ],
        ),
        (
            "7.2 跨倍率、跨樣品與跨系統泛化",
            [
                "本研究的主要量化結果集中於同一HS-100MG樣品及20倍、50倍成像條件。未來應加入不同批次的標準樣品、不同尺寸的pillar與via、線條及方形結構，並保留完整的樣品分組測試，使訓練、驗證與測試資料在實體樣品層級互相獨立。",
                "若模型進一步應用於不同相機、物鏡或照明模組，應先重新建立物體空間尺度與PSF資料，再檢查模型是否需要完整重訓或只需少量校正。此類實驗可明確區分模型對結構幾何的泛化能力與對特定光學系統的依賴程度。",
            ],
        ),
        (
            "7.3 獨立位移真值與量測不確定度",
            [
                "現有位移結果以平台名義指令作為一致性參考，足以比較步間散佈與九顆pillar分歧，但不能取代獨立位移真值。未來可加入雷射干涉儀、外部電容式感測器或具校正證書的編碼器，同步記錄每一步的實際位移，進一步分離平台誤差、影像定位誤差與重建造成的偏差。",
                "在取得獨立真值後，可建立包含重複性、系統偏差、像素尺度、SEM參考與模型變動的量測不確定度預算，並報告不同倍率與不同位置的擴充不確定度。如此才能把『結果較穩定』推進為『結果可追溯且具有明確可信範圍』。",
            ],
        ),
        (
            "7.4 模型選擇與量測指標的整合",
            [
                "後續模型不應只依訓練損失或單張影像品質選擇，而可建立由邊緣寬度、跨門檻CD穩健度、振鈴比例、位移重複性與多顆空間一致性共同構成的驗證準則。各項指標需使用預先固定的資料與門檻，避免模型選擇受到測試結果反向影響。",
                "此外，可讓模型同時輸出重建影像與量測可信度，標示低對比、邊界不完整、超出訓練範圍或可能振鈴的區域。當可信度不足時，系統應保留原始影像、降低修正強度或要求重新取像，使模型成為量測流程中的受控工具，而不是無條件取代原始資料。",
            ],
        ),
        (
            "7.5 自動化流程與實際部署",
            [
                "為提升實務使用效率，未來可將樣品方向判定、pillar與via分區、中心九宮格定位、影像品質檢查、重建及後端量測整合為自動化流程。所有方法仍應保存相同灰階顯示範圍、量測中心與門檻設定，並同步輸出排除影格及排除原因，以維持結果可稽核性。",
                "在運算層面，可進一步評估單張推論時間、記憶體需求及連續取像吞吐量，並以固定模型與固定PSF資料建立部署版本。若能在不犧牲幾何一致性的前提下達到近即時處理，MOSR可望應用於顯微定位、標準片自動檢查及製程量測前處理。",
            ],
        ),
        (
            "7.6 小結",
            [
                "未來研究的核心方向，是將目前已觀察到的局部量測改善擴展為跨位置、跨樣品與跨系統均可驗證的能力。隨著獨立位移真值、完整不確定度預算及自動品質控管加入，量測導向超解析可由影像重建方法進一步發展為具有可追溯性與明確適用界線的顯微量測工具。",
            ],
        ),
    ]
    for heading, bodies in future_sections:
        insert_paragraph_before(p[1099], heading, style="Heading 2", template=heading2_template)
        for body in bodies:
            insert_paragraph_before(p[1099], body, style="Normal", template=body_template)

    # Reference list starts on a new page after Chapter 7.
    add_page_break_before(p[1099], p[1098])

    # Add the literature used to replace placeholders in Chapter 2.
    references = [
        '[30] Z. Wang and A. C. Bovik, "Mean squared error: Love it or leave it? A new look at signal fidelity measures," IEEE Signal Processing Magazine, vol. 26, no. 1, pp. 98–117, 2009.',
        '[31] Y. Blau and T. Michaeli, "The perception-distortion tradeoff," in 2018 IEEE/CVF Conference on Computer Vision and Pattern Recognition, 2018, pp. 6228–6237.',
        '[32] D. Sage et al., "Super-resolution fight club: assessment of 2D and 3D single-molecule localization microscopy software," Nature Methods, vol. 16, no. 5, pp. 387–395, 2019.',
        '[33] A. Speiser et al., "Deep learning enables fast and dense single-molecule localization with high accuracy," Nature Methods, vol. 18, no. 9, pp. 1082–1090, 2021.',
        '[34] M. Raissi, P. Perdikaris, and G. E. Karniadakis, "Physics-informed neural networks: A deep learning framework for solving forward and inverse problems involving nonlinear partial differential equations," Journal of Computational Physics, vol. 378, pp. 686–707, 2019.',
        '[35] M. Weigert et al., "Content-aware image restoration: pushing the limits of fluorescence microscopy," Nature Methods, vol. 15, no. 12, pp. 1090–1097, 2018.',
        '[36] E. Nehme, L. E. Weiss, T. Michaeli, and Y. Shechtman, "Deep-STORM: super-resolution single-molecule microscopy by deep learning," Optica, vol. 5, no. 4, pp. 458–464, 2018.',
        '[37] C. Zuo et al., "Deep learning in optical metrology: a review," Light: Science & Applications, vol. 11, art. no. 39, 2022.',
        '[38] T. S. Cohen and M. Welling, "Group equivariant convolutional networks," in Proceedings of the 33rd International Conference on Machine Learning, 2016, pp. 2990–2999.',
        '[39] R. Zhang, "Making convolutional networks shift-invariant again," in Proceedings of the 36th International Conference on Machine Learning, 2019, pp. 7324–7333.',
        '[40] C. Dong, C. C. Loy, K. He, and X. Tang, "Learning a deep convolutional network for image super-resolution," in Computer Vision – ECCV 2014, 2014, pp. 184–199.',
        '[41] C. Dong, C. C. Loy, and X. Tang, "Accelerating the super-resolution convolutional neural network," in Computer Vision – ECCV 2016, 2016, pp. 391–407.',
        '[42] K. He, X. Zhang, S. Ren, and J. Sun, "Deep residual learning for image recognition," in 2016 IEEE Conference on Computer Vision and Pattern Recognition, 2016, pp. 770–778.',
        '[43] J. Kim, J. K. Lee, and K. M. Lee, "Accurate image super-resolution using very deep convolutional networks," in 2016 IEEE Conference on Computer Vision and Pattern Recognition, 2016, pp. 1646–1654.',
        '[44] B. Lim, S. Son, H. Kim, S. Nah, and K. M. Lee, "Enhanced deep residual networks for single image super-resolution," in 2017 IEEE Conference on Computer Vision and Pattern Recognition Workshops, 2017, pp. 1132–1140.',
        '[45] J. Liang, J. Cao, G. Sun, K. Zhang, L. Van Gool, and R. Timofte, "SwinIR: Image restoration using Swin Transformer," in 2021 IEEE/CVF International Conference on Computer Vision Workshops, 2021, pp. 1833–1844.',
    ]
    anchor = p[1128]
    for reference in references:
        anchor = insert_paragraph_after(
            anchor, reference, style="EndNote Bibliography", template=reference_template
        )

    # Ask Word to refresh all fields on opening as an additional safeguard.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
