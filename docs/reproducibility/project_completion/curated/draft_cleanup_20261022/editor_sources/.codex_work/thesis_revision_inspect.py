from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def main() -> None:
    src = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(src)
    paragraphs = []
    styles = Counter()
    for index, para in enumerate(doc.paragraphs):
        text = clean(para.text)
        style = para.style.name if para.style is not None else ""
        styles[style] += 1
        paragraphs.append(
            {
                "index": index,
                "style": style,
                "text": text,
                "runs": [clean(run.text) for run in para.runs if clean(run.text)],
            }
        )

    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text) for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    payload = {
        "source": str(src),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "section_count": len(doc.sections),
        "styles": styles.most_common(),
        "paragraphs": paragraphs,
        "tables": tables,
    }
    (out_dir / "document.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    patterns = re.compile(
        r"第[一二三四五六七]章|^[1-7](?:\.\d+){0,3}\s|\[REF:|錯誤!|尚未定義書籤|"
        r"V42|epoch\s*120|圖\s*6[-－]?X|本節定位修正|未來展望|參考文獻",
        re.IGNORECASE,
    )
    with (out_dir / "matches.tsv").open("w", encoding="utf-8", newline="") as handle:
        handle.write("index\tstyle\ttext\n")
        for item in paragraphs:
            if patterns.search(item["text"]):
                handle.write(
                    f'{item["index"]}\t{item["style"]}\t{item["text"].replace(chr(9), " ")}\n'
                )

    with (out_dir / "tables.tsv").open("w", encoding="utf-8", newline="") as handle:
        for item in tables:
            handle.write(f'TABLE {item["index"]}\n')
            for row in item["rows"]:
                handle.write("\t".join(cell.replace(chr(9), " ") for cell in row) + "\n")
            handle.write("\n")

    print(json.dumps({
        "paragraphs": len(paragraphs),
        "tables": len(tables),
        "sections": len(doc.sections),
        "top_styles": styles.most_common(15),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
