from __future__ import annotations

import sys
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


CHAPTER_TITLES = {
    "緒論",
    "文獻回顧",
    "理論背景",
    "研究方法",
    "實驗系統與模型實作",
    "實驗結果與分析",
    "未來展望",
    "參考文獻",
}


def remove_highlight_and_shading(element) -> None:
    for tag in (qn("w:highlight"), qn("w:shd")):
        for node in list(element.iter(tag)):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(path)

    found = set()
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in CHAPTER_TITLES:
            paragraph.paragraph_format.page_break_before = True
            found.add(text)
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            remove_highlight_and_shading(paragraph._p)
            heading_runs = [run for run in paragraph.runs if run.text]
            for run_index, run in enumerate(heading_runs):
                if run.text:
                    before = run.text
                    run.text = re.sub(
                        r"^(\d+(?:\.\d+)+)[ \u3000]*",
                        r"\1 ",
                        run.text,
                        count=1,
                    )
                    if run.text != before or re.fullmatch(r"\d+(?:\.\d+)+ ", run.text):
                        if re.fullmatch(r"\d+(?:\.\d+)+ ", run.text):
                            for following in heading_runs[run_index + 1 :]:
                                following.text = following.text.lstrip(" \u3000")
                                if following.text:
                                    break
                    break
        if paragraph.style and paragraph.style.name == "Figure Caption":
            for run in paragraph.runs:
                if run.text:
                    run.text = re.sub(r"^(圖\s+\d+-\d+)(?=\S)", r"\1 ", run.text, count=1)
                    break

    toc_heading_style = doc.styles["TOC Heading"]
    figure_title = next(p for p in doc.paragraphs if p.text.strip() == "圖目錄")
    figure_title.style = toc_heading_style

    main_toc_title = next(
        p
        for p in doc.element.body.xpath(".//w:sdt//w:p")
        if "".join(p.itertext()).strip().startswith("目錄")
    )
    ppr = main_toc_title.get_or_add_pPr()
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        from docx.oxml import OxmlElement

        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), toc_heading_style.style_id)

    missing = CHAPTER_TITLES - found
    if missing:
        raise RuntimeError(f"Missing chapter headings: {sorted(missing)}")

    settings = doc.settings.element
    remove_highlight_and_shading(settings)
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
