from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def delete_paragraph(paragraph) -> None:
    parent = paragraph._p.getparent()
    parent.remove(paragraph._p)


def add_toc_field(paragraph, instruction: str) -> None:
    clear_paragraph(paragraph)
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run.append(begin)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run.append(instr)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run.append(separate)
    placeholder = OxmlElement("w:t")
    placeholder.text = "更新欄位後產生圖目錄"
    run.append(placeholder)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)
    paragraph._p.append(run)


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(path)
    paragraphs = list(doc.paragraphs)

    first = next(i for i, p in enumerate(paragraphs) if p.text.startswith("圖 2-1 Born"))
    last = first
    while last + 1 < len(paragraphs) and paragraphs[last + 1].style.name == "table of figures":
        last += 1
    title = paragraphs[first - 2]
    field = paragraphs[first - 1]

    # Reuse the existing main table-of-contents heading formatting.
    toc_heading = next(
        p
        for p in doc.element.body.xpath(".//w:sdt//w:p")
        if "".join(p.itertext()).strip().startswith("目錄")
    )
    old_ppr = title._p.pPr
    if old_ppr is not None:
        title._p.remove(old_ppr)
    source_ppr = toc_heading.find(qn("w:pPr"))
    if source_ppr is not None:
        title._p.insert(0, copy.deepcopy(source_ppr))
    clear_paragraph(title)
    run = title.add_run("圖目錄")
    source_run = toc_heading.find(qn("w:r"))
    if source_run is not None and source_run.find(qn("w:rPr")) is not None:
        run._r.insert(0, copy.deepcopy(source_run.find(qn("w:rPr"))))

    add_toc_field(field, ' TOC \\h \\z \\t "Caption,1" ')

    for paragraph in paragraphs[first : last + 1]:
        delete_paragraph(paragraph)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
